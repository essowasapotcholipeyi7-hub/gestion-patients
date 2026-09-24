from flask import Flask, render_template, redirect, url_for, flash, request, session, jsonify
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from datetime import datetime
import os
from dotenv import load_dotenv
from scheduler import start_scheduler, stop_scheduler
from datetime import datetime, timedelta
import requests
from bs4 import BeautifulSoup
import csv
from io import StringIO
from decorators import has_permission
import uuid
import logging
import re
from models import db, Engagement, Patient
from models import Consultation, Patient, ExamenType, ExamenPrescrit



load_dotenv()

# Initialisation de l'application
app = Flask(__name__)
app.config.from_object('config.Config')
logger = logging.getLogger(__name__)

# Initialisation de la base de données
from models import db
db.init_app(app)

from routes.engagements import engagements_bp
app.register_blueprint(engagements_bp)

def create_structure_sheets(structure_id):
    """Crée automatiquement les feuilles Google Sheets pour une nouvelle structure"""
    from sheets_sync import GoogleSheetsSync
    import os
    
    SPREADSHEET_ID = "1nCUArOaWgXVFszjEhH1GqNJXGCV7cF754W87vXvQ-lQ"
    CREDENTIALS_FILE = 'credentials.json'
    
    if os.path.exists(CREDENTIALS_FILE):
        syncer = GoogleSheetsSync(SPREADSHEET_ID, CREDENTIALS_FILE)
        if syncer.authenticate():
            syncer.ensure_structure_sheets(structure_id)
            print(f"✅ Feuilles Sheets créées pour structure ID {structure_id}")
            return True
    return False

# Initialisation du login manager
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'
login_manager.login_message = 'Veuillez vous connecter pour accéder à cette page'

@login_manager.user_loader
def load_user(user_id):
    from models import db, Utilisateur
    return db.session.get(Utilisateur, int(user_id))

@app.context_processor
def inject_non_lus():
    from models import Message
    if current_user.is_authenticated:
        non_lus = Message.query.filter_by(id_destinataire=current_user.id, lu=False).count()
        return dict(non_lus=non_lus)
    return dict(non_lus=0)

@app.context_processor
def inject_medicaments_en_retard():
    """Badge rouge dans la sidebar (voir base.html, lien "Médicaments à
    administrer") — doses dont l'heure prévue est dépassée, recalculé à
    chaque page comme inject_non_lus ci-dessus (même limite : pas de temps
    réel, juste un compteur à jour au chargement de la page)."""
    from models import AdministrationMedicament, Patient
    if current_user.is_authenticated and current_user.role in ['admin_structure', 'medecin', 'infirmier']:
        try:
            count = (
                AdministrationMedicament.query
                .join(Patient, AdministrationMedicament.patient_id == Patient.id)
                .filter(
                    Patient.id_structure == current_user.id_structure,
                    AdministrationMedicament.statut == 'a_faire',
                    AdministrationMedicament.heure_prevue <= datetime.utcnow(),
                )
                .count()
            )
            return dict(medicaments_en_retard_count=count)
        except Exception:
            return dict(medicaments_en_retard_count=0)
    return dict(medicaments_en_retard_count=0)

@app.context_processor
def inject_demandes_acces_en_attente():
    """Badge nav "Demandes d'accès" (voir base.html) — même patron que
    inject_non_lus ci-dessus."""
    from models import DemandeAccesHospitalisation, Hospitalisation, Patient
    if not current_user.is_authenticated or current_user.role not in ('medecin', 'admin_structure'):
        return dict(demandes_acces_en_attente_count=0)
    try:
        if current_user.role == 'medecin':
            count = DemandeAccesHospitalisation.query.filter_by(
                destinataire_id=current_user.id, statut='en_attente'
            ).count()
        else:
            count = DemandeAccesHospitalisation.query.filter(
                DemandeAccesHospitalisation.destinataire_id.is_(None),
                DemandeAccesHospitalisation.statut == 'en_attente'
            ).join(Hospitalisation).join(Patient).filter(
                Patient.id_structure == current_user.id_structure
            ).count()
        return dict(demandes_acces_en_attente_count=count)
    except Exception:
        return dict(demandes_acces_en_attente_count=0)

@app.context_processor
def utility_processor():
    from datetime import datetime
    return {
        'now': datetime.now(),
        'fichier_source_info': _fichier_source_info,
    }

@app.template_filter('nl2br')
def nl2br_filter(text):
    """Convertit les sauts de ligne en <br>"""
    if not text:
        return text
    return text.replace('\n', '<br>')

@app.template_filter('from_json')
def from_json_filter(value):
    """Convertit une chaîne JSON en objet Python"""
    import json
    try:
        return json.loads(value) if value else []
    except:
        return []


# ============================================================
# ⭐ MIROIR VERS GHP DEPUIS LES ORDONNANCES/EXAMENS VERSIONNÉS
# ============================================================
# La table Prescription (avec synced_at) + tasks.sync_prescriptions_to_ghp()
# est le SEUL mécanisme qui envoie effectivement vers GHP. Or les écrans de
# prescription réels (ordonnances versionnées avec modèles/protocoles,
# examens prescrits) écrivent dans Ordonnance/ExamenPrescrit — ou, pour
# l'hospitalisation, directement dans Hospitalisation.ordonnance_prescite —
# et ne touchent JAMAIS Prescription. Résultat : rien de ce qui est
# réellement prescrit en pratique ne partait vers GHP. On crée donc ici,
# EN PLUS, une ligne Prescription (miroir) par médicament/acte, pour que
# le pipeline de sync existant (déjà testé, avec rattrapage automatique)
# les prenne en charge sans rien avoir à changer côté GHP.
def _creer_prescriptions_miroir(patient_id, prescripteur_nom, items, type_prescription, id_consultation=None):
    """Crée des lignes Prescription à partir d'une liste d'items (médicaments
    sous forme de dict, ou actes sous forme de chaînes) — ne commit pas,
    l'appelant doit le faire avec le reste de sa transaction."""
    from models import Prescription

    def _duree_en_jours(valeur):
        try:
            return int(valeur)
        except (TypeError, ValueError):
            return 7

    creees = []
    for item in items or []:
        if isinstance(item, dict):
            nom = (item.get('medicament') or item.get('nom') or '').strip()
            if not nom:
                continue
            p = Prescription(
                id_patient=patient_id,
                id_consultation=id_consultation,
                medicament=nom,
                dosage=item.get('dosage', '') or '',
                forme=item.get('forme', '') or '',
                quantite=str(item.get('quantite', '1')) or '1',
                duree_jours=_duree_en_jours(item.get('duree')),
                frequence=item.get('posologie') or item.get('frequence') or '',
                instructions=item.get('instructions', '') or '',
                type_prescription=type_prescription,
                prescripteur=prescripteur_nom,
                statut='active',
                date_prescription=datetime.utcnow(),
                notes=item.get('notes', '') or ''
            )
        else:
            nom = str(item).strip()
            if not nom:
                continue
            p = Prescription(
                id_patient=patient_id,
                id_consultation=id_consultation,
                medicament=nom,
                type_prescription=type_prescription,
                prescripteur=prescripteur_nom,
                statut='active',
                date_prescription=datetime.utcnow()
            )
        db.session.add(p)
        creees.append(p)
    return creees


def _creer_analyses_demandees_depuis_examen_prescrit(examen_prescrit, items, structure_id):
    """Crée une AnalyseDemande par item d'un ExamenPrescrit (consultation ou
    hospitalisation), pour que laborantin/radiologue le voient dans leur
    file (/analyses) — sans ce pont, un examen prescrit pendant une
    consultation/hospitalisation restait invisible pour eux (leur tableau
    de bord ne lit que AnalyseDemande), alors qu'il partait déjà vers les
    prescriptions reçues de GHP via _creer_prescriptions_miroir. Certains
    chemins de création (hospitalisation directe, création d'examen en
    consultation) le faisaient déjà chacun à leur façon sans lien fiable
    vers l'ExamenPrescrit d'origine ; celle-ci pose explicitement
    examen_prescrit_id pour fiabiliser saisir_resultats_examen ensuite.
    `examen_prescrit.id` doit déjà exister (flush l'ExamenPrescrit avant
    d'appeler cette fonction). Ne commit pas, l'appelant s'en charge."""
    from models import AnalyseDemande

    type_analyse = examen_prescrit.nature if examen_prescrit.nature in ('BIOLOGIE', 'IMAGERIE') else 'AUTRE'
    creees = []
    for item in items or []:
        nom = str(item).strip()
        if not nom:
            continue
        demande = AnalyseDemande(
            consultation_id=examen_prescrit.consultation_id,
            hospitalisation_id=examen_prescrit.hospitalisation_id,
            patient_id=examen_prescrit.patient_id,
            structure_id=structure_id,
            type_analyse=type_analyse,
            nom_analyse=nom,
            description=examen_prescrit.motif or examen_prescrit.description,
            prescrit_par=examen_prescrit.medecin_id,
            examen_prescrit_id=examen_prescrit.id,
            statut='EN_ATTENTE',
            date_demande=datetime.utcnow(),
            date_prescription=examen_prescrit.date_prescription or datetime.utcnow(),
        )
        db.session.add(demande)
        creees.append(demande)
    return creees


def _signature_item(item):
    """Signature stable d'un item (médicament dict ou acte chaîne), pour
    comparer une nouvelle version d'ordonnance/examens à l'ancienne et ne
    ré-envoyer vers GHP que les items réellement nouveaux (évite les doublons
    à chaque modification/réimpression)."""
    if isinstance(item, dict):
        nom = (item.get('medicament') or item.get('nom') or '').strip().lower()
        return (nom, str(item.get('dosage') or '').strip().lower(), str(item.get('posologie') or '').strip().lower())
    return (str(item).strip().lower(),)


def _items_nouveaux(items_nouveaux, items_anciens):
    """Retourne les items de `items_nouveaux` absents de `items_anciens`."""
    signatures_anciennes = {_signature_item(i) for i in (items_anciens or [])}
    return [i for i in (items_nouveaux or []) if _signature_item(i) not in signatures_anciennes]


def _date_naissance_depuis_formulaire(date_naissance_str, age_str):
    """Résout la date de naissance d'un patient à partir du formulaire :
    la date exacte si elle est saisie, sinon une date approximative
    déduite d'un âge saisi (même jour/mois que la date du jour, comme le
    fait déjà côté client calculerDateNaissance() dans patients/ajouter.html)
    quand la date de naissance est inconnue — filet de sécurité serveur si
    ce JS n'a pas pu s'exécuter (patron : "il faut prevoir qu'on mette
    l'age directe si on ne connait pas date de naissance"). Retourne un
    datetime.date ou None."""
    if date_naissance_str:
        return datetime.strptime(date_naissance_str, '%Y-%m-%d').date()
    if age_str:
        try:
            age = int(age_str)
        except (TypeError, ValueError):
            return None
        if 0 <= age <= 130:
            aujourdhui = datetime.utcnow()
            try:
                return aujourdhui.replace(year=aujourdhui.year - age).date()
            except ValueError:
                # 29 février sans année bissextile correspondante
                return aujourdhui.replace(year=aujourdhui.year - age, day=28).date()
    return None


def _patient_recherche_conditions(search_term):
    """Conditions de recherche patient communes (nom/prénom/téléphone/email
    + numéro de dossier) — le numéro de dossier affiché partout dans
    l'appli (ex. "P00020", voir "%05d" % patient.id) n'était comparé nulle
    part à une recherche : soit pas du tout (page /recherche, page /patients),
    soit contre l'id brut non-paddé (api_patients_search), qui ne matche
    jamais "P00020" ni même "00020" tel que l'utilisateur le voit et le
    tape. Reconnaît "P00020", "p00020", "00020" ou "20" en retirant le
    préfixe P et les zéros de tête avant de comparer à l'id (exact, comme
    une référence de dossier, pas une recherche floue)."""
    from sqlalchemy import or_

    conditions = [
        Patient.nom.ilike(f'%{search_term}%'),
        Patient.prenom.ilike(f'%{search_term}%'),
        Patient.telephone.ilike(f'%{search_term}%'),
        Patient.email.ilike(f'%{search_term}%'),
    ]

    chiffres = search_term.strip()
    if chiffres[:1] in ('P', 'p'):
        chiffres = chiffres[1:]
    chiffres = chiffres.lstrip('0')
    if chiffres.isdigit():
        conditions.append(Patient.id == int(chiffres))

    return or_(*conditions)


# ============================================================
# ⭐ CONNEXION PAR BIOMÉTRIE DE L'APPAREIL (Face ID / Windows Hello /
# empreinte), via WebAuthn — même principe que côté medilogic_ghp
# (services/webauthn_login_service.py), adapté ici à Flask-Login (un seul
# type de compte, Utilisateur) plutôt qu'à la session Sheets de GHP.
# Deux cérémonies : inscription (compte déjà connecté par mot de passe,
# associe un appareil) et connexion (identifie le compte À PARTIR de la
# clé biométrique elle-même, via une clé résidente/"discoverable
# credential" — pas de saisie d'email au préalable). N'est autorisé par
# les navigateurs que sur localhost ou en HTTPS.
# ============================================================

WEBAUTHN_RP_NAME = "MediLogicConsult"


def _webauthn_rp_id_et_origin(request):
    rp_id = request.host.split(':')[0]
    est_local = rp_id in ('localhost', '127.0.0.1', '::1')
    scheme = request.scheme if est_local else 'https'
    origin = f"{scheme}://{request.host}"
    return rp_id, origin


def _webauthn_b64(data):
    import base64
    return base64.b64encode(data).decode('ascii')


def _webauthn_unb64(s):
    import base64
    return base64.b64decode(s.encode('ascii'))


def _webauthn_options_inscription(request, utilisateur_id, utilisateur_nom):
    import webauthn
    from webauthn.helpers.structs import (
        AuthenticatorAttachment, AuthenticatorSelectionCriteria,
        PublicKeyCredentialDescriptor, ResidentKeyRequirement, UserVerificationRequirement,
    )
    from models import IdentifiantWebauthn

    rp_id, _ = _webauthn_rp_id_et_origin(request)

    existants = IdentifiantWebauthn.query.filter_by(utilisateur_id=utilisateur_id, actif=True).all()
    exclude = [PublicKeyCredentialDescriptor(id=_webauthn_unb64(e.credential_id)) for e in existants]

    options = webauthn.generate_registration_options(
        rp_id=rp_id,
        rp_name=WEBAUTHN_RP_NAME,
        user_id=f"user:{utilisateur_id}".encode('utf-8'),
        user_name=utilisateur_nom or f"utilisateur-{utilisateur_id}",
        user_display_name=utilisateur_nom or f"utilisateur-{utilisateur_id}",
        authenticator_selection=AuthenticatorSelectionCriteria(
            # PLATFORM : force le capteur intégré (Face ID/Windows
            # Hello/empreinte), pas une clé de sécurité USB externe.
            authenticator_attachment=AuthenticatorAttachment.PLATFORM,
            # REQUIRED : la clé doit être "résidente"/découvrable pour que
            # la page de connexion puisse la proposer sans connaître le
            # compte à l'avance.
            resident_key=ResidentKeyRequirement.REQUIRED,
            user_verification=UserVerificationRequirement.REQUIRED,
        ),
        exclude_credentials=exclude,
    )
    challenge_b64 = _webauthn_b64(options.challenge)
    return webauthn.options_to_json(options), challenge_b64


def _webauthn_verifier_inscription(request, utilisateur_id, utilisateur_nom, credential, challenge_b64, libelle_appareil=None):
    import webauthn
    from models import IdentifiantWebauthn

    rp_id, origin = _webauthn_rp_id_et_origin(request)

    verification = webauthn.verify_registration_response(
        credential=credential,
        expected_challenge=_webauthn_unb64(challenge_b64),
        expected_rp_id=rp_id,
        expected_origin=origin,
        require_user_verification=True,
    )

    identifiant = IdentifiantWebauthn(
        utilisateur_id=utilisateur_id,
        credential_id=_webauthn_b64(verification.credential_id),
        public_key=_webauthn_b64(verification.credential_public_key),
        sign_count=verification.sign_count,
        libelle_appareil=libelle_appareil or request.host,
    )
    db.session.add(identifiant)
    db.session.commit()
    return identifiant


def _webauthn_options_connexion(request):
    """Pas de allow_credentials : c'est le principe d'une clé résidente —
    le navigateur retrouve tout seul, sur l'appareil, les clés déjà
    enregistrées pour ce rp_id (Face ID/Windows Hello affiche son propre
    sélecteur de compte)."""
    import webauthn
    from webauthn.helpers.structs import UserVerificationRequirement

    rp_id, _ = _webauthn_rp_id_et_origin(request)
    options = webauthn.generate_authentication_options(
        rp_id=rp_id,
        user_verification=UserVerificationRequirement.REQUIRED,
    )
    challenge_b64 = _webauthn_b64(options.challenge)
    return webauthn.options_to_json(options), challenge_b64


def _webauthn_verifier_connexion(request, credential, challenge_b64):
    """Vérifie la clé biométrique et identifie le compte. Retourne la ligne
    IdentifiantWebauthn — à l'appelant de résoudre l'Utilisateur et de
    poser la session (login_user), comme pour la connexion par mot de
    passe. Lève ValueError avec un message utilisateur en cas d'échec."""
    import json
    import webauthn
    from models import IdentifiantWebauthn

    rp_id, origin = _webauthn_rp_id_et_origin(request)

    cred_dict = credential if isinstance(credential, dict) else json.loads(credential)
    credential_id_b64url = cred_dict.get('id') or cred_dict.get('rawId')
    if not credential_id_b64url:
        raise ValueError("Réponse de l'appareil incomplète.")

    raw_id = webauthn.base64url_to_bytes(credential_id_b64url)
    identifiant = IdentifiantWebauthn.query.filter_by(credential_id=_webauthn_b64(raw_id), actif=True).first()
    if not identifiant:
        raise ValueError("Cet appareil n'est associé à aucun compte (ou a été révoqué).")

    verification = webauthn.verify_authentication_response(
        credential=credential,
        expected_challenge=_webauthn_unb64(challenge_b64),
        expected_rp_id=rp_id,
        expected_origin=origin,
        credential_public_key=_webauthn_unb64(identifiant.public_key),
        credential_current_sign_count=identifiant.sign_count,
        require_user_verification=True,
    )
    identifiant.sign_count = verification.new_sign_count
    identifiant.derniere_utilisation = datetime.utcnow()
    db.session.commit()
    return identifiant


def _envoyer_prescriptions_ghp_immediat():
    """Tente un envoi immédiat vers GHP (best-effort, ne bloque jamais la
    transaction métier) — le scheduler (tasks.py, toutes les 5 min) rattrape
    de toute façon en cas d'échec, comme pour les actes posés."""
    try:
        from tasks import sync_prescriptions_to_ghp
        sync_prescriptions_to_ghp()
    except Exception as e:
        print(f"⚠️ Sync immédiate GHP échouée (rattrapage automatique par le scheduler) : {e}")


def _envoyer_resultats_examens_ghp_immediat():
    """Même principe que _envoyer_prescriptions_ghp_immediat, pour les
    résultats d'analyses/examens et les modèles de résultats (voir
    tasks.sync_resultats_examens_to_ghp)."""
    try:
        from tasks import sync_resultats_examens_to_ghp
        sync_resultats_examens_to_ghp()
    except Exception as e:
        print(f"⚠️ Sync immédiate résultats GHP échouée (rattrapage automatique par le scheduler) : {e}")


def _pousser_rendez_vous_ghp(consultation, patient, medecin_nom):
    """Si la consultation porte une date de suivi (Consultation.prochain_rdv),
    pousse automatiquement un rendez-vous vers GHP au lieu de le laisser
    ressaisir à l'accueil. Best-effort (ne bloque jamais l'enregistrement de
    la consultation) — sans table de rattrapage dédiée : un échec est juste
    signalé à l'utilisateur, à reprogrammer manuellement côté GHP le cas
    échéant (contrairement aux prescriptions, il n'y a pas de conséquence
    financière à un RDV manqué, donc pas de scheduler de retry ici).
    """
    if not consultation.prochain_rdv:
        return

    from models import StructureMapping
    import requests as _requests

    mapping = StructureMapping.query.filter_by(
        local_structure_id=patient.id_structure, actif=True
    ).first()
    if not mapping:
        return

    try:
        resp = _requests.post(
            f"{mapping.api_url}/api/rendez-vous/creer-externe",
            params={'token': mapping.api_key},
            json={
                'patient_nom': patient.nom,
                'patient_prenom': patient.prenom,
                'medecin_nom': medecin_nom,
                'date': consultation.prochain_rdv.strftime('%Y-%m-%d'),
                'heure': '08:00',
                'motif': f"Suivi — {consultation.motif or 'consultation'}"[:255],
                'notes': f"Programmé automatiquement depuis gestion_patients (consultation #{consultation.id})",
                'source_id': consultation.id
            },
            timeout=10
        )
        try:
            data = resp.json()
        except Exception:
            data = {}

        if resp.status_code == 200 and data.get('success'):
            if not data.get('deja_existant'):
                flash(
                    f"📅 Rendez-vous de suivi du {consultation.prochain_rdv.strftime('%d/%m/%Y')} "
                    f"envoyé automatiquement à GHP (pas besoin de le ressaisir à l'accueil).",
                    'info'
                )
        elif data.get('error') == 'medecin_introuvable':
            flash(
                data.get('message') or "Rendez-vous de suivi : médecin non retrouvé côté GHP — à programmer manuellement.",
                'warning'
            )
        # patient_introuvable : le patient n'est probablement pas encore
        # synchronisé côté GHP — pas la peine d'alarmer l'utilisateur pour
        # ça, le rattrapage habituel du patient se fera dans les prochaines
        # minutes et le RDV pourra être repoussé à la prochaine modification.
    except Exception as e:
        print(f"⚠️ Push RDV GHP échoué : {e}")


# ============================================================
# ⭐ MIROIR PROTOCOLES / ORDONNANCES-TYPES / EXAMENS-TYPES VERS GHP
# ============================================================
# gestion_patients reste le côté auteur : c'est ici que les médecins créent
# ET appliquent réellement ces modèles (hospitalisation, consultation).
# Chaque création/modification/suppression est poussée en miroir vers le
# modèle ProtocoleMedical de GHP (plus riche : statut, versioning, historique,
# impression avec en-tête), pour que la même donnée soit visible des deux
# côtés sans jamais la ressaisir. Best-effort, ne bloque jamais la
# sauvegarde du médecin — voir /api/protocoles/sync-externe côté GHP.

def _generer_contenu_protocole_soins(nom, description):
    """Texte imprimable pour un ProtocoleSoins — GHP l'affiche tel quel."""
    return f"{nom}\n\n{description or ''}".strip()


def _generer_contenu_ordonnance(nom, medicaments):
    """Même mise en forme que genererContenuOrdonnance() côté GHP
    (templates/protocoles.html) pour un rendu identique à l'impression."""
    date_jour = datetime.utcnow().strftime('%d/%m/%Y')
    lignes = [f"Date : {date_jour}", ""]
    if not medicaments:
        lignes.append("Aucun médicament ajouté.")
    else:
        for i, med in enumerate(medicaments, start=1):
            med_nom = (med.get('medicament') or med.get('nom') or '').strip()
            dosage = (med.get('dosage') or '').strip()
            posologie = (med.get('posologie') or med.get('frequence') or '').strip()
            duree = (med.get('duree') or '').strip()
            ligne = f"{i}. {med_nom}"
            if dosage:
                ligne += f" ({dosage})"
            espaces = max(50 - len(ligne), 5)
            ligne += " " + ("." * espaces) + " "
            ligne += ", ".join(x for x in (posologie, duree) if x)
            lignes.append(ligne)
    lignes.append("")
    lignes.append("─────────────────────────────────────────")
    lignes.append("")
    lignes.append("Signature : {{medecin}}")
    return "\n".join(lignes)


def _generer_contenu_bulletin(nom, motif, examens):
    """Même mise en forme que genererContenuBulletin() côté GHP."""
    date_jour = datetime.utcnow().strftime('%d/%m/%Y')
    lignes = [f"Date : {date_jour}", ""]
    lignes.append("1. Motif / Diagnostic :")
    lignes.append(f"   {motif or '__________________________'}")
    lignes.append("")
    lignes.append("2. Eléments complémentaires :")
    lignes.append("   __________________________")
    lignes.append("")
    lignes.append("3. Nature d'examen(s) demandé(s) :")
    if examens:
        for ex in examens:
            ex = str(ex).strip()
            if ex:
                lignes.append(f"   - {ex}")
    else:
        lignes.append("   __________________________")
    lignes.append("")
    lignes.append("─────────────────────────────────────────")
    lignes.append("")
    lignes.append("Signature : {{medecin}}")
    return "\n".join(lignes)


def _pousser_protocole_ghp(categorie, source_model, source_id, structure_id,
                            titre, description='', contenu='', medicaments=None,
                            examens=None, actif=True, action='upsert'):
    """Pousse (best-effort) un ProtocoleSoins/OrdonnanceType/ExamenType vers
    GHP en tant que ProtocoleMedical miroir. Boucle sur TOUS les mappings
    actifs de la structure (jamais .first() — une structure peut avoir
    plusieurs cibles GHP), échec non bloquant comme le reste de ce miroir."""
    from models import StructureMapping
    import requests as _requests

    mappings = StructureMapping.query.filter_by(
        local_structure_id=structure_id, actif=True
    ).all()
    if not mappings:
        return

    auteur_nom = None
    try:
        auteur_nom = f"{current_user.prenom} {current_user.nom}".strip()
    except Exception:
        pass

    for mapping in mappings:
        try:
            _requests.post(
                f"{mapping.api_url}/api/protocoles/sync-externe",
                params={'token': mapping.api_key},
                json={
                    'categorie': categorie,
                    'source_app': 'gestion_patients',
                    'source_model': source_model,
                    'source_id': source_id,
                    'titre': titre,
                    'description': description or '',
                    'contenu': contenu or '',
                    'medicaments': medicaments or [],
                    'examens': examens or [],
                    'actif': actif,
                    'action': action,
                    'auteur_nom': auteur_nom,
                },
                timeout=10
            )
        except Exception as e:
            print(f"⚠️ Push protocole GHP échoué (structure {structure_id}) : {e}")


# ============================================================
# ⭐ JOURNAL DE SOINS — actes réellement effectués sur un patient
# ============================================================
# Liste des actes courants proposés en un clic (dossier patient / détail de
# consultation), avec le nom EXACT du catalogue GHP (voir P15x/P160/C104
# dans le catalogue d'actes) pour que le matching (badge vert) fonctionne
# dès l'arrivée dans "Prescriptions reçues" côté GHP. `produit` indique si
# le soin doit demander le produit administré (injections/perfusion).
ACTES_SOINS_HABITUELS = [
    {'nom': 'P152 Injection IM*', 'label': 'Injection IM', 'icone': 'fa-syringe', 'produit': True},
    {'nom': 'P153 Injection IV*', 'label': 'Injection IV', 'icone': 'fa-syringe', 'produit': True},
    {'nom': 'P154 Perfusion*', 'label': 'Perfusion', 'icone': 'fa-droplet', 'produit': True},
    {'nom': 'P155 Pansement*', 'label': 'Pansement', 'icone': 'fa-bandage', 'produit': False},
    {'nom': 'P158 Sutures*', 'label': 'Sutures', 'icone': 'fa-scissors', 'produit': False},
    {'nom': "P157 Incision d'abces*", 'label': "Incision d'abcès", 'icone': 'fa-kit-medical', 'produit': False},
    {'nom': 'P159 POSE DE SONDE URINAIRE*', 'label': 'Pose de sonde urinaire', 'icone': 'fa-notes-medical', 'produit': False},
    {'nom': 'P160 MISE EN OBSERVATION (MEO)', 'label': 'Mise en observation', 'icone': 'fa-eye', 'produit': False},
    {'nom': 'C104 Ponction lombaire avec ou sans injection medicamenteuse*', 'label': 'Ponction lombaire', 'icone': 'fa-syringe', 'produit': False},
    {'nom': 'P156 Circoncision*', 'label': 'Circoncision', 'icone': 'fa-kit-medical', 'produit': False},
]


@app.route('/patient/<int:patient_id>/soins/ajouter', methods=['POST'])
@login_required
def soin_ajouter(patient_id):
    """Consigne un soin réellement effectué (journal de soins) — depuis le
    dossier patient ou le détail d'une consultation. Atterrit en brouillon
    dans "Actes posés" : rien ne part vers GHP tant que ce n'est pas
    vérifié/validé là-bas (voir ACTES_SOINS_HABITUELS et actes_poses_liste)."""
    from models import Patient, ActeType, ActePose

    patient = Patient.query.get_or_404(patient_id)
    if patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('patients_list'))

    nom = (request.form.get('nom') or '').strip()
    if not nom:
        flash('Veuillez choisir ou saisir un acte', 'danger')
        return redirect(request.referrer or url_for('patient_detail', id=patient_id))

    produit_administre = (request.form.get('produit_administre') or '').strip() or None
    quantite = (request.form.get('quantite') or '1').strip() or '1'
    notes = (request.form.get('notes') or '').strip() or None
    heure_str = request.form.get('heure')  # datetime-local : "2026-09-08T15:30"
    consultation_id = request.form.get('consultation_id', type=int)
    hospitalisation_id = request.form.get('hospitalisation_id', type=int)

    try:
        date_pose = datetime.strptime(heure_str, '%Y-%m-%dT%H:%M') if heure_str else datetime.utcnow()
    except ValueError:
        date_pose = datetime.utcnow()

    acte_type = ActeType.query.filter_by(structure_id=current_user.id_structure, nom=nom).first()
    if not acte_type:
        acte_type = ActeType(structure_id=current_user.id_structure, nom=nom, created_by=current_user.id)
        db.session.add(acte_type)
        db.session.flush()

    acte_pose = ActePose(
        patient_id=patient.id,
        consultation_id=consultation_id,
        hospitalisation_id=hospitalisation_id,
        acte_type_id=acte_type.id,
        nom=nom,
        quantite=quantite,
        produit_administre=produit_administre,
        notes=notes,
        date_pose=date_pose,
        pose_par_id=current_user.id,
        statut='actif',
        valide=False
    )
    db.session.add(acte_pose)
    db.session.commit()

    flash(
        f'✅ Soin consigné : {nom}{" — " + produit_administre if produit_administre else ""} '
        f'— à vérifier et valider dans l\'onglet "Actes posés".',
        'success'
    )
    return redirect(request.form.get('retour_url') or url_for('patient_detail', id=patient_id))


# Routes principales
@app.route('/')
def index():
    return render_template('index.html')


def _url_dashboard_pour_role(role):
    """URL du tableau de bord adapté au rôle, après connexion (mot de
    passe ou biométrie WebAuthn) — même dispatch utilisé aux deux endroits."""
    return {
        'super_admin': url_for('admin_dashboard'),
        'admin_structure': url_for('structure_dashboard'),
        'infirmier': url_for('infirmier_dashboard'),
        'medecin': url_for('medecin_dashboard'),
        'laborantin': url_for('laborantin_dashboard'),
        'radiologue': url_for('radiologue_dashboard'),
    }.get(role, url_for('dashboard'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        from models import Utilisateur, Structure
        
        email = request.form.get('email')
        password = request.form.get('password')
        
        user = Utilisateur.query.filter_by(email=email).first()
        
        if user and user.check_password(password):
            # Vérifier que l'utilisateur est actif
            if not user.actif:
                flash('Votre compte a été désactivé. Contactez l\'administrateur.', 'danger')
                return redirect(url_for('login'))
            
            # Vérifier que la structure est active (sauf super_admin)
            if user.role != 'super_admin':
                structure = Structure.query.get(user.id_structure)
                if not structure:
                    flash('Structure non trouvée. Contactez l\'administrateur.', 'danger')
                    return redirect(url_for('login'))
                if structure.statut != 'actif':
                    flash('Votre structure n\'est pas active. Contactez l\'administrateur.', 'warning')
                    return redirect(url_for('login'))
            
            login_user(user)
            user.derniere_connexion = datetime.utcnow()
            db.session.commit()

            return redirect(_url_dashboard_pour_role(user.role))
        else:
            flash('Email ou mot de passe incorrect', 'danger')
    
    return render_template('login.html')


@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('Vous avez été déconnecté', 'info')
    return redirect(url_for('index'))


# ============================================================
# ⭐ CONNEXION PAR BIOMÉTRIE (Face ID / Windows Hello / empreinte) — voir
# les fonctions _webauthn_* plus haut et models.IdentifiantWebauthn.
# ============================================================

@app.route('/api/webauthn/inscription/options', methods=['POST'])
@login_required
def api_webauthn_inscription_options():
    import json
    options_json, challenge = _webauthn_options_inscription(
        request, current_user.id, f"{current_user.prenom} {current_user.nom}"
    )
    session['webauthn_challenge'] = challenge
    return jsonify({'success': True, 'options': json.loads(options_json)})


@app.route('/api/webauthn/inscription/verifier', methods=['POST'])
@login_required
def api_webauthn_inscription_verifier():
    data = request.json or {}
    challenge = session.pop('webauthn_challenge', None)
    if not challenge:
        return jsonify({'success': False, 'error': 'Session expirée, recommencez.'}), 400

    try:
        _webauthn_verifier_inscription(
            request, current_user.id, f"{current_user.prenom} {current_user.nom}",
            data.get('credential'), challenge,
            libelle_appareil=data.get('libelle_appareil'),
        )
        return jsonify({'success': True})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': f"Échec de l'enregistrement : {e}"}), 400


@app.route('/api/webauthn/mes-appareils', methods=['GET'])
@login_required
def api_webauthn_mes_appareils():
    from models import IdentifiantWebauthn
    appareils = IdentifiantWebauthn.query.filter_by(
        utilisateur_id=current_user.id, actif=True
    ).order_by(IdentifiantWebauthn.date_creation.desc()).all()
    return jsonify({'success': True, 'data': [{
        'id': a.id,
        'libelle_appareil': a.libelle_appareil,
        'date_creation': a.date_creation.strftime('%d/%m/%Y %H:%M') if a.date_creation else None,
        'derniere_utilisation': a.derniere_utilisation.strftime('%d/%m/%Y %H:%M') if a.derniere_utilisation else None,
    } for a in appareils]})


@app.route('/api/webauthn/appareils/<int:appareil_id>', methods=['DELETE'])
@login_required
def api_webauthn_revoquer(appareil_id):
    from models import IdentifiantWebauthn
    appareil = IdentifiantWebauthn.query.filter_by(id=appareil_id, utilisateur_id=current_user.id).first()
    if not appareil:
        return jsonify({'success': False, 'error': 'Introuvable'}), 404

    appareil.actif = False
    appareil.date_revocation = datetime.utcnow()
    db.session.commit()
    return jsonify({'success': True})


@app.route('/login/webauthn/options', methods=['POST'])
def login_webauthn_options():
    """Public — page de connexion, personne n'est encore authentifié. Pas
    de allow_credentials (clé résidente) : le navigateur propose lui-même,
    via Face ID/Windows Hello, les comptes déjà enregistrés sur cet
    appareil pour ce site."""
    import json
    options_json, challenge = _webauthn_options_connexion(request)
    session['webauthn_login_challenge'] = challenge
    return jsonify({'success': True, 'options': json.loads(options_json)})


@app.route('/login/webauthn/verifier', methods=['POST'])
def login_webauthn_verifier():
    from models import Utilisateur, Structure

    data = request.json or {}
    challenge = session.pop('webauthn_login_challenge', None)
    if not challenge:
        return jsonify({'success': False, 'error': 'Session expirée, recommencez.'}), 400

    try:
        identifiant = _webauthn_verifier_connexion(request, data.get('credential'), challenge)
    except ValueError as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)}), 401
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': f'Échec de la vérification : {e}'}), 400

    user = Utilisateur.query.get(identifiant.utilisateur_id)
    if not user or not user.actif:
        return jsonify({'success': False, 'error': 'Compte introuvable ou désactivé.'}), 401
    if user.role != 'super_admin':
        structure = Structure.query.get(user.id_structure)
        if not structure or structure.statut != 'actif':
            return jsonify({'success': False, 'error': "Structure introuvable ou inactive."}), 401

    login_user(user)
    user.derniere_connexion = datetime.utcnow()
    db.session.commit()

    return jsonify({'success': True, 'redirect': _url_dashboard_pour_role(user.role)})


# ============================================================
# ⭐ APPLICATION INSTALLABLE (PWA) — "MediLogicConsult", voir
# static/app-manifest.json et static/sw-app.js. Publics (pas de
# login_required) : le navigateur les récupère avant toute session.
# ============================================================

@app.route('/app-manifest.json')
def app_manifest():
    from flask import send_from_directory
    return send_from_directory('static', 'app-manifest.json', mimetype='application/manifest+json')


@app.route('/sw-app.js')
def app_service_worker():
    from flask import send_from_directory
    return send_from_directory('static', 'sw-app.js', mimetype='application/javascript')


@app.route('/register', methods=['GET', 'POST'])
def register_structure():
    if request.method == 'POST':
        from models import Structure, Utilisateur, db
        import hashlib
        
        # Récupération des données
        nom_structure = request.form.get('nom_structure')
        nom_responsable = request.form.get('nom_responsable')
        prenom_responsable = request.form.get('prenom_responsable')
        adresse = request.form.get('adresse')
        email = request.form.get('email')
        telephone = request.form.get('telephone')
        password = request.form.get('password')
        confirm_password = request.form.get('confirm_password')
        reset_question = request.form.get('reset_question')
        reset_answer = request.form.get('reset_answer')
        
        # Validation
        if not nom_structure or not email or not password:
            flash('Tous les champs obligatoires doivent être remplis', 'danger')
            return redirect(url_for('register_structure'))
        
        if password != confirm_password:
            flash('Les mots de passe ne correspondent pas', 'danger')
            return redirect(url_for('register_structure'))
        
        # Validation de la complexité du mot de passe
        import re
        errors = []
        if len(password) < 8:
            errors.append("Minimum 8 caractères")
        if not re.search(r"[A-Z]", password):
            errors.append("Au moins 1 majuscule")
        if not re.search(r"[a-z]", password):
            errors.append("Au moins 1 minuscule")
        if not re.search(r"[0-9]", password):
            errors.append("Au moins 1 chiffre")
        if not re.search(r"[!@#$%^&*(),.?\":{}|<>]", password):
            errors.append("Au moins 1 symbole")
        
        if errors:
            for error in errors:
                flash(error, 'danger')
            return redirect(url_for('register_structure'))
        
        # Vérifier si l'email existe déjà
        existing = Utilisateur.query.filter_by(email=email).first()
        if existing:
            flash('Cet email est déjà utilisé', 'danger')
            return redirect(url_for('register_structure'))
        
        # Création de la structure
        structure = Structure(
            nom=nom_structure,
            email=email,
            telephone=telephone,
            adresse=adresse,
            statut='en_attente',
            reset_question=reset_question,
            reset_answer_hash=hashlib.sha256(reset_answer.lower().strip().encode()).hexdigest()
        )
        db.session.add(structure)
        db.session.flush()
        
        # Création de l'admin de structure (responsable)
        admin = Utilisateur(
            email=email,
            nom=nom_responsable,
            prenom=prenom_responsable,
            role='admin_structure',
            id_structure=structure.id,
            actif=True
        )
        admin.set_password(password)
        db.session.add(admin)
        db.session.commit()
        
        flash('Votre demande d\'inscription a été envoyée. Un administrateur va la valider.', 'success')
        return redirect(url_for('login'))
    
    return render_template('register.html')

@app.route('/forgot-password', methods=['GET', 'POST'])
def forgot_password():
    if request.method == 'POST':
        from models import Utilisateur, Structure
        email = request.form.get('email')
        use_secret_question = request.form.get('use_secret_question') == 'on'
        
        user = Utilisateur.query.filter_by(email=email).first()
        
        if not user:
            flash('Aucun compte trouvé avec cet email', 'danger')
            return redirect(url_for('forgot_password'))
        
        if use_secret_question and user.reset_question:
            session['reset_user_id'] = user.id
            return redirect(url_for('verify_secret_question'))
        else:
            import secrets
            token = secrets.token_urlsafe(32)
            user.reset_token = token
            user.reset_token_expiry = datetime.utcnow() + timedelta(hours=24)  # ✅ CORRIGÉ
            db.session.commit()
            
            flash(f'Lien de réinitialisation (valable 24h) : /reset-password/{token}', 'info')
            return redirect(url_for('login'))
    
    return render_template('forgot_password.html')

@app.route('/verify-secret-question', methods=['GET', 'POST'])
def verify_secret_question():
    user_id = session.get('reset_user_id')
    if not user_id:
        return redirect(url_for('forgot_password'))
    
    from models import Utilisateur
    user = Utilisateur.query.get(user_id)
    
    if request.method == 'POST':
        answer = request.form.get('answer')
        new_password = request.form.get('new_password')
        confirm_password = request.form.get('confirm_password')
        
        if new_password != confirm_password:
            flash('Les mots de passe ne correspondent pas', 'danger')
            return redirect(url_for('verify_secret_question'))
        
        if user.check_reset_answer(answer):
            user.set_password(new_password)
            user.reset_token = None
            user.reset_token_expiry = None
            db.session.commit()
            session.pop('reset_user_id', None)
            flash('Mot de passe réinitialisé avec succès', 'success')
            return redirect(url_for('login'))
        else:
            flash('Réponse incorrecte', 'danger')
    
    return render_template('verify_secret_question.html', question=user.reset_question)

@app.route('/laborantin/dashboard')
@login_required
def laborantin_dashboard():
    """Dashboard pour le laborantin"""
    from models import AnalyseDemande
    
    if current_user.role != 'laborantin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Statistiques pour le laborantin
    en_attente = AnalyseDemande.query.filter_by(
        structure_id=current_user.id_structure,
        statut='EN_ATTENTE'
    ).count()
    
    en_cours = AnalyseDemande.query.filter_by(
        structure_id=current_user.id_structure,
        statut='EN_COURS'
    ).count()
    
    termine = AnalyseDemande.query.filter_by(
        structure_id=current_user.id_structure,
        statut='TERMINE'
    ).count()
    
    total = en_attente + en_cours + termine
    
    return render_template('laborantin/dashboard.html',
                         en_attente=en_attente,
                         en_cours=en_cours,
                         termine=termine,
                         total=total)

@app.route('/radiologue/dashboard')
@login_required
def radiologue_dashboard():
    """Dashboard pour le radiologue"""
    from models import AnalyseDemande
    
    if current_user.role != 'radiologue':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Analyses d'imagerie uniquement
    en_attente = AnalyseDemande.query.filter_by(
        structure_id=current_user.id_structure,
        type_analyse='IMAGERIE',
        statut='EN_ATTENTE'
    ).count()
    
    en_cours = AnalyseDemande.query.filter_by(
        structure_id=current_user.id_structure,
        type_analyse='IMAGERIE',
        statut='EN_COURS'
    ).count()
    
    termine = AnalyseDemande.query.filter_by(
        structure_id=current_user.id_structure,
        type_analyse='IMAGERIE',
        statut='TERMINE'
    ).count()
    
    total = en_attente + en_cours + termine
    
    return render_template('radiologue/dashboard.html',
                         en_attente=en_attente,
                         en_cours=en_cours,
                         termine=termine,
                         total=total)


@app.route('/dashboard')
@login_required
def dashboard():
    from models import Patient, Consultation, Prescription
    from datetime import date, datetime
    
    # ⭐ REDIRECTIONS POUR TOUS LES RÔLES
    if current_user.role == 'super_admin':
        return redirect(url_for('admin_dashboard'))
    elif current_user.role == 'admin_structure':
        return redirect(url_for('structure_dashboard'))
    elif current_user.role == 'infirmier':
        return redirect(url_for('infirmier_dashboard'))
    elif current_user.role == 'laborantin':
        return redirect(url_for('laborantin_dashboard'))  # ⭐ AJOUTÉ
    elif current_user.role == 'radiologue':
        return redirect(url_for('radiologue_dashboard'))  # ⭐ AJOUTÉ
    else:
        # Dashboard médecin
        patients_actifs = Patient.query.filter_by(
            id_structure=current_user.id_structure,
            id_medecin_referent=current_user.id,
            archived=False
        ).count()
        
        consultations_aujourdhui = Consultation.query.filter(
            Consultation.id_medecin == current_user.id,
            Consultation.date_consultation >= datetime.now().replace(hour=0, minute=0, second=0)
        ).count()
        
        prescriptions_actives = Prescription.query.filter_by(
            id_medecin=current_user.id,
            statut='active'
        ).count()
        
        patients_gueris = Patient.query.filter_by(
            id_structure=current_user.id_structure,
            id_medecin_referent=current_user.id,
            statut_medical='GUERI',
            archived=False
        ).count()
        
        derniers_patients = Patient.query.filter_by(
            id_structure=current_user.id_structure,
            id_medecin_referent=current_user.id,
            archived=False
        ).order_by(Patient.date_creation.desc()).limit(5).all()
        
        prochains_rdv = Consultation.query.filter(
            Consultation.id_medecin == current_user.id,
            Consultation.prochain_rdv >= datetime.now()
        ).order_by(Consultation.prochain_rdv.asc()).limit(5).all()
        
        return render_template('medecin_dashboard.html',
                             patients_actifs=patients_actifs,
                             consultations_aujourdhui=consultations_aujourdhui,
                             prescriptions_actives=prescriptions_actives,
                             patients_gueris=patients_gueris,
                             derniers_patients=derniers_patients,
                             prochains_rdv=prochains_rdv,
                             today=date.today(),
                             now=datetime.now())

# Routes admin super admin
@app.route('/admin')
@login_required
def admin_dashboard():
    if current_user.role != 'super_admin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Structure, Utilisateur
    structures_en_attente = Structure.query.filter_by(statut='en_attente').count()
    structures_actives = Structure.query.filter_by(statut='actif').count()
    total_utilisateurs = Utilisateur.query.count()
    
    return render_template('admin/dashboard.html',
                         structures_en_attente=structures_en_attente,
                         structures_actives=structures_actives,
                         total_utilisateurs=total_utilisateurs)


@app.route('/radiologie')
@login_required
def liste_radiologie():
    """Liste des examens d'imagerie pour le radiologue"""
    from models import AnalyseDemande, Patient
    from sqlalchemy import or_
    
    if current_user.role not in ['admin_structure', 'radiologue', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    statut = request.args.get('statut', '')
    search = request.args.get('search', '')
    
    query = AnalyseDemande.query.filter_by(
        structure_id=current_user.id_structure,
        type_analyse='IMAGERIE'  # ⭐ UNIQUEMENT L'IMAGERIE
    )
    
    if statut:
        query = query.filter_by(statut=statut)
    
    if search:
        search = search.strip()
        filters = [Patient.nom.ilike(f'%{search}%'), Patient.prenom.ilike(f'%{search}%')]
        patients_trouves = Patient.query.filter(or_(*filters)).all()
        patient_ids = [p.id for p in patients_trouves]
        if patient_ids:
            query = query.filter(AnalyseDemande.patient_id.in_(patient_ids))
        else:
            query = query.filter(AnalyseDemande.patient_id == -1)
    
    analyses = query.order_by(AnalyseDemande.date_demande.desc()).all()
    
    patients_dict = {}
    for analyse in analyses:
        patient_id = analyse.patient_id
        if patient_id not in patients_dict:
            patients_dict[patient_id] = {
                'patient': analyse.patient,
                'analyses': []
            }
        patients_dict[patient_id]['analyses'].append(analyse)
    
    patients = list(patients_dict.values())
    statuts = ['EN_ATTENTE', 'EN_COURS', 'TERMINE']
    
    return render_template('radiologue/liste.html',
                         patients=patients,
                         statut_actuel=statut,
                         statuts=statuts,
                         search=search)


# ==================== DASHBOARD INFIRMIER ====================
@app.route('/infirmier/dashboard')
@login_required
def infirmier_dashboard():
    from models import Patient, Utilisateur
    from datetime import datetime
    
    # Vérifier que l'utilisateur est un infirmier
    if current_user.role != 'infirmier':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # ⭐ PATIENTS EN ATTENTE - TRIÉS PAR DATE DE CRÉATION (les plus récents d'abord)
    # Exclut les patients marqués "pas de pré-consultation nécessaire" (patron :
    # "c'est pas tous les patients que l'infirmier aura à preconsulter") —
    # voir infirmier_exclure_pre_consultation / infirmier_reinclure_pre_consultation.
    patients_attente = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        archived=False
    ).filter(
        db.or_(
            Patient.pre_consultation_faite == False,
            Patient.pre_consultation_faite.is_(None)
        )
    ).filter(
        db.or_(
            Patient.pre_consultation_non_requise == False,
            Patient.pre_consultation_non_requise.is_(None)
        )
    ).order_by(Patient.date_creation.desc()).all()  # ⭐ CHANGÉ

    # ⭐ PATIENTS DÉSÉLECTIONNÉS - pas encore pré-consultés mais marqués comme
    # n'en ayant pas besoin (pour pouvoir les réinclure en cas d'erreur).
    patients_exclus = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        archived=False,
        pre_consultation_non_requise=True
    ).filter(
        db.or_(
            Patient.pre_consultation_faite == False,
            Patient.pre_consultation_faite.is_(None)
        )
    ).order_by(Patient.date_creation.desc()).all()

    # ⭐ PATIENTS DÉJÀ PRÉPARÉS - TRIÉS PAR DATE DE PRÉ-CONSULTATION
    patients_prets = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        archived=False,
        pre_consultation_faite=True
    ).order_by(Patient.pre_consultation_date.desc()).all()

    # Médecins actifs
    medecins = Utilisateur.query.filter_by(
        id_structure=current_user.id_structure,
        role='medecin',
        actif=True
    ).all()

    # Total patients
    total_patients = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        archived=False
    ).count()

    return render_template('infirmier/dashboard.html',
                         patients_attente=patients_attente,
                         patients_exclus=patients_exclus,
                         patients_prets=patients_prets,
                         medecins=medecins,
                         total_patients=total_patients,
                         now=datetime.now())


@app.route('/infirmier/patient/<int:patient_id>/preconsultation/exclure', methods=['POST'])
@login_required
def infirmier_exclure_pre_consultation(patient_id):
    """Retire un patient de la file d'attente de pré-consultation sans la
    faire — tous les patients non préparés n'en ont pas forcément besoin."""
    from models import Patient

    if current_user.role != 'infirmier':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    patient = Patient.query.get_or_404(patient_id)
    if patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé à ce patient', 'danger')
        return redirect(url_for('infirmier_dashboard'))

    patient.pre_consultation_non_requise = True
    db.session.commit()
    flash(f'{patient.prenom} {patient.nom} retiré(e) de la file de pré-consultation', 'success')
    return redirect(url_for('infirmier_dashboard'))


@app.route('/infirmier/patient/<int:patient_id>/preconsultation/reinclure', methods=['POST'])
@login_required
def infirmier_reinclure_pre_consultation(patient_id):
    """Annule l'exclusion et remet le patient dans la file d'attente."""
    from models import Patient

    if current_user.role != 'infirmier':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    patient = Patient.query.get_or_404(patient_id)
    if patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé à ce patient', 'danger')
        return redirect(url_for('infirmier_dashboard'))

    patient.pre_consultation_non_requise = False
    db.session.commit()
    flash(f'{patient.prenom} {patient.nom} remis(e) dans la file de pré-consultation', 'success')
    return redirect(url_for('infirmier_dashboard'))

@app.route('/infirmier/pre_consultation/<int:patient_id>', methods=['GET', 'POST'])
@login_required
def infirmier_pre_consultation(patient_id):
    from models import Patient
    from datetime import datetime, timezone
    
    patient = Patient.query.get_or_404(patient_id)
    
    if current_user.role != 'infirmier':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    if patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé à ce patient', 'danger')
        return redirect(url_for('infirmier_dashboard'))
    
    # ⭐ RÉCUPÉRER LES PATIENTS EN ATTENTE
    patients_attente = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        archived=False
    ).filter(
        db.or_(
            Patient.pre_consultation_faite == False,
            Patient.pre_consultation_faite.is_(None)
        )
    ).order_by(Patient.date_creation.desc()).all()

    step = request.args.get('step', 1, type=int)
    if step < 1 or step > 4:
        step = 1
    
    if request.method == 'POST':
        # Motif
        patient.motif_pre_consultation = request.form.get('motif')
        
        # Constantes avec gestion des valeurs vides
        tension = request.form.get('tension')
        temperature = request.form.get('temperature')
        pouls = request.form.get('pouls')
        saturation = request.form.get('saturation')
        poids = request.form.get('poids')
        taille = request.form.get('taille')
        imc = request.form.get('imc')
        
        patient.tension_arterielle = tension if tension and tension.strip() else None
        
        try:
            patient.temperature_c = float(temperature) if temperature and temperature.strip() else None
        except (ValueError, TypeError):
            patient.temperature_c = None
        
        try:
            patient.pulse_bpm = int(pouls) if pouls and pouls.strip() else None
        except (ValueError, TypeError):
            patient.pulse_bpm = None
        
        try:
            patient.oxygene_saturation = int(saturation) if saturation and saturation.strip() else None
        except (ValueError, TypeError):
            patient.oxygene_saturation = None
        
        try:
            patient.poids_kg = float(poids) if poids and poids.strip() else None
        except (ValueError, TypeError):
            patient.poids_kg = None
        
        try:
            patient.taille_cm = float(taille) if taille and taille.strip() else None
        except (ValueError, TypeError):
            patient.taille_cm = None
        
        try:
            patient.imc = float(imc) if imc and imc.strip() else None
        except (ValueError, TypeError):
            patient.imc = None
        
        # Habitudes de vie
        patient.tabac = request.form.get('tabac')
        patient.alcool = request.form.get('alcool')
        patient.allaitement = request.form.get('allaitement') == 'Oui'
        patient.grossesse = request.form.get('grossesse') == 'Oui'
        patient.groupe_sanguin = request.form.get('groupe_sanguin')
        patient.mutuelle = request.form.get('mutuelle')
        patient.medecin_traitant = request.form.get('medecin_traitant')
        
        # Marquer la pré-consultation comme faite
        patient.pre_consultation_faite = True
        patient.pre_consultation_par = current_user.id
        patient.pre_consultation_date = datetime.now(timezone.utc)
        
        db.session.commit()

        flash(f'✅ Pré-consultation de {patient.prenom} {patient.nom} enregistrée avec succès !', 'success')
        # ⭐ Retour au tableau (liste des patients en attente/prêts) plutôt que
        # de rester sur la fiche du patient qui vient d'être traité — demandé
        # explicitement, l'infirmier enchaîne sur le patient suivant.
        return redirect(url_for('infirmier_dashboard'))
    
    return render_template('infirmier/pre_consultation.html', 
                         patient=patient,
                         patients_attente=patients_attente,
                         step=step)

# ==================== SAUVEGARDE MOTIF (ÉTAPE 1) ====================
@app.route('/infirmier/pre_consultation/<int:patient_id>/save-motif', methods=['POST'])
@login_required
def save_motif_pre_consultation(patient_id):
    from models import Patient
    from flask import jsonify
    
    patient = Patient.query.get_or_404(patient_id)
    
    if current_user.role != 'infirmier':
        return jsonify({'success': False, 'message': 'Accès non autorisé'}), 403
    
    motif = request.form.get('motif', '').strip()
    patient.motif_pre_consultation = motif
    db.session.commit()
    
    return jsonify({'success': True})


# ==================== SAUVEGARDE CONSTANTES (ÉTAPE 2) ====================
@app.route('/infirmier/pre_consultation/<int:patient_id>/save-constantes', methods=['POST'])
@login_required
def save_constantes_pre_consultation(patient_id):
    from models import Patient
    from flask import jsonify
    
    patient = Patient.query.get_or_404(patient_id)
    
    if current_user.role != 'infirmier':
        return jsonify({'success': False, 'message': 'Accès non autorisé'}), 403
    
    # Récupérer les constantes
    tension = request.form.get('tension')
    temperature = request.form.get('temperature')
    pouls = request.form.get('pouls')
    saturation = request.form.get('saturation')
    poids = request.form.get('poids')
    taille = request.form.get('taille')
    imc = request.form.get('imc')
    
    # Sauvegarder avec gestion des valeurs vides
    patient.tension_arterielle = tension if tension and tension.strip() else None
    
    try:
        patient.temperature_c = float(temperature) if temperature and temperature.strip() else None
    except (ValueError, TypeError):
        patient.temperature_c = None
    
    try:
        patient.pulse_bpm = int(pouls) if pouls and pouls.strip() else None
    except (ValueError, TypeError):
        patient.pulse_bpm = None
    
    try:
        patient.oxygene_saturation = int(saturation) if saturation and saturation.strip() else None
    except (ValueError, TypeError):
        patient.oxygene_saturation = None
    
    try:
        patient.poids_kg = float(poids) if poids and poids.strip() else None
    except (ValueError, TypeError):
        patient.poids_kg = None
    
    try:
        patient.taille_cm = float(taille) if taille and taille.strip() else None
    except (ValueError, TypeError):
        patient.taille_cm = None
    
    try:
        patient.imc = float(imc) if imc and imc.strip() else None
    except (ValueError, TypeError):
        patient.imc = None
    
    db.session.commit()
    
    return jsonify({'success': True})

@app.route('/medecin/dashboard')
@login_required
def medecin_dashboard():
    from models import Patient, Consultation, Prescription
    from datetime import date, datetime
    
    if current_user.role != 'medecin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Statistiques pour le médecin
    patients_actifs = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        id_medecin_referent=current_user.id,
        archived=False
    ).count()
    
    consultations_aujourdhui = Consultation.query.filter(
        Consultation.id_medecin == current_user.id,
        Consultation.date_consultation >= datetime.now().replace(hour=0, minute=0, second=0)
    ).count()
    
    prescriptions_actives = Prescription.query.filter_by(
        id_medecin=current_user.id,
        statut='active'
    ).count()
    
    patients_gueris = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        id_medecin_referent=current_user.id,
        statut_medical='GUERI',
        archived=False
    ).count()
    
    derniers_patients = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        id_medecin_referent=current_user.id,
        archived=False
    ).order_by(Patient.date_creation.desc()).limit(5).all()
    
    prochains_rdv = Consultation.query.filter(
        Consultation.id_medecin == current_user.id,
        Consultation.prochain_rdv >= datetime.now()
    ).order_by(Consultation.prochain_rdv.asc()).limit(5).all()
    
    return render_template('medecin_dashboard.html',
                         patients_actifs=patients_actifs,
                         consultations_aujourdhui=consultations_aujourdhui,
                         prescriptions_actives=prescriptions_actives,
                         patients_gueris=patients_gueris,
                         derniers_patients=derniers_patients,
                         prochains_rdv=prochains_rdv,
                         today=date.today())

@app.route('/api/patient/<int:patient_id>/pre_consultation')
@login_required
def api_patient_pre_consultation(patient_id):
    from models import Patient
    
    patient = Patient.query.get_or_404(patient_id)
    
    return jsonify({
        'motif': patient.motif_pre_consultation or '',
        'pre_consultation_faite': patient.pre_consultation_faite or False,
        'tabac': patient.tabac or '',
        'alcool': patient.alcool or '',
        'allaitement': patient.allaitement or False,
        'grossesse': patient.grossesse or False,
        'groupe_sanguin': patient.groupe_sanguin or '',
        'mutuelle': patient.mutuelle or '',
        'medecin_traitant': patient.medecin_traitant or ''
    })

# ==================== GESTION DES UTILISATEURS PAR ADMIN STRUCTURE ====================

@app.route('/structure/utilisateurs')
@login_required
def structure_utilisateurs():
    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Utilisateur
    utilisateurs = Utilisateur.query.filter_by(id_structure=current_user.id_structure).all()
    return render_template('structure/utilisateurs.html', utilisateurs=utilisateurs)


@app.route('/structure/parametrage-amu', methods=['GET', 'POST'])
@login_required
def parametrage_amu():
    """Paramétrage AMU hospitalisation : seuils de jours par palier (semaine 1
    / semaine 2 / 15j et plus), commun à toute la structure.

    ⭐ Le mapping par salle vers le catalogue GHP (Salle.acte_ghp_semaine1/2/3)
    ne se fait plus ici — il se configure directement à la création/
    modification de la salle (voir ajouter_salle/modifier_salle), avec un
    sélecteur qui choisit les 3 paliers d'un coup depuis le vrai catalogue
    GHP (fini la saisie libre risquant un mismatch). Cet écran ne montre
    plus qu'un résumé (salles non configurées) qui pointe vers leur fiche.
    """
    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    from models import ParametrageAMU, Salle, Service

    parametrage = ParametrageAMU.query.filter_by(structure_id=current_user.id_structure).first()
    if not parametrage:
        parametrage = ParametrageAMU(structure_id=current_user.id_structure)
        db.session.add(parametrage)
        db.session.commit()

    if request.method == 'POST':
        try:
            s1 = int(request.form.get('seuil_jours_semaine1', 7))
            s2 = int(request.form.get('seuil_jours_semaine2', 14))
            taux = float(request.form.get('taux_amu_info', 90))
        except (TypeError, ValueError):
            flash('Valeurs invalides', 'danger')
            return redirect(url_for('parametrage_amu'))

        if s1 < 1 or s2 <= s1:
            flash('Le palier 2 doit se terminer après le palier 1', 'danger')
            return redirect(url_for('parametrage_amu'))

        parametrage.seuil_jours_semaine1 = s1
        parametrage.seuil_jours_semaine2 = s2
        parametrage.taux_amu_info = taux
        db.session.commit()
        flash('Paramétrage AMU mis à jour', 'success')
        return redirect(url_for('parametrage_amu'))

    salles = Salle.query.join(Service).filter(
        Service.structure_id == current_user.id_structure
    ).order_by(Salle.nom).all()

    return render_template('structure/parametrage_amu.html', parametrage=parametrage, salles=salles)


@app.route('/structure/utilisateur/ajouter', methods=['GET', 'POST'])
@login_required
def structure_ajouter_utilisateur():
    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Utilisateur
    
    if request.method == 'POST':
        email = request.form.get('email')
        nom = request.form.get('nom')
        prenom = request.form.get('prenom')
        role = request.form.get('role')
        password = request.form.get('password')
        confirm_password = request.form.get('confirm_password')
        
        # Vérifications
        if password != confirm_password:
            flash('Les mots de passe ne correspondent pas', 'danger')
            return redirect(url_for('structure_ajouter_utilisateur'))
        
        # Vérifier si l'email existe déjà
        existing = Utilisateur.query.filter_by(email=email).first()
        if existing:
            flash('Cet email est déjà utilisé', 'danger')
            return redirect(url_for('structure_ajouter_utilisateur'))
        
        # Créer l'utilisateur
        new_user = Utilisateur(
            email=email,
            nom=nom,
            prenom=prenom,
            role=role,
            id_structure=current_user.id_structure,
            actif=True
        )
        new_user.set_password(password)
        db.session.add(new_user)
        db.session.commit()
        
        flash(f'Utilisateur {prenom} {nom} créé avec succès', 'success')
        return redirect(url_for('structure_utilisateurs'))
    
    return render_template('structure/ajouter_utilisateur.html')

@app.route('/structure/utilisateur/<int:id>/modifier', methods=['GET', 'POST'])
@login_required
def structure_modifier_utilisateur(id):
    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Utilisateur
    user = Utilisateur.query.get_or_404(id)
    
    # Vérifier que l'utilisateur appartient à la structure
    if user.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('structure_utilisateurs'))
    
    if request.method == 'POST':
        user.nom = request.form.get('nom')
        user.prenom = request.form.get('prenom')
        # ⭐ Garde-fou : n'accepter que les rôles réels de l'application — sans
        # ça, un <select> incomplet ou un champ manipulé peut écraser
        # silencieusement le rôle d'un utilisateur (voir le commentaire sur
        # modifier_utilisateur.html : c'est exactement ce qui arrivait pour
        # le rôle "infirmier", absent de l'ancienne liste déroulante).
        role_soumis = request.form.get('role')
        roles_valides = [
            'medecin', 'infirmier', 'sage-femme', 'assistant_medical',
            'technicien_superieur', 'secretaire', 'laborantin', 'radiologue',
            'kinésithérapeute', 'psychologue', 'nutritionniste', 'pharmacien',
            'ambulancier', 'accueil', 'comptable',
        ]
        if role_soumis in roles_valides:
            user.role = role_soumis
        elif role_soumis:
            flash(f'Rôle "{role_soumis}" invalide — rôle inchangé.', 'danger')
            return redirect(url_for('structure_modifier_utilisateur', id=id))
        user.actif = request.form.get('actif') == 'on'
        
        # Changement de mot de passe optionnel
        new_password = request.form.get('new_password')
        if new_password:
            confirm = request.form.get('confirm_password')
            if new_password == confirm:
                user.set_password(new_password)
                flash('Mot de passe modifié', 'success')
            else:
                flash('Les mots de passe ne correspondent pas', 'danger')
                return redirect(url_for('structure_modifier_utilisateur', id=id))
        
        db.session.commit()
        flash(f'Utilisateur {user.prenom} {user.nom} modifié', 'success')
        return redirect(url_for('structure_utilisateurs'))
    
    return render_template('structure/modifier_utilisateur.html', user=user)

@app.route('/structure/utilisateur/<int:id>/supprimer')
@login_required
def structure_supprimer_utilisateur(id):
    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Utilisateur
    user = Utilisateur.query.get_or_404(id)
    
    # Ne pas permettre de supprimer son propre compte
    if user.id == current_user.id:
        flash('Vous ne pouvez pas supprimer votre propre compte', 'danger')
        return redirect(url_for('structure_utilisateurs'))
    
    if user.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('structure_utilisateurs'))
    
    db.session.delete(user)
    db.session.commit()
    flash(f'Utilisateur {user.prenom} {user.nom} supprimé', 'success')
    return redirect(url_for('structure_utilisateurs'))

@app.route('/admin/structures')
@login_required
def admin_structures():
    if current_user.role != 'super_admin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Structure
    structures = Structure.query.all()
    return render_template('admin/structures.html', structures=structures)

@app.route('/admin/structure/<int:id>/activate')
@login_required
def admin_activate_structure(id):
    if current_user.role != 'super_admin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Structure
    structure = Structure.query.get_or_404(id)
    structure.statut = 'actif'
    structure.date_activation = datetime.utcnow()
    db.session.commit()
    
    # 👇 AJOUTEZ CETTE LIGNE
    create_structure_sheets(structure.id)
    
    flash(f'Structure {structure.nom} activée avec succès', 'success')
    return redirect(url_for('admin_structures'))

@app.route('/admin/structure/<int:id>/reset-password', methods=['GET', 'POST'])
@login_required
def admin_reset_structure_password(id):
    if current_user.role != 'super_admin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Structure, Utilisateur
    structure = Structure.query.get_or_404(id)
    admin = Utilisateur.query.filter_by(id_structure=id, role='admin_structure').first()
    
    if request.method == 'POST':
        new_password = request.form.get('new_password')
        confirm_password = request.form.get('confirm_password')
        
        if new_password != confirm_password:
            flash('Les mots de passe ne correspondent pas', 'danger')
        else:
            admin.set_password(new_password)
            db.session.commit()
            flash(f'Mot de passe réinitialisé pour {structure.nom}', 'success')
            return redirect(url_for('admin_structures'))
    
    return render_template('admin/reset_password.html', structure=structure, admin=admin)

# Route pour tableau de bord structure admin
@app.route('/structure')
@login_required
def structure_dashboard():
    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Patient, Utilisateur, Consultation
    total_patients = Patient.query.filter_by(id_structure=current_user.id_structure, archived=False).count()
    total_medecins = Utilisateur.query.filter_by(id_structure=current_user.id_structure, role='medecin').count()
    consultations_mois = Consultation.query.filter(
        Consultation.id_patient.in_(
            db.session.query(Patient.id).filter_by(id_structure=current_user.id_structure)
        ),
        Consultation.date_consultation >= datetime.utcnow().replace(day=1)
    ).count()
    
    structure = current_user.structure
    zone_key = _resoudre_zone_climatique(structure)

    return render_template('structure/dashboard.html',
                         total_patients=total_patients,
                         total_medecins=total_medecins,
                         consultations_mois=consultations_mois,
                         pays_liste=_PAYS_AFRIQUE_LISTE,
                         zone_climatique_label=_ZONES_CLIMATIQUES[zone_key]['label'] if zone_key else None)


@app.route('/structure/localisation', methods=['POST'])
@login_required
def structure_localisation():
    """Enregistre le pays/ville de la structure — sert à déterminer sa zone
    climatique pour contextualiser les observations épidémiologiques des
    statistiques (voir _resoudre_zone_climatique). L'appli n'est pas
    utilisée qu'à Lomé/au Togo : sans ce réglage, aucune interprétation
    locale (saison des pluies, harmattan...) n'est générée, par choix."""
    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    structure = current_user.structure
    pays = (request.form.get('pays') or '').strip()
    if pays == '__autre__':
        pays = (request.form.get('pays_autre') or '').strip()
    ville = (request.form.get('ville') or '').strip()

    structure.pays = pays or None
    structure.ville = ville or None
    db.session.commit()
    flash('Localisation de la structure mise à jour.', 'success')
    return redirect(url_for('structure_dashboard'))

@app.route('/patients')
@login_required
@has_permission('PATIENTS')
def patients_list():
    # Super Admin n'a pas accès aux patients
    if current_user.role == 'super_admin':
        flash('Accès non autorisé. Zone réservée aux structures médicales.', 'danger')
        return redirect(url_for('admin_dashboard'))
    
    from models import Patient
    from sqlalchemy import or_
    
    # ⭐ RÉCUPÉRER LES FILTRES
    statut_filter = request.args.get('statut', '')
    search = request.args.get('search', '')
    
    # ⭐ CONSTRUIRE LA REQUÊTE DE BASE
    if current_user.role == 'admin_structure':
        query = Patient.query.filter_by(id_structure=current_user.id_structure, archived=False)
    elif current_user.role == 'medecin':
        query = Patient.query.filter_by(
            id_structure=current_user.id_structure, 
            id_medecin_referent=current_user.id,
            archived=False
        )
    else:  # secretaire
        query = Patient.query.filter_by(id_structure=current_user.id_structure, archived=False)
    
    # ⭐ APPLIQUER LE FILTRE STATUT
    if statut_filter:
        if statut_filter == 'GUERI':
            query = query.filter(Patient.statut_medical == 'GUERI')
        elif statut_filter == 'EN_TRAITEMENT':
            query = query.filter(Patient.statut_medical == 'EN_TRAITEMENT')
        elif statut_filter == 'PREMIERE_VISITE':
            query = query.filter(Patient.statut_medical == 'PREMIERE_VISITE')
        elif statut_filter == 'TRANSFERE':
            query = query.filter(Patient.statut_medical == 'TRANSFERE')
        elif statut_filter == 'PERDU_VUE':
            query = query.filter(Patient.statut_medical == 'PERDU_VUE')
    
    # ⭐ APPLIQUER LA RECHERCHE
    if search:
        search = search.strip()
        query = query.filter(_patient_recherche_conditions(search))
    
    patients = query.order_by(Patient.date_creation.desc()).all()
    
    # ⭐ COMPTER LES STATUTS POUR LE FILTRE
    total_gueris = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        statut_medical='GUERI',
        archived=False
    ).count()
    
    total_traitement = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        statut_medical='EN_TRAITEMENT',
        archived=False
    ).count()
    
    total_attente = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        statut_medical='PREMIERE_VISITE',
        archived=False
    ).count()
    
    return render_template('patients/list.html', 
                         patients=patients,
                         statut_actuel=statut_filter,
                         search=search,
                         total_gueris=total_gueris,
                         total_traitement=total_traitement,
                         total_attente=total_attente)

@app.route('/patients/ajouter', methods=['GET', 'POST'])
@login_required
@has_permission('PATIENTS')
def patient_ajouter():
    from models import Patient, Utilisateur, Consultation
    from datetime import datetime
    
    # Récupérer la liste des médecins
    medecins = []
    if current_user.role in ['admin_structure', 'secretaire']:
        medecins = Utilisateur.query.filter_by(
            id_structure=current_user.id_structure, 
            role='medecin',
            actif=True
        ).all()
    elif current_user.role == 'medecin':
        medecins = [current_user]
    
    if request.method == 'POST':
        # Identité
        nom = request.form.get('nom')
        prenom = request.form.get('prenom')
        date_naissance = request.form.get('date_naissance')
        sexe = request.form.get('sexe')
        telephone = request.form.get('telephone')
        email = request.form.get('email')
        adresse = request.form.get('adresse')
        profession = request.form.get('profession')
        
        # Assurance principale
        type_assurance = request.form.get('type_assurance')
        autre_assurance_nom = request.form.get('autre_assurance_nom')
        num_assure = request.form.get('num_assure')
        
        # ⭐ ASSURANCE 2 (NOUVEAU)
        assurance2_nom = request.form.get('assurance2_nom')
        taux_assurance2 = request.form.get('taux_assurance2')
        numero_assure2 = request.form.get('numero_assure2')
        
        # ⭐ PERSONNE À PRÉVENIR (NOUVEAU)
        personne_a_prevenir_nom = request.form.get('personne_a_prevenir_nom')
        personne_a_prevenir_telephone = request.form.get('personne_a_prevenir_telephone')
        personne_a_prevenir_relation = request.form.get('personne_a_prevenir_relation')
        
        # ⭐ TAUX DE PRISE EN CHARGE (NOUVEAU)
        taux_prise_charge = request.form.get('taux_prise_charge')

        groupe_sanguin = request.form.get('groupe_sanguin')

        # Médecin référent
        id_medecin_referent = request.form.get('id_medecin_referent')
        
        # Constantes vitales
        temperature = request.form.get('temperature')
        tension = request.form.get('tension')
        pouls = request.form.get('pouls')
        saturation = request.form.get('saturation')
        poids = request.form.get('poids')
        taille = request.form.get('taille')
        imc = request.form.get('imc')
        
        # Notes
        notes = request.form.get('notes')
        
        # Validation
        if not nom or not prenom:
            flash('Le nom et le prénom sont obligatoires', 'danger')
            return redirect(url_for('patient_ajouter'))
        
        # Création du patient avec TOUS les champs
        patient = Patient(
            id_structure=current_user.id_structure if current_user.id_structure else 1,
            nom=nom,
            prenom=prenom,
            date_naissance=datetime.strptime(date_naissance, '%Y-%m-%d') if date_naissance else None,
            sexe=sexe,
            telephone=telephone,
            email=email,
            adresse=adresse,
            profession=profession,
            groupe_sanguin=groupe_sanguin,
            # Assurance principale
            type_assurance=type_assurance,
            autre_assurance_nom=autre_assurance_nom if type_assurance == 'AUTRE_ASSURANCE' else None,
            num_assure=num_assure,
            
            # ⭐ ASSURANCE 2
            assurance2_nom=assurance2_nom,
            taux_assurance2=float(taux_assurance2) if taux_assurance2 else None,
            numero_assure2=numero_assure2,
            
            # ⭐ PERSONNE À PRÉVENIR
            personne_a_prevenir_nom=personne_a_prevenir_nom,
            personne_a_prevenir_telephone=personne_a_prevenir_telephone,
            personne_a_prevenir_relation=personne_a_prevenir_relation,
            
            # ⭐ TAUX DE PRISE EN CHARGE
            taux_prise_charge=float(taux_prise_charge) if taux_prise_charge else None,
            
            # Médecin référent
            id_medecin_referent=int(id_medecin_referent) if id_medecin_referent else None,
            
            # Constantes
            temperature_c=float(temperature) if temperature else None,
            tension_arterielle=tension,
            pulse_bpm=int(pouls) if pouls else None,
            oxygene_saturation=int(saturation) if saturation else None,
            poids_kg=float(poids) if poids else None,
            taille_cm=float(taille) if taille else None,
            imc=float(imc) if imc else None,
            
            # Notes
            notes=notes,
            
            # Statut
            statut_medical='PREMIERE_VISITE',
            date_premiere_visite=datetime.utcnow(),
            archived=False
        )
        
        db.session.add(patient)
        db.session.flush()
        
        # Créer une première consultation avec les constantes
        consultation = Consultation(
            id_patient=patient.id,
            id_medecin=current_user.id if current_user.role == 'medecin' else int(id_medecin_referent) if id_medecin_referent else None,
            motif="Première consultation - Enregistrement initial",
            temperature_c=float(temperature) if temperature else None,
            tension_arterielle=tension,
            pulse_bpm=int(pouls) if pouls else None,
            oxygene_saturation=int(saturation) if saturation else None,
            poids_kg=float(poids) if poids else None,
            taille_cm=float(taille) if taille else None,
            imc=float(imc) if imc else None,
            date_consultation=datetime.utcnow()
        )
        db.session.add(consultation)
        
        db.session.commit()
        
        flash(f'✅ Patient {prenom} {nom} créé avec succès !', 'success')
        flash('📝 Renseignez maintenant les antécédents du patient.', 'info')
        return redirect(url_for('patient_antecedents', patient_id=patient.id))
    
    return render_template('patients/ajouter.html', medecins=medecins)

@app.route('/patient/<int:id>')
@login_required
@has_permission('PATIENTS')
def patient_detail(id):
    from models import Patient, Consultation, Prescription, ExamenPhysique, SectionExamenPhysique, Hospitalisation
    from datetime import datetime
    import json
    
    patient = Patient.query.get_or_404(id)
    
    # Vérification pour le médecin
    if current_user.role == 'medecin' and patient.id_medecin_referent is not None and patient.id_medecin_referent != current_user.id:
        flash('Accès non autorisé - Ce patient n\'est pas votre patient référent', 'danger')
        return redirect(url_for('patients_list'))
    
    consultations = Consultation.query.filter_by(id_patient=patient.id).order_by(Consultation.date_consultation.desc()).all()
    prescriptions = Prescription.query.filter_by(id_patient=patient.id).order_by(Prescription.date_prescription.desc()).all()

    from models import ActePose
    soins_poses = ActePose.query.filter_by(patient_id=patient.id).order_by(ActePose.date_pose.desc()).all()

    # ⭐ Récupérer les sections standard pour référence
    sections_standard = SectionExamenPhysique.query.filter_by(actif=True).order_by(SectionExamenPhysique.ordre).all()
    sections_standard_dict = {s.nom: s.texte_fr for s in sections_standard}
    
    # ⭐ Pour chaque consultation, récupérer les sections modifiées
    for consultation in consultations:
        examen = ExamenPhysique.query.filter_by(consultation_id=consultation.id).first()
        if examen and examen.sections_modifiees:
            try:
                consultation.sections_modifiees = json.loads(examen.sections_modifiees)
                consultation.examen_complet = examen.examen_complet
            except:
                consultation.sections_modifiees = {}
                consultation.examen_complet = None
        else:
            consultation.sections_modifiees = {}
            consultation.examen_complet = None

        # ⭐ Ajouter les sections standard pour comparaison
        consultation.sections_standard = sections_standard_dict

        # ⭐ FIX : un résultat saisi via l'éditeur en ligne ou un fichier
        # importé (voir saisir_resultats_analyse) ne met jamais à jour
        # consultation.resultats_biologie/imagerie (texte libre uniquement,
        # voir le commentaire à cet endroit) — sans ceci, ce résultat
        # n'apparaissait nulle part dans le dossier du patient, alors qu'il
        # est bien "Terminé" et visible dans la file labo/radio.
        consultation.resultats_riches = [
            a for a in consultation.analyses_demandees
            if a.statut == 'TERMINE' and (a.contenu_html or a.fichier_data)
        ]

    # ⭐ NOUVEAU : Historique des ordonnances, patient-wide — jusqu'ici
    # aucune vue d'ensemble n'existait dans le dossier (contrairement aux
    # soins posés, voir soins/_widget.html) : une ordonnance de consultation
    # n'apparaissait que noyée dans la carte de SA consultation, et une
    # ordonnance d'hospitalisation n'apparaissait NULLE PART dans le
    # dossier patient (les hospitalisations elles-mêmes n'y sont même pas
    # listées). Un item par consultation/hospitalisation ayant une
    # ordonnance active ; le détail complet (toutes les versions) reste à
    # un clic via "Voir historique" (pages déjà existantes).
    hospitalisations = Hospitalisation.query.filter_by(
        patient_id=patient.id
    ).order_by(Hospitalisation.date_debut.desc()).all()

    historique_ordonnances = []
    for consultation in consultations:
        o = consultation.ordonnance_active
        if o and o.get_medicaments_list():
            historique_ordonnances.append({
                'origine': 'Consultation',
                'origine_label': f"Consultation du {consultation.date_consultation.strftime('%d/%m/%Y')}",
                'version': o.version,
                'medicaments': o.get_medicaments_list(),
                'date': o.date_redaction,
                'prescrit_par_nom': f"{o.redacteur.prenom} {o.redacteur.nom}" if o.redacteur else '',
                'imprimer_url': url_for('imprimer_ordonnance_consultation', id=consultation.id),
                'historique_url': url_for('historique_ordonnances_consultation', id=consultation.id),
                'fichier': _fichier_source_info(o.source_type, o.source_id, 'ordonnance'),
            })

    for hosp in hospitalisations:
        meds = []
        if hosp.ordonnance_prescite:
            try:
                meds = json.loads(hosp.ordonnance_prescite)
            except Exception:
                meds = []
        if meds:
            historique_ordonnances.append({
                'origine': 'Hospitalisation',
                'origine_label': f"Hospitalisation du {hosp.date_debut.strftime('%d/%m/%Y')} ({hosp.service})",
                'version': hosp.ordonnance_version or 1,
                'medicaments': meds,
                'date': hosp.updated_at,
                'prescrit_par_nom': f"{hosp.createur.prenom} {hosp.createur.nom}" if hosp.createur else '',
                'imprimer_url': url_for('imprimer_ordonnance', id=hosp.id),
                'historique_url': url_for('historique_ordonnances_hospitalisation', id=hosp.id),
                # ⭐ Hospitalisation.ordonnance_prescite n'a pas son propre
                # source_type/source_id (juste le JSON copié) — seul
                # protocole_id trace sa provenance, traité comme
                # source_type='protocole' pour retrouver le fichier importé.
                'fichier': _fichier_source_info('protocole', hosp.protocole_id, 'ordonnance') if hosp.protocole_id else None,
            })

    historique_ordonnances.sort(key=lambda x: x['date'] or datetime.min, reverse=True)

    # ⭐ La note d'admission n'est plus reprise ici dans le dossier patient
    # général (patron : pas nécessaire sur cette page) — elle reste
    # consultable sur la fiche hospitalisation elle-même (detail_hospitalisation).

    return render_template('patients/detail.html',
                         patient=patient,
                         consultations=consultations,
                         prescriptions=prescriptions,
                         soins_poses=soins_poses,
                         actes_soins_habituels=ACTES_SOINS_HABITUELS,
                         historique_ordonnances=historique_ordonnances,
                         hospitalisations=hospitalisations,
                         now=datetime.now())

@app.route('/consultation/ajouter', methods=['GET', 'POST'])
@login_required
def consultation_ajouter():
    from models import Patient, Consultation, AnalyseDemande, Prescription
    from datetime import datetime
    import json
    
    # Récupération des patients
    if current_user.role == 'medecin':
        patients = Patient.query.filter_by(
            id_structure=current_user.id_structure,
            id_medecin_referent=current_user.id,
            archived=False
        ).all()
    else:
        patients = Patient.query.filter_by(
            id_structure=current_user.id_structure,
            archived=False
        ).all()
    
    if request.method == 'POST':
        id_patient = request.form.get('id_patient')
        motif = request.form.get('motif')
        diagnostic = request.form.get('diagnostic')
        orientation = request.form.get('orientation') or None
        if orientation not in ('ambulatoire', 'hospitalisation'):
            orientation = None

        # ============================================================
        # RÉCUPÉRATION DES CHAMPS HPI
        # ============================================================
        hpi_date_debut = request.form.get('hpi_date_debut')
        hpi_debut_type = request.form.get('hpi_debut_type')
        hpi_circonstances = request.form.get('hpi_circonstances')
        hpi_evolution = request.form.get('hpi_evolution')
        hpi_facteurs = request.form.get('hpi_facteurs')
        hpi_traitements = request.form.get('hpi_traitements')
        hpi_signes = request.form.get('hpi_signes')
        
        # ⭐ RÉCUPÉRER LES CHECKBOXES DES SIGNES ASSOCIÉS
        hpi_fievre = request.form.get('hpi_fievre') == 'on'
        hpi_nausees = request.form.get('hpi_nausees') == 'on'
        hpi_douleur = request.form.get('hpi_douleur') == 'on'
        hpi_cephalées = request.form.get('hpi_cephalées') == 'on'
        hpi_vertiges = request.form.get('hpi_vertiges') == 'on'
        hpi_dyspnee = request.form.get('hpi_dyspnee') == 'on'
        
        hpi_trauma_mecanisme = request.form.get('hpi_trauma_mecanisme')
        hpi_trauma_heure = request.form.get('hpi_trauma_heure')
        hpi_trauma_pc = request.form.get('hpi_trauma_pc')
        hpi_trauma_description = request.form.get('hpi_trauma_description')
        
        hpi_morsure_type = request.form.get('hpi_morsure_type')
        hpi_morsure_espece = request.form.get('hpi_morsure_espece')
        hpi_morsure_siege = request.form.get('hpi_morsure_siege')
        hpi_morsure_signes = request.form.get('hpi_morsure_signes')
        
        hpi_intox_substance = request.form.get('hpi_intox_substance')
        hpi_intox_heure = request.form.get('hpi_intox_heure')
        hpi_intox_circonstances = request.form.get('hpi_intox_circonstances')
        
        hpi_autres_signes = request.form.get('hpi_autres_signes')
        hpi_autre_infos = request.form.get('hpi_autre_infos')
        
        hpi_complements_medecin = request.form.get('hpi_complements_medecin')
        histoire_maladie = request.form.get('histoire_maladie')

        # Constantes
        tension = request.form.get('tension')
        temperature = request.form.get('temperature')
        pouls = request.form.get('pouls')
        saturation = request.form.get('saturation')
        poids = request.form.get('poids')
        taille = request.form.get('taille')
        imc = request.form.get('imc')
        
        # Examens
        examens_cliniques = request.form.get('examens_cliniques')
        examens_biologie = request.form.get('examens_biologie')
        examens_imagerie = request.form.get('examens_imagerie')
        
        traitement = request.form.get('traitement')
        notes = request.form.get('notes')
        prochain_rdv = request.form.get('prochain_rdv')
        arret_travail = request.form.get('arret_travail') == 'on'
        arret_jours = request.form.get('arret_jours')
        statut_medical = request.form.get('statut_medical')
        cim10 = request.form.get('cim10')
        
        allergies = request.form.get('allergies')
        traitements_en_cours = request.form.get('traitements_en_cours')
        antecedents_medicaux = request.form.get('antecedents_medicaux')
        antecedents_chirurgicaux = request.form.get('antecedents_chirurgicaux')
        
        medicaments_prescrits = request.form.get('medicaments_prescrits')
        
        # ⭐⭐⭐ RÉCUPÉRER LE PATIENT ⭐⭐⭐
        patient = Patient.query.get(id_patient)
        if not patient:
            flash('Patient non trouvé', 'danger')
            return redirect(url_for('consultation_ajouter'))
        
        # ⭐⭐⭐ LE MÉDECIN QUI CONSULTE DEVIENT LE RÉFÉRENT ⭐⭐⭐
        patient.id_medecin_referent = current_user.id
        print(f"✅ Médecin référent : Dr {current_user.nom} {current_user.prenom} pour le patient {patient.nom} {patient.prenom}")
        
        # ═══════════════════════════════════════════
        # ⭐ 2. RÉCUPÉRATION DE LA CONSULTATION TEMPORAIRE
        # ═══════════════════════════════════════════
        
        consultation = Consultation.query.filter_by(
            id_patient=patient.id,
            is_temporary=True
        ).first()
        
        if consultation:
            # ⭐ METTRE À JOUR LA CONSULTATION TEMPORAIRE
            print(f"✅ Consultation temporaire #{consultation.id} trouvée !")
            
            consultation.motif = motif
            consultation.diagnostic = diagnostic
            consultation.tension_arterielle = tension
            consultation.temperature_c = float(temperature) if temperature else None
            consultation.pulse_bpm = int(pouls) if pouls else None
            consultation.oxygene_saturation = int(saturation) if saturation else None
            consultation.poids_kg = float(poids) if poids else None
            consultation.taille_cm = float(taille) if taille else None
            consultation.imc = float(imc) if imc else None
            consultation.examens_cliniques = examens_cliniques
            consultation.examens_biologie = examens_biologie
            consultation.examens_imagerie = examens_imagerie
            consultation.traitement_prescrit = traitement
            consultation.notes_cliniques = notes
            consultation.cim10 = cim10
            consultation.arret_travail = arret_travail
            consultation.arret_jours = int(arret_jours) if arret_jours else None
            consultation.prochain_rdv = datetime.strptime(prochain_rdv, '%Y-%m-%d') if prochain_rdv else None
            consultation.allergies = allergies
            consultation.traitements_en_cours = traitements_en_cours
            consultation.antecedents_medicaux = antecedents_medicaux
            consultation.antecedents_chirurgicaux = antecedents_chirurgicaux
            
            # HPI - Champs texte
            consultation.hpi_date_debut = datetime.strptime(hpi_date_debut, '%Y-%m-%d').date() if hpi_date_debut else None
            consultation.hpi_debut_type = hpi_debut_type
            consultation.hpi_circonstances = hpi_circonstances
            consultation.hpi_evolution = hpi_evolution
            consultation.hpi_facteurs = hpi_facteurs
            consultation.hpi_traitements = hpi_traitements
            consultation.hpi_signes = hpi_signes
            consultation.hpi_trauma_mecanisme = hpi_trauma_mecanisme
            consultation.hpi_trauma_heure = datetime.fromisoformat(hpi_trauma_heure) if hpi_trauma_heure else None
            consultation.hpi_trauma_pc = hpi_trauma_pc
            consultation.hpi_trauma_description = hpi_trauma_description
            consultation.hpi_morsure_type = hpi_morsure_type
            consultation.hpi_morsure_espece = hpi_morsure_espece
            consultation.hpi_morsure_siege = hpi_morsure_siege
            consultation.hpi_morsure_signes = hpi_morsure_signes
            consultation.hpi_intox_substance = hpi_intox_substance
            consultation.hpi_intox_heure = datetime.fromisoformat(hpi_intox_heure) if hpi_intox_heure else None
            consultation.hpi_intox_circonstances = hpi_intox_circonstances
            consultation.hpi_autres_signes = hpi_autres_signes
            consultation.hpi_autre_infos = hpi_autre_infos
            consultation.hpi_complements_medecin = hpi_complements_medecin
            consultation.histoire_maladie = histoire_maladie
            
            # ⭐⭐⭐ AJOUTER LES SIGNES ASSOCIÉS ⭐⭐⭐
            consultation.hpi_fievre = hpi_fievre
            consultation.hpi_nausees = hpi_nausees
            consultation.hpi_douleur = hpi_douleur
            consultation.hpi_cephalées = hpi_cephalées
            consultation.hpi_vertiges = hpi_vertiges
            consultation.hpi_dyspnee = hpi_dyspnee
            
            # ⭐ MARQUER COMME DÉFINITIVE
            consultation.is_temporary = False
            consultation.statut = 'terminee'
            consultation.orientation = orientation
            consultation.id_medecin = current_user.id
            consultation.date_consultation = datetime.utcnow()

        else:
            # ⭐ PAS DE CONSULTATION TEMPORAIRE : EN CRÉER UNE
            print(f"⚠️ Aucune consultation temporaire trouvée pour patient #{patient.id}")

            consultation = Consultation(
                id_patient=int(id_patient),
                id_medecin=current_user.id,
                motif=motif,
                diagnostic=diagnostic,
                orientation=orientation,
                tension_arterielle=tension,
                temperature_c=float(temperature) if temperature else None,
                pulse_bpm=int(pouls) if pouls else None,
                oxygene_saturation=int(saturation) if saturation else None,
                poids_kg=float(poids) if poids else None,
                taille_cm=float(taille) if taille else None,
                imc=float(imc) if imc else None,
                examens_cliniques=examens_cliniques,
                examens_biologie=examens_biologie,
                examens_imagerie=examens_imagerie,
                traitement_prescrit=traitement,
                notes_cliniques=notes,
                arret_travail=arret_travail,
                arret_jours=int(arret_jours) if arret_jours else None,
                prochain_rdv=datetime.strptime(prochain_rdv, '%Y-%m-%d') if prochain_rdv else None,
                allergies=allergies,
                traitements_en_cours=traitements_en_cours,
                antecedents_medicaux=antecedents_medicaux,
                antecedents_chirurgicaux=antecedents_chirurgicaux,
                cim10=cim10,
                hpi_date_debut=datetime.strptime(hpi_date_debut, '%Y-%m-%d').date() if hpi_date_debut else None,
                hpi_debut_type=hpi_debut_type,
                hpi_circonstances=hpi_circonstances,
                hpi_evolution=hpi_evolution,
                hpi_facteurs=hpi_facteurs,
                hpi_traitements=hpi_traitements,
                hpi_signes=hpi_signes,
                hpi_trauma_mecanisme=hpi_trauma_mecanisme,
                hpi_trauma_heure=datetime.fromisoformat(hpi_trauma_heure) if hpi_trauma_heure else None,
                hpi_trauma_pc=hpi_trauma_pc,
                hpi_trauma_description=hpi_trauma_description,
                hpi_morsure_type=hpi_morsure_type,
                hpi_morsure_espece=hpi_morsure_espece,
                hpi_morsure_siege=hpi_morsure_siege,
                hpi_morsure_signes=hpi_morsure_signes,
                hpi_intox_substance=hpi_intox_substance,
                hpi_intox_heure=datetime.fromisoformat(hpi_intox_heure) if hpi_intox_heure else None,
                hpi_intox_circonstances=hpi_intox_circonstances,
                hpi_autres_signes=hpi_autres_signes,
                hpi_autre_infos=hpi_autre_infos,
                hpi_complements_medecin=hpi_complements_medecin,
                histoire_maladie=histoire_maladie,
                date_consultation=datetime.utcnow(),
                is_temporary=False,
                # ⭐⭐⭐ AJOUTER LES SIGNES ASSOCIÉS ⭐⭐⭐
                hpi_fievre=hpi_fievre,
                hpi_nausees=hpi_nausees,
                hpi_douleur=hpi_douleur,
                hpi_cephalées=hpi_cephalées,
                hpi_vertiges=hpi_vertiges,
                hpi_dyspnee=hpi_dyspnee
            )
            db.session.add(consultation)
            db.session.flush()
            print(f"✅ Nouvelle consultation #{consultation.id} créée")
        
        # ═══════════════════════════════════════════
        # 3. CRÉATION DES ANALYSES
        # ═══════════════════════════════════════════
        
        if examens_biologie:
            for ligne in examens_biologie.split('\n'):
                nom = ligne.strip()
                if nom:
                    analyse = AnalyseDemande(
                        consultation_id=consultation.id,
                        patient_id=consultation.id_patient,
                        structure_id=current_user.id_structure,
                        type_analyse='BIOLOGIE',
                        nom_analyse=nom,
                        prescrit_par=current_user.id,
                        statut='EN_ATTENTE'
                    )
                    db.session.add(analyse)
        
        if examens_imagerie:
            for ligne in examens_imagerie.split('\n'):
                nom = ligne.strip()
                if nom:
                    analyse = AnalyseDemande(
                        consultation_id=consultation.id,
                        patient_id=consultation.id_patient,
                        structure_id=current_user.id_structure,
                        type_analyse='IMAGERIE',
                        nom_analyse=nom,
                        prescrit_par=current_user.id,
                        statut='EN_ATTENTE'
                    )
                    db.session.add(analyse)
        
        # ═══════════════════════════════════════════
        # 4. CRÉATION DES PRESCRIPTIONS
        # ═══════════════════════════════════════════
        
        prescriptions_creees = 0
        
        if medicaments_prescrits:
            try:
                meds_data = json.loads(medicaments_prescrits)
                
                for med in meds_data:
                    prescription = Prescription(
                        id_patient=int(id_patient),
                        id_consultation=consultation.id,
                        id_medecin=current_user.id,
                        medicament=med.get('nom', ''),
                        dosage=med.get('dosage', ''),
                        forme=med.get('forme', ''),
                        quantite=str(med.get('quantite', 1)),
                        duree_jours=med.get('duree', 7),
                        frequence=med.get('posologie', ''),
                        instructions=med.get('instructions', ''),
                        renouvelable=med.get('renouvelable', False),
                        type_prescription='medicament',
                        prescripteur=f"{current_user.prenom} {current_user.nom}",
                        statut='active',
                        date_prescription=datetime.utcnow(),
                        notes=med.get('notes', ''),
                        source_id=med.get('source_id'),
                        stock_disponible=med.get('stock')
                    )
                    db.session.add(prescription)
                    prescriptions_creees += 1
                    
                print(f"✅ {len(meds_data)} prescription(s) médicamenteuse(s) enregistrée(s)")
                
            except Exception as e:
                print(f"❌ Erreur sauvegarde prescriptions médicaments: {e}")
                import traceback
                traceback.print_exc()
        
        if examens_biologie:
            for ligne in examens_biologie.split('\n'):
                nom = ligne.strip()
                if nom:
                    prescription = Prescription(
                        id_patient=int(id_patient),
                        id_consultation=consultation.id,
                        id_medecin=current_user.id,
                        medicament=nom,
                        type_prescription='acte',
                        prescripteur=f"{current_user.prenom} {current_user.nom}",
                        statut='active',
                        date_prescription=datetime.utcnow()
                    )
                    db.session.add(prescription)
                    prescriptions_creees += 1
        
        if examens_imagerie:
            for ligne in examens_imagerie.split('\n'):
                nom = ligne.strip()
                if nom:
                    prescription = Prescription(
                        id_patient=int(id_patient),
                        id_consultation=consultation.id,
                        id_medecin=current_user.id,
                        medicament=nom,
                        type_prescription='acte',
                        prescripteur=f"{current_user.prenom} {current_user.nom}",
                        statut='active',
                        date_prescription=datetime.utcnow()
                    )
                    db.session.add(prescription)
                    prescriptions_creees += 1
        
        # ═══════════════════════════════════════════
        # 5. MISE À JOUR DU PATIENT
        # ═══════════════════════════════════════════
        
        if temperature:
            patient.temperature_c = float(temperature)
        if tension:
            patient.tension_arterielle = tension
        if pouls:
            patient.pulse_bpm = int(pouls)
        if saturation:
            patient.oxygene_saturation = int(saturation)
        if poids:
            patient.poids_kg = float(poids)
        if taille:
            patient.taille_cm = float(taille)
        if imc:
            patient.imc = float(imc)
        
        patient.date_derniere_consultation = datetime.utcnow()
        
        if statut_medical:
            patient.statut_medical = statut_medical
            if statut_medical == 'GUERI':
                patient.date_guerison = datetime.utcnow()
        elif patient.statut_medical == 'PREMIERE_VISITE':
            patient.statut_medical = 'EN_TRAITEMENT'
        
        # ═══════════════════════════════════════════
        # 6. COMMIT FINAL
        # ═══════════════════════════════════════════
        
        db.session.commit()
        
        if prescriptions_creees > 0:
            try:
                from tasks import sync_prescriptions_to_ghp
                result = sync_prescriptions_to_ghp()
                if result.get('success'):
                    print(f"✅ {result.get('message')}")
                else:
                    print(f"⚠️ {result.get('message')}")
            except Exception as e:
                print(f"⚠️ Erreur sync auto: {e}")

        _pousser_rendez_vous_ghp(consultation, patient, f"{current_user.prenom} {current_user.nom}")

        flash('Consultation enregistrée avec succès', 'success')
        if orientation == 'hospitalisation':
            flash('🛏️ Patient orienté vers une hospitalisation — visible dans "Hospitalisations en attente".', 'info')
        return redirect(url_for('patient_detail', id=id_patient))

    empty_consultation = Consultation()
    return render_template('consultations/ajouter.html', patients=patients, consultation=empty_consultation)


@app.route('/consultation/<int:id>')
@login_required
def consultation_detail(id):
    from models import Consultation, Patient, ExamenType
    
    consultation = Consultation.query.get_or_404(id)
    patient = Patient.query.get(consultation.id_patient)
    
    # ⭐ RÉCUPÉRER LES EXAMENS TYPES DISPONIBLES POUR LA STRUCTURE
    examens_types = ExamenType.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    # ⭐ RÉCUPÉRER LES EXAMENS PRESCRITS (si la relation existe)
    examens_prescrits = []
    if hasattr(consultation, 'examens_prescrits'):
        examens_prescrits = consultation.examens_prescrits.order_by(
            ExamenPrescrit.date_prescription.desc()
        ).all() if consultation.examens_prescrits else []

    from models import ActePose
    soins_poses = ActePose.query.filter_by(consultation_id=consultation.id).order_by(ActePose.date_pose.desc()).all()

    # ⭐ EXAMEN PHYSIQUE — sections en lecture seule, celles qui diffèrent
    # du texte par défaut du catalogue apparaissent grisées (patron :
    # "fait de même pour consultation details ... là où l'examen
    # physique apparaît").
    from models import ExamenPhysique
    examen_physique_obj = ExamenPhysique.query.filter_by(consultation_id=consultation.id).first()
    examen_physique_sections = _rendre_sections_examen_physique(examen_physique_obj)

    return render_template('consultations/detail.html',
                         consultation=consultation,
                         patient=patient,
                         examens_types=examens_types,
                         examens_prescrits=examens_prescrits,
                         soins_poses=soins_poses,
                         actes_soins_habituels=ACTES_SOINS_HABITUELS,
                         examen_physique_sections=examen_physique_sections,
                         now=datetime.utcnow())


# ==================== PRESCRIPTIONS ====================

@app.route('/prescription/ajouter', methods=['GET', 'POST'])
@login_required
def prescription_ajouter():
    from models import Patient, Prescription
    from datetime import datetime, timezone
    import json
    
    try:
        # ✅ FORCER UNE NOUVELLE SESSION
        db.session.expire_all()
        
        if current_user.role == 'medecin':
            patients = Patient.query.filter_by(
                id_structure=current_user.id_structure,
                id_medecin_referent=current_user.id,
                archived=False
            ).all()
        else:
            patients = Patient.query.filter_by(
                id_structure=current_user.id_structure,
                archived=False
            ).all()
    except Exception as e:
        print(f"❌ Erreur récupération patients: {e}")
        patients = []
        flash('Erreur de chargement des patients', 'danger')
    
    if request.method == 'POST':
        try:
            id_patient = request.form.get('id_patient')
            notes = request.form.get('notes', '')
            medicaments_prescrits = request.form.get('medicaments_prescrits')
            
            if not medicaments_prescrits:
                flash('Veuillez ajouter au moins un médicament', 'danger')
                return redirect(url_for('prescription_ajouter'))
            
            meds_data = json.loads(medicaments_prescrits)
            
            if not meds_data:
                flash('Aucun médicament valide', 'danger')
                return redirect(url_for('prescription_ajouter'))
            
            patient = db.session.get(Patient, int(id_patient))
            if not patient:
                flash('Patient non trouvé', 'danger')
                return redirect(url_for('prescription_ajouter'))
            
            prescriptions = []
            
            for med in meds_data:
                nom_med = med.get('nom', '').strip()
                if not nom_med:
                    continue
                
                prescription = Prescription(
                    id_patient=int(id_patient),
                    id_medecin=current_user.id,
                    medicament=nom_med,
                    dosage=med.get('dosage', ''),
                    quantite=str(med.get('quantite', 1)),
                    duree_jours=int(med.get('duree', 7) or 7),
                    frequence=med.get('posologie', ''),
                    instructions=med.get('instructions', ''),
                    prescripteur=f"{current_user.prenom} {current_user.nom}",
                    statut='active',
                    date_prescription=datetime.now(timezone.utc),
                    notes=notes
                )
                prescriptions.append(prescription)
                print(f"✅ Prescription préparée: {nom_med}")
            
            if len(prescriptions) == 0:
                flash('Aucun médicament valide à enregistrer', 'danger')
                return redirect(url_for('prescription_ajouter'))
            
            # ✅ AJOUTER UN PAR UN AVEC FLUSH
            for p in prescriptions:
                db.session.add(p)
                db.session.flush()  # ← Flush après chaque ajout
            
            # ✅ COMMIT FINAL
            db.session.commit()
            
            flash(f'✅ {len(prescriptions)} prescription(s) enregistrée(s)', 'success')
            return redirect(url_for('patients_list'))
            
        except Exception as e:
            print(f"❌ Erreur: {e}")
            import traceback
            traceback.print_exc()
            db.session.rollback()
            flash(f'Erreur: {str(e)}', 'danger')
            return redirect(url_for('prescription_ajouter'))
    
    return render_template('prescriptions/ajouter.html', patients=patients)


# ==================== ACTES POSÉS ====================
# ⭐ Un acte posé est réalisé directement par le médecin/infirmier
# (pansement, injection, suture, petit soin...) — pas un examen prescrit à
# faire réaliser ailleurs (ça reste Prescription). Synchronisé vers GHP
# pour facturation, comme les prescriptions.

@app.route('/actes-poses/ajouter', methods=['GET', 'POST'])
@login_required
def acte_pose_ajouter():
    from models import Patient, ActeType, ActePose
    from datetime import datetime
    import json

    try:
        db.session.expire_all()
        if current_user.role == 'medecin':
            patients = Patient.query.filter_by(
                id_structure=current_user.id_structure,
                id_medecin_referent=current_user.id,
                archived=False
            ).all()
        else:
            patients = Patient.query.filter_by(
                id_structure=current_user.id_structure,
                archived=False
            ).all()
    except Exception as e:
        print(f"❌ Erreur récupération patients: {e}")
        patients = []
        flash('Erreur de chargement des patients', 'danger')

    if request.method == 'POST':
        try:
            id_patient = request.form.get('id_patient')
            notes = request.form.get('notes', '')
            actes_poses_json = request.form.get('actes_poses_json')
            date_pose_str = request.form.get('date_pose')
            heure_pose_str = request.form.get('heure_pose')

            if not id_patient:
                flash('Veuillez sélectionner un patient', 'danger')
                return redirect(url_for('acte_pose_ajouter'))

            # ⭐ Date/heure RÉELLES de l'acte — obligatoires, jamais déduites en
            # silence côté serveur (patron : "je veux qu'on ait la trace de ce
            # que l'infirmier a posé comme acte à quel moment quelle heure").
            if not date_pose_str or not heure_pose_str:
                flash('La date et l\'heure de l\'acte sont obligatoires', 'danger')
                return redirect(url_for('acte_pose_ajouter'))
            try:
                date_pose = datetime.strptime(f"{date_pose_str} {heure_pose_str}", '%Y-%m-%d %H:%M')
            except ValueError:
                flash('Date ou heure invalide', 'danger')
                return redirect(url_for('acte_pose_ajouter'))

            if not actes_poses_json:
                flash('Veuillez ajouter au moins un acte', 'danger')
                return redirect(url_for('acte_pose_ajouter'))

            actes_data = json.loads(actes_poses_json)
            actes_valides = [a for a in actes_data if a.get('nom', '').strip()]

            if not actes_valides:
                flash('Aucun acte valide', 'danger')
                return redirect(url_for('acte_pose_ajouter'))

            actes_crees = []
            for a in actes_valides:
                nom = a['nom'].strip()

                # ⭐ Si l'acte n'existe pas encore dans le catalogue de la
                # structure, on le crée à la volée (comme pour un
                # médicament ajouté manuellement).
                acte_type = ActeType.query.filter_by(
                    structure_id=current_user.id_structure, nom=nom
                ).first()
                if not acte_type:
                    acte_type = ActeType(
                        structure_id=current_user.id_structure,
                        nom=nom,
                        created_by=current_user.id
                    )
                    db.session.add(acte_type)
                    db.session.flush()

                acte_pose = ActePose(
                    patient_id=int(id_patient),
                    acte_type_id=acte_type.id,
                    nom=nom,
                    quantite=str(a.get('quantite', 1)),
                    notes=notes,
                    date_pose=date_pose,
                    pose_par_id=current_user.id,
                    statut='actif',
                    valide=False  # ⭐ atterrit en brouillon dans "Actes posés" — voir _actes_poses_liste()
                )
                db.session.add(acte_pose)
                actes_crees.append(acte_pose)

            db.session.commit()

            flash(
                f'✅ {len(actes_crees)} acte(s) posé(s) enregistré(s) — à vérifier et valider '
                f'dans l\'onglet "Actes posés" avant l\'envoi à GHP.',
                'success'
            )
            return redirect(url_for('actes_poses_liste'))

        except Exception as e:
            print(f"❌ Erreur: {e}")
            import traceback
            traceback.print_exc()
            db.session.rollback()
            flash(f'Erreur: {str(e)}', 'danger')
            return redirect(url_for('acte_pose_ajouter'))

    return render_template('actes_poses/ajouter.html', patients=patients)


@app.route('/actes-poses')
@login_required
def actes_poses_liste():
    from models import ActePose, Patient

    base_query = (
        ActePose.query
        .join(Patient, ActePose.patient_id == Patient.id)
        .filter(Patient.id_structure == current_user.id_structure)
    )

    # ⭐ Deux blocs bien distincts : ce qui reste à vérifier/valider (issu du
    # journal de soins ou de la saisie manuelle) en premier, bien visible ;
    # l'historique (déjà validé/envoyé/annulé) ensuite, pour référence.
    actes_a_valider = (
        base_query.filter(ActePose.valide == False, ActePose.statut == 'actif')
        .order_by(ActePose.date_pose.desc())
        .all()
    )
    historique = (
        base_query.filter(db.or_(ActePose.valide == True, ActePose.statut == 'annule'))
        .order_by(ActePose.date_pose.desc())
        .limit(200)
        .all()
    )
    return render_template('actes_poses/liste.html', actes_a_valider=actes_a_valider, historique=historique)


@app.route('/actes-poses/<int:id>/valider', methods=['POST'])
@login_required
def acte_pose_valider(id):
    """Vérifie/ajuste puis valide un acte posé — déclenche son envoi à GHP.
    Rien ne part vers GHP tant que cette étape n'a pas eu lieu (voir
    tasks.sync_actes_poses_to_ghp, filtré sur valide=True)."""
    from models import ActePose

    acte = ActePose.query.get_or_404(id)
    if acte.patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('actes_poses_liste'))
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('actes_poses_liste'))

    quantite = (request.form.get('quantite') or '').strip()
    if quantite:
        acte.quantite = quantite

    acte.valide = True
    acte.valide_par_id = current_user.id
    acte.date_validation = datetime.utcnow()
    db.session.commit()

    try:
        from tasks import sync_actes_poses_to_ghp
        result = sync_actes_poses_to_ghp()
        if result.get('success'):
            print(f"✅ {result.get('message')}")
    except Exception as e:
        print(f"⚠️ Erreur sync auto actes posés : {e}")

    flash(f'✅ "{acte.nom}" validé et envoyé à GHP', 'success')
    return redirect(url_for('actes_poses_liste'))


@app.route('/actes-poses/<int:id>/annuler', methods=['POST'])
@login_required
def acte_pose_annuler(id):
    """Annule un acte posé consigné par erreur — ne part jamais à GHP."""
    from models import ActePose

    acte = ActePose.query.get_or_404(id)
    if acte.patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('actes_poses_liste'))
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('actes_poses_liste'))

    acte.statut = 'annule'
    db.session.commit()
    flash(f'Acte "{acte.nom}" annulé', 'info')
    return redirect(url_for('actes_poses_liste'))


# ==================== ADMINISTRATION DES MÉDICAMENTS ====================
# ⭐ Suivi infirmier des médicaments prescrits (Consultation ET Hospitalisation
# — les deux créent des lignes Prescription, voir _creer_prescriptions_miroir)
# — planning, rappel, écart prévu/réel. Ne part JAMAIS vers GHP : lecture
# infirmier/médecin/admin_structure, écriture réservée à l'infirmier(-ère)
# (et admin_structure en dépannage), comme pour les Actes posés.
def _calculer_prochaine_echeance(prescription, heure_prevue_precedente, intervalle_heures):
    """Renvoie la prochaine heure_prevue, ancrée sur l'heure PRÉVUE (pas
    l'heure réelle) pour que des retards ponctuels ne décalent pas tout le
    planning restant — ou None si la prescription est arrivée à son terme
    (date_fin dépassée) OU si son administration a été arrêtée manuellement
    (voir infirmier_medicament_terminer) — sans ce garde-fou, un traitement
    sans date de fin connue continuerait à réclamer une dose à l'infini."""
    from datetime import timedelta
    if prescription.administration_arretee:
        return None
    prochaine = heure_prevue_precedente + timedelta(hours=intervalle_heures or 24)
    date_fin = prescription.date_fin
    if not date_fin and prescription.date_debut and prescription.duree_jours:
        date_fin = prescription.date_debut + timedelta(days=prescription.duree_jours)
    if date_fin and prochaine.date() > date_fin:
        return None
    return prochaine


@app.route('/infirmier/medicaments')
@login_required
def infirmier_medicaments():
    """Écran de suivi d'administration — lecture pour infirmier/médecin/
    admin_structure (patron : "permets aux médecins et à l'administrateur de
    voir cet onglet de l'infirmier"), écriture réservée à l'infirmier dans
    le template (mêmes boutons masqués pour les autres rôles, comme pour
    Actes posés)."""
    from models import Prescription, AdministrationMedicament, Patient

    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    # Prescriptions de médicaments actives, sans AUCUNE administration
    # planifiée pour l'instant — à planifier (1ère dose) par l'infirmier.
    deja_planifiees = db.session.query(AdministrationMedicament.prescription_id).distinct()
    a_planifier = (
        Prescription.query
        .join(Patient, Prescription.id_patient == Patient.id)
        .filter(
            Patient.id_structure == current_user.id_structure,
            Prescription.type_prescription == 'medicament',
            Prescription.statut == 'active',
            ~Prescription.id.in_(deja_planifiees)
        )
        .order_by(Prescription.date_prescription.desc())
        .all()
    )

    base_admin = (
        AdministrationMedicament.query
        .join(Patient, AdministrationMedicament.patient_id == Patient.id)
        .filter(Patient.id_structure == current_user.id_structure)
    )
    a_faire = (
        base_admin.filter(AdministrationMedicament.statut == 'a_faire')
        .order_by(AdministrationMedicament.heure_prevue.asc())
        .all()
    )

    patients = (
        Patient.query.filter_by(id_structure=current_user.id_structure, archived=False)
        .order_by(Patient.nom).all()
    )

    # ⭐ Protocoles de soins actifs — jusqu'ici l'infirmier ne voyait que les
    # lignes de médicaments/examens qui en découlent, sans jamais voir le
    # protocole lui-même (patron : "il doit voir le protocole des soins à
    # un protocole on peut assigner une ordonnance et examens à faire").
    from models import Hospitalisation
    hospitalisations_avec_protocole = (
        Hospitalisation.query
        .join(Patient, Hospitalisation.patient_id == Patient.id)
        .filter(
            Patient.id_structure == current_user.id_structure,
            Hospitalisation.statut == 'actif',
            Hospitalisation.protocole_id.isnot(None),
        )
        .order_by(Hospitalisation.date_debut.desc())
        .all()
    )

    return render_template('infirmier/medicaments.html', a_planifier=a_planifier, a_faire=a_faire, patients=patients,
                         hospitalisations_avec_protocole=hospitalisations_avec_protocole)


@app.route('/infirmier/medicaments/<int:prescription_id>/planifier', methods=['POST'])
@login_required
def infirmier_medicament_planifier(prescription_id):
    """Fixe la 1ère dose (heure prévue + fréquence + dose) d'une prescription
    — les doses suivantes seront ensuite créées automatiquement à mesure
    qu'on marque chaque dose comme faite."""
    from models import Prescription, AdministrationMedicament, Patient

    if current_user.role not in ['admin_structure', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('infirmier_medicaments'))

    prescription = Prescription.query.get_or_404(prescription_id)
    if prescription.patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('infirmier_medicaments'))

    date_str = request.form.get('date_prevue')
    heure_str = request.form.get('heure_prevue')
    intervalle = request.form.get('intervalle_heures', type=float)
    dose = (request.form.get('dose') or prescription.dosage or '').strip()

    if not date_str or not heure_str or not intervalle:
        flash('Date, heure et fréquence (en heures) sont obligatoires', 'danger')
        return redirect(url_for('infirmier_medicaments'))

    try:
        heure_prevue = datetime.strptime(f"{date_str} {heure_str}", '%Y-%m-%d %H:%M')
    except ValueError:
        flash('Date ou heure invalide', 'danger')
        return redirect(url_for('infirmier_medicaments'))

    admin = AdministrationMedicament(
        prescription_id=prescription.id,
        patient_id=prescription.id_patient,
        medicament=prescription.medicament,
        dose=dose,
        numero_dose=1,
        intervalle_heures=intervalle,
        heure_prevue=heure_prevue,
        statut='a_faire',
    )
    db.session.add(admin)
    db.session.commit()
    flash(f'Planning démarré pour "{prescription.medicament}"', 'success')
    return redirect(url_for('infirmier_medicaments'))


@app.route('/infirmier/medicaments/administration/<int:id>/marquer-fait', methods=['POST'])
@login_required
def infirmier_medicament_marquer_fait(id):
    """Marque une dose comme administrée — heure_reelle posée UNE SEULE FOIS
    ici (jamais modifiable ensuite), puis crée automatiquement la dose
    suivante de la même prescription si elle n'est pas arrivée à son terme."""
    from models import AdministrationMedicament

    if current_user.role not in ['admin_structure', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('infirmier_medicaments'))

    admin = AdministrationMedicament.query.get_or_404(id)
    if admin.patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('infirmier_medicaments'))
    if admin.statut == 'fait':
        flash('Cette dose est déjà marquée comme faite', 'warning')
        return redirect(url_for('infirmier_medicaments'))

    admin.statut = 'fait'
    admin.heure_reelle = datetime.utcnow()
    admin.fait_par_id = current_user.id
    db.session.flush()

    prochaine = _calculer_prochaine_echeance(admin.prescription, admin.heure_prevue, admin.intervalle_heures)
    if prochaine:
        suivante = AdministrationMedicament(
            prescription_id=admin.prescription_id,
            patient_id=admin.patient_id,
            medicament=admin.medicament,
            dose=admin.dose,
            numero_dose=admin.numero_dose + 1,
            intervalle_heures=admin.intervalle_heures,
            heure_prevue=prochaine,
            statut='a_faire',
        )
        db.session.add(suivante)

    db.session.commit()
    flash(f'"{admin.medicament}" administré — enregistré', 'success')
    return redirect(url_for('infirmier_medicaments'))


@app.route('/infirmier/medicaments/prescription/<int:prescription_id>/terminer', methods=['POST'])
@login_required
def infirmier_medicament_terminer(prescription_id):
    """Marque un traitement comme terminé — patron : "il faut prévoir qu'on
    marque fin à un traitement, pour ne pas que la machine considère que le
    traitement continue et ne cesse d'alerter". N'annule QUE le planning
    d'administration (administration_arretee, local) : la prescription
    elle-même n'est pas touchée et reste synchronisable vers GHP comme
    avant. Toute dose encore "à faire" pour ce traitement est annulée —
    elle n'apparaîtra plus dans les alertes ni dans "Doses à faire"."""
    from models import Prescription, AdministrationMedicament

    if current_user.role not in ['admin_structure', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('infirmier_medicaments'))

    prescription = Prescription.query.get_or_404(prescription_id)
    if prescription.patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('infirmier_medicaments'))

    prescription.administration_arretee = True
    AdministrationMedicament.query.filter_by(
        prescription_id=prescription_id, statut='a_faire'
    ).update({'statut': 'annule'})
    db.session.commit()
    flash(f'Traitement "{prescription.medicament}" marqué comme terminé — plus aucune alerte ne sera générée.', 'info')
    return redirect(url_for('infirmier_medicaments'))


@app.route('/infirmier/medicaments/ajouter-adhoc', methods=['POST'])
@login_required
def infirmier_medicament_ajouter_adhoc():
    """Ajoute un médicament non prescrit par un médecin — la prescription
    créée est marquée origine_prescripteur='infirmier' pour qu'on sache
    qu'elle sort du circuit normal (patron : "cette prescription sera
    marquée comme prescrite par l'infirmier")."""
    from models import Prescription, AdministrationMedicament, Patient

    if current_user.role not in ['admin_structure', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('infirmier_medicaments'))

    id_patient = request.form.get('id_patient', type=int)
    medicament = (request.form.get('medicament') or '').strip()
    dosage = (request.form.get('dosage') or '').strip()
    date_str = request.form.get('date_prevue')
    heure_str = request.form.get('heure_prevue')
    intervalle = request.form.get('intervalle_heures', type=float)

    if not id_patient or not medicament or not date_str or not heure_str or not intervalle:
        flash('Patient, médicament, date, heure et fréquence sont obligatoires', 'danger')
        return redirect(url_for('infirmier_medicaments'))

    patient = Patient.query.get_or_404(id_patient)
    if patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('infirmier_medicaments'))

    try:
        heure_prevue = datetime.strptime(f"{date_str} {heure_str}", '%Y-%m-%d %H:%M')
    except ValueError:
        flash('Date ou heure invalide', 'danger')
        return redirect(url_for('infirmier_medicaments'))

    prescription = Prescription(
        id_patient=id_patient,
        medicament=medicament,
        dosage=dosage,
        type_prescription='medicament',
        prescripteur=f"{current_user.prenom} {current_user.nom}",
        origine_prescripteur='infirmier',
        statut='active',
        date_debut=heure_prevue.date(),
        date_prescription=datetime.utcnow(),
        notes="Ajouté par l'infirmier(-ère) — non prescrit par un médecin.",
    )
    db.session.add(prescription)
    db.session.flush()

    admin = AdministrationMedicament(
        prescription_id=prescription.id,
        patient_id=id_patient,
        medicament=medicament,
        dose=dosage,
        numero_dose=1,
        intervalle_heures=intervalle,
        heure_prevue=heure_prevue,
        statut='a_faire',
    )
    db.session.add(admin)
    db.session.commit()
    flash(f'"{medicament}" ajouté (non prescrit) et planifié', 'success')
    return redirect(url_for('infirmier_medicaments'))


@app.route('/api/infirmier/medicaments/dues')
@login_required
def api_infirmier_medicaments_dues():
    """Interrogé régulièrement en JS depuis l'écran médicaments (et la
    sidebar) pour le badge/l'alerte — pas d'infrastructure de notification
    temps réel existante dans l'appli (voir exploration), donc scrutation
    simple côté client."""
    from models import AdministrationMedicament, Patient

    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        return jsonify({'count': 0, 'items': []})

    dues = (
        AdministrationMedicament.query
        .join(Patient, AdministrationMedicament.patient_id == Patient.id)
        .filter(
            Patient.id_structure == current_user.id_structure,
            AdministrationMedicament.statut == 'a_faire',
            AdministrationMedicament.heure_prevue <= datetime.utcnow(),
        )
        .all()
    )
    return jsonify({
        'count': len(dues),
        'items': [
            {
                'id': d.id,
                'medicament': d.medicament,
                'dose': d.dose,
                'patient': f"{d.patient.prenom} {d.patient.nom}" if d.patient else '',
                'heure_prevue': d.heure_prevue.strftime('%H:%M'),
            }
            for d in dues
        ],
    })


@app.route('/medicaments/historique')
@login_required
def medicaments_historique():
    """Historique complet des administrations — accessible en lecture à
    infirmier/médecin/admin_structure, filtrable par patient et par statut
    (patron : "on doit revoir l'historique et tout")."""
    from models import AdministrationMedicament, Patient

    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    statut = request.args.get('statut', '')

    query = (
        AdministrationMedicament.query
        .join(Patient, AdministrationMedicament.patient_id == Patient.id)
        .filter(Patient.id_structure == current_user.id_structure)
    )
    if statut in ('a_faire', 'fait', 'annule'):
        query = query.filter(AdministrationMedicament.statut == statut)

    # ⭐ La recherche patient est filtrée en direct côté client sur ces lignes
    # déjà chargées (voir historique.html) : elle ne porte donc jamais que
    # sur des patients ayant réellement une administration dans cette liste,
    # sans requête serveur supplémentaire ni bouton à cliquer (patron :
    # "xa devrait juste concerner aux patients qui sont dans l'historique
    # des administrations... et la recherche doit etre en temps réel").
    administrations = query.order_by(AdministrationMedicament.heure_prevue.desc()).limit(300).all()

    return render_template('medicaments/historique.html', administrations=administrations,
                            statut_filtre=statut)


@app.route('/api/actes-types/rechercher')
@login_required
def api_actes_types_rechercher():
    """Recherche dans le catalogue LOCAL déjà créé pour cette structure
    (complète la recherche live sur GHP côté client — utile pour retrouver
    tout de suite un acte créé manuellement la veille, avant qu'il ait pu
    être resynchronisé)."""
    from models import ActeType

    terme = request.args.get('q', '').strip()
    if len(terme) < 2:
        return jsonify([])

    actes = ActeType.query.filter(
        ActeType.structure_id == current_user.id_structure,
        ActeType.actif == True,
        ActeType.nom.ilike(f'%{terme}%')
    ).order_by(ActeType.nom).limit(20).all()

    return jsonify([{'id': a.id, 'nom': a.nom} for a in actes])


# ==================== RECHERCHE ====================

@app.route('/recherche', methods=['GET', 'POST'])
@login_required
def recherche_patients():
    from models import Patient
    
    patients = []
    search_term = ''
    
    if request.method == 'POST':
        search_term = request.form.get('search_term', '').strip()
        
        if search_term:
            query = Patient.query.filter(
                Patient.id_structure == current_user.id_structure,
                Patient.archived == False
            )
            
            # Filtrer selon le rôle médecin
            if current_user.role == 'medecin':
                query = query.filter(Patient.id_medecin_referent == current_user.id)
            
            # Recherche multi-champs
            patients = query.filter(
                _patient_recherche_conditions(search_term)
            ).limit(50).all()

    return render_template('recherche.html', patients=patients, search_term=search_term)

# ==================== MODIFIER PATIENT ====================

@app.route('/patient/<int:id>/modifier', methods=['GET', 'POST'])
@login_required
@has_permission('PATIENTS')
def patient_modifier(id):
    from models import Patient, Utilisateur
    from datetime import datetime
    
    patient = Patient.query.get_or_404(id)

    # ⭐ FIX : même bug que consultation_ajouter_avec_patient — comparaison
    # stricte contre id_medecin_referent=None (patient sans référent, ex.
    # tout juste synchronisé depuis GHP) bloquait à tort TOUT médecin.
    if current_user.role == 'medecin' and patient.id_medecin_referent is not None and patient.id_medecin_referent != current_user.id:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('patients_list'))

    # ⭐ Retour vers une autre page que le dossier complet (ex. la
    # pré-consultation infirmier) après enregistrement — voir le bouton
    # "Compléter les informations du patient" dans
    # infirmier/pre_consultation.html. Sans ce paramètre, comportement
    # inchangé (retour au dossier patient).
    retour = request.values.get('retour') or None
    # ⭐ N'accepte qu'un chemin interne relatif (jamais une URL externe /
    # protocole-relative type "//autre-site") — évite une redirection
    # ouverte via ce paramètre.
    if retour and (not retour.startswith('/') or retour.startswith('//')):
        retour = None

    # Récupérer les médecins
    medecins = Utilisateur.query.filter_by(
        id_structure=current_user.id_structure if current_user.id_structure else patient.id_structure,
        role='medecin',
        actif=True
    ).all()
    
    if request.method == 'POST':
        # ═══════════════════════════════════════════
        # IDENTITÉ
        # ═══════════════════════════════════════════
        patient.nom = request.form.get('nom')
        patient.prenom = request.form.get('prenom')
        
        date_naissance = request.form.get('date_naissance')
        age = request.form.get('age')
        patient.date_naissance = _date_naissance_depuis_formulaire(date_naissance, age)
        
        patient.lieu_naissance = request.form.get('lieu_naissance')
        patient.sexe = request.form.get('sexe')
        patient.telephone = request.form.get('telephone')
        patient.email = request.form.get('email')
        patient.profession = request.form.get('profession')
        patient.code_postal = request.form.get('code_postal')
        patient.ville = request.form.get('ville')
        patient.adresse = request.form.get('adresse')
        
        # ═══════════════════════════════════════════
        # ASSURANCE PRINCIPALE
        # ═══════════════════════════════════════════
        patient.type_assurance = request.form.get('type_assurance')
        
        autre_assurance_nom = request.form.get('autre_assurance_nom')
        if patient.type_assurance == 'AUTRE_ASSURANCE':
            patient.autre_assurance_nom = autre_assurance_nom
        else:
            patient.autre_assurance_nom = None
        
        patient.num_assure = request.form.get('num_assure')
        
        # ═══════════════════════════════════════════
        # ASSURANCE 2
        # ═══════════════════════════════════════════
        patient.assurance2_nom = request.form.get('assurance2_nom')
        
        taux_assurance2 = request.form.get('taux_assurance2')
        patient.taux_assurance2 = float(taux_assurance2) if taux_assurance2 else None
        
        patient.numero_assure2 = request.form.get('numero_assure2')
        
        # ═══════════════════════════════════════════
        # TAUX DE PRISE EN CHARGE
        # ═══════════════════════════════════════════
        taux_prise_charge = request.form.get('taux_prise_charge')
        patient.taux_prise_charge = taux_prise_charge if taux_prise_charge else None
        
        # ═══════════════════════════════════════════
        # PERSONNE À PRÉVENIR
        # ═══════════════════════════════════════════
        patient.personne_a_prevenir_nom = request.form.get('personne_a_prevenir_nom')
        patient.personne_a_prevenir_telephone = request.form.get('personne_a_prevenir_telephone')
        patient.personne_a_prevenir_relation = request.form.get('personne_a_prevenir_relation')
        
        # ═══════════════════════════════════════════
        # INFORMATIONS MÉDICALES
        # ═══════════════════════════════════════════
        patient.groupe_sanguin = request.form.get('groupe_sanguin')
        patient.mutuelle = request.form.get('mutuelle')
        patient.medecin_traitant = request.form.get('medecin_traitant')
        
        # ═══════════════════════════════════════════
        # HABITUDES DE VIE
        # ═══════════════════════════════════════════
        patient.tabac = request.form.get('tabac')
        patient.alcool = request.form.get('alcool')
        patient.allaitement = request.form.get('allaitement') == 'on'
        patient.grossesse = request.form.get('grossesse') == 'on'
        
        # ═══════════════════════════════════════════
        # MÉDECIN RÉFÉRENT
        # ═══════════════════════════════════════════
        id_medecin_referent = request.form.get('id_medecin_referent')
        patient.id_medecin_referent = int(id_medecin_referent) if id_medecin_referent else None
        
        # ═══════════════════════════════════════════
        # NOTES
        # ═══════════════════════════════════════════
        patient.notes = request.form.get('notes')
        
        # ═══════════════════════════════════════════
        # ALLERGIES ET ANTÉCÉDENTS (si présents dans le formulaire)
        # ═══════════════════════════════════════════
        # Si tu as ces champs dans le formulaire, décommente :
        # patient.allergies = request.form.get('allergies')
        # patient.antecedents_medicaux = request.form.get('antecedents')
        
        db.session.commit()
        flash('✅ Patient modifié avec succès', 'success')
        return redirect(retour or url_for('patient_detail', id=patient.id))

    return render_template('patients/modifier.html', patient=patient, medecins=medecins, retour=retour)

# ==================== CONSULTATION AVEC PATIENT SPECIFIQUE ====================

@app.route('/patient/<int:id>/consultation/ajouter', methods=['GET', 'POST'])
@login_required
def consultation_ajouter_avec_patient(id):
    from models import Patient, Consultation, Prescription, AnalyseDemande
    from datetime import datetime
    import json
    
    patient = Patient.query.get_or_404(id)

    # ⭐ FIX E2E : un patient tout juste synchronisé depuis GHP (ou sans
    # consultation antérieure) n'a pas encore de médecin référent
    # (id_medecin_referent=None, voir sync_patients_from_ghp) — la
    # comparaison stricte bloquait ici TOUT médecin voulant faire la toute
    # première consultation (même en suivant le lien "Nouvelle
    # consultation" depuis la fiche patient, qui reste accessible tant
    # qu'aucun référent n'est encore assigné, voir patient_detail()
    # ci-dessus). On n'interdit donc que l'accès à un patient déjà
    # référencé par un AUTRE médecin.
    if current_user.role == 'medecin' and patient.id_medecin_referent is not None and patient.id_medecin_referent != current_user.id:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('patients_list'))
    
    if request.method == 'POST':
        # ═══════════════════════════════════════════
        # 1. RÉCUPÉRATION DES DONNÉES DU FORMULAIRE
        # ═══════════════════════════════════════════

        motif = request.form.get('motif')
        diagnostic = request.form.get('diagnostic')
        orientation = request.form.get('orientation') or None
        if orientation not in ('ambulatoire', 'hospitalisation'):
            orientation = None

        # RÉCUPÉRATION DES CHAMPS HPI
        hpi_date_debut = request.form.get('hpi_date_debut')
        hpi_debut_type = request.form.get('hpi_debut_type')
        hpi_circonstances = request.form.get('hpi_circonstances')
        hpi_evolution = request.form.get('hpi_evolution')
        hpi_facteurs = request.form.get('hpi_facteurs')
        hpi_traitements = request.form.get('hpi_traitements')
        hpi_signes = request.form.get('hpi_signes')
        
        # ⭐ RÉCUPÉRER LES CHECKBOXES (MAIS PAS ENCORE LES APPLIQUER)
        hpi_fievre = request.form.get('hpi_fievre') == 'on'
        hpi_nausees = request.form.get('hpi_nausees') == 'on'
        hpi_douleur = request.form.get('hpi_douleur') == 'on'
        hpi_cephalées = request.form.get('hpi_cephalées') == 'on'
        hpi_vertiges = request.form.get('hpi_vertiges') == 'on'
        hpi_dyspnee = request.form.get('hpi_dyspnee') == 'on'
        
        hpi_trauma_mecanisme = request.form.get('hpi_trauma_mecanisme')
        hpi_trauma_heure = request.form.get('hpi_trauma_heure')
        hpi_trauma_pc = request.form.get('hpi_trauma_pc')
        hpi_trauma_description = request.form.get('hpi_trauma_description')
        
        hpi_morsure_type = request.form.get('hpi_morsure_type')
        hpi_morsure_espece = request.form.get('hpi_morsure_espece')
        hpi_morsure_siege = request.form.get('hpi_morsure_siege')
        hpi_morsure_signes = request.form.get('hpi_morsure_signes')
        
        hpi_intox_substance = request.form.get('hpi_intox_substance')
        hpi_intox_heure = request.form.get('hpi_intox_heure')
        hpi_intox_circonstances = request.form.get('hpi_intox_circonstances')
        
        hpi_autres_signes = request.form.get('hpi_autres_signes')
        hpi_autre_infos = request.form.get('hpi_autre_infos')
        
        hpi_complements_medecin = request.form.get('hpi_complements_medecin')
        histoire_maladie = request.form.get('histoire_maladie')

        # Constantes
        tension = request.form.get('tension')
        temperature = request.form.get('temperature')
        pouls = request.form.get('pouls')
        saturation = request.form.get('saturation')
        poids = request.form.get('poids')
        taille = request.form.get('taille')
        imc = request.form.get('imc')
        
        # Examens
        examens_cliniques = request.form.get('examens_cliniques')
        examens_biologie = request.form.get('examens_biologie')
        examens_imagerie = request.form.get('examens_imagerie')
        
        traitement = request.form.get('traitement')
        notes = request.form.get('notes')
        cim10 = request.form.get('cim10')
        
        arret_travail = request.form.get('arret_travail') == 'on'
        arret_jours = request.form.get('arret_jours')
        
        prochain_rdv = request.form.get('prochain_rdv')
        statut_medical = request.form.get('statut_medical')
        
        allergies = request.form.get('allergies')
        traitements_en_cours = request.form.get('traitements_en_cours')
        antecedents_medicaux = request.form.get('antecedents_medicaux')
        antecedents_chirurgicaux = request.form.get('antecedents_chirurgicaux')
        
        medicaments_prescrits = request.form.get('medicaments_prescrits')
        
        # ═══════════════════════════════════════════
        # ⭐ 2. RÉCUPÉRATION DE LA CONSULTATION TEMPORAIRE
        # ═══════════════════════════════════════════
        
        consultation = Consultation.query.filter_by(
            id_patient=patient.id,
            is_temporary=True
        ).first()
        
        if consultation:
            # ⭐ METTRE À JOUR LA CONSULTATION TEMPORAIRE
            print(f"✅ Consultation temporaire #{consultation.id} trouvée !")
            
            consultation.motif = motif
            consultation.diagnostic = diagnostic
            consultation.tension_arterielle = tension
            consultation.temperature_c = float(temperature) if temperature else None
            consultation.pulse_bpm = int(pouls) if pouls else None
            consultation.oxygene_saturation = int(saturation) if saturation else None
            consultation.poids_kg = float(poids) if poids else None
            consultation.taille_cm = float(taille) if taille else None
            consultation.imc = float(imc) if imc else None
            consultation.examens_cliniques = examens_cliniques
            consultation.examens_biologie = examens_biologie
            consultation.examens_imagerie = examens_imagerie
            consultation.traitement_prescrit = traitement
            consultation.notes_cliniques = notes
            consultation.cim10 = cim10
            consultation.arret_travail = arret_travail
            consultation.arret_jours = int(arret_jours) if arret_jours else None
            consultation.prochain_rdv = datetime.strptime(prochain_rdv, '%Y-%m-%d') if prochain_rdv else None
            consultation.allergies = allergies
            consultation.traitements_en_cours = traitements_en_cours
            consultation.antecedents_medicaux = antecedents_medicaux
            consultation.antecedents_chirurgicaux = antecedents_chirurgicaux
            
            # HPI - Champs texte
            consultation.hpi_date_debut = datetime.strptime(hpi_date_debut, '%Y-%m-%d').date() if hpi_date_debut else None
            consultation.hpi_debut_type = hpi_debut_type
            consultation.hpi_circonstances = hpi_circonstances
            consultation.hpi_evolution = hpi_evolution
            consultation.hpi_facteurs = hpi_facteurs
            consultation.hpi_traitements = hpi_traitements
            consultation.hpi_signes = hpi_signes
            consultation.hpi_trauma_mecanisme = hpi_trauma_mecanisme
            consultation.hpi_trauma_heure = datetime.fromisoformat(hpi_trauma_heure) if hpi_trauma_heure else None
            consultation.hpi_trauma_pc = hpi_trauma_pc
            consultation.hpi_trauma_description = hpi_trauma_description
            consultation.hpi_morsure_type = hpi_morsure_type
            consultation.hpi_morsure_espece = hpi_morsure_espece
            consultation.hpi_morsure_siege = hpi_morsure_siege
            consultation.hpi_morsure_signes = hpi_morsure_signes
            consultation.hpi_intox_substance = hpi_intox_substance
            consultation.hpi_intox_heure = datetime.fromisoformat(hpi_intox_heure) if hpi_intox_heure else None
            consultation.hpi_intox_circonstances = hpi_intox_circonstances
            consultation.hpi_autres_signes = hpi_autres_signes
            consultation.hpi_autre_infos = hpi_autre_infos
            consultation.hpi_complements_medecin = hpi_complements_medecin
            consultation.histoire_maladie = histoire_maladie
            
            # ⭐⭐⭐ AJOUTER LES SIGNES ASSOCIÉS ICI ⭐⭐⭐
            consultation.hpi_fievre = hpi_fievre
            consultation.hpi_nausees = hpi_nausees
            consultation.hpi_douleur = hpi_douleur
            consultation.hpi_cephalées = hpi_cephalées
            consultation.hpi_vertiges = hpi_vertiges
            consultation.hpi_dyspnee = hpi_dyspnee
            
            # ⭐ MARQUER COMME DÉFINITIVE
            consultation.is_temporary = False
            consultation.statut = 'terminee'
            consultation.orientation = orientation
            consultation.id_medecin = current_user.id if current_user.role == 'medecin' else None
            consultation.date_consultation = datetime.utcnow()

        else:
            # ⭐ PAS DE CONSULTATION TEMPORAIRE : EN CRÉER UNE
            print(f"⚠️ Aucune consultation temporaire trouvée pour patient #{patient.id}")

            consultation = Consultation(
                id_patient=patient.id,
                id_medecin=current_user.id if current_user.role == 'medecin' else None,
                motif=motif,
                diagnostic=diagnostic,
                orientation=orientation,
                tension_arterielle=tension,
                temperature_c=float(temperature) if temperature else None,
                pulse_bpm=int(pouls) if pouls else None,
                oxygene_saturation=int(saturation) if saturation else None,
                poids_kg=float(poids) if poids else None,
                taille_cm=float(taille) if taille else None,
                imc=float(imc) if imc else None,
                examens_cliniques=examens_cliniques,
                examens_biologie=examens_biologie,
                examens_imagerie=examens_imagerie,
                traitement_prescrit=traitement,
                notes_cliniques=notes,
                cim10=cim10,
                arret_travail=arret_travail,
                arret_jours=int(arret_jours) if arret_jours else None,
                prochain_rdv=datetime.strptime(prochain_rdv, '%Y-%m-%d') if prochain_rdv else None,
                allergies=allergies,
                traitements_en_cours=traitements_en_cours,
                antecedents_medicaux=antecedents_medicaux,
                antecedents_chirurgicaux=antecedents_chirurgicaux,
                hpi_date_debut=datetime.strptime(hpi_date_debut, '%Y-%m-%d').date() if hpi_date_debut else None,
                hpi_debut_type=hpi_debut_type,
                hpi_circonstances=hpi_circonstances,
                hpi_evolution=hpi_evolution,
                hpi_facteurs=hpi_facteurs,
                hpi_traitements=hpi_traitements,
                hpi_signes=hpi_signes,
                hpi_trauma_mecanisme=hpi_trauma_mecanisme,
                hpi_trauma_heure=datetime.fromisoformat(hpi_trauma_heure) if hpi_trauma_heure else None,
                hpi_trauma_pc=hpi_trauma_pc,
                hpi_trauma_description=hpi_trauma_description,
                hpi_morsure_type=hpi_morsure_type,
                hpi_morsure_espece=hpi_morsure_espece,
                hpi_morsure_siege=hpi_morsure_siege,
                hpi_morsure_signes=hpi_morsure_signes,
                hpi_intox_substance=hpi_intox_substance,
                hpi_intox_heure=datetime.fromisoformat(hpi_intox_heure) if hpi_intox_heure else None,
                hpi_intox_circonstances=hpi_intox_circonstances,
                hpi_autres_signes=hpi_autres_signes,
                hpi_autre_infos=hpi_autre_infos,
                hpi_complements_medecin=hpi_complements_medecin,
                histoire_maladie=histoire_maladie,
                date_consultation=datetime.utcnow(),
                is_temporary=False,
                # ⭐⭐⭐ AJOUTER LES SIGNES ASSOCIÉS ICI ⭐⭐⭐
                hpi_fievre=hpi_fievre,
                hpi_nausees=hpi_nausees,
                hpi_douleur=hpi_douleur,
                hpi_cephalées=hpi_cephalées,
                hpi_vertiges=hpi_vertiges,
                hpi_dyspnee=hpi_dyspnee
            )
            db.session.add(consultation)
            db.session.flush()
            print(f"✅ Nouvelle consultation #{consultation.id} créée")
        
        # ═══════════════════════════════════════════
        # 3. CRÉATION DES ANALYSES
        # ═══════════════════════════════════════════
        
        if examens_biologie:
            for ligne in examens_biologie.split('\n'):
                nom = ligne.strip()
                if nom:
                    analyse = AnalyseDemande(
                        consultation_id=consultation.id,
                        patient_id=consultation.id_patient,
                        structure_id=current_user.id_structure,
                        type_analyse='BIOLOGIE',
                        nom_analyse=nom,
                        prescrit_par=current_user.id,
                        statut='EN_ATTENTE'
                    )
                    db.session.add(analyse)
        
        if examens_imagerie:
            for ligne in examens_imagerie.split('\n'):
                nom = ligne.strip()
                if nom:
                    analyse = AnalyseDemande(
                        consultation_id=consultation.id,
                        patient_id=consultation.id_patient,
                        structure_id=current_user.id_structure,
                        type_analyse='IMAGERIE',
                        nom_analyse=nom,
                        prescrit_par=current_user.id,
                        statut='EN_ATTENTE'
                    )
                    db.session.add(analyse)
        
        # ═══════════════════════════════════════════
        # 4. CRÉATION DES PRESCRIPTIONS
        # ═══════════════════════════════════════════
        
        prescriptions_creees = 0
        
        if medicaments_prescrits:
            try:
                meds_data = json.loads(medicaments_prescrits)
                print(f"📦 Données reçues: {medicaments_prescrits}")

                for med in meds_data:
                    print(f"   Médicament: {med.get('nom')}, Durée: {med.get('duree')}")
                    prescription = Prescription(
                        id_patient=patient.id,
                        id_consultation=consultation.id,
                        id_medecin=current_user.id,
                        medicament=med.get('nom', ''),
                        dosage=med.get('dosage', ''),
                        forme=med.get('forme', ''),
                        quantite=str(med.get('quantite', 1)),
                        duree_jours=med.get('duree', 7),
                        frequence=med.get('posologie', ''),
                        instructions=med.get('instructions', ''),
                        renouvelable=med.get('renouvelable', False),
                        type_prescription='medicament',
                        prescripteur=f"{current_user.prenom} {current_user.nom}",
                        statut='active',
                        date_prescription=datetime.utcnow(),
                        notes=med.get('notes', ''),
                        source_id=med.get('source_id'),
                        stock_disponible=med.get('stock')
                    )
                    db.session.add(prescription)
                    prescriptions_creees += 1
                    
                print(f"✅ {len(meds_data)} prescription(s) médicamenteuse(s) enregistrée(s)")
                
            except Exception as e:
                print(f"❌ Erreur sauvegarde prescriptions médicaments: {e}")
                import traceback
                traceback.print_exc()
        
        if examens_biologie:
            for ligne in examens_biologie.split('\n'):
                nom = ligne.strip()
                if nom:
                    prescription = Prescription(
                        id_patient=patient.id,
                        id_consultation=consultation.id,
                        id_medecin=current_user.id,
                        medicament=nom,
                        type_prescription='acte',
                        prescripteur=f"{current_user.prenom} {current_user.nom}",
                        statut='active',
                        date_prescription=datetime.utcnow()
                    )
                    db.session.add(prescription)
                    prescriptions_creees += 1
                    print(f"📋 Prescription d'acte (biologie) ajoutée: {nom}")
        
        if examens_imagerie:
            for ligne in examens_imagerie.split('\n'):
                nom = ligne.strip()
                if nom:
                    prescription = Prescription(
                        id_patient=patient.id,
                        id_consultation=consultation.id,
                        id_medecin=current_user.id,
                        medicament=nom,
                        type_prescription='acte',
                        prescripteur=f"{current_user.prenom} {current_user.nom}",
                        statut='active',
                        date_prescription=datetime.utcnow()
                    )
                    db.session.add(prescription)
                    prescriptions_creees += 1
                    print(f"📋 Prescription d'acte (imagerie) ajoutée: {nom}")
        
        # ═══════════════════════════════════════════
        # 5. MISE À JOUR DU PATIENT
        # ═══════════════════════════════════════════
        
        if temperature:
            patient.temperature_c = float(temperature)
        if tension:
            patient.tension_arterielle = tension
        if pouls:
            patient.pulse_bpm = int(pouls)
        if saturation:
            patient.oxygene_saturation = int(saturation)
        if poids:
            patient.poids_kg = float(poids)
        if taille:
            patient.taille_cm = float(taille)
        if imc:
            patient.imc = float(imc)
        
        patient.date_derniere_consultation = datetime.utcnow()
        patient.id_medecin_referent = current_user.id  # ⭐ LE MÉDECIN QUI CONSULTE DEVIENT RÉFÉRENT
        
        if statut_medical:
            patient.statut_medical = statut_medical
            if statut_medical == 'GUERI':
                patient.date_guerison = datetime.utcnow()
        elif patient.statut_medical == 'PREMIERE_VISITE':
            patient.statut_medical = 'EN_TRAITEMENT'
        
        # ═══════════════════════════════════════════
        # 6. COMMIT FINAL
        # ═══════════════════════════════════════════
        
        db.session.commit()
        
        if prescriptions_creees > 0:
            try:
                from tasks import sync_prescriptions_to_ghp
                result = sync_prescriptions_to_ghp()
                if result.get('success'):
                    print(f"✅ {result.get('message')}")
                else:
                    print(f"⚠️ {result.get('message')}")
            except Exception as e:
                print(f"⚠️ Erreur sync auto: {e}")

        _pousser_rendez_vous_ghp(consultation, patient, f"{current_user.prenom} {current_user.nom}")

        flash(f'Consultation pour {patient.prenom} {patient.nom} enregistrée avec succès', 'success')
        if orientation == 'hospitalisation':
            flash('🛏️ Patient orienté vers une hospitalisation — visible dans "Hospitalisations en attente".', 'info')
        return redirect(url_for('patient_detail', id=patient.id))
    
    empty_consultation = Consultation()

    return render_template('consultations/ajouter_avec_patient.html', patient=patient, consultation=empty_consultation)

# ==================== STATISTIQUES ====================

# ⭐ Un diagnostic saisi via le sélecteur CIM-10 (templates/consultations/
# ajouter.html) est une ligne "CODE Libellé" par pathologie retenue (ex.
# "B50 Paludisme à Plasmodium falciparum"), plusieurs lignes possibles par
# consultation — jusqu'ici, "Pathologies les plus fréquentes" regroupait le
# texte ENTIER du diagnostic (tel-quel, multi-ligne) : deux consultations
# avec les 2 mêmes pathologies dans un ordre différent, ou une de plus/moins,
# ne se regroupaient jamais ensemble. On regroupe maintenant PATHOLOGIE PAR
# PATHOLOGIE (une ligne = une pathologie), par code CIM-10 quand reconnu
# (fiable même si le libellé varie légèrement), repli sur le texte brut de
# la ligne sinon (saisie manuelle sans passer par le sélecteur).
_RE_CIM10_LIGNE = re.compile(r'^([A-Z]\d{2}(?:\.\d+)?)\s+(.+)$')


def _extraire_pathologies(diagnostic_text):
    """Découpe un texte de diagnostic (une ligne = une pathologie) en liste
    de (cle, libelle) — cle = code CIM-10 si reconnu, sinon le texte de la
    ligne lui-même (normalisé) pour un regroupement au moins cohérent."""
    if not diagnostic_text:
        return []
    resultats = []
    for ligne in diagnostic_text.split('\n'):
        ligne = ligne.strip()
        if not ligne or ligne in ('-', '—'):
            continue
        m = _RE_CIM10_LIGNE.match(ligne)
        if m:
            resultats.append((m.group(1), m.group(2).strip()))
        else:
            resultats.append((ligne.lower(), ligne))
    return resultats


# ⭐ Grille de lecture épidémiologique — GÉNÉRIQUE PAR ZONE CLIMATIQUE, pas
# figée sur Lomé : l'appli est utilisée par des structures dans plusieurs
# pays, avec des saisons différentes (l'Afrique australe est même dans
# l'hémisphère sud — pluies en décembre, pas en juillet). Chaque catégorie
# ne connaît que son TYPE de période ('pluies' ou 'seche'), résolu en mois
# concrets pour la zone de LA STRUCTURE via _resoudre_zone_climatique
# ci-dessous (déterminée à partir de son pays/ville, renseignés dans
# /structure). Si la zone de la structure est inconnue, aucune explication
# climatique n'est inventée — seul le constat statistique brut est gardé
# (voir _construire_analyse_pathologies).
_CATEGORIES_EPIDEMIO = [
    {
        'icone': '🦟', 'nom': 'paludisme',
        'codes': ('B50', 'B51', 'B52', 'B53', 'B54'),
        'mots_cles': ('paludisme', 'palu', 'malaria'),
        'periode_type': 'pluies',
        'explication': "la prolifération des moustiques anophèles vecteurs après les pluies",
    },
    {
        'icone': '💧', 'nom': 'maladies hydriques',
        'codes': ('A00', 'A01', 'A02', 'A03', 'A04', 'A05', 'A06', 'A07', 'A08', 'A09'),
        'mots_cles': ('choléra', 'cholera', 'diarrh', 'typho', 'gastro-entérite', 'gastro entérite', 'intoxication alimentaire'),
        'periode_type': 'pluies',
        'explication': "la contamination des points d'eau et les risques d'inondation",
    },
    {
        'icone': '🌬️', 'nom': 'infections respiratoires',
        'codes': ('J',),
        'mots_cles': ('ira', 'grippe', 'rhume', 'bronchite', 'pneumonie', 'rhinopharyngite', 'toux', 'respiratoire'),
        'periode_type': 'seche',
        'explication': "l'air sec (souvent chargé de poussière en saison sèche), qui fragilise les voies respiratoires",
    },
    {
        'icone': '🌡️', 'nom': 'méningite',
        'codes': ('A39', 'G00', 'G01', 'G02', 'G03'),
        'mots_cles': ('méningite', 'meningite'),
        'periode_type': 'seche',
        'explication': "la saison sèche, propice à la transmission du méningocoque",
    },
    {
        'icone': '👁️', 'nom': 'conjonctivite',
        'codes': ('H10', 'H11'),
        'mots_cles': ('conjonctiv',),
        'periode_type': 'seche',
        'explication': "la poussière de la saison sèche, irritante pour les yeux",
    },
    {
        'icone': '🔴', 'nom': 'rougeole',
        'codes': ('B05',),
        'mots_cles': ('rougeole',),
        'periode_type': 'seche',
        'explication': "la saison sèche, période de circulation accrue du virus",
    },
    {
        'icone': '❤️', 'nom': 'maladie chronique non transmissible', 'chronique': True,
        'codes': ('I10', 'I11', 'I12', 'I13', 'I14', 'I15', 'I20', 'I21', 'I22', 'I23', 'I24', 'I25', 'E10', 'E11', 'E12', 'E13', 'E14'),
        'mots_cles': ('hypertension', 'diabète', 'diabete', 'cardiopathie', 'cardiovasculaire'),
        'periode_type': None, 'explication': None,
    },
]


def _categorie_epidemiologique(code, label):
    """Retourne la catégorie épidémiologique connue correspondant à ce code
    CIM-10/libellé, ou None si la pathologie n'est pas dans la grille de
    lecture (dans ce cas, on n'invente aucune explication — seul le constat
    statistique brut est gardé)."""
    code = (code or '').upper()
    texte = (label or '').lower()
    for cat in _CATEGORIES_EPIDEMIO:
        if any(code.startswith(c) for c in cat['codes']):
            return cat
        if any(m in texte for m in cat['mots_cles']):
            return cat
    return None


# ⭐ Zones climatiques d'Afrique subsaharienne, simplifiées pour l'usage
# statistique ci-dessus (pas une classification climatologique complète —
# juste de quoi situer "saison des pluies" / "saison sèche" par grande
# région). L'Afrique australe est en hémisphère sud : ses pluies tombent en
# été austral (nov-mars), pas en même temps que l'Afrique de l'Ouest.
_ZONES_CLIMATIQUES = {
    'sahel': {
        'label': "zone sahélienne",
        'pluies': {'Juin', 'Juil', 'Août', 'Sep'},
        'seche': {'Nov', 'Déc', 'Jan', 'Fév', 'Mar'},
    },
    'soudanien': {
        'label': "zone soudanienne",
        'pluies': {'Mai', 'Juin', 'Juil', 'Août', 'Sep', 'Oct'},
        'seche': {'Nov', 'Déc', 'Jan', 'Fév', 'Mar', 'Avr'},
    },
    'guineen': {
        'label': "zone guinéenne/côtière",
        'pluies': {'Mar', 'Avr', 'Mai', 'Juin', 'Juil', 'Sep', 'Oct'},
        'seche': {'Nov', 'Déc', 'Jan', 'Fév'},
    },
    'equatorial': {
        'label': "zone équatoriale",
        'pluies': {'Mar', 'Avr', 'Mai', 'Sep', 'Oct', 'Nov'},
        'seche': {'Juin', 'Juil', 'Août', 'Déc', 'Jan', 'Fév'},
    },
    'est_africain': {
        'label': "zone d'Afrique de l'Est (pluies bimodales)",
        'pluies': {'Mar', 'Avr', 'Mai', 'Oct', 'Nov'},
        'seche': {'Juin', 'Juil', 'Août', 'Sep', 'Déc', 'Jan', 'Fév'},
    },
    'australe': {
        'label': "zone d'Afrique australe (hémisphère sud)",
        'pluies': {'Nov', 'Déc', 'Jan', 'Fév', 'Mar'},
        'seche': {'Avr', 'Mai', 'Juin', 'Juil', 'Août', 'Sep', 'Oct'},
    },
}

# Villes reconnues avec leur zone précise (prioritaire sur le pays — utile
# pour les pays qui chevauchent plusieurs zones, ex. Nigeria, Cameroun).
_ZONE_PAR_VILLE = {
    'lome': 'guineen', 'lomé': 'guineen', 'cotonou': 'guineen', 'lagos': 'guineen',
    'abuja': 'soudanien', 'kano': 'sahel', 'accra': 'guineen', 'abidjan': 'guineen',
    'conakry': 'guineen', 'freetown': 'guineen', 'monrovia': 'guineen', 'banjul': 'guineen',
    'dakar': 'sahel', 'bamako': 'sahel', 'niamey': 'sahel', "n'djamena": 'sahel',
    'nouakchott': 'sahel', 'ouagadougou': 'soudanien', 'bissau': 'soudanien',
    'douala': 'equatorial', 'yaounde': 'equatorial', 'yaoundé': 'equatorial',
    'kinshasa': 'equatorial', 'brazzaville': 'equatorial', 'libreville': 'equatorial',
    'bangui': 'equatorial', 'malabo': 'equatorial',
    'nairobi': 'est_africain', 'dar es salaam': 'est_africain', 'kampala': 'est_africain',
    'kigali': 'est_africain', 'bujumbura': 'est_africain', 'addis-abeba': 'est_africain',
    'addis abeba': 'est_africain', 'mogadiscio': 'est_africain',
    'johannesburg': 'australe', 'pretoria': 'australe', 'le cap': 'australe',
    'harare': 'australe', 'lusaka': 'australe', 'maputo': 'australe', 'luanda': 'australe',
    'antananarivo': 'australe', 'windhoek': 'australe', 'gaborone': 'australe',
    'lilongwe': 'australe', 'mbabane': 'australe', 'maseru': 'australe',
}

# Pays reconnus, zone par défaut (climat dominant/capitale) si la ville
# n'est pas renseignée ou pas reconnue ci-dessus.
_ZONE_PAR_PAYS = {
    'togo': 'guineen', 'bénin': 'guineen', 'benin': 'guineen', 'ghana': 'guineen',
    "côte d'ivoire": 'guineen', "cote d'ivoire": 'guineen', 'nigeria': 'guineen',
    'sierra leone': 'guineen', 'liberia': 'guineen', 'guinée': 'guineen', 'guinee': 'guineen',
    'gambie': 'guineen',
    'mali': 'sahel', 'niger': 'sahel', 'mauritanie': 'sahel', 'tchad': 'sahel',
    'sénégal': 'sahel', 'senegal': 'sahel', 'soudan': 'sahel',
    'burkina faso': 'soudanien', 'guinée-bissau': 'soudanien', 'guinee-bissau': 'soudanien',
    'cameroun': 'equatorial', 'gabon': 'equatorial', 'congo': 'equatorial', 'rdc': 'equatorial',
    'république démocratique du congo': 'equatorial', 'republique democratique du congo': 'equatorial',
    'centrafrique': 'equatorial', 'république centrafricaine': 'equatorial',
    'guinée équatoriale': 'equatorial',
    'kenya': 'est_africain', 'tanzanie': 'est_africain', 'ouganda': 'est_africain',
    'rwanda': 'est_africain', 'burundi': 'est_africain', 'éthiopie': 'est_africain',
    'ethiopie': 'est_africain', 'somalie': 'est_africain', 'djibouti': 'est_africain',
    'afrique du sud': 'australe', 'zimbabwe': 'australe', 'zambie': 'australe',
    'mozambique': 'australe', 'angola': 'australe', 'namibie': 'australe',
    'botswana': 'australe', 'madagascar': 'australe', 'malawi': 'australe',
    'lesotho': 'australe', 'eswatini': 'australe',
}


def _resoudre_zone_climatique(structure):
    """Détermine la zone climatique d'une structure à partir de sa ville
    (précis, prioritaire) ou son pays (repli), tels que renseignés dans son
    profil (voir /structure/localisation). Retourne None si rien n'est
    renseigné ou reconnu — dans ce cas, aucune observation climatique
    n'est fabriquée ailleurs, par choix (mieux vaut une structure qui
    configure son pays qu'une supposition fausse)."""
    if not structure:
        return None
    ville = (structure.ville or '').strip().lower()
    if ville in _ZONE_PAR_VILLE:
        return _ZONE_PAR_VILLE[ville]
    pays = (structure.pays or '').strip().lower()
    if pays in _ZONE_PAR_PAYS:
        return _ZONE_PAR_PAYS[pays]
    return None


# Liste affichée dans le sélecteur "Pays" de /structure (voir
# structure_localisation ci-dessous) — chaque libellé, en minuscules, doit
# être une clé de _ZONE_PAR_PAYS pour que la zone soit reconnue. Une
# structure hors de cette liste peut quand même saisir son pays via
# "Autre" : il sera enregistré tel quel, la zone restera simplement
# inconnue (pas d'observation climatique inventée pour elle).
_PAYS_AFRIQUE_LISTE = [
    'Togo', 'Bénin', 'Ghana', "Côte d'Ivoire", 'Nigeria', 'Sierra Leone', 'Liberia', 'Guinée', 'Gambie',
    'Mali', 'Niger', 'Mauritanie', 'Tchad', 'Sénégal', 'Soudan',
    'Burkina Faso', 'Guinée-Bissau',
    'Cameroun', 'Gabon', 'Congo', 'République démocratique du Congo', 'Centrafrique', 'Guinée équatoriale',
    'Kenya', 'Tanzanie', 'Ouganda', 'Rwanda', 'Burundi', 'Éthiopie', 'Somalie', 'Djibouti',
    'Afrique du Sud', 'Zimbabwe', 'Zambie', 'Mozambique', 'Angola', 'Namibie', 'Botswana',
    'Madagascar', 'Malawi', 'Lesotho', 'Eswatini',
]


def _construire_analyse_pathologies(consultations, zone_key=None, lieu_label=None):
    """Construit, à partir d'une liste d'objets Consultation (avec
    date_consultation et diagnostic déjà chargés), un classement des
    pathologies avec répartition mensuelle et un pic saisonnier détecté
    par pathologie — voir _extraire_pathologies ci-dessus pour le
    découpage. `consultations` doit être une vraie liste (pas une requête
    encore lazy), un seul passage suffit.

    `zone_key` (voir _ZONES_CLIMATIQUES/_resoudre_zone_climatique) et
    `lieu_label` (ville ou pays à citer dans le texte) pilotent
    l'interprétation contextuelle — si `zone_key` est None (structure sans
    pays/ville renseigné, ou zone non reconnue), aucune explication
    climatique n'est ajoutée, seul le constat statistique brut est gardé."""
    NOMS_MOIS = ['Jan', 'Fév', 'Mar', 'Avr', 'Mai', 'Juin', 'Juil', 'Août', 'Sep', 'Oct', 'Nov', 'Déc']
    zone = _ZONES_CLIMATIQUES.get(zone_key)
    pathologies = {}  # cle -> {'label', 'code', 'total', 'par_mois': [0]*12}

    for c in consultations:
        if not c.date_consultation:
            continue
        mois_idx = c.date_consultation.month - 1
        for code, label in _extraire_pathologies(c.diagnostic):
            cle = code or label
            if cle not in pathologies:
                pathologies[cle] = {'label': label, 'code': code, 'total': 0, 'par_mois': [0] * 12}
            pathologies[cle]['total'] += 1
            pathologies[cle]['par_mois'][mois_idx] += 1
            # ⭐ Garder le libellé le plus long vu pour ce code (plus
            # descriptif — la première rencontre n'est pas toujours la
            # plus complète, ex. "Paludisme" puis "Paludisme grave").
            if len(label) > len(pathologies[cle]['label']):
                pathologies[cle]['label'] = label

    total_mentions = sum(p['total'] for p in pathologies.values())

    resultat = []
    for cle, p in pathologies.items():
        moyenne_mensuelle = p['total'] / 12
        # ⭐ Pic saisonnier : mois dont le compte dépasse nettement la
        # moyenne mensuelle de CETTE pathologie (pas la moyenne globale) —
        # seuil × 1.5 et au moins 3 cas au total pour éviter de qualifier
        # de "pic" un simple bruit statistique sur 1-2 cas.
        pic_mois = []
        if p['total'] >= 3 and moyenne_mensuelle > 0:
            pic_mois = [NOMS_MOIS[i] for i, n in enumerate(p['par_mois']) if n >= moyenne_mensuelle * 1.5 and n >= 2]

        insight_parts = []
        if pic_mois:
            ratio = max(p['par_mois']) / moyenne_mensuelle if moyenne_mensuelle else 0
            insight_parts.append(f"Nettement plus fréquent en {', '.join(pic_mois)} (jusqu'à {ratio:.1f}× la moyenne mensuelle de cette pathologie).")

        # ⭐ Interprétation contextuelle (voir _CATEGORIES_EPIDEMIO) : compare
        # le pic détecté à la saisonnalité habituellement attendue dans LA
        # ZONE CLIMATIQUE DE LA STRUCTURE (zone=None → pays/ville non
        # renseigné ou non reconnu → pas d'explication climatique inventée,
        # constat statistique brut uniquement). Un pic HORS saison attendue
        # est signalé "atypique" — potentiellement l'observation la plus
        # utile pour la structure.
        atypique = False
        categorie_icone = None
        cat = _categorie_epidemiologique(p['code'], p['label'])
        if cat:
            categorie_icone = cat['icone']
            if cat.get('chronique'):
                if pic_mois:
                    insight_parts.append(f"{cat['icone']} Pic ponctuel malgré une pathologie habituellement non saisonnière — probablement lié au suivi médical (renouvellements d'ordonnance, rendez-vous programmés) plutôt qu'à un facteur épidémiologique.")
                elif p['total'] >= 3:
                    insight_parts.append(f"{cat['icone']} Maladie chronique non transmissible : répartition stable attendue, cohérente avec un suivi régulier plutôt que saisonnier.")
            elif zone and cat.get('periode_type'):
                saison_mois = zone[cat['periode_type']]
                saison_label = "la saison des pluies" if cat['periode_type'] == 'pluies' else "la saison sèche"
                ou = f" à {lieu_label}" if lieu_label else f" en {zone['label']}"
                chevauche = bool(set(pic_mois) & saison_mois)
                if pic_mois and chevauche:
                    insight_parts.append(f"{cat['icone']} Cohérent avec {saison_label}{ou} : {cat['explication']}.")
                elif pic_mois and not chevauche:
                    atypique = True
                    insight_parts.append(f"{cat['icone']} Pic atypique : survient hors de {saison_label} habituellement associée à cette pathologie{ou} — à surveiller (foyer localisé, ou effectif encore faible).")
                elif not pic_mois and p['total'] >= 3:
                    insight_parts.append(f"{cat['icone']} Habituellement plus marqué en {saison_label}{ou}, mais réparti ici de façon régulière sur la période analysée.")

        resultat.append({
            'code': p['code'],
            'label': p['label'],
            'total': p['total'],
            'pourcentage': round(p['total'] / total_mentions * 100, 1) if total_mentions else 0,
            'par_mois': p['par_mois'],
            'pic_mois': pic_mois,
            'insight': ' '.join(insight_parts) or None,
            'categorie_icone': categorie_icone,
            'atypique': atypique,
        })

    resultat.sort(key=lambda x: x['total'], reverse=True)
    return resultat, NOMS_MOIS


def _calculer_statistiques():
    """Calcule toutes les statistiques (KPI, pathologies, médecins,
    infirmiers, hospitalisations, analyses...) selon les filtres de la
    requête courante (request.args) et le rôle de l'utilisateur connecté.
    Factorisé hors de statistiques() pour être réutilisé tel quel par les
    exports Excel/TXT — mêmes chiffres partout, un seul endroit à faire
    évoluer. Retourne un dict prêt à passer à render_template(**d)."""
    from models import Patient, Consultation, Utilisateur, Hospitalisation, ConstanteVitale, AnalyseDemande, HospitalisationInfirmier
    from datetime import datetime, timedelta
    from sqlalchemy import func, extract

    # Récupérer les filtres
    date_debut = request.args.get('date_debut', '')
    date_fin = request.args.get('date_fin', '')
    periode = request.args.get('periode', 'mois')
    medecin_id = request.args.get('medecin_id', '')
    type_assurance = request.args.get('type_assurance', '')
    
    # ========== Construction de la requête de base ==========
    base_query = Consultation.query.join(Patient, Consultation.id_patient == Patient.id)
    
    # Filtrer selon le rôle
    if current_user.role == 'medecin':
        base_query = base_query.filter(Consultation.id_medecin == current_user.id)
    elif current_user.role == 'admin_structure':
        base_query = base_query.filter(Patient.id_structure == current_user.id_structure)
    
    # Application des filtres de date
    if date_debut:
        base_query = base_query.filter(Consultation.date_consultation >= datetime.strptime(date_debut, '%Y-%m-%d'))
    if date_fin:
        base_query = base_query.filter(Consultation.date_consultation <= datetime.strptime(date_fin, '%Y-%m-%d') + timedelta(days=1))
    
    # Filtre par médecin
    if medecin_id and medecin_id != '':
        base_query = base_query.filter(Consultation.id_medecin == int(medecin_id))
    
    # Filtre par type d'assurance
    if type_assurance and type_assurance != '':
        base_query = base_query.filter(Patient.type_assurance == type_assurance)
    
    # ========== KPI principaux ==========
    total_consultations = base_query.count()
    total_patients = db.session.query(Patient).filter(Patient.id.in_(
        base_query.with_entities(Consultation.id_patient).distinct()
    )).count()
    
    # ========== Patients par période ==========
    patients_par_periode = []
    
    if periode == 'jour':
        periodes = base_query.with_entities(
            func.date(Consultation.date_consultation).label('jour'),
            func.count(func.distinct(Consultation.id_patient)).label('nb')
        ).group_by('jour').order_by('jour').limit(30).all()
        patients_par_periode = [{'periode': p.jour.strftime('%d/%m/%Y'), 'nb': p.nb} for p in periodes]
    elif periode == 'semaine':
        periodes = base_query.with_entities(
            extract('year', Consultation.date_consultation).label('annee'),
            extract('week', Consultation.date_consultation).label('semaine'),
            func.count(func.distinct(Consultation.id_patient)).label('nb')
        ).group_by('annee', 'semaine').order_by('annee', 'semaine').limit(30).all()
        patients_par_periode = [{'periode': f"S{int(p.semaine)}-{int(p.annee)}", 'nb': p.nb} for p in periodes]
    elif periode == 'mois':
        periodes = base_query.with_entities(
            func.to_char(Consultation.date_consultation, 'YYYY-MM').label('mois'),
            func.count(func.distinct(Consultation.id_patient)).label('nb')
        ).group_by('mois').order_by('mois').limit(12).all()
        patients_par_periode = [{'periode': p.mois, 'nb': p.nb} for p in periodes]
    else:
        periodes = base_query.with_entities(
            extract('year', Consultation.date_consultation).label('annee'),
            func.count(func.distinct(Consultation.id_patient)).label('nb')
        ).group_by('annee').order_by('annee').all()
        patients_par_periode = [{'periode': str(int(p.annee)), 'nb': p.nb} for p in periodes]
    
    # ========== Top pathologies ==========
    top_pathologies = base_query.with_entities(
        Consultation.diagnostic,
        func.count(Consultation.id).label('total')
    ).filter(
        Consultation.diagnostic.isnot(None),
        Consultation.diagnostic != '',
        Consultation.diagnostic != '-'
    ).group_by(Consultation.diagnostic).order_by(func.count(Consultation.id).desc()).limit(10).all()

    # ========== ⭐ ANALYSE INTELLIGENTE DES PATHOLOGIES ==========
    # Une pathologie par ligne de diagnostic (voir _construire_analyse_
    # pathologies ci-dessus), avec répartition mensuelle et pic saisonnier
    # détecté — complète top_pathologies (gardé pour le KPI "guérison" plus
    # haut) sans le remplacer.
    consultations_diag = base_query.with_entities(
        Consultation.date_consultation, Consultation.diagnostic
    ).filter(
        Consultation.diagnostic.isnot(None),
        Consultation.diagnostic != ''
    ).all()
    # ⭐ Zone climatique de LA STRUCTURE (pas figée sur Lomé — voir
    # _resoudre_zone_climatique) : None si pays/ville non renseigné dans
    # /structure, auquel cas aucune explication climatique locale n'est
    # inventée dans l'analyse ci-dessous.
    zone_key = _resoudre_zone_climatique(current_user.structure)
    lieu_label = None
    if current_user.structure:
        lieu_label = current_user.structure.ville or current_user.structure.pays
    analyse_pathologies, noms_mois = _construire_analyse_pathologies(consultations_diag, zone_key, lieu_label)

    # ========== Répartition assurances ==========
    # ⭐ FIX : comptait Patient.id une fois PAR CONSULTATION (jointure
    # Patient-Consultation, un patient avec plusieurs consultations sur la
    # période était compté plusieurs fois) — les pourcentages dépassaient
    # largement 100% dès qu'un patient avait plus d'une consultation.
    # func.count(distinct(...)) aligne ce total sur total_patients (patients
    # UNIQUES), qui sert de dénominateur du pourcentage dans le template.
    assurances = db.session.query(
        Patient.type_assurance,
        func.count(func.distinct(Patient.id)).label('total')
    ).join(Consultation, Patient.id == Consultation.id_patient).filter(
        Consultation.id.in_(base_query.with_entities(Consultation.id))
    ).group_by(Patient.type_assurance).all()
    
    # ========== 📊 PERFORMANCE DES MÉDECINS ==========
    if current_user.role == 'admin_structure':
        stats_medecins = db.session.query(
            Utilisateur.id,
            Utilisateur.nom,
            Utilisateur.prenom,
            func.count(Consultation.id).label('nb_consultations'),
            func.count(func.distinct(Consultation.id_patient)).label('nb_patients')
        ).join(Consultation, Utilisateur.id == Consultation.id_medecin).filter(
            Consultation.id.in_(base_query.with_entities(Consultation.id))
        ).group_by(Utilisateur.id).all()
    else:
        stats_medecins = []
    
    # ========== 🩺 PERFORMANCE DES INFIRMIERS ==========
    stats_infirmiers = []
    if current_user.role == 'admin_structure':
        infirmiers = Utilisateur.query.filter_by(
            id_structure=current_user.id_structure,
            role='infirmier',
            actif=True
        ).all()
        
        for inf in infirmiers:
            nb_constantes = ConstanteVitale.query.filter_by(infirmier_id=inf.id).count()
            nb_hospitalisations = HospitalisationInfirmier.query.filter_by(
                infirmier_id=inf.id,
                actif=True
            ).count()
            derniere_constante = ConstanteVitale.query.filter_by(
                infirmier_id=inf.id
            ).order_by(ConstanteVitale.date_prise.desc()).first()
            
            stats_infirmiers.append({
                'id': inf.id,
                'nom': inf.nom,
                'prenom': inf.prenom,
                'nb_constantes': nb_constantes,
                'nb_hospitalisations': nb_hospitalisations,
                'derniere_constante': derniere_constante.date_prise if derniere_constante else None
            })
        
        stats_infirmiers.sort(key=lambda x: x['nb_constantes'], reverse=True)
    
    # ========== 📊 STATISTIQUES HOSPITALISATIONS ==========
    stats_hospitalisations = {}
    if current_user.role == 'admin_structure' or current_user.role == 'medecin':
        structure_id = current_user.id_structure if current_user.id_structure else 1
        
        # Récupérer les IDs des patients de la structure
        patient_ids = db.session.query(Patient.id).filter(
            Patient.id_structure == structure_id
        ).all()
        patient_ids = [p[0] for p in patient_ids]
        
        if patient_ids:
            # ⭐ Respecte désormais les mêmes filtres de dates que le reste de
            # la page (avant : toujours toutes dates confondues, impossible
            # de voir l'évolution des hospitalisations sur une période).
            hosp_query = Hospitalisation.query.filter(Hospitalisation.patient_id.in_(patient_ids))
            if date_debut:
                hosp_query = hosp_query.filter(Hospitalisation.date_debut >= datetime.strptime(date_debut, '%Y-%m-%d'))
            if date_fin:
                hosp_query = hosp_query.filter(Hospitalisation.date_debut <= datetime.strptime(date_fin, '%Y-%m-%d') + timedelta(days=1))
            hosp_ids = hosp_query.with_entities(Hospitalisation.id)

            total_hosp = hosp_query.count()
            hosp_actives = hosp_query.filter(Hospitalisation.statut == 'actif').count()

            hosp_par_service = db.session.query(
                Hospitalisation.service,
                func.count(Hospitalisation.id).label('total')
            ).filter(
                Hospitalisation.id.in_(hosp_ids)
            ).group_by(Hospitalisation.service).all()

            # ⭐ Admissions par mois — pour voir les périodes de forte
            # affluence (ex. pics saisonniers de paludisme grave nécessitant
            # une hospitalisation), demandé explicitement par la structure.
            hosp_par_periode_rows = db.session.query(
                func.to_char(Hospitalisation.date_debut, 'YYYY-MM').label('mois'),
                func.count(Hospitalisation.id).label('nb')
            ).filter(
                Hospitalisation.id.in_(hosp_ids)
            ).group_by('mois').order_by('mois').limit(12).all()

            # ⭐ Durée moyenne de séjour (jours), sur les hospitalisations
            # déjà clôturées (date_fin renseignée) de la période filtrée.
            duree_moyenne = db.session.query(
                func.avg(func.extract('epoch', Hospitalisation.date_fin - Hospitalisation.date_debut) / 86400.0)
            ).filter(
                Hospitalisation.id.in_(hosp_ids),
                Hospitalisation.date_fin.isnot(None)
            ).scalar()

            stats_hospitalisations = {
                'total': total_hosp,
                'actives': hosp_actives,
                'par_service': [{'service': s[0], 'total': s[1]} for s in hosp_par_service],
                'par_periode': [{'periode': h.mois, 'nb': h.nb} for h in hosp_par_periode_rows],
                'duree_moyenne_sejour': round(duree_moyenne, 1) if duree_moyenne else None,
            }
        else:
            stats_hospitalisations = {'total': 0, 'actives': 0, 'par_service': [], 'par_periode': [], 'duree_moyenne_sejour': None}
    
    # ========== 📊 STATISTIQUES ANALYSES ==========
    stats_analyses = {}
    if current_user.role == 'admin_structure' or current_user.role == 'medecin' or current_user.role == 'laborantin':
        structure_id = current_user.id_structure if current_user.id_structure else 1
        
        # Récupérer les IDs des patients de la structure
        patient_ids = db.session.query(Patient.id).filter(
            Patient.id_structure == structure_id
        ).all()
        patient_ids = [p[0] for p in patient_ids]
        
        if patient_ids:
            total_analyses = AnalyseDemande.query.filter(
                AnalyseDemande.patient_id.in_(patient_ids)
            ).count()
            
            analyses_par_statut = db.session.query(
                AnalyseDemande.statut,
                func.count(AnalyseDemande.id).label('total')
            ).filter(
                AnalyseDemande.patient_id.in_(patient_ids)
            ).group_by(AnalyseDemande.statut).all()
            
            analyses_par_type = db.session.query(
                AnalyseDemande.type_analyse,
                func.count(AnalyseDemande.id).label('total')
            ).filter(
                AnalyseDemande.patient_id.in_(patient_ids)
            ).group_by(AnalyseDemande.type_analyse).all()
            
            stats_analyses = {
                'total': total_analyses,
                'par_statut': [{'statut': s[0], 'total': s[1]} for s in analyses_par_statut],
                'par_type': [{'type': t[0], 'total': t[1]} for t in analyses_par_type]
            }
        else:
            stats_analyses = {'total': 0, 'par_statut': [], 'par_type': []}
    
    # ========== Évolution quotidienne ==========
    evolution = base_query.with_entities(
        func.date_trunc('day', Consultation.date_consultation).label('date'),
        func.count(Consultation.id).label('total')
    ).group_by('date').order_by('date').limit(60).all()
    
    evolution_labels = [e.date.strftime('%d/%m') if e.date else '-' for e in evolution]
    evolution_data = [e.total for e in evolution]
    
    # Liste des médecins pour le filtre (admin structure)
    if current_user.role == 'admin_structure':
        medecins = Utilisateur.query.filter_by(id_structure=current_user.id_structure, role='medecin', actif=True).all()
    else:
        medecins = []
    
    types_assurance = ['AMU-CNSS', 'AMU-INAM', 'AUTRE_ASSURANCE', 'NON_ASSURÉ']
    
    return {
        'total_consultations': total_consultations,
        'total_patients': total_patients,
        'patients_par_periode': patients_par_periode,
        'top_pathologies': top_pathologies,
        'analyse_pathologies': analyse_pathologies,
        'noms_mois': noms_mois,
        'zone_climatique_label': _ZONES_CLIMATIQUES[zone_key]['label'] if zone_key else None,
        'zone_climatique_lieu': lieu_label,
        'assurances': assurances,
        'stats_medecins': stats_medecins,
        'stats_infirmiers': stats_infirmiers,
        'stats_hospitalisations': stats_hospitalisations,
        'stats_analyses': stats_analyses,
        'medecins': medecins,
        'types_assurance': types_assurance,
        'evolution_labels': evolution_labels,
        'evolution_data': evolution_data,
        'periode': periode,
        'date_debut': date_debut,
        'date_fin': date_fin,
        'medecin_id': medecin_id,
        'type_assurance': type_assurance,
    }


@app.route('/statistiques')
@login_required
def statistiques():
    return render_template('statistiques.html', **_calculer_statistiques())


def _nom_periode_filtre(donnees):
    """Texte lisible décrivant la période filtrée, pour l'en-tête des
    exports."""
    if donnees['date_debut'] and donnees['date_fin']:
        return f"du {donnees['date_debut']} au {donnees['date_fin']}"
    if donnees['date_debut']:
        return f"depuis le {donnees['date_debut']}"
    if donnees['date_fin']:
        return f"jusqu'au {donnees['date_fin']}"
    return "toutes dates confondues"


@app.route('/statistiques/export/excel')
@login_required
@has_permission('STATISTIQUES')
def export_statistiques_excel():
    """Export Excel multi-feuilles des statistiques — mêmes filtres/chiffres
    que la page (voir _calculer_statistiques). Feuilles : Résumé, Pathologies
    (avec répartition mensuelle), Évolution, Médecins, Hospitalisations."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from flask import Response
    from datetime import datetime
    from io import BytesIO

    d = _calculer_statistiques()
    entete_fill = PatternFill(start_color='1a3a5c', end_color='1a3a5c', fill_type='solid')
    entete_font = Font(color='FFFFFF', bold=True)

    def _entete(ws, colonnes):
        ws.append(colonnes)
        for cell in ws[1]:
            cell.fill = entete_fill
            cell.font = entete_font
            cell.alignment = Alignment(horizontal='center')

    wb = Workbook()

    # ---- Résumé ----
    ws = wb.active
    ws.title = 'Résumé'
    ws.append([f"Statistiques — {current_user.structure.nom if current_user.structure else ''}"])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([f"Période : {_nom_periode_filtre(d)} — généré le {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    if d['zone_climatique_label']:
        lieu = f" ({d['zone_climatique_lieu']})" if d['zone_climatique_lieu'] else ''
        ws.append([f"Zone climatique : {d['zone_climatique_label']}{lieu} — utilisée pour les observations épidémiologiques"])
    ws.append([])
    _entete(ws, ['Indicateur', 'Valeur'])
    ws.append(['Total consultations', d['total_consultations']])
    ws.append(['Patients uniques', d['total_patients']])
    if d['stats_hospitalisations']:
        ws.append(['Hospitalisations (total)', d['stats_hospitalisations'].get('total', 0)])
        ws.append(['Hospitalisations actives', d['stats_hospitalisations'].get('actives', 0)])
        if d['stats_hospitalisations'].get('duree_moyenne_sejour') is not None:
            ws.append(['Durée moyenne de séjour (jours)', d['stats_hospitalisations']['duree_moyenne_sejour']])
    if d['stats_analyses']:
        ws.append(['Analyses labo/radio (total)', d['stats_analyses'].get('total', 0)])
    for col, width in (('A', 32), ('B', 14)):
        ws.column_dimensions[col].width = width

    # ---- Pathologies ----
    ws2 = wb.create_sheet('Pathologies')
    _entete(ws2, ['#', 'Code CIM-10', 'Pathologie', 'Cas', '%'] + d['noms_mois'] + ['Pic saisonnier', 'Observation'])
    for i, p in enumerate(d['analyse_pathologies'], start=1):
        ws2.append([
            i, p['code'] or '-', p['label'], p['total'], p['pourcentage'],
            *p['par_mois'],
            ', '.join(p['pic_mois']) or '-',
            p['insight'] or '-',
        ])
    largeurs2 = [5, 12, 40] + [8] * (2 + 12) + [16, 55]
    for idx, largeur in enumerate(largeurs2, start=1):
        ws2.column_dimensions[get_column_letter(idx)].width = largeur

    # ---- Évolution (patients par période) ----
    ws3 = wb.create_sheet('Évolution')
    _entete(ws3, ['Période', 'Patients (uniques)'])
    for p in d['patients_par_periode']:
        ws3.append([p['periode'], p['nb']])
    ws3.column_dimensions['A'].width = 18
    ws3.column_dimensions['B'].width = 18

    # ---- Médecins ----
    if d['stats_medecins']:
        ws4 = wb.create_sheet('Médecins')
        _entete(ws4, ['Médecin', 'Consultations', 'Patients suivis'])
        for m in d['stats_medecins']:
            ws4.append([f"Dr {m.prenom} {m.nom}", m.nb_consultations, m.nb_patients])
        ws4.column_dimensions['A'].width = 28

    # ---- Hospitalisations ----
    if d['stats_hospitalisations'] and d['stats_hospitalisations'].get('par_service'):
        ws5 = wb.create_sheet('Hospitalisations')
        _entete(ws5, ['Service', 'Nombre'])
        for s in d['stats_hospitalisations']['par_service']:
            ws5.append([s['service'], s['total']])
        ws5.column_dimensions['A'].width = 28

        if d['stats_hospitalisations'].get('par_periode'):
            ws5.append([])
            row_entete_periode = ws5.max_row + 1
            ws5.append(['Période (admissions)', 'Nombre'])
            for cell in ws5[row_entete_periode]:
                cell.fill = entete_fill
                cell.font = entete_font
                cell.alignment = Alignment(horizontal='center')
            for h in d['stats_hospitalisations']['par_periode']:
                ws5.append([h['periode'], h['nb']])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename=statistiques_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'}
    )


@app.route('/statistiques/export/txt')
@login_required
@has_permission('STATISTIQUES')
def export_statistiques_txt():
    """Export texte brut, lisible — mêmes filtres/chiffres que la page."""
    from flask import Response
    from datetime import datetime

    d = _calculer_statistiques()
    nom_structure = current_user.structure.nom if current_user.structure else 'Structure'
    lignes = []
    lignes.append(f"STATISTIQUES — {nom_structure}")
    lignes.append(f"Période : {_nom_periode_filtre(d)}")
    lignes.append(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}")
    if d['zone_climatique_label']:
        lieu = f" ({d['zone_climatique_lieu']})" if d['zone_climatique_lieu'] else ''
        lignes.append(f"Zone climatique : {d['zone_climatique_label']}{lieu} — utilisée pour les observations épidémiologiques")
    lignes.append('=' * 70)
    lignes.append('')
    lignes.append('RÉSUMÉ')
    lignes.append('-' * 70)
    lignes.append(f"Total consultations : {d['total_consultations']}")
    lignes.append(f"Patients uniques : {d['total_patients']}")
    if d['stats_hospitalisations']:
        lignes.append(f"Hospitalisations : {d['stats_hospitalisations'].get('total', 0)} (dont {d['stats_hospitalisations'].get('actives', 0)} en cours)")
        if d['stats_hospitalisations'].get('duree_moyenne_sejour') is not None:
            lignes.append(f"Durée moyenne de séjour : {d['stats_hospitalisations']['duree_moyenne_sejour']} jours")
    if d['stats_analyses']:
        lignes.append(f"Analyses labo/radio : {d['stats_analyses'].get('total', 0)}")
    lignes.append('')

    lignes.append('PATHOLOGIES LES PLUS FRÉQUENTES')
    lignes.append('-' * 70)
    if d['analyse_pathologies']:
        for i, p in enumerate(d['analyse_pathologies'][:30], start=1):
            code = f"[{p['code']}] " if p['code'] else ''
            lignes.append(f"{i:>2}. {code}{p['label']} — {p['total']} cas ({p['pourcentage']}%)")
            if p['insight']:
                lignes.append(f"    → {p['insight']}")
    else:
        lignes.append("Aucune pathologie enregistrée sur cette période.")
    lignes.append('')

    lignes.append('ÉVOLUTION')
    lignes.append('-' * 70)
    for p in d['patients_par_periode']:
        lignes.append(f"{p['periode']:<15} {p['nb']} patient(s)")
    lignes.append('')

    if d['stats_medecins']:
        lignes.append('MÉDECINS')
        lignes.append('-' * 70)
        for m in d['stats_medecins']:
            lignes.append(f"Dr {m.prenom} {m.nom} — {m.nb_consultations} consultation(s), {m.nb_patients} patient(s)")
        lignes.append('')

    if d['stats_hospitalisations'] and d['stats_hospitalisations'].get('par_service'):
        lignes.append('HOSPITALISATIONS PAR SERVICE')
        lignes.append('-' * 70)
        for s in d['stats_hospitalisations']['par_service']:
            lignes.append(f"{s['service']:<30} {s['total']}")
        lignes.append('')

    if d['stats_hospitalisations'] and d['stats_hospitalisations'].get('par_periode'):
        lignes.append('ADMISSIONS PAR PÉRIODE')
        lignes.append('-' * 70)
        for h in d['stats_hospitalisations']['par_periode']:
            lignes.append(f"{h['periode']:<15} {h['nb']} admission(s)")
        lignes.append('')

    contenu = '\n'.join(lignes)
    return Response(
        contenu,
        mimetype='text/plain; charset=utf-8',
        headers={'Content-Disposition': f'attachment; filename=statistiques_{datetime.now().strftime("%Y%m%d_%H%M")}.txt'}
    )

@app.route('/admin/sync-sheets', methods=['POST'])
@login_required
def admin_sync_sheets():
    """Synchronisation manuelle vers Google Sheets (Super Admin uniquement)"""
    from flask import jsonify
    import traceback
    
    if current_user.role != 'super_admin':
        return jsonify({'success': False, 'error': 'Non autorisé'}), 403
    
    try:
        from sheets_sync import GoogleSheetsSync
        SPREADSHEET_ID = "1nCUArOaWgXVFszjEhH1GqNJXGCV7cF754W87vXvQ-lQ"
        syncer = GoogleSheetsSync(SPREADSHEET_ID)
        syncer.sync_all()
        return jsonify({'success': True, 'message': '✅ Synchronisation terminée vers Google Sheets'})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)})

# ==================== MESSAGERIE INTERNE ====================

@app.route('/messages')
@login_required
def messages_liste():
    from models import Message
    
    # Messages reçus
    messages_recus = Message.query.filter_by(id_destinataire=current_user.id).order_by(Message.date_envoi.desc()).all()
    
    # Messages envoyés
    messages_envoyes = Message.query.filter_by(id_expediteur=current_user.id).order_by(Message.date_envoi.desc()).all()
    
    # Nombre de messages non lus
    non_lus = Message.query.filter_by(id_destinataire=current_user.id, lu=False).count()
    
    return render_template('messages/liste.html', 
                         messages_recus=messages_recus,
                         messages_envoyes=messages_envoyes,
                         non_lus=non_lus)


@app.route('/messages/nouveau', methods=['GET', 'POST'])
@login_required
def messages_nouveau():
    from models import Message, Utilisateur
    
    # Récupérer les destinataires possibles (même structure)
    if current_user.role == 'super_admin':
        destinataires = Utilisateur.query.filter(Utilisateur.id != current_user.id).all()
    else:
        destinataires = Utilisateur.query.filter(
            Utilisateur.id_structure == current_user.id_structure,
            Utilisateur.id != current_user.id,
            Utilisateur.actif == True
        ).all()
    
    if request.method == 'POST':
        id_destinataire = request.form.get('id_destinataire')
        sujet = request.form.get('sujet')
        contenu = request.form.get('contenu')
        
        if not id_destinataire or not sujet or not contenu:
            flash('Tous les champs sont obligatoires', 'danger')
            return redirect(url_for('messages_nouveau'))
        
        message = Message(
            id_expediteur=current_user.id,
            id_destinataire=int(id_destinataire),
            id_structure=current_user.id_structure if current_user.id_structure else 1,
            sujet=sujet,
            contenu=contenu
        )
        
        db.session.add(message)
        db.session.commit()
        
        flash('Message envoyé avec succès', 'success')
        return redirect(url_for('messages_liste'))
    
    return render_template('messages/nouveau.html', destinataires=destinataires)


@app.route('/messages/lire/<int:id>')
@login_required
def messages_lire(id):
    from models import Message
    from datetime import datetime
    
    message = Message.query.get_or_404(id)
    
    # Vérifier que l'utilisateur est concerné
    if message.id_destinataire != current_user.id and message.id_expediteur != current_user.id:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('messages_liste'))
    
    # Marquer comme lu si c'est le destinataire
    if message.id_destinataire == current_user.id and not message.lu:
        message.lu = True
        message.lu_at = datetime.utcnow()
        db.session.commit()
    
    return render_template('messages/lire.html', message=message)


@app.route('/messages/supprimer/<int:id>')
@login_required
def messages_supprimer(id):
    from models import Message
    
    message = Message.query.get_or_404(id)
    
    # Vérifier que l'utilisateur est l'expéditeur ou le destinataire
    if message.id_expediteur != current_user.id and message.id_destinataire != current_user.id:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('messages_liste'))
    
    db.session.delete(message)
    db.session.commit()
    
    flash('Message supprimé', 'success')
    return redirect(url_for('messages_liste'))


# Démarrer le scheduler au lancement de l'application
start_scheduler()

# Arrêter le scheduler proprement à la fermeture
import atexit
atexit.register(stop_scheduler)

@app.route('/admin/profil', methods=['GET', 'POST'])
@login_required
def admin_profil():
    if current_user.role != 'super_admin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    if request.method == 'POST':
        ancien_mdp = request.form.get('ancien_mdp')
        nouveau_mdp = request.form.get('nouveau_mdp')
        confirmer_mdp = request.form.get('confirmer_mdp')
        
        if not current_user.check_password(ancien_mdp):
            flash('Ancien mot de passe incorrect', 'danger')
            return redirect(url_for('admin_profil'))
        
        if nouveau_mdp != confirmer_mdp:
            flash('Les nouveaux mots de passe ne correspondent pas', 'danger')
            return redirect(url_for('admin_profil'))
        
        # Validation de la complexité
        import re
        if len(nouveau_mdp) < 8:
            flash('Minimum 8 caractères', 'danger')
        elif not re.search(r"[A-Z]", nouveau_mdp):
            flash('Au moins 1 majuscule', 'danger')
        elif not re.search(r"[a-z]", nouveau_mdp):
            flash('Au moins 1 minuscule', 'danger')
        elif not re.search(r"[0-9]", nouveau_mdp):
            flash('Au moins 1 chiffre', 'danger')
        elif not re.search(r"[!@#$%^&*(),.?\":{}|<>]", nouveau_mdp):
            flash('Au moins 1 symbole', 'danger')
        else:
            current_user.set_password(nouveau_mdp)
            db.session.commit()
            flash('Mot de passe modifié avec succès', 'success')
            return redirect(url_for('admin_profil'))
    
    return render_template('admin/profil.html')

@app.route('/admin/nouveau-super-admin', methods=['GET', 'POST'])
@login_required
def admin_nouveau_super_admin():
    if current_user.role != 'super_admin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Utilisateur
    from datetime import datetime
    
    if request.method == 'POST':
        email = request.form.get('email')
        nom = request.form.get('nom')
        prenom = request.form.get('prenom')
        password = request.form.get('password')
        
        # Vérifier si l'email existe déjà
        existing = Utilisateur.query.filter_by(email=email).first()
        if existing:
            flash('Cet email est déjà utilisé', 'danger')
            return redirect(url_for('admin_nouveau_super_admin'))
        
        # Créer le Super Admin
        new_admin = Utilisateur(
            email=email,
            nom=nom,
            prenom=prenom,
            role='super_admin',
            actif=True
        )
        new_admin.set_password(password)
        db.session.add(new_admin)
        db.session.commit()
        
        flash(f'Super Admin {prenom} {nom} créé avec succès', 'success')
        return redirect(url_for('admin_structures'))
    
    return render_template('admin/nouveau_super_admin.html')

@app.route('/reset-password/<token>', methods=['GET', 'POST'])
def reset_password_token(token):
    from models import Utilisateur
    from datetime import datetime
    
    user = Utilisateur.query.filter_by(reset_token=token).first()
    
    if not user or not user.reset_token_expiry or user.reset_token_expiry < datetime.utcnow():
        flash('Le lien de réinitialisation est invalide ou a expiré.', 'danger')
        return redirect(url_for('forgot_password'))
    
    if request.method == 'POST':
        new_password = request.form.get('new_password')
        confirm_password = request.form.get('confirm_password')
        
        if new_password != confirm_password:
            flash('Les mots de passe ne correspondent pas.', 'danger')
        else:
            user.set_password(new_password)
            user.reset_token = None
            user.reset_token_expiry = None
            db.session.commit()
            flash('Mot de passe réinitialisé avec succès.', 'success')
            return redirect(url_for('login'))
    
    return render_template('reset_password_token.html', token=token)

@app.route('/admin/structure/<int:id>/desactivate', methods=['GET', 'POST'])
@login_required
def admin_desactivate_structure(id):
    if current_user.role != 'super_admin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Structure
    structure = Structure.query.get_or_404(id)
    
    if request.method == 'POST':
        structure.statut = 'desactive'
        db.session.commit()
        flash(f'Structure {structure.nom} désactivée avec succès', 'success')
        return redirect(url_for('admin_structures'))
    
    return render_template('admin/desactiver_structure.html', structure=structure)

@app.route('/admin/structure/<int:id>/delete', methods=['GET', 'POST'])
@login_required
def admin_delete_structure(id):
    if current_user.role != 'super_admin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Structure, Utilisateur, Patient, Consultation, Prescription
    
    structure = Structure.query.get_or_404(id)
    
    if request.method == 'POST':
        # Compter les données avant suppression
        nb_users = Utilisateur.query.filter_by(id_structure=id).count()
        nb_patients = Patient.query.filter_by(id_structure=id).count()
        nb_consultations = Consultation.query.join(Patient).filter(Patient.id_structure == id).count()
        nb_prescriptions = Prescription.query.join(Patient).filter(Patient.id_structure == id).count()
        
        # Supprimer en cascade
        # 1. Supprimer les prescriptions
        Prescription.query.filter(Prescription.id_patient.in_(
            db.session.query(Patient.id).filter_by(id_structure=id)
        )).delete(synchronize_session=False)
        
        # 2. Supprimer les consultations
        Consultation.query.filter(Consultation.id_patient.in_(
            db.session.query(Patient.id).filter_by(id_structure=id)
        )).delete(synchronize_session=False)
        
        # 3. Supprimer les patients
        Patient.query.filter_by(id_structure=id).delete()
        
        # 4. Supprimer les utilisateurs
        Utilisateur.query.filter_by(id_structure=id).delete()
        
        # 5. Supprimer la structure
        db.session.delete(structure)
        db.session.commit()
        
        flash(f'Structure {structure.nom} supprimée avec succès. Données supprimées : {nb_users} utilisateurs, {nb_patients} patients, {nb_consultations} consultations, {nb_prescriptions} prescriptions.', 'success')
        return redirect(url_for('admin_structures'))
    
    return render_template('admin/supprimer_structure.html', structure=structure)

# ==================== HOSPITALISATIONS ====================

@app.route('/hospitalisations')
@login_required
@has_permission('HOSPITALISATION')
def liste_hospitalisations():
    """Liste des hospitalisations"""
    from models import Hospitalisation, HospitalisationInfirmier, Patient, Consultation

    if current_user.role not in ['admin_structure', 'medecin', 'infirmier', 'secretaire', 'super_admin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    # ⭐ File d'attente d'admission : consultations orientées "à
    # hospitaliser" (Consultation.orientation) pour lesquelles aucune
    # Hospitalisation n'a encore été créée (Hospitalisation.consultation_id)
    # — patron : "si c'est à hospitaliser il doit apparaitre dans
    # hospitalisation (hospitalisation en attente par exemple)".
    consultations_query = (
        Consultation.query
        .join(Patient, Consultation.id_patient == Patient.id)
        .outerjoin(Hospitalisation, Hospitalisation.consultation_id == Consultation.id)
        .filter(
            Consultation.orientation == 'hospitalisation',
            Hospitalisation.id.is_(None),
        )
    )
    if current_user.role != 'super_admin':
        consultations_query = consultations_query.filter(Patient.id_structure == current_user.id_structure)
    consultations_en_attente = consultations_query.order_by(Consultation.date_consultation.desc()).all()
    
    statut = request.args.get('statut', 'tous')
    service = request.args.get('service', '')
    page = request.args.get('page', 1, type=int)
    per_page = 20
    
    # Base de la requête
    if current_user.role == 'super_admin':
        query = Hospitalisation.query
    else:
        query = Hospitalisation.query.join(Patient).filter(Patient.id_structure == current_user.id_structure)
    
    # Filtrer par statut
    if statut == 'actif':
        query = query.filter(Hospitalisation.statut == 'actif')
    elif statut == 'sorti':
        query = query.filter(Hospitalisation.statut == 'sorti')
    elif statut == 'transfere':
        query = query.filter(Hospitalisation.statut == 'transfere')
    
    if service:
        query = query.filter(Hospitalisation.service.ilike(f'%{service}%'))
    
    # Filtrer selon le rôle
    # ⭐ Un médecin voit maintenant TOUTES les hospitalisations de sa
    # structure (plus seulement les siennes) — condition nécessaire pour
    # pouvoir repérer un patient d'un autre service et lui demander l'accès
    # (voir demander_acces_hospitalisation) ; le détail clinique complet
    # reste gated par _a_acces_hospitalisation (le nom/service/médecin
    # assigné affichés dans cette liste sont l'équivalent d'un tableau de
    # service hospitalier classique, pas d'information clinique).
    if current_user.role == 'infirmier':
        query = query.join(HospitalisationInfirmier).filter(
            HospitalisationInfirmier.infirmier_id == current_user.id,
            HospitalisationInfirmier.actif == True
        )
    
    hospitalisations = query.order_by(Hospitalisation.date_debut.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    
    # Liste des services
    if current_user.role == 'super_admin':
        services_query = db.session.query(Hospitalisation.service).distinct()
    else:
        services_query = db.session.query(Hospitalisation.service).join(Patient).filter(
            Patient.id_structure == current_user.id_structure
        ).distinct()
    services = [s[0] for s in services_query.all() if s[0]]
    
    return render_template('hospitalisations/liste.html',
                         hospitalisations=hospitalisations,
                         statut_actuel=statut,
                         services=services,
                         consultations_en_attente=consultations_en_attente)


@app.route('/hospitalisation/nouvelle', methods=['GET', 'POST'])
@login_required
def nouvelle_hospitalisation():
    """Créer une nouvelle hospitalisation avec note d'admission structurée.

    ⭐ Un infirmier peut aussi admettre un patient (patron : "on va aussi
    permettre à l'infirmier d'admettre un patient en hospitalisation mais
    les informations seront complétées par le médecin") — dans ce cas la
    note d'admission (réservée au médecin) n'est PAS exigée ici : elle
    reste à compléter via /hospitalisation/<id>/note/ajouter (route déjà
    existante, gérée par ajouter_note_admission), et son absence
    (hospitalisation.note_admission_active_id NULL) sert de signal
    "à compléter" affiché en rouge (voir detail_hospitalisation/liste.html)."""
    from models import Patient, Utilisateur, Hospitalisation, HospitalisationMedecin, HospitalisationInfirmier, Service, Salle, Lit, NoteAdmission, Consultation, Message

    if current_user.role not in ['admin_structure', 'medecin', 'secretaire', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    if request.method == 'POST':
        # ============================================================
        # 1. RÉCUPÉRATION DES DONNÉES
        # ============================================================
        patient_id = request.form.get('patient_id')
        consultation_id = request.form.get('consultation_id', type=int)
        motif = request.form.get('motif')
        service = request.form.get('service')
        chambre = request.form.get('chambre')
        lit = request.form.get('lit')  # ⭐ Gardé pour compatibilité
        lit_id = request.form.get('lit_id', type=int)
        medecins_ids = request.form.getlist('medecins_ids')
        infirmiers_ids = request.form.getlist('infirmiers_ids')

        # ⭐ Le médecin de la consultation d'origine (s'il y en a une) est
        # automatiquement assigné comme médecin traitant — continuité de
        # la prise en charge, même s'il n'a pas été coché explicitement
        # dans le formulaire.
        if consultation_id:
            consultation_origine_assign = Consultation.query.get(consultation_id)
            if consultation_origine_assign and consultation_origine_assign.id_medecin and \
                    str(consultation_origine_assign.id_medecin) not in medecins_ids:
                medecins_ids.append(str(consultation_origine_assign.id_medecin))

        # ⭐ Un infirmier qui admet le patient est automatiquement inclus
        # dans les infirmiers assignés — sinon il perdrait l'accès à
        # l'hospitalisation qu'il vient lui-même de créer (voir
        # _a_acces_hospitalisation, qui exige une assignation active).
        if current_user.role == 'infirmier' and str(current_user.id) not in infirmiers_ids:
            infirmiers_ids.append(str(current_user.id))

        # ===== NOTE D'ADMISSION =====
        note_motif = request.form.get('note_motif', '').strip()
        note_contexte = request.form.get('note_contexte', '').strip()
        note_examen_clinique = request.form.get('note_examen_clinique', '').strip()
        note_diagnostic = request.form.get('note_diagnostic', '').strip()
        note_examens = request.form.get('note_examens', '').strip()
        note_traitement = request.form.get('note_traitement', '').strip()
        note_evolution_prevue = request.form.get('note_evolution_prevue', '').strip()
        note_conclusion = request.form.get('note_conclusion', '').strip()
        note_constantes = request.form.get('note_constantes', '').strip()
        
        # ============================================================
        # 2. VALIDATIONS
        # ============================================================
        
        if not patient_id:
            flash('Veuillez sélectionner un patient.', 'danger')
            return redirect(url_for('nouvelle_hospitalisation'))
        
        if not motif:
            flash('Le motif d\'hospitalisation est obligatoire.', 'danger')
            return redirect(url_for('nouvelle_hospitalisation'))
        
        if not service:
            flash('Le service est obligatoire.', 'danger')
            return redirect(url_for('nouvelle_hospitalisation'))
        
        # ⭐ Note d'admission — réservée au médecin (voir docstring), pas
        # exigée quand c'est un infirmier qui admet le patient.
        if current_user.role != 'infirmier':
            champs_obligatoires = {
                'Motif d\'hospitalisation': note_motif,
                'Contexte et antécédents': note_contexte,
                'Examen clinique initial': note_examen_clinique,
                'Diagnostic présumé': note_diagnostic,
                'Traitement initial': note_traitement,
                'Conclusion du médecin référent': note_conclusion
            }

            champs_manquants = []
            for nom, valeur in champs_obligatoires.items():
                if not valeur:
                    champs_manquants.append(nom)

            if champs_manquants:
                flash(f'La note d\'admission est incomplète. Champs obligatoires : {", ".join(champs_manquants)}', 'danger')
                return redirect(url_for('nouvelle_hospitalisation'))
        
        # ============================================================
        # 3. CRÉATION
        # ============================================================
        try:
            # --- Création de l'hospitalisation ---
            hospitalisation = Hospitalisation(
                patient_id=int(patient_id),
                consultation_id=consultation_id,
                motif=motif,
                service=service,
                chambre=chambre,
                lit=lit,  # ⭐ Gardé comme avant
                notes_admission=None,
                statut='actif',
                note_admission_a_completer=(current_user.role == 'infirmier'),
                created_by=current_user.id,
                created_at=datetime.utcnow()
            )
            db.session.add(hospitalisation)
            db.session.flush()

            # --- Copier l'examen physique de la consultation d'origine ---
            # (patron : "examen physique préremplie à modifier exactement
            # comme dans consultations") — sections_origine fige la copie
            # pour permettre, ensuite, de griser ce que le médecin change
            # depuis l'admission (voir hospitalisation_examen_physique).
            if consultation_id:
                from models import ExamenPhysique
                examen_source = ExamenPhysique.query.filter_by(consultation_id=consultation_id).first()
                if examen_source:
                    db.session.add(ExamenPhysique(
                        hospitalisation_id=hospitalisation.id,
                        sections_modifiees=examen_source.sections_modifiees,
                        sections_origine=examen_source.sections_modifiees or '{}',
                        examen_complet=examen_source.examen_complet,
                        created_by=current_user.id
                    ))

            # --- Assigner le lit ---
            if lit_id:
                lit_obj = Lit.query.get(lit_id)  # ⭐ Comme avant
                if lit_obj and lit_obj.statut == 'disponible':
                    lit_obj.occuper(hospitalisation.id)
                    hospitalisation.lit_id = lit_obj.id
                    hospitalisation.chambre = lit_obj.salle.nom
                    hospitalisation.lit = lit_obj.numero  # ⭐ Comme avant
            
            # --- Assigner les médecins ---
            for medecin_id in medecins_ids:
                hm = HospitalisationMedecin(
                    hospitalisation_id=hospitalisation.id,
                    medecin_id=int(medecin_id),
                    date_assignation=datetime.utcnow(),
                    actif=True
                )
                db.session.add(hm)
            
            # --- Assigner les infirmiers ---
            for infirmier_id in infirmiers_ids:
                hi = HospitalisationInfirmier(
                    hospitalisation_id=hospitalisation.id,
                    infirmier_id=int(infirmier_id),
                    date_assignation=datetime.utcnow(),
                    actif=True
                )
                db.session.add(hi)
            
            # --- Création de la note d'admission (médecin/admin/secrétaire
            #     uniquement — voir docstring de la route) ---
            if current_user.role != 'infirmier':
                note = NoteAdmission(
                    hospitalisation_id=hospitalisation.id,
                    version=1,
                    est_initial=True,
                    est_verrouillee=True,
                    motif_admission=note_motif,
                    contexte_admission=note_contexte,
                    examen_clinique_admission=note_examen_clinique,
                    diagnostic_admission=note_diagnostic,
                    examens_admission=note_examens if note_examens else None,
                    traitement_admission=note_traitement,
                    evolution_prevue=note_evolution_prevue if note_evolution_prevue else None,
                    conclusion_admission=note_conclusion,
                    constantes_admission=note_constantes if note_constantes else None,
                    redige_par=current_user.id,
                    date_redaction=datetime.utcnow(),
                    valide_par=current_user.id,
                    date_validation=datetime.utcnow()
                )
                db.session.add(note)
                db.session.flush()

                # --- Lier la note active ---
                hospitalisation.note_admission_active_id = note.id

            db.session.commit()

            flash(f'✅ Hospitalisation de {hospitalisation.patient.nom} {hospitalisation.patient.prenom} créée avec succès !', 'success')
            if current_user.role != 'infirmier':
                flash('📋 Note d\'admission verrouillée - Elle servira de référence pour le suivi.', 'info')
            else:
                # ⭐ Signale aux médecins assignés (ou, à défaut, à l'admin
                # de la structure) qu'une note d'admission reste à
                # compléter — même mécanisme de notification que la
                # demande d'accès inter-service (messagerie interne).
                flash('⚠️ Note d\'admission non renseignée — un médecin doit la compléter.', 'warning')
                destinataires_notif = [Utilisateur.query.get(int(mid)) for mid in medecins_ids] or Utilisateur.query.filter_by(
                    id_structure=hospitalisation.patient.id_structure, role='admin_structure', actif=True
                ).all()
                for cible in destinataires_notif:
                    if not cible:
                        continue
                    db.session.add(Message(
                        id_expediteur=current_user.id,
                        id_destinataire=cible.id,
                        id_structure=hospitalisation.patient.id_structure,
                        sujet=f"Note d'admission à compléter — {hospitalisation.patient.prenom} {hospitalisation.patient.nom}",
                        contenu=f"{current_user.prenom} {current_user.nom} (infirmier) a admis {hospitalisation.patient.prenom} {hospitalisation.patient.nom} en {hospitalisation.service}.\nLa note d'admission reste à rédiger — voir la fiche hospitalisation.",
                    ))
                db.session.commit()

            return redirect(url_for('detail_hospitalisation', id=hospitalisation.id))
            
        except Exception as e:
            db.session.rollback()
            print(f"❌ Erreur: {e}")
            import traceback
            traceback.print_exc()
            flash(f'❌ Erreur lors de l\'enregistrement : {str(e)}', 'danger')
            return redirect(url_for('nouvelle_hospitalisation'))
    
    # ============================================================
    # GET : Afficher le formulaire
    # ============================================================
    if current_user.role == 'super_admin':
        patients = Patient.query.filter_by(archived=False).all()
        medecins = Utilisateur.query.filter_by(role='medecin', actif=True).all()
        infirmiers = Utilisateur.query.filter_by(role='infirmier', actif=True).all()
    else:
        patients = Patient.query.filter_by(
            id_structure=current_user.id_structure,
            archived=False
        ).all()
        medecins = Utilisateur.query.filter_by(
            id_structure=current_user.id_structure,
            role='medecin',
            actif=True
        ).all()
        infirmiers = Utilisateur.query.filter_by(
            id_structure=current_user.id_structure,
            role='infirmier',
            actif=True
        ).all()
    
    services = Service.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()

    # ⭐ Arrivée depuis "Hospitalisations en attente" (liste_hospitalisations)
    # — pré-sélectionne le patient et prépare le lien vers la consultation
    # d'origine (motif pré-rempli, examen physique copié à la création).
    from models import Consultation
    consultation_origine = None
    patient_id_prerempli = request.args.get('patient_id', type=int)
    consultation_id_prerempli = request.args.get('consultation_id', type=int)
    if consultation_id_prerempli:
        consultation_origine = Consultation.query.get(consultation_id_prerempli)

    return render_template('hospitalisations/nouvelle.html',
                         patients=patients,
                         medecins=medecins,
                         infirmiers=infirmiers,
                         services=services,
                         consultation_origine=consultation_origine,
                         patient_id_prerempli=patient_id_prerempli)


def _a_acces_hospitalisation(user, hospitalisation):
    """⭐ Un médecin/infirmier a-t-il accès à CETTE hospitalisation ? Factorise
    la vérification dupliquée dans detail_hospitalisation/ajouter_evolution/
    ajouter_constante, et l'étend : un médecin non assigné a aussi accès s'il
    a une DemandeAccesHospitalisation acceptée et encore dans sa fenêtre de
    validité (voir /hospitalisation/<id>/demande-acces). super_admin/
    admin_structure : accès toujours vrai (inchangé)."""
    from models import HospitalisationMedecin, HospitalisationInfirmier, DemandeAccesHospitalisation
    from datetime import datetime

    if user.role in ('super_admin', 'admin_structure'):
        return True

    if user.role == 'medecin':
        assigne = HospitalisationMedecin.query.filter_by(
            hospitalisation_id=hospitalisation.id,
            medecin_id=user.id,
            actif=True
        ).first()
        if assigne:
            return True
        # ⭐ Hospitalisation "orpheline" (aucun médecin assigné, ex. admise
        # par un infirmier sans en choisir un — voir nouvelle_hospitalisation)
        # : n'importe quel médecin de la structure peut la prendre en
        # charge, pas besoin d'une demande d'accès (qui suppose un médecin
        # traitant existant à qui demander).
        if not HospitalisationMedecin.query.filter_by(hospitalisation_id=hospitalisation.id, actif=True).first() \
                and user.id_structure == hospitalisation.patient.id_structure:
            return True
        acces_temp = DemandeAccesHospitalisation.query.filter_by(
            hospitalisation_id=hospitalisation.id,
            demandeur_id=user.id,
            statut='acceptee'
        ).filter(DemandeAccesHospitalisation.date_fin > datetime.utcnow()).first()
        return acces_temp is not None

    if user.role == 'infirmier':
        assigne = HospitalisationInfirmier.query.filter_by(
            hospitalisation_id=hospitalisation.id,
            infirmier_id=user.id,
            actif=True
        ).first()
        return assigne is not None

    return False


@app.route('/hospitalisation/<int:id>')
@login_required
def detail_hospitalisation(id):
    """Détails d'une hospitalisation avec note d'admission structurée"""
    from models import Hospitalisation, HospitalisationMedecin, HospitalisationInfirmier, ConstanteVitale, EvolutionPatient, NoteAdmission, ProtocoleSoins, ExamenType, ExamenPrescrit
    import json

    hospitalisation = Hospitalisation.query.get_or_404(id)

    # Vérifier les permissions — voir _a_acces_hospitalisation (inclut
    # l'accès temporaire accordé via une demande d'accès inter-service).
    if not _a_acces_hospitalisation(current_user, hospitalisation):
        # ⭐ Un médecin non assigné (patient d'un autre service) est envoyé
        # vers le formulaire de demande d'accès plutôt qu'un simple refus —
        # c'est précisément le cas d'usage de cette fonctionnalité.
        if current_user.role == 'medecin':
            flash('Vous n\'êtes pas assigné à cette hospitalisation — vous pouvez demander l\'accès.', 'info')
            return redirect(url_for('demander_acces_hospitalisation', id=id))
        flash('Vous n\'etes pas assigne a cette hospitalisation', 'danger')
        return redirect(url_for('dashboard'))

    # ============================================================
    # 1. RÉCUPÉRATION DES DONNÉES ASSOCIÉES
    # ============================================================
    
    medecins = hospitalisation.medecins.filter_by(actif=True).all()
    infirmiers = hospitalisation.infirmiers.filter_by(actif=True).all()
    constantes = hospitalisation.constantes.order_by(ConstanteVitale.date_prise.desc()).limit(50).all()
    evolutions = hospitalisation.evolutions.order_by(EvolutionPatient.date_evolution.desc()).all()
    
    # ============================================================
    # 2. NOTES D'ADMISSION
    # ============================================================
    
    toutes_notes = hospitalisation.notes_admission_list.order_by(
        NoteAdmission.version.asc()
    ).all()
    
    note_active = None
    if hospitalisation.note_admission_active_id:
        note_active = NoteAdmission.query.get(hospitalisation.note_admission_active_id)
    
    if not note_active and toutes_notes:
        for note in reversed(toutes_notes):
            if note.est_verrouillee:
                note_active = note
                break
    
    if not note_active and toutes_notes:
        note_active = toutes_notes[0]
    
    nb_versions = len(toutes_notes)
    a_plusieurs_versions = nb_versions > 1
    
    # ============================================================
    # 3. COMPARAISON DES NOTES
    # ============================================================
    
    comparaison = None
    if a_plusieurs_versions and note_active:
        note_initiale = toutes_notes[0] if toutes_notes else None
        
        if note_initiale and note_initiale.id != note_active.id:
            diag_initial = note_initiale.diagnostic_admission or ''
            diag_actuel = note_active.diagnostic_admission or ''
            diag_change = diag_initial != diag_actuel
            
            trait_initial = note_initiale.traitement_admission or ''
            trait_actuel = note_active.traitement_admission or ''
            trait_change = trait_initial != trait_actuel
            
            comparaison = {
                'note_initiale': note_initiale,
                'note_active': note_active,
                'diag_change': diag_change,
                'trait_change': trait_change,
                'nb_versions': nb_versions
            }
    
    # ============================================================
    # 4. ⭐ PROTOCOLES DISPONIBLES
    # ============================================================
    protocoles_disponibles = ProtocoleSoins.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    # ============================================================
    # 5. ⭐ PROTOCOLE ACTIF
    # ============================================================
    protocole_actif = None
    if hospitalisation.protocole_id:
        try:
            protocole_actif = ProtocoleSoins.query.get(hospitalisation.protocole_id)
        except:
            protocole_actif = None
    
    # ============================================================
    # 6. ⭐ EXAMENS TYPES DISPONIBLES
    # ============================================================
    examens_types = ExamenType.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    # ============================================================
    # 7. ⭐ EXAMENS PRESCRITS (CORRIGÉ - Requête directe)
    # ============================================================
    examens_prescrits = []
    nb_examens_prescrits = 0
    
    try:
        examens_prescrits = ExamenPrescrit.query.filter_by(
            hospitalisation_id=hospitalisation.id
        ).order_by(ExamenPrescrit.date_prescription.desc()).all()
        nb_examens_prescrits = len(examens_prescrits)
    except Exception as e:
        print(f"Erreur chargement examens prescrits: {e}")
        examens_prescrits = []
        nb_examens_prescrits = 0
    
    # ============================================================
    # 8. ⭐ ORDONNANCE
    # ============================================================
    ordonnance_medicaments = []
    if hospitalisation.ordonnance_prescite:
        try:
            ordonnance_medicaments = json.loads(hospitalisation.ordonnance_prescite)
            if not isinstance(ordonnance_medicaments, list):
                ordonnance_medicaments = []
        except:
            ordonnance_medicaments = []
    
    # ============================================================
    # 10. ⭐ EXAMEN PHYSIQUE (pré-rempli depuis la consultation d'origine,
    #     sections modifiées depuis l'admission grisées à l'affichage)
    # ============================================================
    from models import ExamenPhysique
    examen_physique_obj = ExamenPhysique.query.filter_by(hospitalisation_id=hospitalisation.id).first()
    examen_physique_sections = _rendre_sections_examen_physique(examen_physique_obj)

    # ============================================================
    # 11. ⭐ SOINS ADMINISTRÉS (journal de soins — widget partagé avec
    #     patients/detail.html et consultations/detail.html, jusqu'ici
    #     jamais branché sur la page hospitalisation)
    # ============================================================
    from models import ActePose
    soins_poses = ActePose.query.filter_by(
        hospitalisation_id=hospitalisation.id
    ).order_by(ActePose.date_pose.desc()).all()

    # ============================================================
    # 12. ⭐ VISITE INFIRMIÈRE DU JOUR + CONSIGNE MÉDICALE ACTIVE
    # ============================================================
    from models import VisiteInfirmiere, ConsigneMedicale, DemandeAccesHospitalisation
    visites_infirmieres_recentes = hospitalisation.visites_infirmieres.order_by(
        VisiteInfirmiere.date_visite.desc()
    ).limit(3).all()
    consigne_active = hospitalisation.consignes_medicales.order_by(
        ConsigneMedicale.date_consigne.desc()
    ).first()

    # ============================================================
    # 13. ⭐ DEMANDE D'ACCÈS INTER-SERVICE — pour un médecin non assigné,
    #     savoir s'il a déjà une demande en cours (pour ne pas en proposer
    #     une nouvelle) ; il ne peut de toute façon pas arriver jusqu'ici
    #     sans accès déjà accordé (voir _a_acces_hospitalisation), donc ce
    #     bloc ne sert qu'aux admin/medecin déjà légitimement sur la page.
    # ============================================================
    demande_acces_en_cours = None
    if current_user.role == 'medecin':
        demande_acces_en_cours = DemandeAccesHospitalisation.query.filter_by(
            hospitalisation_id=hospitalisation.id,
            demandeur_id=current_user.id,
            statut='en_attente'
        ).first()

    # ============================================================
    # 14. ⭐ ÉQUIPE SOIGNANTE — médecins/infirmiers de la structure pas
    #     encore activement assignés à CETTE hospitalisation, pour les
    #     sélecteurs "ajouter" (voir equipe_soignante_ajouter/retirer).
    # ============================================================
    from models import Utilisateur
    medecins_assignes_ids = [m.medecin_id for m in medecins] or [0]
    infirmiers_assignes_ids = [i.infirmier_id for i in infirmiers] or [0]
    medecins_disponibles = Utilisateur.query.filter_by(
        id_structure=hospitalisation.patient.id_structure, role='medecin', actif=True
    ).filter(~Utilisateur.id.in_(medecins_assignes_ids)).order_by(Utilisateur.nom).all()
    infirmiers_disponibles = Utilisateur.query.filter_by(
        id_structure=hospitalisation.patient.id_structure, role='infirmier', actif=True
    ).filter(~Utilisateur.id.in_(infirmiers_assignes_ids)).order_by(Utilisateur.nom).all()

    # ============================================================
    # 9. RENDU
    # ============================================================

    return render_template('hospitalisations/detail.html',
                         hospitalisation=hospitalisation,
                         patient=hospitalisation.patient,
                         medecins=medecins,
                         infirmiers=infirmiers,
                         constantes=constantes,
                         evolutions=evolutions,
                         notes=toutes_notes,
                         note_active=note_active,
                         nb_versions=nb_versions,
                         a_plusieurs_versions=a_plusieurs_versions,
                         comparaison=comparaison,
                         protocoles_disponibles=protocoles_disponibles,
                         protocole_actif=protocole_actif,
                         examens_types=examens_types,
                         examens_prescrits=examens_prescrits,
                         nb_examens_prescrits=nb_examens_prescrits,
                         ordonnance_medicaments=ordonnance_medicaments,
                         examen_physique_sections=examen_physique_sections,
                         soins_poses=soins_poses,
                         actes_soins_habituels=ACTES_SOINS_HABITUELS,
                         visites_infirmieres_recentes=visites_infirmieres_recentes,
                         consigne_active=consigne_active,
                         demande_acces_en_cours=demande_acces_en_cours,
                         medecins_disponibles=medecins_disponibles,
                         infirmiers_disponibles=infirmiers_disponibles,
                         now=datetime.utcnow())


# ==================== ÉQUIPE SOIGNANTE (AJOUT/RETRAIT) ====================
# ⭐ Jusqu'ici l'équipe (médecins/infirmiers) n'était fixée qu'à la création
# de l'hospitalisation (nouvelle_hospitalisation), sans aucun moyen de la
# modifier ensuite — ni pour remplacer un médecin traitant en cours de
# séjour (rotation, congé...), ni pour ajouter un infirmier. Ces deux
# routes le permettent, réservées à l'admin et au médecin déjà assigné
# (voir peut_gerer_equipe, hospitalisations/detail.html).

def _peut_gerer_equipe_soignante(user, hospitalisation):
    if hospitalisation.statut != 'actif':
        return False
    if user.role in ('admin_structure', 'secretaire'):
        return True
    if user.role == 'medecin':
        from models import HospitalisationMedecin
        return HospitalisationMedecin.query.filter_by(
            hospitalisation_id=hospitalisation.id, medecin_id=user.id, actif=True
        ).first() is not None
    return False


@app.route('/hospitalisation/<int:id>/equipe/ajouter', methods=['POST'])
@login_required
def equipe_soignante_ajouter(id):
    from models import Hospitalisation, HospitalisationMedecin, HospitalisationInfirmier, Utilisateur

    hospitalisation = Hospitalisation.query.get_or_404(id)
    if not _peut_gerer_equipe_soignante(current_user, hospitalisation):
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))

    type_membre = request.form.get('type')
    user_id = request.form.get('user_id', type=int)
    membre = Utilisateur.query.get(user_id) if user_id else None
    if not membre or membre.id_structure != hospitalisation.patient.id_structure:
        flash('Utilisateur invalide.', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))

    if type_membre == 'medecin' and membre.role == 'medecin':
        existant = HospitalisationMedecin.query.filter_by(hospitalisation_id=id, medecin_id=membre.id).first()
        if existant:
            existant.actif = True
            existant.date_assignation = datetime.utcnow()
        else:
            db.session.add(HospitalisationMedecin(
                hospitalisation_id=id, medecin_id=membre.id, role='medecin_traitant',
                date_assignation=datetime.utcnow(), actif=True
            ))
        db.session.commit()
        flash(f'✅ Dr {membre.prenom} {membre.nom} ajouté à l\'équipe.', 'success')
    elif type_membre == 'infirmier' and membre.role == 'infirmier':
        existant = HospitalisationInfirmier.query.filter_by(hospitalisation_id=id, infirmier_id=membre.id).first()
        if existant:
            existant.actif = True
            existant.date_assignation = datetime.utcnow()
        else:
            db.session.add(HospitalisationInfirmier(
                hospitalisation_id=id, infirmier_id=membre.id,
                date_assignation=datetime.utcnow(), actif=True
            ))
        db.session.commit()
        flash(f'✅ {membre.prenom} {membre.nom} ajouté(e) à l\'équipe.', 'success')
    else:
        flash('Type ou rôle invalide.', 'danger')

    return redirect(url_for('detail_hospitalisation', id=id))


@app.route('/hospitalisation/<int:id>/equipe/retirer', methods=['POST'])
@login_required
def equipe_soignante_retirer(id):
    from models import Hospitalisation, HospitalisationMedecin, HospitalisationInfirmier

    hospitalisation = Hospitalisation.query.get_or_404(id)
    if not _peut_gerer_equipe_soignante(current_user, hospitalisation):
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))

    type_membre = request.form.get('type')
    user_id = request.form.get('user_id', type=int)

    if type_membre == 'medecin':
        nb_actifs = HospitalisationMedecin.query.filter_by(hospitalisation_id=id, actif=True).count()
        if nb_actifs <= 1:
            flash('Impossible de retirer le dernier médecin assigné — assignez-en un autre d\'abord.', 'danger')
            return redirect(url_for('detail_hospitalisation', id=id))
        ligne = HospitalisationMedecin.query.filter_by(hospitalisation_id=id, medecin_id=user_id, actif=True).first()
        if ligne:
            ligne.actif = False
            db.session.commit()
            flash('Médecin retiré de l\'équipe.', 'info')
    elif type_membre == 'infirmier':
        ligne = HospitalisationInfirmier.query.filter_by(hospitalisation_id=id, infirmier_id=user_id, actif=True).first()
        if ligne:
            ligne.actif = False
            db.session.commit()
            flash('Infirmier retiré de l\'équipe.', 'info')

    return redirect(url_for('detail_hospitalisation', id=id))


# ============================================================
# AJOUTER UNE NOUVELLE NOTE D'ADMISSION (RÉÉVALUATION)
# ============================================================
@app.route('/hospitalisation/<int:id>/note/ajouter', methods=['POST'])
@login_required
def ajouter_note_admission(id):
    """
    Ajouter une nouvelle version de la note d'admission
    (Réévaluation du patient par le médecin)
    """
    from models import Hospitalisation, NoteAdmission, HospitalisationMedecin
    from datetime import datetime

    hospitalisation = Hospitalisation.query.get_or_404(id)

    # ============================================================
    # 1. VÉRIFICATION DES PERMISSIONS
    # ============================================================

    # Seuls les médecins et admins peuvent ajouter une note
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Seuls les médecins peuvent ajouter une note de réévaluation.', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))

    # Vérifier que l'hospitalisation est active
    if hospitalisation.statut != 'actif':
        flash('Impossible d\'ajouter une note sur une hospitalisation clôturée.', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))

    # ⭐ Vérifier que le médecin est assigné — SAUF si l'hospitalisation n'a
    # encore AUCUN médecin assigné (cas d'une admission par un infirmier
    # sans médecin sélectionné, voir nouvelle_hospitalisation) : dans ce
    # cas, n'importe quel médecin de la structure peut rédiger la première
    # note, ce qui l'assigne du même coup comme médecin traitant — sinon
    # personne ne pourrait jamais la compléter.
    if current_user.role == 'medecin':
        aucun_medecin_assigne = not HospitalisationMedecin.query.filter_by(
            hospitalisation_id=id, actif=True
        ).first()
        if aucun_medecin_assigne:
            db.session.add(HospitalisationMedecin(
                hospitalisation_id=id,
                medecin_id=current_user.id,
                role='medecin_traitant',
                date_assignation=datetime.utcnow(),
                actif=True
            ))
        else:
            assigne = HospitalisationMedecin.query.filter_by(
                hospitalisation_id=id,
                medecin_id=current_user.id,
                actif=True
            ).first()
            if not assigne:
                flash('Vous n\'êtes pas assigné à cette hospitalisation.', 'danger')
                return redirect(url_for('detail_hospitalisation', id=id))
    
    # ============================================================
    # 2. RÉCUPÉRATION DES DONNÉES DU FORMULAIRE
    # ============================================================
    
    note_motif = request.form.get('note_motif', '').strip()
    note_contexte = request.form.get('note_contexte', '').strip()
    note_examen_clinique = request.form.get('note_examen_clinique', '').strip()
    note_diagnostic = request.form.get('note_diagnostic', '').strip()
    note_examens = request.form.get('note_examens', '').strip()
    note_traitement = request.form.get('note_traitement', '').strip()
    note_evolution_prevue = request.form.get('note_evolution_prevue', '').strip()
    note_conclusion = request.form.get('note_conclusion', '').strip()
    note_constantes = request.form.get('note_constantes', '').strip()
    
    # ============================================================
    # 3. VALIDATION
    # ============================================================
    
    champs_obligatoires = {
        'Motif d\'hospitalisation': note_motif,
        'Contexte et antécédents': note_contexte,
        'Examen clinique initial': note_examen_clinique,
        'Diagnostic présumé': note_diagnostic,
        'Traitement initial': note_traitement,
        'Conclusion du médecin': note_conclusion
    }
    
    champs_manquants = []
    for nom, valeur in champs_obligatoires.items():
        if not valeur:
            champs_manquants.append(nom)
    
    if champs_manquants:
        flash(f'La note est incomplète. Champs obligatoires : {", ".join(champs_manquants)}', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    # ============================================================
    # 4. CALCUL DU NUMÉRO DE VERSION
    # ============================================================
    
    notes_existantes = hospitalisation.notes_admission_list.count()
    nouvelle_version = notes_existantes + 1
    
    # Récupérer la note précédente pour référence
    note_precedente = None
    if hospitalisation.note_admission_active_id:
        note_precedente = NoteAdmission.query.get(hospitalisation.note_admission_active_id)
    
    # ============================================================
    # 5. CRÉATION DE LA NOUVELLE NOTE
    # ============================================================
    
    try:
        note = NoteAdmission(
            hospitalisation_id=hospitalisation.id,
            version=nouvelle_version,
            # ⭐ Vraiment initiale seulement s'il n'existe encore aucune
            # note — cas d'une hospitalisation admise par un infirmier
            # (voir nouvelle_hospitalisation), où cette route sert à
            # rédiger la toute première note plutôt qu'une réévaluation.
            est_initial=(notes_existantes == 0),
            est_verrouillee=True,  # Verrouillée immédiatement
            motif_admission=note_motif,
            contexte_admission=note_contexte,
            examen_clinique_admission=note_examen_clinique,
            diagnostic_admission=note_diagnostic,
            examens_admission=note_examens if note_examens else None,
            traitement_admission=note_traitement,
            evolution_prevue=note_evolution_prevue if note_evolution_prevue else None,
            conclusion_admission=note_conclusion,
            constantes_admission=note_constantes if note_constantes else None,
            redige_par=current_user.id,
            date_redaction=datetime.utcnow(),
            valide_par=current_user.id,
            date_validation=datetime.utcnow()
        )
        
        db.session.add(note)
        db.session.flush()
        
        # ============================================================
        # 6. METTRE À JOUR LA NOTE ACTIVE
        # ============================================================
        
        hospitalisation.note_admission_active_id = note.id
        # ⭐ Lève l'alerte "à compléter" si c'était une admission infirmière
        # sans note — voir nouvelle_hospitalisation.
        hospitalisation.note_admission_a_completer = False

        # ============================================================
        # 7. LOG DANS LES NOTES CLINIQUES (optionnel)
        # ============================================================
        
        # Ajouter une entrée dans les notes d'admission (texte libre) pour traçabilité
        ancienne_note = hospitalisation.notes_admission or ''
        nouvelle_entree = f"""
--- RÉÉVALUATION v{ nouvelle_version} - {datetime.utcnow().strftime('%d/%m/%Y %H:%M')} ---
Médecin : Dr {current_user.prenom} {current_user.nom}
Motif : {note_motif[:100]}...
Diagnostic : {note_diagnostic[:100]}...
Traitement : {note_traitement[:100]}...
Conclusion : {note_conclusion[:100]}...
"""
        hospitalisation.notes_admission = (ancienne_note + nouvelle_entree) if ancienne_note else nouvelle_entree
        
        # ============================================================
        # 8. COMMIT FINAL
        # ============================================================
        
        db.session.commit()
        
        flash(f'✅ Nouvelle note de réévaluation (version {nouvelle_version}) enregistrée avec succès !', 'success')
        
        if note_precedente:
            flash(f'📋 Cette note remplace la version {note_precedente.version} comme référence active.', 'info')
        else:
            flash('📋 Cette note est désormais la référence active pour le suivi.', 'info')
        
        return redirect(url_for('detail_hospitalisation', id=id))
        
    except Exception as e:
        db.session.rollback()
        flash(f'❌ Erreur lors de l\'enregistrement : {str(e)}', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))


# ============================================================
# COMPARER LES NOTES D'ADMISSION
# ============================================================
@app.route('/hospitalisation/<int:id>/notes/comparer')
@login_required
def comparer_notes_admission(id):
    """
    Comparer les différentes versions de la note d'admission
    Affiche un tableau comparatif des versions
    """
    from models import Hospitalisation, NoteAdmission, HospitalisationMedecin
    from datetime import datetime
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    
    # ============================================================
    # 1. VÉRIFICATION DES PERMISSIONS
    # ============================================================
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé.', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    # Vérifier que le médecin est assigné à cette hospitalisation
    if current_user.role == 'medecin':
        assigne = HospitalisationMedecin.query.filter_by(
            hospitalisation_id=id,
            medecin_id=current_user.id,
            actif=True
        ).first()
        if not assigne:
            flash('Vous n\'êtes pas assigné à cette hospitalisation.', 'danger')
            return redirect(url_for('detail_hospitalisation', id=id))
    
    # ============================================================
    # 2. RÉCUPÉRATION DES NOTES
    # ============================================================
    
    toutes_notes = hospitalisation.notes_admission_list.order_by(
        NoteAdmission.version.asc()
    ).all()
    
    if len(toutes_notes) < 2:
        flash('Il n\'y a pas assez de notes pour faire une comparaison (minimum 2).', 'warning')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    # ============================================================
    # 3. ANALYSE DES DIFFÉRENCES
    # ============================================================
    
    # Préparer les données pour le tableau comparatif
    notes_data = []
    for note in toutes_notes:
        # Récupérer le nom du médecin
        medecin_nom = "Inconnu"
        if note.redacteur:
            medecin_nom = f"Dr {note.redacteur.prenom} {note.redacteur.nom}"
        
        notes_data.append({
            'id': note.id,
            'version': note.version,
            'est_initial': note.est_initial,
            'est_verrouillee': note.est_verrouillee,
            'date_redaction': note.date_redaction,
            'redacteur': medecin_nom,
            'motif': note.motif_admission or '-',
            'contexte': note.contexte_admission or '-',
            'examen_clinique': note.examen_clinique_admission or '-',
            'diagnostic': note.diagnostic_admission or '-',
            'examens': note.examens_admission or '-',
            'traitement': note.traitement_admission or '-',
            'evolution_prevue': note.evolution_prevue or '-',
            'conclusion': note.conclusion_admission or '-',
            'constantes': note.constantes_admission or '-'
        })
    
    # ============================================================
    # 4. DÉTECTION DES CHANGEMENTS MAJEURS
    # ============================================================
    
    changements = []
    
    # Comparer chaque note avec la précédente
    for i in range(1, len(notes_data)):
        note_prec = notes_data[i-1]
        note_act = notes_data[i]
        
        changements_note = {
            'version': note_act['version'],
            'changements': []
        }
        
        # Vérifier les changements dans les champs clés
        if note_prec['diagnostic'] != note_act['diagnostic']:
            changements_note['changements'].append({
                'champ': 'Diagnostic',
                'ancien': note_prec['diagnostic'],
                'nouveau': note_act['diagnostic']
            })
        
        if note_prec['traitement'] != note_act['traitement']:
            changements_note['changements'].append({
                'champ': 'Traitement',
                'ancien': note_prec['traitement'],
                'nouveau': note_act['traitement']
            })
        
        if note_prec['examen_clinique'] != note_act['examen_clinique']:
            changements_note['changements'].append({
                'champ': 'Examen clinique',
                'ancien': note_prec['examen_clinique'],
                'nouveau': note_act['examen_clinique']
            })
        
        if note_prec['conclusion'] != note_act['conclusion']:
            changements_note['changements'].append({
                'champ': 'Conclusion',
                'ancien': note_prec['conclusion'],
                'nouveau': note_act['conclusion']
            })
        
        if changements_note['changements']:
            changements.append(changements_note)
    
    # ============================================================
    # 5. RENDU DU TEMPLATE
    # ============================================================
    
    return render_template('hospitalisations/comparaison_notes.html',
                         hospitalisation=hospitalisation,
                         notes_data=notes_data,
                         changements=changements,
                         nb_notes=len(toutes_notes),
                         now=datetime.utcnow())


@app.route('/hospitalisation/<int:id>/evolution', methods=['GET', 'POST'])
@login_required
def ajouter_evolution(id):
    """
    Ajouter une évolution pour un patient hospitalisé
    Avec affichage de la note d'admission en référence
    """
    from models import Hospitalisation, HospitalisationMedecin, HospitalisationInfirmier, EvolutionPatient, NoteAdmission
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    
    # ============================================================
    # 1. VÉRIFICATION DES PERMISSIONS
    # ============================================================
    
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    if not _a_acces_hospitalisation(current_user, hospitalisation):
        flash('Vous n\'êtes pas assigné à cette hospitalisation', 'danger')
        return redirect(url_for('dashboard'))

    # ============================================================
    # 2. RÉCUPÉRATION DE LA NOTE D'ADMISSION ACTIVE (RÉFÉRENCE)
    # ============================================================

    note_active = None
    if hospitalisation.note_admission_active_id:
        note_active = NoteAdmission.query.get(hospitalisation.note_admission_active_id)
    
    # Si pas de note active, prendre la dernière verrouillée
    if not note_active:
        toutes_notes = hospitalisation.notes_admission_list.order_by(
            NoteAdmission.version.desc()
        ).all()
        for note in toutes_notes:
            if note.est_verrouillee:
                note_active = note
                break
    
    # Si toujours pas, prendre la première
    if not note_active:
        note_active = hospitalisation.notes_admission_list.first()
    
    # ============================================================
    # 3. TRAITEMENT DU POST
    # ============================================================
    
    if request.method == 'POST':
        etat_echelle = request.form.get('etat_echelle', type=int)
        temperature = request.form.get('temperature', type=float)
        pression = request.form.get('pression')
        fc = request.form.get('fc', type=int)
        symptomes = request.form.get('symptomes')
        traitement_administre = request.form.get('traitement_administre')
        observations = request.form.get('observations')
        prochaines_etapes = request.form.get('prochaines_etapes')
        
        # ⭐ NOUVEAU : Évolution par rapport à l'admission
        evolution_par_rapport = request.form.get('evolution_par_rapport')
        
        if etat_echelle is None or etat_echelle < 0 or etat_echelle > 10:
            flash('L\'état doit être entre 0 et 10', 'danger')
            return redirect(url_for('ajouter_evolution', id=id))
        
        try:
            evolution = EvolutionPatient(
                hospitalisation_id=id,
                etat_echelle=etat_echelle,
                temperature=temperature,
                pression=pression,
                fc=fc,
                symptomes=symptomes,
                traitement_administre=traitement_administre,
                observations=observations,
                prochaines_etapes=prochaines_etapes,
                evolution_par_rapport=evolution_par_rapport,  # ⭐ NOUVEAU
                redige_par=current_user.id,
                date_evolution=datetime.utcnow()
            )
            db.session.add(evolution)
            db.session.commit()
            
            # ⭐ Message avec rappel de la note de référence
            flash('✅ Évolution enregistrée avec succès', 'success')
            if note_active:
                flash(f'📋 Référence : Note d\'admission v{note_active.version} du {note_active.date_redaction.strftime("%d/%m/%Y")}', 'info')
            
            return redirect(url_for('detail_hospitalisation', id=id))
            
        except Exception as e:
            db.session.rollback()
            flash(f'❌ Erreur : {str(e)}', 'danger')
            return redirect(url_for('ajouter_evolution', id=id))
    
    # ============================================================
    # 4. AFFICHAGE DU FORMULAIRE (GET)
    # ============================================================
    
    # Récupérer les évolutions précédentes pour référence
    evolutions_precedentes = hospitalisation.evolutions.order_by(
        EvolutionPatient.date_evolution.desc()
    ).limit(5).all()
    
    # Calculer la tendance
    tendance = None
    if len(evolutions_precedentes) >= 2:
        dernier_etat = evolutions_precedentes[0].etat_echelle if evolutions_precedentes else None
        avant_dernier = evolutions_precedentes[1].etat_echelle if len(evolutions_precedentes) > 1 else None
        
        if dernier_etat is not None and avant_dernier is not None:
            if dernier_etat > avant_dernier:
                tendance = 'amelioration'
            elif dernier_etat < avant_dernier:
                tendance = 'aggravation'
            else:
                tendance = 'stable'
    
    return render_template('hospitalisations/evolution.html',
                         hospitalisation=hospitalisation,
                         note_active=note_active,
                         evolutions_precedentes=evolutions_precedentes,
                         tendance=tendance,
                         now=datetime.utcnow())


# ==================== VISITE INFIRMIÈRE DU JOUR ====================
# ⭐ Pendant du suivi médecin (EvolutionPatient/ajouter_evolution) côté
# infirmier, mais avec un vrai examen physique système par système (comme
# en consultation) et une proposition de décision explicitement INDICATIVE
# — seule ConsigneMedicale (ci-dessous) est appliquée. Voir plan :
# "l'infirmier fait la visite du jour, il peut mettre une décision mais
# c'est juste indicationnel, c'est pour le médecin qui sera appliqué".

@app.route('/hospitalisation/<int:id>/visite-infirmiere/ajouter')
@login_required
def nouvelle_visite_infirmiere(id):
    from models import Hospitalisation, ExamenPhysique, VisiteInfirmiere

    hospitalisation = Hospitalisation.query.get_or_404(id)
    patient = hospitalisation.patient

    if current_user.role not in ('admin_structure', 'infirmier'):
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    if not _a_acces_hospitalisation(current_user, hospitalisation):
        flash('Vous n\'êtes pas assigné à cette hospitalisation', 'danger')
        return redirect(url_for('dashboard'))

    # ⭐ Formulaire pré-rempli avec l'examen physique COURANT de
    # l'hospitalisation (dossier officiel tenu par le médecin) — jamais
    # persisté ici (voir enregistrer_visite_infirmiere pour la création
    # réelle). sections_origine = référence pour le surlignage de ce que
    # L'INFIRMIER changera pendant SA visite.
    examen_actuel = ExamenPhysique.query.filter_by(hospitalisation_id=id).first()
    reference = examen_actuel.sections_modifiees if examen_actuel else '{}'
    visite = VisiteInfirmiere(
        sections_modifiees=reference,
        sections_origine=reference or '{}',
    )

    return render_template('hospitalisations/visite_infirmiere_form.html',
                         patient=patient,
                         hospitalisation=hospitalisation,
                         visite=visite)


@app.route('/hospitalisation/<int:id>/visite-infirmiere/enregistrer', methods=['POST'])
@login_required
def enregistrer_visite_infirmiere(id):
    from models import Hospitalisation, ExamenPhysique, VisiteInfirmiere

    hospitalisation = Hospitalisation.query.get_or_404(id)

    if current_user.role not in ('admin_structure', 'infirmier'):
        return jsonify({'success': False, 'message': 'Accès non autorisé'}), 403
    if not _a_acces_hospitalisation(current_user, hospitalisation):
        return jsonify({'success': False, 'message': 'Vous n\'êtes pas assigné à cette hospitalisation'}), 403

    try:
        examen_actuel = ExamenPhysique.query.filter_by(hospitalisation_id=id).first()
        sections_origine = examen_actuel.sections_modifiees if examen_actuel else '{}'

        visite = VisiteInfirmiere(
            hospitalisation_id=id,
            infirmier_id=current_user.id,
            plaintes_patient=request.form.get('plaintes_patient'),
            etat_general=request.form.get('etat_general'),
            examen_complet=nettoyer_examen_complet(request.form.get('examen_complet', '')),
            sections_modifiees=request.form.get('sections_modifiees', '{}'),
            sections_origine=sections_origine or '{}',
            appareil_dysfonctionnel=bool(request.form.get('appareil_dysfonctionnel')),
            appareil_dysfonctionnel_detail=request.form.get('appareil_dysfonctionnel_detail'),
            decision_suggeree=request.form.get('decision_suggeree'),
        )
        db.session.add(visite)
        db.session.commit()

        return jsonify({
            'success': True,
            'message': 'Visite enregistrée',
            'redirect': url_for('detail_hospitalisation', id=id),
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500


@app.route('/hospitalisation/<int:id>/visites-infirmieres')
@login_required
def visites_infirmieres_liste(id):
    from models import Hospitalisation, VisiteInfirmiere

    hospitalisation = Hospitalisation.query.get_or_404(id)

    if not _a_acces_hospitalisation(current_user, hospitalisation):
        flash('Vous n\'êtes pas assigné à cette hospitalisation', 'danger')
        return redirect(url_for('dashboard'))

    visites = hospitalisation.visites_infirmieres.order_by(VisiteInfirmiere.date_visite.desc()).all()

    return render_template('hospitalisations/visites_infirmieres_liste.html',
                         hospitalisation=hospitalisation,
                         patient=hospitalisation.patient,
                         visites=visites)


# ==================== CONSIGNE MÉDICALE ====================
# ⭐ La seule chose que l'infirmier doit réellement appliquer — distincte
# de VisiteInfirmiere.decision_suggeree (indicative, voir ci-dessus).

@app.route('/hospitalisation/<int:id>/consigne/ajouter', methods=['GET', 'POST'])
@login_required
def ajouter_consigne(id):
    from models import Hospitalisation, ConsigneMedicale, VisiteInfirmiere

    hospitalisation = Hospitalisation.query.get_or_404(id)
    patient = hospitalisation.patient

    if current_user.role not in ('admin_structure', 'medecin'):
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    if not _a_acces_hospitalisation(current_user, hospitalisation):
        flash('Vous n\'êtes pas assigné à cette hospitalisation', 'danger')
        return redirect(url_for('dashboard'))

    if request.method == 'POST':
        type_decision = request.form.get('type_decision')
        instructions = request.form.get('instructions', '').strip()
        visite_infirmiere_id = request.form.get('visite_infirmiere_id', type=int)

        if not type_decision or not instructions:
            flash('Merci de préciser la décision et les instructions', 'danger')
            return redirect(url_for('ajouter_consigne', id=id, visite_infirmiere_id=visite_infirmiere_id))

        consigne = ConsigneMedicale(
            hospitalisation_id=id,
            medecin_id=current_user.id,
            type_decision=type_decision,
            instructions=instructions,
            visite_infirmiere_id=visite_infirmiere_id,
        )
        db.session.add(consigne)
        db.session.commit()

        flash('✅ Consigne enregistrée — visible côté infirmier', 'success')
        return redirect(url_for('detail_hospitalisation', id=id))

    visite_ref = None
    visite_ref_id = request.args.get('visite_infirmiere_id', type=int)
    if visite_ref_id:
        visite_ref = VisiteInfirmiere.query.get(visite_ref_id)

    return render_template('hospitalisations/consigne_form.html',
                         patient=patient,
                         hospitalisation=hospitalisation,
                         visite_ref=visite_ref)


# ==================== DEMANDE D'ACCÈS INTER-SERVICE ====================
# ⭐ Un médecin non assigné à une hospitalisation (patient d'un autre
# service) peut demander l'accès — routée automatiquement au médecin
# traitant s'il y en a un, sinon à l'administrateur de la structure.

def _resoudre_destinataire_demande_acces(hospitalisation, demandeur):
    """Retourne (destinataire_utilisateur_ou_None, libelle_lisible). None =
    routée vers l'administrateur de la structure du patient (pas de
    médecin traitant assigné, ou seulement le demandeur lui-même)."""
    from models import HospitalisationMedecin
    assigne = HospitalisationMedecin.query.filter_by(
        hospitalisation_id=hospitalisation.id,
        role='medecin_traitant',
        actif=True
    ).filter(HospitalisationMedecin.medecin_id != demandeur.id).first()
    if assigne:
        return assigne.medecin, f"Dr {assigne.medecin.prenom} {assigne.medecin.nom} (médecin traitant)"
    return None, "l'administrateur de la structure"


@app.route('/hospitalisation/<int:id>/demande-acces', methods=['GET', 'POST'])
@login_required
def demander_acces_hospitalisation(id):
    from models import Hospitalisation, DemandeAccesHospitalisation, Message, Utilisateur

    hospitalisation = Hospitalisation.query.get_or_404(id)
    patient = hospitalisation.patient

    if current_user.role != 'medecin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    if current_user.id_structure != patient.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    # Déjà accès (assigné ou demande acceptée encore valide) → rien à faire ici.
    if _a_acces_hospitalisation(current_user, hospitalisation):
        return redirect(url_for('detail_hospitalisation', id=id))

    demande_existante = DemandeAccesHospitalisation.query.filter_by(
        hospitalisation_id=id,
        demandeur_id=current_user.id,
        statut='en_attente'
    ).first()

    destinataire, destinataire_label = _resoudre_destinataire_demande_acces(hospitalisation, current_user)

    if request.method == 'POST':
        if demande_existante:
            flash('Une demande est déjà en attente pour ce patient.', 'warning')
            return redirect(url_for('liste_hospitalisations'))

        motif = request.form.get('motif', '').strip()
        if not motif:
            flash('Merci de préciser le motif de la demande', 'danger')
            return redirect(url_for('demander_acces_hospitalisation', id=id))

        demande = DemandeAccesHospitalisation(
            hospitalisation_id=id,
            demandeur_id=current_user.id,
            destinataire_id=destinataire.id if destinataire else None,
            motif=motif,
        )
        db.session.add(demande)

        # ⭐ Notification — réutilise la messagerie interne existante,
        # pas de nouveau canal. Si routée admin (destinataire None), tous
        # les admin_structure de la structure sont notifiés.
        cible_messages = [destinataire] if destinataire else Utilisateur.query.filter_by(
            id_structure=patient.id_structure, role='admin_structure', actif=True
        ).all()
        for cible in cible_messages:
            db.session.add(Message(
                id_expediteur=current_user.id,
                id_destinataire=cible.id,
                id_structure=patient.id_structure,
                sujet=f"Demande d'accès — {patient.prenom} {patient.nom}",
                contenu=f"Dr {current_user.prenom} {current_user.nom} demande l'accès au dossier de {patient.prenom} {patient.nom} ({hospitalisation.service}).\nMotif : {motif}\n\nVoir : /demandes-acces",
            ))

        db.session.commit()
        flash(f'✅ Demande envoyée à {destinataire_label}.', 'success')
        return redirect(url_for('liste_hospitalisations'))

    return render_template('hospitalisations/demande_acces_form.html',
                         patient=patient,
                         hospitalisation=hospitalisation,
                         destinataire_label=destinataire_label,
                         demande_existante=demande_existante)


@app.route('/demandes-acces')
@login_required
def demandes_acces_inbox():
    from models import DemandeAccesHospitalisation, Hospitalisation, Patient

    if current_user.role not in ('medecin', 'admin_structure'):
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    if current_user.role == 'medecin':
        query = DemandeAccesHospitalisation.query.filter_by(destinataire_id=current_user.id)
    else:
        query = DemandeAccesHospitalisation.query.filter(
            DemandeAccesHospitalisation.destinataire_id.is_(None)
        ).join(Hospitalisation).join(Patient).filter(
            Patient.id_structure == current_user.id_structure
        )

    toutes = query.order_by(DemandeAccesHospitalisation.created_at.desc()).all()
    en_attente = [d for d in toutes if d.statut == 'en_attente']
    traitees = [d for d in toutes if d.statut != 'en_attente']

    return render_template('hospitalisations/demandes_acces_liste.html',
                         en_attente=en_attente,
                         traitees=traitees)


@app.route('/demande-acces/<int:id>/traiter', methods=['POST'])
@login_required
def traiter_demande_acces(id):
    from models import DemandeAccesHospitalisation, Message
    from datetime import timedelta

    demande = DemandeAccesHospitalisation.query.get_or_404(id)

    autorise = (
        (demande.destinataire_id and demande.destinataire_id == current_user.id) or
        (not demande.destinataire_id and current_user.role == 'admin_structure' and
         current_user.id_structure == demande.hospitalisation.patient.id_structure) or
        current_user.role == 'super_admin'
    )
    if not autorise:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('demandes_acces_inbox'))

    if demande.statut != 'en_attente':
        flash('Cette demande a déjà été traitée.', 'warning')
        return redirect(url_for('demandes_acces_inbox'))

    action = request.form.get('action')
    patient = demande.hospitalisation.patient

    if action == 'accepter':
        duree_heures = request.form.get('duree_heures', 24, type=int)
        demande.statut = 'acceptee'
        demande.date_debut = datetime.utcnow()
        demande.date_fin = datetime.utcnow() + timedelta(hours=duree_heures)
        demande.traite_par = current_user.id
        demande.date_traitement = datetime.utcnow()
        message_contenu = f"Votre demande d'accès au dossier de {patient.prenom} {patient.nom} a été acceptée pour {duree_heures}h."
        flash_msg = ('✅ Accès accordé', 'success')
    elif action == 'refuser':
        motif_refus = request.form.get('motif_refus', '').strip()
        if not motif_refus:
            flash('Merci de préciser un motif de refus', 'danger')
            return redirect(url_for('demandes_acces_inbox'))
        demande.statut = 'refusee'
        demande.motif_refus = motif_refus
        demande.traite_par = current_user.id
        demande.date_traitement = datetime.utcnow()
        message_contenu = f"Votre demande d'accès au dossier de {patient.prenom} {patient.nom} a été refusée.\nMotif : {motif_refus}"
        flash_msg = ('Demande refusée', 'info')
    else:
        flash('Action inconnue', 'danger')
        return redirect(url_for('demandes_acces_inbox'))

    db.session.add(Message(
        id_expediteur=current_user.id,
        id_destinataire=demande.demandeur_id,
        id_structure=patient.id_structure,
        sujet=f"Demande d'accès — {patient.prenom} {patient.nom}",
        contenu=message_contenu,
    ))
    db.session.commit()

    flash(*flash_msg)
    return redirect(url_for('demandes_acces_inbox'))


@app.route('/hospitalisation/<int:id>/constante', methods=['GET', 'POST'])
@login_required
def ajouter_constante(id):
    """
    Ajouter des constantes vitales pour un patient hospitalisé
    Avec affichage de la note d'admission en référence
    """
    from models import Hospitalisation, HospitalisationInfirmier, HospitalisationMedecin, ConstanteVitale, NoteAdmission
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    
    # ============================================================
    # 1. VÉRIFICATION DES PERMISSIONS
    # ============================================================
    
    if current_user.role not in ['super_admin', 'admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    if current_user.id_structure and hospitalisation.patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_hospitalisations'))

    if not _a_acces_hospitalisation(current_user, hospitalisation):
        flash('Vous n\'êtes pas assigné à cette hospitalisation', 'danger')
        return redirect(url_for('liste_hospitalisations'))

    # ============================================================
    # 2. RÉCUPÉRATION DE LA NOTE D'ADMISSION ACTIVE (RÉFÉRENCE)
    # ============================================================
    
    note_active = None
    if hospitalisation.note_admission_active_id:
        note_active = NoteAdmission.query.get(hospitalisation.note_admission_active_id)
    
    # Si pas de note active, prendre la dernière verrouillée
    if not note_active:
        toutes_notes = hospitalisation.notes_admission_list.order_by(
            NoteAdmission.version.desc()
        ).all()
        for note in toutes_notes:
            if note.est_verrouillee:
                note_active = note
                break
    
    # Si toujours pas, prendre la première
    if not note_active:
        note_active = hospitalisation.notes_admission_list.first()
    
    # ============================================================
    # 3. RÉCUPÉRATION DES DERNIÈRES CONSTANTES POUR COMPARAISON
    # ============================================================
    
    dernieres_constantes = hospitalisation.constantes.order_by(
        ConstanteVitale.date_prise.desc()
    ).first()
    
    # ============================================================
    # 4. TRAITEMENT DU POST
    # ============================================================
    
    if request.method == 'POST':
        temperature = request.form.get('temperature', type=float)
        pression_arterielle = request.form.get('pression_arterielle')
        frequence_cardiaque = request.form.get('frequence_cardiaque', type=int)
        frequence_respiratoire = request.form.get('frequence_respiratoire', type=int)
        saturation_oxygene = request.form.get('saturation_oxygene', type=float)
        glycemie = request.form.get('glycemie', type=float)
        poids = request.form.get('poids', type=float)
        taille = request.form.get('taille', type=float)
        
        # Nouveaux champs
        diurese = request.form.get('diurese')
        emission_gaz = request.form.get('emission_gaz')
        selles = request.form.get('selles')
        vomissements = request.form.get('vomissements')
        douleur = request.form.get('douleur', type=int)
        conscience = request.form.get('conscience')
        pouls_peripherique = request.form.get('pouls_peripherique')
        temperature_cutanee = request.form.get('temperature_cutanee')
        
        autres_constantes = request.form.get('autres_constantes')
        notes = request.form.get('notes')
        
        # ⭐ NOUVEAU : Évolution par rapport à l'admission
        evolution_par_rapport = request.form.get('evolution_par_rapport')
        
        # Calcul de l'IMC
        imc = None
        if poids and taille and taille > 0:
            imc = round(poids / ((taille/100) ** 2), 1)
        
        try:
            constante = ConstanteVitale(
                hospitalisation_id=id,
                infirmier_id=current_user.id,
                temperature=temperature,
                pression_arterielle=pression_arterielle,
                frequence_cardiaque=frequence_cardiaque,
                frequence_respiratoire=frequence_respiratoire,
                saturation_oxygene=saturation_oxygene,
                glycemie=glycemie,
                poids=poids,
                taille=taille,
                imc=imc,
                # Nouveaux champs
                diurese=diurese,
                emission_gaz=emission_gaz,
                selles=selles,
                vomissements=vomissements,
                douleur=douleur,
                conscience=conscience,
                pouls_peripherique=pouls_peripherique,
                temperature_cutanee=temperature_cutanee,
                # ⭐ NOUVEAU
                evolution_par_rapport=evolution_par_rapport,
                autres_constantes=autres_constantes,
                notes=notes,
                date_prise=datetime.utcnow()
            )
            db.session.add(constante)
            db.session.commit()
            
            # ⭐ Message avec rappel de la note de référence
            flash('✅ Constantes vitales enregistrées avec succès', 'success')
            if note_active:
                flash(f'📋 Référence : Note d\'admission v{note_active.version} du {note_active.date_redaction.strftime("%d/%m/%Y")}', 'info')
            
            return redirect(url_for('detail_hospitalisation', id=id))
            
        except Exception as e:
            db.session.rollback()
            flash(f'❌ Erreur : {str(e)}', 'danger')
            return redirect(url_for('ajouter_constante', id=id))
    
    # ============================================================
    # 5. AFFICHAGE DU FORMULAIRE (GET)
    # ============================================================
    
    return render_template('hospitalisations/constante.html',
                         hospitalisation=hospitalisation,
                         note_active=note_active,
                         dernieres_constantes=dernieres_constantes,
                         now=datetime.utcnow())

def _taux_amu_pour_acte(nom_acte, taux_defaut):
    """⭐ FIX : l'acte P160 (hospitalisation) est remboursé par l'AMU à 90%,
    alors que le taux général du patient (souvent 80%) s'applique à tous
    les autres actes — même règle que côté GHP (app.py:taux_amu_pour_article(),
    templates/actes_vente.html:tauxAMUPourArticle()). Les actes de
    facturation d'hospitalisation portent le nom exact du catalogue GHP
    (ex: "P160 Hospi Cabine ventillee..."), donc la même détection par nom
    s'applique ici."""
    return 90 if (nom_acte and 'P160' in nom_acte) else taux_defaut


def _calculer_paliers_hospitalisation(jours_total, structure_id):
    """Répartit un nombre de jours en paliers (semaine 1 / semaine 2 / 15j
    et plus) selon le paramétrage AMU de la structure. Retourne la liste des
    paliers avec un nombre de jours > 0, dans l'ordre."""
    from models import ParametrageAMU
    param = ParametrageAMU.get_ou_defaut(structure_id)
    s1, s2 = param.seuil_jours_semaine1, param.seuil_jours_semaine2

    jours_p1 = min(jours_total, s1)
    jours_p2 = min(max(jours_total - s1, 0), max(s2 - s1, 0))
    jours_p3 = max(jours_total - s2, 0)

    paliers = []
    if jours_p1 > 0:
        paliers.append({'palier': 'semaine1', 'label': f'Semaine 1 (jours 1-{s1})', 'jours': jours_p1})
    if jours_p2 > 0:
        paliers.append({'palier': 'semaine2', 'label': f'Semaine 2 (jours {s1+1}-{s2})', 'jours': jours_p2})
    if jours_p3 > 0:
        paliers.append({'palier': 'semaine3plus', 'label': f'À partir du jour {s2+1}', 'jours': jours_p3})

    return paliers, param


def _salle_hospitalisation(hospitalisation, salle_id_manuel=None):
    """Retrouve la salle occupée pendant l'hospitalisation, pour la
    facturation. Doit être appelé AVANT de libérer le lit à la clôture.

    ⭐ Dans les faits, l'assignation formelle d'un lit (écran "Assigner un
    lit") n'est utilisée que pour une minorité des hospitalisations —
    beaucoup de dossiers ne renseignent qu'un numéro de chambre en texte
    libre (Hospitalisation.chambre), sans lien avec la table Salle. Sans
    repli, la facturation automatique échouait silencieusement pour la
    plupart des clôtures réelles ("aucun lit n'était assigné"). On essaie
    donc, dans l'ordre :
    1. Un choix manuel explicite (sélectionné par l'utilisateur à la
       clôture, quand aucune des méthodes automatiques n'a abouti).
    2. Le lit formellement assigné (le cas fiable).
    3. Une correspondance de nom entre Hospitalisation.chambre et
       Salle.nom, pour la même structure (repli best-effort).
    """
    from models import Lit, Salle, Service

    if salle_id_manuel:
        return Salle.query.get(int(salle_id_manuel))

    if hospitalisation.lit_id:
        lit = Lit.query.get(hospitalisation.lit_id)
        if lit and lit.salle:
            return lit.salle

    if hospitalisation.chambre:
        chambre = hospitalisation.chambre.strip().lower()
        salle = Salle.query.join(Service).filter(
            Service.structure_id == hospitalisation.patient.id_structure,
            db.func.lower(Salle.nom) == chambre
        ).first()
        if salle:
            return salle

    return None


def _tarifs_ghp_structure(structure_id):
    """Récupère en direct (par nom) le catalogue d'actes GHP de la
    structure — utilisé pour estimer le montant avant envoi. Retourne un
    dict {nom_en_minuscule: {'nom', 'prix', 'pbr'}}, vide si indisponible
    (ce n'est qu'une estimation, GHP recalcule de toute façon à la vente)."""
    from models import StructureMapping
    import requests

    mapping = StructureMapping.query.filter_by(local_structure_id=structure_id, actif=True).first()
    if not mapping:
        return {}
    try:
        resp = requests.get(
            f"{mapping.api_url}/api/actes/disponibles",
            params={'token': mapping.api_key},
            timeout=10
        )
        if resp.status_code != 200:
            return {}
        return {a['nom'].lower().strip(): a for a in resp.json().get('actes', []) if a.get('nom')}
    except Exception:
        return {}


# ⭐ Regex des 3 paliers du catalogue GHP hospitalisation, ex :
# "P160 Hospi Cabine climatisee avec 1 lit Premiere Semaine"
# "P160 Hospi Cabine climatisee avec 1 lit 8e jour au 14e jour"
# "P160 Hospi Cabine climatisee avec 1 lit 15 jours et plus"
# — même préfixe "P160 Hospi <type de salle>", seul le palier change. Le
# prix (cash, clinique) est identique sur les 3 paliers d'une même famille —
# seul le PBR (base de remboursement AMU) diffère par palier. Voir
# _familles_ghp_hospitalisation() : on regroupe donc par famille pour que
# choisir UN type de salle GHP renseigne les 3 paliers d'un coup, sans
# jamais pouvoir les dépareiller (fini la saisie libre au risque de
# mismatch/typo).
import re as _re
_RE_PALIER_GHP = _re.compile(
    r'^P160\s+Hospi(?:talisation)?\s+(.+?)\s+'
    r'(Premiere\s+Semaine|8e\s+jour\s+au\s+14e\s+jour|15\s+jours\s+et\s+plus)$',
    _re.IGNORECASE
)
_PALIER_LABEL_VERS_CLE = {
    'premiere semaine': 'semaine1',
    '8e jour au 14e jour': 'semaine2',
    '15 jours et plus': 'semaine3',
}


def _familles_ghp_hospitalisation(structure_id):
    """Regroupe le catalogue GHP en 'familles' de tarifs d'hospitalisation :
    un type de salle (ex: 'Cabine climatisee avec 1 lit') -> ses 3 actes
    GHP (un par palier), ne retient que les familles où les 3 paliers
    existent réellement dans le catalogue. Retourne une liste triée par nom
    de famille : [{'famille', 'semaine1', 'semaine2', 'semaine3', 'prix',
    'pbr_semaine1', 'pbr_semaine2', 'pbr_semaine3'}, ...]."""
    tarifs = _tarifs_ghp_structure(structure_id)

    familles = {}
    for acte in tarifs.values():
        nom = (acte.get('nom') or '').strip()
        m = _RE_PALIER_GHP.match(nom)
        if not m:
            continue
        famille_nom = m.group(1).strip()
        palier_cle = _PALIER_LABEL_VERS_CLE.get(m.group(2).strip().lower())
        if not palier_cle:
            continue
        entry = familles.setdefault(famille_nom, {'famille': famille_nom})
        entry[palier_cle] = nom
        entry[f'prix_{palier_cle}'] = acte.get('prix')
        entry[f'pbr_{palier_cle}'] = acte.get('pbr')

    resultat = []
    for famille_nom, entry in familles.items():
        # Ne garder que les familles complètes (les 3 paliers présents) —
        # une famille incomplète produirait un mapping partiel silencieux.
        if not all(entry.get(p) for p in ('semaine1', 'semaine2', 'semaine3')):
            continue
        entry['prix'] = entry.get('prix_semaine1')  # flat sur les 3 paliers
        resultat.append(entry)

    return sorted(resultat, key=lambda e: e['famille'].lower())


@app.route('/api/hospitalisation/familles-salles-ghp')
@login_required
def api_familles_salles_ghp():
    """Types de salle d'hospitalisation disponibles dans le catalogue GHP
    (groupés par famille, 3 paliers). Alimente le sélecteur unique de
    Paramétrage tarif GHP dans la création/modification d'une salle."""
    familles = _familles_ghp_hospitalisation(current_user.id_structure)
    return jsonify({'familles': familles})


@app.route('/hospitalisation/<int:id>/facturation/apercu')
@login_required
def apercu_facturation_hospitalisation(id):
    """Aperçu (lecture seule) du calcul de facturation d'une hospitalisation
    en cours — jours par palier, tarifs GHP, estimation AMU/complémentaire.
    Appelé en AJAX à l'ouverture de la modale de clôture."""
    from models import Hospitalisation
    import math

    hospitalisation = Hospitalisation.query.get_or_404(id)
    if current_user.id_structure and hospitalisation.patient.id_structure != current_user.id_structure:
        return jsonify({'success': False, 'error': 'Accès non autorisé'}), 403

    salle_id_manuel = request.args.get('salle_id', type=int)
    salle = _salle_hospitalisation(hospitalisation, salle_id_manuel=salle_id_manuel)
    if not salle:
        # ⭐ Ni lit assigné, ni correspondance trouvée avec le nom de chambre
        # ("{{ chambre }}" tapé en texte libre) — on propose la liste des
        # salles de la structure pour un choix manuel, au lieu d'abandonner.
        from models import Salle, Service
        salles = Salle.query.join(Service).filter(
            Service.structure_id == hospitalisation.patient.id_structure
        ).order_by(Salle.nom).all()
        return jsonify({
            'success': False,
            'error': (
                f"Aucune salle retrouvée automatiquement pour la chambre "
                f"\"{hospitalisation.chambre or 'non renseignée'}\" — sélectionnez la salle facturée ci-dessous."
            ),
            'chambre_saisie': hospitalisation.chambre,
            'salles_disponibles': [{'id': s.id, 'nom': s.nom} for s in salles]
        })

    jours_total = max(1, math.ceil((datetime.utcnow() - hospitalisation.date_debut).total_seconds() / 86400))
    structure_id = hospitalisation.patient.id_structure
    paliers, param = _calculer_paliers_hospitalisation(jours_total, structure_id)

    noms_actes = {
        'semaine1': salle.acte_ghp_semaine1,
        'semaine2': salle.acte_ghp_semaine2,
        'semaine3plus': salle.acte_ghp_semaine3,
    }
    tarifs = _tarifs_ghp_structure(structure_id)

    patient = hospitalisation.patient
    try:
        taux_amu_patient = float(patient.taux_prise_charge) if patient.taux_prise_charge else 0
    except (TypeError, ValueError):
        taux_amu_patient = 0
    est_assure = (
        bool(patient.type_assurance)
        and 'non_assur' not in patient.type_assurance.lower().replace('é', 'e')
        and taux_amu_patient > 0
    )
    taux_cac = float(patient.taux_assurance2 or 0) if patient.assurance2_nom else 0

    lignes = []
    total_brut = 0
    total_pbr_base = 0
    prise_en_charge_amu = 0
    mapping_manquant = False

    for p in paliers:
        acte_nom = noms_actes.get(p['palier'])
        tarif = tarifs.get((acte_nom or '').lower().strip())
        if not acte_nom:
            mapping_manquant = True
        prix = float(tarif['prix']) if tarif else 0
        pbr = float(tarif['pbr']) if tarif else 0
        sous_total = prix * p['jours']
        pbr_base = min(prix, pbr) * p['jours'] if tarif else 0
        total_brut += sous_total
        total_pbr_base += pbr_base
        # ⭐ FIX : taux AMU par palier (P160 = 90%, le reste = taux du
        # patient) au lieu d'un taux unique appliqué à tout le total —
        # voir _taux_amu_pour_acte() ci-dessus.
        taux_ligne = _taux_amu_pour_acte(acte_nom, taux_amu_patient)
        if est_assure and tarif:
            prise_en_charge_amu += pbr_base * taux_ligne / 100
        lignes.append({
            'palier': p['palier'],
            'label': p['label'],
            'jours': p['jours'],
            'acte_nom': acte_nom or '(non configuré — voir Paramétrage AMU)',
            'trouve_dans_catalogue': tarif is not None,
            'prix_unitaire': prix,
            'sous_total': sous_total,
            'taux_amu': taux_ligne if est_assure else 0,
        })

    reste_apres_amu = max(total_brut - prise_en_charge_amu, 0)
    prise_en_charge_cac = (reste_apres_amu * taux_cac / 100) if taux_cac > 0 else 0
    net_estime = max(reste_apres_amu - prise_en_charge_cac, 0)

    return jsonify({
        'success': True,
        'jours_total': jours_total,
        'salle_id': salle.id,
        'salle_nom': salle.nom,
        'lignes': lignes,
        'mapping_manquant': mapping_manquant,
        'total_brut': total_brut,
        'est_assure': est_assure,
        'taux_amu_patient': taux_amu_patient,
        'prise_en_charge_amu': prise_en_charge_amu,
        'assurance2_nom': patient.assurance2_nom,
        'taux_cac': taux_cac,
        'prise_en_charge_cac': prise_en_charge_cac,
        'net_estime': net_estime,
    })


@app.route('/hospitalisation/<int:id>/cloturer', methods=['POST'])
@login_required
def cloturer_hospitalisation(id):
    """Clôturer une hospitalisation (sortie du patient)"""
    from models import Hospitalisation, HospitalisationMedecin, Lit
    from datetime import datetime
    import math

    hospitalisation = Hospitalisation.query.get_or_404(id)
    
    if current_user.role not in ['super_admin', 'admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    # Vérifier l'appartenance à la structure
    if current_user.id_structure and hospitalisation.patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_hospitalisations'))
    
    if current_user.role == 'medecin':
        assigne = HospitalisationMedecin.query.filter_by(
            hospitalisation_id=id,
            medecin_id=current_user.id,
            actif=True
        ).first()
        if not assigne:
            flash('Vous n\'êtes pas assigné à cette hospitalisation', 'danger')
            return redirect(url_for('liste_hospitalisations'))
    
    # Récupération des données du formulaire
    type_sortie = request.form.get('type_sortie', 'sortie')
    motif_sortie = request.form.get('motif_sortie', '')
    
    # --- Gestion du type de sortie ---
    if type_sortie == 'transfere':
        centre_transfert = request.form.get('centre_transfert')
        motif_transfert = request.form.get('motif_transfert')
        date_transfert_str = request.form.get('date_transfert')
        
        hospitalisation.statut = 'transfere'
        hospitalisation.centre_transfert = centre_transfert
        hospitalisation.motif_transfert = motif_transfert
        if date_transfert_str:
            hospitalisation.date_transfert = datetime.fromisoformat(date_transfert_str)
            
    elif type_sortie == 'deces':
        hospitalisation.statut = 'sorti'
        motif_sortie = f"DÉCÈS - {motif_sortie}" if motif_sortie else "DÉCÈS"
        
    elif type_sortie == 'autres':
        hospitalisation.statut = 'sorti'
        motif_autres = request.form.get('motif_autres')
        if motif_autres:
            motif_sortie = f"Autre motif: {motif_autres}"
        else:
            motif_sortie = "Autre motif non spécifié"
            
    else:  # sortie normale
        hospitalisation.statut = 'sorti'
    
    # ⭐ Facturation : retrouver la salle occupée AVANT de libérer le lit
    # (sinon l'info est perdue — lit_id est mis à None juste après).
    # salle_id (choix manuel) n'est fourni que si l'aperçu n'a pas pu
    # résoudre automatiquement une salle (voir apercu_facturation_hospitalisation).
    salle_id_manuel = request.form.get('salle_id_facturation', type=int)
    salle_facturation = _salle_hospitalisation(hospitalisation, salle_id_manuel=salle_id_manuel)

    # ⭐ LIBÉRER LE LIT (corrigé)
    if hospitalisation.lit_id:
        lit = Lit.query.get(hospitalisation.lit_id)
        if lit:
            lit.liberer()
            hospitalisation.lit_id = None

    # --- Gestion des avis externes ---
    medecins_externes = request.form.get('medecins_externes')
    demandes_avis = request.form.get('demandes_avis')
    avis_externes = request.form.get('avis_externes')
    
    if medecins_externes:
        hospitalisation.medecins_externes = medecins_externes
    if demandes_avis:
        hospitalisation.demandes_avis = demandes_avis
    if avis_externes:
        hospitalisation.avis_externes = avis_externes
    
    hospitalisation.date_fin = datetime.utcnow()
    
    # --- Construction des notes de sortie ---
    notes_completes = f"\n--- SORTIE DU PATIENT ---\n"
    notes_completes += f"Date de sortie: {hospitalisation.date_fin.strftime('%d/%m/%Y %H:%M')}\n"
    notes_completes += f"Type: {type_sortie}\n"
    notes_completes += f"Motif: {motif_sortie}\n"
    
    if type_sortie == 'transfere':
        notes_completes += f"Transfert vers: {centre_transfert or 'Non spécifié'}\n"
        notes_completes += f"Motif du transfert: {motif_transfert or 'Non spécifié'}\n"
        if date_transfert_str:
            notes_completes += f"Date du transfert: {date_transfert_str}\n"
    
    if medecins_externes:
        notes_completes += f"Médecins externes consultés: {medecins_externes}\n"
    if demandes_avis:
        notes_completes += f"Demandes d'avis: {demandes_avis}\n"
    if avis_externes:
        notes_completes += f"Avis reçus: {avis_externes}\n"
    
    notes_completes += "---\n"
    
    hospitalisation.notes_admission = (hospitalisation.notes_admission or '') + notes_completes

    db.session.commit()

    # ⭐ Facturation automatique vers GHP (jours × tarif par palier)
    if not salle_facturation:
        flash(
            f'Hospitalisation clôturée. Salle non identifiée (chambre "{hospitalisation.chambre or "?"}" '
            f'sans lit assigné, et sans salle du même nom en Paramétrage AMU) : la facturation '
            f'n\'a pas pu être calculée automatiquement — à saisir manuellement côté GHP si besoin.',
            'warning'
        )
    else:
        from models import HospitalisationFacturation
        jours_total = max(1, math.ceil((hospitalisation.date_fin - hospitalisation.date_debut).total_seconds() / 86400))
        structure_id = hospitalisation.patient.id_structure
        paliers, _param = _calculer_paliers_hospitalisation(jours_total, structure_id)

        noms_actes = {
            'semaine1': salle_facturation.acte_ghp_semaine1,
            'semaine2': salle_facturation.acte_ghp_semaine2,
            'semaine3plus': salle_facturation.acte_ghp_semaine3,
        }
        tarifs = _tarifs_ghp_structure(structure_id)

        lignes_creees = 0
        paliers_sans_mapping = []
        for p in paliers:
            acte_nom = noms_actes.get(p['palier'])
            if not acte_nom:
                paliers_sans_mapping.append(p['label'])
                continue
            tarif = tarifs.get(acte_nom.lower().strip())
            ligne = HospitalisationFacturation(
                hospitalisation_id=hospitalisation.id,
                patient_id=hospitalisation.patient_id,
                palier=p['palier'],
                acte_nom=acte_nom,
                nombre_jours=p['jours'],
                prix_unitaire_estime=float(tarif['prix']) if tarif else None
            )
            db.session.add(ligne)
            lignes_creees += 1

        if lignes_creees:
            db.session.commit()
            try:
                from tasks import sync_hospitalisations_to_ghp
                resultat_sync = sync_hospitalisations_to_ghp()
            except Exception as e:
                resultat_sync = {'success': False, 'message': str(e)}

            if resultat_sync.get('success') and resultat_sync.get('count'):
                flash(
                    f'Hospitalisation clôturée ({jours_total} jour(s), {salle_facturation.nom}) '
                    f'— facturation envoyée à GHP (onglet Prescriptions reçues).',
                    'success'
                )
            else:
                flash(
                    f'Hospitalisation clôturée ({jours_total} jour(s)). La facturation sera '
                    f'envoyée à GHP automatiquement dans les prochaines minutes (rattrapage).',
                    'warning'
                )
        else:
            flash(
                f'Hospitalisation clôturée. Aucun acte GHP n\'est configuré pour la salle '
                f'"{salle_facturation.nom}" ({", ".join(paliers_sans_mapping)}) — configurez-le '
                f'dans Paramétrage AMU, puis facturez manuellement côté GHP.',
                'warning'
            )

    return redirect(url_for('liste_hospitalisations'))

@app.route('/patient/<int:patient_id>/hospitalisations')
@login_required
def patient_hospitalisations(patient_id):
    """Voir toutes les hospitalisations d'un patient"""
    from models import Patient, Hospitalisation
    
    patient = Patient.query.get_or_404(patient_id)
    
    # Vérifier les permissions
    if current_user.role not in ['super_admin']:
        if current_user.id_structure and patient.id_structure != current_user.id_structure:
            flash('Accès non autorisé', 'danger')
            return redirect(url_for('dashboard'))
        
        if current_user.role == 'medecin' and patient.id_medecin_referent != current_user.id:
            flash('Accès non autorisé', 'danger')
            return redirect(url_for('dashboard'))
    
    hospitalisations = Hospitalisation.query.filter_by(
        patient_id=patient_id
    ).order_by(Hospitalisation.date_debut.desc()).all()
    
    return render_template('hospitalisations/patient_hospitalisations.html',
                         patient=patient,
                         hospitalisations=hospitalisations)
@app.route('/hospitalisation/<int:id>/avis-externe', methods=['POST'])
@login_required
def ajouter_avis_externe(id):
    """Ajouter un nouvel avis de médecin externe"""
    from models import Hospitalisation, AvisExterne
    from datetime import datetime
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    
    # Vérifier les permissions
    if current_user.role not in ['super_admin', 'admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Vérifier l'appartenance à la structure
    if current_user.id_structure and hospitalisation.patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_hospitalisations'))
    
    # Récupération des données
    medecin_nom = request.form.get('medecin_nom', '').strip()
    specialite = request.form.get('specialite', '').strip()
    etablissement = request.form.get('etablissement', '').strip()
    demande_avis = request.form.get('demande_avis', '').strip()
    avis_recu = request.form.get('avis_recu', '').strip()
    date_demande_str = request.form.get('date_demande')
    date_reception_str = request.form.get('date_reception')
    
    # Validation
    if not medecin_nom or not avis_recu:
        flash('Le nom du médecin et l\'avis reçu sont obligatoires', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    # Création de l'avis
    avis = AvisExterne(
        hospitalisation_id=id,
        medecin_nom=medecin_nom,
        specialite=specialite if specialite else None,
        etablissement=etablissement if etablissement else None,
        demande_avis=demande_avis if demande_avis else None,
        avis_recu=avis_recu,
        created_by=current_user.id
    )
    
    if date_demande_str:
        avis.date_demande = datetime.fromisoformat(date_demande_str)
    if date_reception_str:
        avis.date_reception = datetime.fromisoformat(date_reception_str)
    
    db.session.add(avis)
    db.session.commit()
    
    flash('✅ Avis externe enregistré avec succès', 'success')
    return redirect(url_for('detail_hospitalisation', id=id))

@app.route('/consultation/<int:id>/resultats', methods=['POST'])
@login_required
def ajouter_resultats(id):
    """Ajouter ou mettre à jour les résultats des examens"""
    from models import Consultation
    from datetime import datetime
    
    consultation = Consultation.query.get_or_404(id)
    
    # Vérifier les permissions
    if current_user.role not in ['super_admin', 'admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Récupération des données
    resultats_biologie = request.form.get('resultats_biologie', '').strip()
    resultats_imagerie = request.form.get('resultats_imagerie', '').strip()
    date_resultats_str = request.form.get('date_resultats')
    
    # Mise à jour
    if resultats_biologie:
        consultation.resultats_biologie = resultats_biologie
    if resultats_imagerie:
        consultation.resultats_imagerie = resultats_imagerie
    
    if date_resultats_str:
        consultation.date_resultats = datetime.fromisoformat(date_resultats_str)
    else:
        consultation.date_resultats = datetime.utcnow()
    
    db.session.commit()
    
    flash('✅ Résultats enregistrés avec succès', 'success')
    return redirect(url_for('consultation_detail', id=id))


# ==================== CIM-10 ====================

import csv
import os

_cim10_cache = None
_cim10_last_update = None

def get_cim10_list():
    """Récupère la liste des codes CIM-10 depuis le fichier local"""
    global _cim10_cache, _cim10_last_update
    
    if _cim10_cache and _cim10_last_update:
        from datetime import datetime
        if (datetime.now() - _cim10_last_update).seconds < 3600:
            return _cim10_cache
    
    try:
        cim10_file = os.path.join(os.path.dirname(__file__), 'cim10.csv')
        
        if not os.path.exists(cim10_file):
            print(f"❌ Fichier {cim10_file} non trouvé !")
            return []
        
        cim10_list = []
        with open(cim10_file, 'r', encoding='utf-8') as f:
            csv_reader = csv.reader(f)
            for i, row in enumerate(csv_reader):
                if i == 0:  # Ignorer l'en-tête "id,code"
                    continue
                if row and len(row) > 1:
                    code = row[1].strip()  # Colonne "code"
                    if code:
                        code = code.strip('"')
                        code = code.replace('""', '"')
                        cim10_list.append(code)
        
        _cim10_cache = cim10_list
        from datetime import datetime
        _cim10_last_update = datetime.now()
        
        return cim10_list
    except Exception as e:
        print(f"❌ Erreur chargement CIM-10: {e}")
        return []

def search_cim10(search_term, limit=30):
    """Recherche dans les codes CIM-10"""
    if not search_term or len(search_term) < 2:
        return []
    
    all_codes = get_cim10_list()
    search_term = search_term.lower().strip()
    
    results = []
    for code in all_codes:
        if search_term in code.lower():
            results.append({'nom': code})
            if len(results) >= limit:
                break
    
    return results

@app.route('/api/cim10/search')
@login_required
def api_cim10_search():
    """API de recherche CIM-10"""
    term = request.args.get('term', '')
    if len(term) < 2:
        return jsonify([])
    
    results = search_cim10(term)
    return jsonify(results)
# ==================== ANALYSES ====================

@app.route('/analyses')
@login_required
@has_permission('ANALYSES')  # ⭐ AJOUTER LE DÉCORATEUR
def liste_analyses():
    """Liste des analyses regroupées par patient"""
    from models import AnalyseDemande, Patient
    from sqlalchemy import or_
    
    # ⭐ ACCÈS POUR MÉDECIN, LABORANTIN ET ADMIN
    if current_user.role not in ['admin_structure', 'laborantin', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Récupérer les paramètres
    statut = request.args.get('statut', '')
    search = request.args.get('search', '')
    
    # Requête de base
    # ⭐ FIX : cette page "Laboratoire Analyses" ne filtrait sur BIOLOGIE
    # que pour le rôle laborantin — un admin_structure/médecin (qui voit
    # les deux filières) y voyait donc aussi l'imagerie mélangée, alors
    # que la page s'appelle "Analyses" (biologie). Filtre systématique,
    # quel que soit le rôle.
    query = AnalyseDemande.query.filter_by(
        structure_id=current_user.id_structure,
        type_analyse='BIOLOGIE',
    )

    # Filtrer par statut
    if statut:
        query = query.filter_by(statut=statut)
    
    # Recherche
    if search:
        search = search.strip()
        filters = []
        filters.append(Patient.nom.ilike(f'%{search}%'))
        filters.append(Patient.prenom.ilike(f'%{search}%'))
        
        if search.upper().startswith('P'):
            try:
                num = int(search[1:])
                filters.append(Patient.id == num)
            except ValueError:
                pass
        elif search.isdigit():
            filters.append(Patient.id == int(search))
        
        if filters:
            patients_trouves = Patient.query.filter(or_(*filters)).all()
            patient_ids = [p.id for p in patients_trouves]
            if patient_ids:
                query = query.filter(AnalyseDemande.patient_id.in_(patient_ids))
            else:
                query = query.filter(AnalyseDemande.patient_id == -1)
    
    # Regrouper par patient
    analyses = query.order_by(AnalyseDemande.date_demande.desc()).all()
    
    patients_dict = {}
    for analyse in analyses:
        patient_id = analyse.patient_id
        if patient_id not in patients_dict:
            patients_dict[patient_id] = {
                'patient': analyse.patient,
                'analyses': []
            }
        patients_dict[patient_id]['analyses'].append(analyse)
    
    patients = list(patients_dict.values())
    statuts = ['EN_ATTENTE', 'EN_COURS', 'TERMINE']
    
    return render_template('analyses/liste.html',
                         patients=patients,
                         statut_actuel=statut,
                         statuts=statuts,
                         search=search)

@app.route('/analyse/<int:id>')
@login_required
def detail_analyse(id):
    """Détail d'une analyse demandée"""
    from models import AnalyseDemande
    
    analyse = AnalyseDemande.query.get_or_404(id)
    
    # ⭐ PERMISSIONS - Médecin, Laborantin, Radiologue, Admin, Super Admin
    if current_user.role not in ['super_admin', 'admin_structure', 'laborantin', 'medecin', 'radiologue']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Vérifier la structure
    if current_user.role not in ['super_admin']:
        if analyse.structure_id != current_user.id_structure:
            flash('Accès non autorisé', 'danger')
            return redirect(url_for('liste_analyses'))
    
    return render_template('analyses/detail.html', analyse=analyse)

@app.route('/analyse/<int:id>/resultats', methods=['POST'])
@login_required
def saisir_resultats_analyse(id):
    """Saisir les résultats d'une analyse (Laborantin ou Radiologue) — texte
    libre (historique, toujours possible), OU fichier importé, OU rédigé
    dans l'éditeur en ligne, avec signature électronique pré-enregistrée
    optionnelle (voir /signatures-intervenants) — même principe que GHP
    (api_enregistrer_resultat, ResultatExamen), porté directement par
    AnalyseDemande. Tout est optionnel/additif : la saisie texte simple
    d'avant continue de fonctionner à l'identique si rien d'autre n'est
    envoyé."""
    from models import AnalyseDemande, Consultation, SignatureIntervenant
    from datetime import datetime

    analyse = AnalyseDemande.query.get_or_404(id)

    # ⭐ PERMISSIONS : Laborantin, Radiologue, Admin, Super Admin
    if current_user.role not in ['laborantin', 'radiologue', 'admin_structure', 'super_admin']:
        flash('Accès non autorisé - réservé au laborantin ou radiologue', 'danger')
        return redirect(url_for('dashboard'))

    # ⭐ VÉRIFIER QUE LE RÔLE CORRESPOND AU TYPE D'ANALYSE
    if current_user.role == 'laborantin' and analyse.type_analyse != 'BIOLOGIE':
        flash('Accès non autorisé - vous ne pouvez saisir que les analyses de biologie', 'danger')
        return redirect(url_for('liste_analyses'))

    if current_user.role == 'radiologue' and analyse.type_analyse != 'IMAGERIE':
        flash('Accès non autorisé - vous ne pouvez saisir que les examens d\'imagerie', 'danger')
        return redirect(url_for('liste_radiologie'))

    if analyse.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_analyses'))

    resultats = (request.form.get('resultats') or '').strip()
    statut = request.form.get('statut', 'TERMINE')
    contenu_html = (request.form.get('contenu_html') or '').strip()
    fichier = request.files.get('fichier')
    modele_id = request.form.get('modele_utilise_id')
    signature_choice = request.form.get('signature_id')  # id numérique, '__autre__', ou vide
    nom_interprete_libre = (request.form.get('nom_interprete') or '').strip()

    if statut == 'TERMINE' and not (resultats or contenu_html or (fichier and fichier.filename)):
        flash('Veuillez saisir les résultats (texte, fichier ou contenu rédigé)', 'danger')
        return redirect(url_for('detail_analyse', id=id))

    # ⭐ Signature pré-enregistrée sélectionnée : sa signature s'appose
    # automatiquement, COPIÉE (jamais relue plus tard) — voir le
    # commentaire sur AnalyseDemande.signature_data (models.py).
    nom_interprete = nom_interprete_libre or None
    titre_interprete = None
    signature_intervenant_id = None
    signature_data = None
    signature_mime = None
    if signature_choice and signature_choice != '__autre__':
        signature = SignatureIntervenant.query.filter_by(
            id=signature_choice, structure_id=current_user.id_structure,
            filiere=analyse.type_analyse, actif=True,
        ).first()
        if signature:
            nom_interprete = signature.nom
            titre_interprete = signature.titre
            signature_intervenant_id = signature.id
            signature_data = signature.signature_data
            signature_mime = signature.signature_mime

    # Mettre à jour l'analyse
    analyse.resultats = resultats or None
    analyse.statut = statut
    analyse.date_resultats = datetime.utcnow()
    analyse.resultats_par = current_user.id
    if fichier and fichier.filename:
        analyse.fichier_nom = fichier.filename
        analyse.fichier_mime = fichier.mimetype
        analyse.fichier_data = fichier.read()
        analyse.contenu_html = None
    elif contenu_html:
        analyse.contenu_html = contenu_html
        analyse.fichier_nom = None
        analyse.fichier_mime = None
        analyse.fichier_data = None
    if modele_id:
        try:
            analyse.modele_utilise_id = int(modele_id)
        except ValueError:
            pass
    if nom_interprete:
        analyse.nom_interprete = nom_interprete
        analyse.titre_interprete = titre_interprete
        analyse.signature_intervenant_id = signature_intervenant_id
        analyse.signature_data = signature_data
        analyse.signature_mime = signature_mime

    # Mettre à jour les résultats de la consultation — uniquement pour le
    # texte libre (comportement historique inchangé) ; un fichier/contenu
    # riche reste consultable via /analyse/<id> ou l'impression, pas
    # dupliqué en texte brut ici.
    consultation = Consultation.query.get(analyse.consultation_id)
    if consultation and resultats:
        if analyse.type_analyse == 'BIOLOGIE':
            if consultation.resultats_biologie:
                consultation.resultats_biologie += f"\n\n--- {analyse.nom_analyse} ---\n{resultats}"
            else:
                consultation.resultats_biologie = f"--- {analyse.nom_analyse} ---\n{resultats}"
        elif analyse.type_analyse == 'IMAGERIE':
            if consultation.resultats_imagerie:
                consultation.resultats_imagerie += f"\n\n--- {analyse.nom_analyse} ---\n{resultats}"
            else:
                consultation.resultats_imagerie = f"--- {analyse.nom_analyse} ---\n{resultats}"

        consultation.date_resultats = datetime.utcnow()

    db.session.commit()

    # ⭐ Miroir vers GHP dès qu'un résultat est disponible (voir
    # _envoyer_resultats_examens_ghp_immediat) — pas avant, un résultat non
    # terminé n'a rien à partager côté GHP.
    if analyse.statut == 'TERMINE':
        _envoyer_resultats_examens_ghp_immediat()

    flash('✅ Résultats enregistrés avec succès', 'success')

    # ⭐ REDIRECTION SELON LA FILIÈRE DE CETTE ANALYSE — pas le rôle : un
    # admin_structure qui saisit un résultat d'imagerie était toujours
    # renvoyé vers la page biologie (liste_analyses), jamais radiologie.
    if analyse.type_analyse == 'IMAGERIE':
        return redirect(url_for('liste_radiologie'))
    else:
        return redirect(url_for('liste_analyses'))


def _type_analyse_depuis_ghp(type_prestation):
    """analyse/examen (vocabulaire GHP) -> BIOLOGIE/IMAGERIE (vocabulaire
    gestion_patients) — inverse de _type_prestation_ghp (tasks.py)."""
    return 'IMAGERIE' if type_prestation == 'examen' else 'BIOLOGIE'


@app.route('/api/resultats-examens/sync-externe', methods=['POST'])
def api_recevoir_resultat_examen_externe():
    """
    Reçoit, en miroir, un résultat d'analyse/examen (ou un modèle de
    résultat) saisi côté GHP — voir _pousser_resultat_examen_gestion_patients
    / _pousser_modele_resultat_gestion_patients dans medilogic_ghp/app.py.
    Même schéma que le sens gestion_patients -> GHP (voir
    _envoyer_resultats_examens_ghp_immediat/tasks.py ici, et
    /api/resultats-examens/sync-externe côté GHP) : token =
    StructureMapping.api_key, upsert idempotent sur
    (structure_id, source_app='ghp', source_model, source_id).

    ⭐⭐ RÈGLE ANTI-BOUCLE ⭐⭐ : `source_synced_at` est posé IMMÉDIATEMENT sur
    toute ligne créée/mise à jour ici (jamais laissé NULL) — sans ça, le job
    sortant de tasks.sync_resultats_examens_to_ghp (qui ne filtre QUE sur
    `source_synced_at IS NULL`) la repousserait vers GHP au prochain cycle,
    créant une boucle infinie. Volontairement PAS de secret statique façon
    /api/webhook/patient-created (WEBHOOK_SECRET) — même token que partout
    ailleurs dans cette intégration, seule source d'auth fiable pour
    identifier la structure appelante.
    """
    from models import Patient, StructureMapping, AnalyseDemande, ModeleResultat
    import base64

    token = request.args.get('token')
    if not token:
        return jsonify({'success': False, 'error': 'Token manquant'}), 401

    mapping = StructureMapping.query.filter_by(api_key=token, actif=True).first()
    if not mapping:
        return jsonify({'success': False, 'error': 'Token invalide'}), 401

    try:
        data = request.json or {}
        categorie = data.get('categorie')
        if categorie not in ('resultat', 'modele'):
            return jsonify({'success': False, 'error': 'Catégorie non synchronisable'}), 400

        source_app = data.get('source_app') or 'ghp'
        source_model = data.get('source_model')
        source_id = data.get('source_id')
        if not source_model or not source_id:
            return jsonify({'success': False, 'error': 'source_model/source_id manquants'}), 400

        structure_id = mapping.local_structure_id
        auteur_nom = data.get('auteur_nom') or 'Sync GHP'

        if categorie == 'modele':
            modele = ModeleResultat.query.filter_by(
                structure_id=structure_id, source_app=source_app,
                source_model=source_model, source_id=source_id,
            ).first()
            fichier_data_b64 = data.get('fichier_data_b64')
            if not modele:
                modele = ModeleResultat(
                    structure_id=structure_id, source_app=source_app,
                    source_model=source_model, source_id=source_id,
                )
                db.session.add(modele)
            modele.type_analyse = _type_analyse_depuis_ghp(data.get('type_prestation'))
            modele.nom = data.get('nom') or 'Sans titre'
            modele.fichier_nom = data.get('fichier_nom')
            modele.fichier_mime = data.get('fichier_mime')
            modele.fichier_data = base64.b64decode(fichier_data_b64) if fichier_data_b64 else None
            modele.contenu_html = data.get('contenu_html') or None
            modele.created_by = auteur_nom
            modele.source_synced_at = datetime.utcnow()  # ⭐ anti-boucle, voir docstring
            db.session.commit()
            return jsonify({'success': True, 'modele_id': modele.id})

        # categorie == 'resultat'
        patient_source_id = data.get('patient_source_id')
        patient_nom = data.get('patient_nom') or ''
        patient_prenom = data.get('patient_prenom') or ''

        # ⭐ Même logique de correspondance que sync_patients_from_ghp
        # (patient_source_id + source_structure_id d'abord), repli nom/prénom.
        patient = None
        if patient_source_id:
            patient = Patient.query.filter_by(
                patient_source_id=str(patient_source_id),
                source_structure_id=mapping.source_structure_id,
                id_structure=structure_id,
            ).first()
        if not patient and patient_nom and patient_prenom:
            patient = Patient.query.filter(
                db.func.lower(Patient.nom) == patient_nom.strip().lower(),
                db.func.lower(Patient.prenom) == patient_prenom.strip().lower(),
                Patient.id_structure == structure_id,
            ).first()
        if not patient:
            return jsonify({'success': False, 'error': 'patient_introuvable'}), 404

        analyse = AnalyseDemande.query.filter_by(
            structure_id=structure_id, source_app=source_app,
            source_model=source_model, source_id=source_id,
        ).first()
        if not analyse:
            analyse = AnalyseDemande(
                structure_id=structure_id, patient_id=patient.id,
                source_app=source_app, source_model=source_model, source_id=source_id,
            )
            db.session.add(analyse)

        analyse.type_analyse = _type_analyse_depuis_ghp(data.get('type_prestation'))
        analyse.nom_analyse = data.get('nom_analyse') or data.get('acte_nom') or 'Examen'
        analyse.description = data.get('motif') or data.get('description') or ''
        analyse.statut = 'TERMINE'
        if not analyse.date_prescription:
            analyse.date_prescription = datetime.utcnow()
        if not analyse.date_demande:
            analyse.date_demande = datetime.utcnow()
        analyse.date_resultats = datetime.utcnow()

        fichier_data_b64 = data.get('fichier_data_b64')
        signature_data_b64 = data.get('signature_data_b64')
        analyse.resultats = data.get('resultats_texte') or None
        analyse.fichier_nom = data.get('fichier_nom')
        analyse.fichier_mime = data.get('fichier_mime')
        analyse.fichier_data = base64.b64decode(fichier_data_b64) if fichier_data_b64 else None
        analyse.contenu_html = data.get('contenu_html') or None
        analyse.nom_interprete = data.get('nom_interprete') or auteur_nom
        analyse.titre_interprete = data.get('titre_interprete')
        analyse.signature_data = base64.b64decode(signature_data_b64) if signature_data_b64 else None
        analyse.signature_mime = data.get('signature_mime')
        analyse.source_synced_at = datetime.utcnow()  # ⭐ anti-boucle, voir docstring

        db.session.commit()
        return jsonify({'success': True, 'analyse_id': analyse.id})

    except Exception as e:
        db.session.rollback()
        print(f"⚠️ Erreur sync résultat depuis GHP : {e}")
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/analyse/<int:id>/fichier')
@login_required
def telecharger_fichier_analyse(id):
    """Sert le fichier de résultat joint à une AnalyseDemande (PDF de
    préférence, Word/Excel accepté)."""
    from models import AnalyseDemande
    from flask import Response

    analyse = AnalyseDemande.query.get_or_404(id)
    if current_user.role != 'super_admin' and analyse.structure_id != current_user.id_structure:
        return "Accès non autorisé", 403
    if not analyse.fichier_data:
        return "Ce résultat a été rédigé en ligne, pas de fichier joint.", 404
    return Response(
        analyse.fichier_data, mimetype=analyse.fichier_mime or 'application/octet-stream',
        headers={'Content-Disposition': f'inline; filename="{analyse.fichier_nom or "resultat"}"'}
    )


@app.route('/analyse/<int:id>/signature-image')
@login_required
def image_signature_analyse(id):
    """Sert l'image de signature FIGÉE sur cette analyse précise (jamais la
    signature actuelle du registre — voir AnalyseDemande.signature_data)."""
    from models import AnalyseDemande
    from flask import Response

    analyse = AnalyseDemande.query.get_or_404(id)
    if current_user.role != 'super_admin' and analyse.structure_id != current_user.id_structure:
        return "Accès non autorisé", 403
    if not analyse.signature_data:
        return "Introuvable", 404
    return Response(analyse.signature_data, mimetype=analyse.signature_mime or 'image/png')


# ================================================================
# LABORATOIRE / RADIOLOGIE — Modèles de résultats (parité GHP)
# ================================================================
# ⭐ Un modèle (Word/Excel importé, ou rédigé en ligne) sert de point de
# départ réutilisable pour la saisie d'un résultat — voir
# saisir_resultats_analyse ci-dessus et le formulaire dans analyses/detail.html.
def _acces_module_resultats():
    return (current_user.role in ('laborantin', 'radiologue', 'secretaire', 'admin_structure'))


@app.route('/modeles-resultats')
@login_required
def page_modeles_resultats():
    if not _acces_module_resultats():
        flash('Accès non autorisé pour votre rôle.', 'danger')
        return redirect(url_for('dashboard'))
    return render_template('analyses/modeles_liste.html')


@app.route('/api/modeles-resultats', methods=['GET'])
@login_required
def api_lister_modeles_resultats():
    from models import ModeleResultat
    q = ModeleResultat.query.filter_by(structure_id=current_user.id_structure)
    type_analyse = request.args.get('type_analyse')
    if type_analyse:
        q = q.filter_by(type_analyse=type_analyse)
    lignes = q.order_by(ModeleResultat.type_analyse, ModeleResultat.nom).all()
    return jsonify([{
        'id': l.id, 'nom': l.nom, 'type_analyse': l.type_analyse,
        'fichier_nom': l.fichier_nom, 'a_contenu_html': bool(l.contenu_html),
        'created_at': l.created_at.strftime('%d/%m/%Y') if l.created_at else '',
        'source_app': l.source_app,  # ⭐ 'ghp' si synchronisé depuis GHP, sinon None (natif)
    } for l in lignes])


@app.route('/api/modeles-resultats/<int:modele_id>/contenu', methods=['GET'])
@login_required
def api_contenu_modele_resultat(modele_id):
    from models import ModeleResultat
    modele = ModeleResultat.query.filter_by(id=modele_id, structure_id=current_user.id_structure).first()
    if not modele:
        return jsonify({'success': False, 'error': 'Introuvable'}), 404
    return jsonify({'success': True, 'contenu_html': modele.contenu_html or ''})


def _aplatir_tableaux_html(html):
    """Remplace chaque <table> par un paragraphe par ligne — l'éditeur en
    ligne (Quill) n'a pas de module tableau et aplatit silencieusement tout
    <table> en un seul bloc de texte collé sans espaces ni sauts de ligne."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, 'html.parser')
    for table in soup.find_all('table'):
        remplacement = soup.new_tag('div')
        for tr in table.find_all('tr'):
            valeurs = [c.get_text(strip=True) for c in tr.find_all(['td', 'th'])]
            valeurs = [v for v in valeurs if v]
            if valeurs:
                p = soup.new_tag('p')
                p.string = ' — '.join(valeurs)
                remplacement.append(p)
        table.replace_with(remplacement)
    return str(soup)


def _convertir_fichier_en_html(fichier_data, fichier_nom, fichier_mime):
    """Convertit un modèle Word (.docx) ou Excel (.xlsx) importé en HTML
    éditable. Les anciens formats binaires .doc/.xls (pré-2007) ne sont pas
    lisibles ainsi — message clair invitant à réenregistrer en .docx/.xlsx."""
    nom = (fichier_nom or '').lower()
    mime = (fichier_mime or '').lower()

    if nom.endswith('.docx') or 'wordprocessingml' in mime:
        import mammoth
        from io import BytesIO
        resultat = mammoth.convert_to_html(BytesIO(fichier_data))
        return _aplatir_tableaux_html(resultat.value), None

    if nom.endswith('.xlsx') or 'spreadsheetml' in mime:
        import openpyxl
        from io import BytesIO
        from markupsafe import escape
        wb = openpyxl.load_workbook(BytesIO(fichier_data), data_only=True)
        ws = wb.active
        lignes_html = []
        for row in ws.iter_rows():
            valeurs = [str(c.value) for c in row if c.value is not None and str(c.value).strip()]
            if valeurs:
                lignes_html.append(f'<p>{escape(" — ".join(valeurs))}</p>')
        html = ''.join(lignes_html) or '<p></p>'
        return html, None

    return None, "Conversion non prise en charge pour ce type de fichier (.doc/.xls ancien format, PDF...) — réenregistrez-le en .docx/.xlsx depuis Word/Excel, ou retapez le contenu directement."


@app.route('/api/modeles-resultats/<int:modele_id>/convertir-html', methods=['GET'])
@login_required
def api_convertir_modele_html(modele_id):
    from models import ModeleResultat
    modele = ModeleResultat.query.filter_by(id=modele_id, structure_id=current_user.id_structure).first()
    if not modele:
        return jsonify({'success': False, 'error': 'Introuvable'}), 404
    if modele.contenu_html:
        return jsonify({'success': True, 'contenu_html': modele.contenu_html})
    if not modele.fichier_data:
        return jsonify({'success': False, 'error': 'Aucun fichier à convertir'}), 400
    try:
        html, erreur = _convertir_fichier_en_html(modele.fichier_data, modele.fichier_nom, modele.fichier_mime)
        if erreur:
            return jsonify({'success': False, 'error': erreur}), 400
        return jsonify({'success': True, 'contenu_html': html})
    except Exception as e:
        return jsonify({'success': False, 'error': f'Échec de la conversion : {e}'}), 500


def _extraire_texte_fichier(fichier_data, fichier_nom, fichier_mime):
    """Extrait le texte brut d'un fichier Word (.docx) ou PDF pour
    pré-remplir un formulaire de création (protocole/ordonnance/examen déjà
    préparé sur l'ordinateur) — best-effort, l'utilisateur reste libre de
    corriger avant d'enregistrer. Les anciens formats binaires .doc
    (pré-2007) ne sont pas lisibles ainsi, comme pour _convertir_fichier_en_html."""
    nom = (fichier_nom or '').lower()
    mime = (fichier_mime or '').lower()

    if nom.endswith('.docx') or 'wordprocessingml' in mime:
        import mammoth
        from io import BytesIO
        resultat = mammoth.extract_raw_text(BytesIO(fichier_data))
        return resultat.value, None

    if nom.endswith('.pdf') or 'pdf' in mime:
        from pypdf import PdfReader
        from io import BytesIO
        lecteur = PdfReader(BytesIO(fichier_data))
        pages = [(p.extract_text() or '') for p in lecteur.pages]
        return '\n'.join(pages), None

    return None, "Import non pris en charge pour ce type de fichier (.doc ancien format...) — réenregistrez-le en .docx ou PDF, ou retapez le contenu directement."


@app.route('/api/import-fichier-texte', methods=['POST'])
@login_required
def api_import_fichier_texte():
    """Extrait le texte d'un fichier Word/PDF envoyé pour pré-remplir un
    formulaire de création de protocole/ordonnance type/examen type — le
    fichier n'est ni stocké ni enregistré, seul le texte extrait est
    renvoyé pour que l'utilisateur le complète/corrige avant d'enregistrer."""
    if current_user.role not in ['admin_structure', 'medecin']:
        return jsonify({'success': False, 'error': 'Accès non autorisé'}), 403

    fichier = request.files.get('fichier')
    if not fichier or not fichier.filename:
        return jsonify({'success': False, 'error': 'Aucun fichier reçu'}), 400

    donnees = fichier.read()
    if len(donnees) > 10 * 1024 * 1024:
        return jsonify({'success': False, 'error': 'Fichier trop volumineux (max 10 Mo)'}), 400

    try:
        texte, erreur = _extraire_texte_fichier(donnees, fichier.filename, fichier.mimetype)
    except Exception as e:
        return jsonify({'success': False, 'error': f'Erreur de lecture du fichier : {e}'}), 400

    if erreur:
        return jsonify({'success': False, 'error': erreur}), 400

    lignes = [l.strip() for l in (texte or '').splitlines() if l.strip()]
    nom_suggere = lignes[0][:200] if lignes else ''

    return jsonify({
        'success': True,
        'nom_suggere': nom_suggere,
        'texte': texte or '',
        'lignes': lignes,
    })


@app.route('/api/modeles-resultats/<int:modele_id>/contenu', methods=['PUT'])
@login_required
def api_definir_contenu_modele(modele_id):
    from models import ModeleResultat
    try:
        modele = ModeleResultat.query.filter_by(id=modele_id, structure_id=current_user.id_structure).first()
        if not modele:
            return jsonify({'success': False, 'error': 'Introuvable'}), 404
        contenu_html = (request.json.get('contenu_html') or '').strip()
        if not contenu_html:
            return jsonify({'success': False, 'error': 'Contenu vide'}), 400
        modele.contenu_html = contenu_html
        db.session.commit()
        return jsonify({'success': True})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/modeles-resultats', methods=['POST'])
@login_required
def api_creer_modele_resultat():
    from models import ModeleResultat
    try:
        nom = (request.form.get('nom') or '').strip()
        type_analyse = request.form.get('type_analyse')
        fichier = request.files.get('fichier')
        contenu_html = (request.form.get('contenu_html') or '').strip()

        if not nom or not (fichier or contenu_html):
            return jsonify({'success': False, 'error': 'Nom et (fichier ou contenu) requis'}), 400
        if type_analyse not in ('BIOLOGIE', 'IMAGERIE'):
            return jsonify({'success': False, 'error': "type_analyse doit être 'BIOLOGIE' ou 'IMAGERIE'"}), 400

        modele = ModeleResultat(
            structure_id=current_user.id_structure, type_analyse=type_analyse, nom=nom,
            fichier_nom=fichier.filename if fichier else None,
            fichier_mime=fichier.mimetype if fichier else None,
            fichier_data=fichier.read() if fichier else None,
            contenu_html=contenu_html or None,
            created_by=f"{current_user.prenom} {current_user.nom}",
        )
        db.session.add(modele)
        db.session.commit()
        _envoyer_resultats_examens_ghp_immediat()
        return jsonify({'success': True, 'id': modele.id})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/modeles-resultats/<int:modele_id>/fichier', methods=['GET'])
@login_required
def api_telecharger_modele_resultat(modele_id):
    from models import ModeleResultat
    from flask import Response
    modele = ModeleResultat.query.filter_by(id=modele_id, structure_id=current_user.id_structure).first()
    if not modele:
        return "Modèle introuvable", 404
    if not modele.fichier_data:
        return "Ce modèle est rédigé en ligne, pas de fichier à télécharger.", 404
    return Response(
        modele.fichier_data, mimetype=modele.fichier_mime or 'application/octet-stream',
        headers={'Content-Disposition': f'attachment; filename="{modele.fichier_nom or "modele"}"'}
    )


@app.route('/api/modeles-resultats/<int:modele_id>', methods=['DELETE'])
@login_required
def api_supprimer_modele_resultat(modele_id):
    from models import ModeleResultat
    try:
        modele = ModeleResultat.query.filter_by(id=modele_id, structure_id=current_user.id_structure).first()
        if not modele:
            return jsonify({'success': False, 'error': 'Introuvable'}), 404
        db.session.delete(modele)
        db.session.commit()
        return jsonify({'success': True})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500


# ================================================================
# LABORATOIRE / RADIOLOGIE — Signatures électroniques (parité GHP)
# ================================================================
TITRES_LABORATOIRE = [
    'Ingénieur de laboratoire',
    'Technicien supérieur de laboratoire',
    'Biologiste',
    'Médecin Biologiste',
]


@app.route('/signatures-intervenants')
@login_required
def page_signatures_intervenants():
    if not _acces_module_resultats():
        flash('Accès non autorisé pour votre rôle.', 'danger')
        return redirect(url_for('dashboard'))
    return render_template('analyses/signatures_liste.html', titres_laboratoire=TITRES_LABORATOIRE)


@app.route('/api/signatures-intervenants', methods=['GET'])
@login_required
def api_lister_signatures_intervenants():
    from models import SignatureIntervenant
    q = SignatureIntervenant.query.filter_by(structure_id=current_user.id_structure)
    filiere = request.args.get('filiere')
    if filiere:
        q = q.filter_by(filiere=filiere)
    if request.args.get('actif_seulement'):
        q = q.filter_by(actif=True)
    lignes = q.order_by(SignatureIntervenant.filiere, SignatureIntervenant.nom).all()
    return jsonify([{
        'id': l.id, 'filiere': l.filiere, 'nom': l.nom, 'titre': l.titre,
        'actif': l.actif, 'created_at': l.created_at.strftime('%d/%m/%Y') if l.created_at else '',
    } for l in lignes])


@app.route('/api/signatures-intervenants', methods=['POST'])
@login_required
def api_creer_signature_intervenant():
    from models import SignatureIntervenant
    try:
        filiere = request.form.get('filiere')
        nom = (request.form.get('nom') or '').strip()
        titre = (request.form.get('titre') or '').strip()
        fichier = request.files.get('fichier')

        if filiere not in ('BIOLOGIE', 'IMAGERIE'):
            return jsonify({'success': False, 'error': "filiere doit être 'BIOLOGIE' ou 'IMAGERIE'"}), 400
        if not nom:
            return jsonify({'success': False, 'error': 'Le nom est obligatoire'}), 400
        if filiere == 'BIOLOGIE' and titre not in TITRES_LABORATOIRE:
            return jsonify({'success': False, 'error': 'Titre invalide pour un biologiste'}), 400
        if not fichier:
            return jsonify({'success': False, 'error': 'Image de signature requise'}), 400

        signature = SignatureIntervenant(
            structure_id=current_user.id_structure, filiere=filiere, nom=nom,
            titre=titre if filiere == 'BIOLOGIE' else None,
            signature_data=fichier.read(), signature_mime=fichier.mimetype,
            created_by=f"{current_user.prenom} {current_user.nom}",
        )
        db.session.add(signature)
        db.session.commit()
        return jsonify({'success': True, 'id': signature.id})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/signatures-intervenants/<int:signature_id>/image', methods=['GET'])
@login_required
def api_image_signature_intervenant(signature_id):
    from models import SignatureIntervenant
    from flask import Response
    signature = SignatureIntervenant.query.filter_by(id=signature_id, structure_id=current_user.id_structure).first()
    if not signature:
        return "Introuvable", 404
    return Response(signature.signature_data, mimetype=signature.signature_mime or 'image/png')


@app.route('/api/signatures-intervenants/<int:signature_id>/toggle', methods=['POST'])
@login_required
def api_toggle_signature_intervenant(signature_id):
    from models import SignatureIntervenant
    try:
        signature = SignatureIntervenant.query.filter_by(id=signature_id, structure_id=current_user.id_structure).first()
        if not signature:
            return jsonify({'success': False, 'error': 'Introuvable'}), 404
        signature.actif = not signature.actif
        db.session.commit()
        return jsonify({'success': True, 'actif': signature.actif})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/signatures-intervenants/<int:signature_id>', methods=['DELETE'])
@login_required
def api_supprimer_signature_intervenant(signature_id):
    from models import SignatureIntervenant
    try:
        signature = SignatureIntervenant.query.filter_by(id=signature_id, structure_id=current_user.id_structure).first()
        if not signature:
            return jsonify({'success': False, 'error': 'Introuvable'}), 404
        db.session.delete(signature)
        db.session.commit()
        return jsonify({'success': True})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/analyse/<int:id>/imprimer')
@login_required
def imprimer_resultat_analyse(id):
    from models import AnalyseDemande, Structure
    from datetime import datetime
    import re
    
    analyse = AnalyseDemande.query.get_or_404(id)
    structure = Structure.query.get(current_user.id_structure)
    
    # ⭐ NETTOYER LE NOM DE L'ANALYSE
    nom_analyse = analyse.nom_analyse
    
    # Règle : supprimer le code qui commence par une ou plusieurs lettres suivies de chiffres
    # Exemples: 
    #   "Q426 - TDM CEREBRALE" → "TDM CEREBRALE"
    #   "Q426 TDM CEREBRALE" → "TDM CEREBRALE"
    #   "Q100 Examen radiologique du doigt" → "Examen radiologique du doigt"
    #   "B12 - NFS" → "NFS"
    #   "TDM CEREBRALE" (sans code) → "TDM CEREBRALE"
    #
    # Pattern: une ou plusieurs lettres majuscules + un ou plusieurs chiffres 
    #          + optionnel espace + optionnel tiret + optionnel espace
    pattern = r'^[A-Z]+[0-9]+\s*-?\s*'
    nom_analyse = re.sub(pattern, '', nom_analyse)

    # ⭐ FIX : jusqu'ici, un résultat sous forme de fichier joint (Word/Excel/
    # PDF) affichait juste "Fichier joint" + un lien "Ouvrir" à l'impression —
    # le contenu lui-même n'était jamais visible/imprimable directement dans
    # la page. Pour Word/Excel : conversion en HTML à la volée (même fonction
    # que pour les modèles de résultats, voir _convertir_fichier_en_html
    # ci-dessus) et rendu en flux normal, comme un résultat rédigé en ligne —
    # pagination correcte à l'impression. Pour un PDF : rendu page par page
    # via PDF.js directement dans le flux (jamais un <iframe>, qui n'imprime
    # que ce qui est visible à l'instant du clic), voir le script en bas du
    # template.
    fichier_converti_html = None
    fichier_est_pdf = False
    if analyse.fichier_data and not analyse.contenu_html:
        mime = (analyse.fichier_mime or '').lower()
        nom_fichier = (analyse.fichier_nom or '').lower()
        if 'pdf' in mime or nom_fichier.endswith('.pdf'):
            fichier_est_pdf = True
        else:
            try:
                html, erreur = _convertir_fichier_en_html(analyse.fichier_data, analyse.fichier_nom, analyse.fichier_mime)
                if not erreur:
                    fichier_converti_html = html
            except Exception:
                fichier_converti_html = None

    return render_template('impressions/resultat.html',
                         analyse=analyse,
                         structure=structure,
                         nom_analyse=nom_analyse,  # ⭐ NOM NETTOYÉ
                         fichier_converti_html=fichier_converti_html,
                         fichier_est_pdf=fichier_est_pdf,
                         now=datetime.utcnow())

@app.route('/consultation/<int:id>/analyse/ajouter', methods=['POST'])
@login_required
def ajouter_analyse_demande(id):
    """Le médecin ajoute une demande d'analyse"""
    from models import Consultation, AnalyseDemande
    
    consultation = Consultation.query.get_or_404(id)
    
    if current_user.role not in ['medecin', 'admin_structure']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    type_analyse = request.form.get('type_analyse')
    nom_analyse = request.form.get('nom_analyse')
    description = request.form.get('description')
    
    if not type_analyse or not nom_analyse:
        flash('Le type et le nom de l\'analyse sont obligatoires', 'danger')
        return redirect(url_for('consultation_detail', id=id))
    
    analyse = AnalyseDemande(
        consultation_id=consultation.id,
        patient_id=consultation.id_patient,
        structure_id=current_user.id_structure,
        type_analyse=type_analyse,
        nom_analyse=nom_analyse,
        description=description,
        prescrit_par=current_user.id,
        statut='EN_ATTENTE'
    )
    
    db.session.add(analyse)
    db.session.commit()
    
    flash(f'✅ Demande d\'analyse "{nom_analyse}" envoyée au laborantin', 'success')
    return redirect(url_for('consultation_detail', id=id))

@app.route('/patient/<int:patient_id>/analyses')
@login_required
def patient_analyses(patient_id):
    from models import Patient, AnalyseDemande
    
    patient = Patient.query.get_or_404(patient_id)
    
    if current_user.role not in ['super_admin']:
        if patient.id_structure != current_user.id_structure:
            flash('Accès non autorisé', 'danger')
            return redirect(url_for('liste_analyses'))
    
    # ⭐ REQUÊTE DE BASE
    query = AnalyseDemande.query.filter_by(
        patient_id=patient_id,
        structure_id=current_user.id_structure
    )

    # ⭐ FILTRER SELON LE RÔLE
    if current_user.role == 'laborantin':
        query = query.filter(AnalyseDemande.type_analyse == 'BIOLOGIE')
    elif current_user.role == 'radiologue':
        query = query.filter(AnalyseDemande.type_analyse == 'IMAGERIE')
    else:
        # ⭐ FIX : un admin_structure/médecin (qui voit les deux filières)
        # arrivait ici sans AUCUN filtre, peu importe qu'il vienne de
        # "Voir toutes les analyses" (liste_analyses, biologie) ou "Voir
        # tous les examens" (liste_radiologie, imagerie) — les deux
        # listes distinctes menaient au même mélange biologie+imagerie.
        # Ces deux pages passent maintenant explicitement leur filière
        # (?type_analyse=BIOLOGIE|IMAGERIE) ; sans ce paramètre (ex. lien
        # "Résultats labo/radio" du dossier patient), la vue reste
        # volontairement complète, comme avant.
        type_analyse_filtre = request.args.get('type_analyse')
        if type_analyse_filtre in ('BIOLOGIE', 'IMAGERIE'):
            query = query.filter(AnalyseDemande.type_analyse == type_analyse_filtre)

    analyses = query.order_by(AnalyseDemande.date_demande.desc()).all()

    return render_template('analyses/patient_analyses.html',
                         patient=patient,
                         analyses=analyses,
                         type_analyse_filtre=request.args.get('type_analyse'))

@app.route('/consultation/<int:id>/reference/ajouter', methods=['GET', 'POST'])
@login_required
def ajouter_reference(id):
    """Ajouter une référence depuis une consultation"""
    from models import Consultation, Reference, Patient
    
    consultation = Consultation.query.get_or_404(id)
    patient = Patient.query.get(consultation.id_patient)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    if request.method == 'POST':
        motif = request.form.get('motif')
        diagnostic = request.form.get('diagnostic')
        centre_reference = request.form.get('centre_reference')
        service_reference = request.form.get('service_reference')
        medecin_referent = request.form.get('medecin_referent')
        resume_clinique = request.form.get('resume_clinique')
        examens_realises = request.form.get('examens_realises')
        traitements_en_cours = request.form.get('traitements_en_cours')
        
        if not motif or not centre_reference:
            flash('Le motif et le centre de référence sont obligatoires', 'danger')
            return redirect(url_for('ajouter_reference', id=id))
        
        # Créer la référence avec les dernières constantes du patient
        reference = Reference(
            patient_id=patient.id,
            consultation_id=consultation.id,
            structure_id=current_user.id_structure,
            motif=motif,
            diagnostic=diagnostic or consultation.diagnostic,
            centre_reference=centre_reference,
            service_reference=service_reference,
            medecin_referent=medecin_referent,
            derniere_tension=patient.tension_arterielle,
            derniere_temperature=patient.temperature_c,
            derniere_pulse=patient.pulse_bpm,
            derniere_saturation=patient.oxygene_saturation,
            dernier_poids=patient.poids_kg,
            derniere_taille=patient.taille_cm,
            dernier_imc=patient.imc,
            resume_clinique=resume_clinique,
            examens_realises=examens_realises,
            traitements_en_cours=traitements_en_cours,
            statut='ENVOYE',
            created_by=current_user.id
        )
        
        db.session.add(reference)
        db.session.commit()
        
        flash('✅ Référence créée avec succès', 'success')
        return redirect(url_for('imprimer_reference', id=reference.id))
    
    return render_template('consultations/ajouter_reference.html',
                         consultation=consultation,
                         patient=patient)
@app.route('/reference/<int:id>/imprimer')
@login_required
def imprimer_reference(id):
    """Imprimer une fiche de référence"""
    from models import Reference, Structure
    from datetime import datetime
    
    reference = Reference.query.get_or_404(id)
    
    if current_user.role not in ['super_admin', 'admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    if reference.structure_id != current_user.id_structure and current_user.role != 'super_admin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # ⭐ RÉCUPÉRER LA STRUCTURE POUR LE LOGO
    structure = Structure.query.get(current_user.id_structure)
    
    return render_template('impressions/reference.html',
                         reference=reference,
                         structure=structure,
                         now=datetime.utcnow())

# ==================== GESTION DES RÉFÉRENCES ====================

@app.route('/references')
@login_required
@has_permission('REFERENCE')  # ⭐ AJOUTER LE DÉCORATEUR
def liste_references():
    """Liste des références effectuées"""
    from models import Reference, Patient
    from sqlalchemy import or_
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Filtres
    search = request.args.get('search', '')
    statut = request.args.get('statut', '')
    
    query = Reference.query.filter_by(structure_id=current_user.id_structure)
    
    # ⭐ RECHERCHE CORRIGÉE
    if search:
        search = search.strip()
        filters = []
        
        # Recherche par nom ou prénom
        filters.append(Patient.nom.ilike(f'%{search}%'))
        filters.append(Patient.prenom.ilike(f'%{search}%'))
        
        # Recherche par numéro de dossier (P00001)
        if search.upper().startswith('P'):
            try:
                num = int(search[1:])
                filters.append(Patient.id == num)
            except ValueError:
                pass
        elif search.isdigit():
            filters.append(Patient.id == int(search))
        
        # Recherche par numéro de référence (REF-00001)
        if search.upper().startswith('REF'):
            try:
                num = int(search[3:])
                filters.append(Reference.id == num)
            except ValueError:
                pass
        
        # Appliquer les filtres
        if filters:
            query = query.join(Patient).filter(or_(*filters))
    
    # Filtre par statut
    if statut:
        query = query.filter_by(statut=statut)
    
    references = query.order_by(Reference.date_reference.desc()).all()
    statuts = ['ENVOYE', 'ACCEPTE', 'REFUSE', 'EN_ATTENTE']
    
    return render_template('references/liste.html',
                         references=references,
                         search=search,
                         statut_actuel=statut,
                         statuts=statuts)

@app.route('/reference/<int:id>/reimprimer')
@login_required
def reimprimer_reference(id):
    """Réimprimer une référence existante"""
    from models import Reference
    
    reference = Reference.query.get_or_404(id)
    
    if current_user.role not in ['super_admin', 'admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    if reference.structure_id != current_user.id_structure and current_user.role != 'super_admin':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_references'))
    
    return render_template('references/imprimer.html', reference=reference)
@app.route('/reference/<int:id>/suivi', methods=['POST'])
@login_required
def suivi_reference(id):
    """Mettre à jour le suivi d'une référence"""
    from models import Reference
    from datetime import datetime
    
    reference = Reference.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    if reference.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_references'))
    
    statut = request.form.get('statut')
    retour_info = request.form.get('retour_info')
    
    if statut:
        reference.statut = statut
        if statut in ['ACCEPTE', 'REFUSE']:
            reference.date_retour = datetime.utcnow()
    
    if retour_info:
        reference.retour_info = retour_info
    
    db.session.commit()
    
    flash('✅ Suivi mis à jour', 'success')
    return redirect(url_for('liste_references'))
# ==================== PERMISSIONS TEMPORAIRES ====================

@app.route('/structure/permissions')
@login_required
def gestion_permissions():
    """Gestion des permissions temporaires"""
    if current_user.role not in ['admin_structure']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Utilisateur, PermissionTemp
    from datetime import datetime
    
    # Liste des utilisateurs de la structure (sauf admin)
    users = Utilisateur.query.filter(
        Utilisateur.id_structure == current_user.id_structure,
        Utilisateur.role != 'admin_structure',
        Utilisateur.actif == True
    ).all()
    
    # Permissions actives
    permissions_actives = PermissionTemp.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    # Historique
    historique = PermissionTemp.query.filter_by(
        structure_id=current_user.id_structure
    ).order_by(PermissionTemp.created_at.desc()).limit(50).all()
    
    # Liste des permissions disponibles
    permissions_list = [
        {'key': 'ANALYSES', 'label': '🧪 Accès Analyses', 'description': 'Voir et saisir les résultats d\'analyses'},
        {'key': 'REFERENCE', 'label': '🚑 Accès Références', 'description': 'Créer et gérer les références'},
        {'key': 'HOSPITALISATION', 'label': '🏥 Accès Hospitalisations', 'description': 'Gérer les hospitalisations'},
        {'key': 'STATISTIQUES', 'label': '📊 Accès Statistiques', 'description': 'Voir les statistiques'},
        {'key': 'PATIENTS', 'label': '👤 Accès Patients', 'description': 'Voir et modifier les patients'},
    ]
    
    # ⭐ PASSER now AU TEMPLATE
    return render_template('structure/permissions.html',
                         users=users,
                         permissions_actives=permissions_actives,
                         historique=historique,
                         permissions_list=permissions_list,
                         now=datetime.utcnow())  # ⭐ AJOUTER CETTE LIGNE

@app.route('/structure/permissions/ajouter', methods=['POST'])
@login_required
def ajouter_permission():
    """Ajouter une permission temporaire"""
    if current_user.role not in ['admin_structure']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import Utilisateur, PermissionTemp
    from datetime import datetime, timedelta
    
    user_id = request.form.get('user_id')
    permission = request.form.get('permission')
    duree = request.form.get('duree', 24)  # Heures par défaut
    motif = request.form.get('motif', '')
    
    if not user_id or not permission:
        flash('Veuillez sélectionner un utilisateur et une permission', 'danger')
        return redirect(url_for('gestion_permissions'))
    
    # Vérifier que l'utilisateur est dans la structure
    user = Utilisateur.query.get(user_id)
    if not user or user.id_structure != current_user.id_structure:
        flash('Utilisateur non trouvé', 'danger')
        return redirect(url_for('gestion_permissions'))
    
    # Vérifier si une permission active existe déjà
    existing = PermissionTemp.query.filter_by(
        user_id=user_id,
        permission=permission,
        actif=True
    ).first()
    
    if existing:
        flash(f'{user.prenom} {user.nom} a déjà cette permission active', 'warning')
        return redirect(url_for('gestion_permissions'))
    
    # Créer la permission
    permission_temp = PermissionTemp(
        user_id=user_id,
        granted_by=current_user.id,
        structure_id=current_user.id_structure,
        permission=permission,
        date_debut=datetime.utcnow(),
        date_fin=datetime.utcnow() + timedelta(hours=int(duree)),
        motif=motif
    )
    
    db.session.add(permission_temp)
    db.session.commit()
    
    flash(f'✅ Permission "{permission}" accordée à {user.prenom} {user.nom} pour {duree}h', 'success')
    return redirect(url_for('gestion_permissions'))


@app.route('/structure/permissions/revoke/<int:id>', methods=['POST'])
@login_required
def revoke_permission(id):
    """Révoquer une permission temporaire"""
    if current_user.role not in ['admin_structure']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    from models import PermissionTemp
    
    permission = PermissionTemp.query.get_or_404(id)
    
    if permission.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('gestion_permissions'))
    
    motif = request.form.get('motif_revocation', 'Révoquée par admin')
    
    permission.actif = False
    permission.date_revocation = datetime.utcnow()
    permission.revoked_by = current_user.id
    permission.motif_revocation = motif
    
    db.session.commit()
    
    flash(f'✅ Permission révoquée avec succès', 'success')
    return redirect(url_for('gestion_permissions'))


@app.route('/api/permissions/check')
@login_required
def check_permission():
    """Vérifier si l'utilisateur a une permission (pour les routes)"""
    permission = request.args.get('permission', '')
    
    if not permission:
        return jsonify({'has_permission': False})
    
    from models import PermissionTemp
    from datetime import datetime
    
    has_permission = PermissionTemp.query.filter_by(
        user_id=current_user.id,
        permission=permission,
        actif=True
    ).filter(
        PermissionTemp.date_fin > datetime.utcnow()
    ).first()
    
    # Vérifier le rôle de base
    role_permissions = {
        'medecin': ['PATIENTS', 'REFERENCE', 'HOSPITALISATION', 'STATISTIQUES'],
        'infirmier': ['PATIENTS', 'HOSPITALISATION'],
        'laborantin': ['ANALYSES'],
        'admin_structure': ['PATIENTS', 'REFERENCE', 'HOSPITALISATION', 'STATISTIQUES', 'ANALYSES']
    }
    
    base_permissions = role_permissions.get(current_user.role, [])
    
    return jsonify({
        'has_permission': bool(has_permission) or permission in base_permissions,
        'permission': permission
    })
@app.route('/reference/nouvelle', methods=['GET', 'POST'])
@login_required
def nouvelle_reference():
    """Créer une référence directement depuis l'onglet Références"""
    from models import Patient, Consultation, Reference
    from sqlalchemy import or_
    from datetime import datetime
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Récupérer les paramètres de recherche
    search = request.args.get('search', '')
    
    # Requête de base
    query = Patient.query.filter_by(
        id_structure=current_user.id_structure,
        archived=False
    )
    
    # Filtrer par recherche
    if search:
        query = query.filter(
            or_(
                Patient.nom.ilike(f'%{search}%'),
                Patient.prenom.ilike(f'%{search}%'),
                Patient.id.cast().ilike(f'%{search}%')
            )
        )
    
    patients = query.order_by(Patient.nom).all()
    
    if request.method == 'POST':
        patient_id = request.form.get('patient_id')
        motif = request.form.get('motif')
        diagnostic = request.form.get('diagnostic')
        centre_reference = request.form.get('centre_reference')
        service_reference = request.form.get('service_reference')
        medecin_referent = request.form.get('medecin_referent')
        resume_clinique = request.form.get('resume_clinique')
        examens_realises = request.form.get('examens_realises')
        traitements_en_cours = request.form.get('traitements_en_cours')
        
        if not patient_id or not motif or not centre_reference:
            flash('Le patient, le motif et le centre de référence sont obligatoires', 'danger')
            return redirect(url_for('nouvelle_reference'))
        
        patient = Patient.query.get(patient_id)
        
        # Créer une consultation automatique
        consultation = Consultation(
            id_patient=patient.id,
            id_medecin=current_user.id,
            motif=f"Référence vers {centre_reference}",
            diagnostic=diagnostic,
            date_consultation=datetime.utcnow()
        )
        db.session.add(consultation)
        db.session.flush()
        
        # ⭐ RÉCUPÉRER LE DERNIER DIAGNOSTIC CORRECTEMENT
        dernier_diagnostic = None
        if patient.consultations:
            # Trier les consultations par date et prendre la plus récente
            consultations_triees = sorted(patient.consultations, key=lambda c: c.date_consultation, reverse=True)
            if consultations_triees and consultations_triees[0].diagnostic:
                dernier_diagnostic = consultations_triees[0].diagnostic
        
        # Créer la référence
        reference = Reference(
            patient_id=patient.id,
            consultation_id=consultation.id,
            structure_id=current_user.id_structure,
            motif=motif,
            diagnostic=diagnostic or dernier_diagnostic,
            centre_reference=centre_reference,
            service_reference=service_reference,
            medecin_referent=medecin_referent,
            derniere_tension=patient.tension_arterielle,
            derniere_temperature=patient.temperature_c,
            derniere_pulse=patient.pulse_bpm,
            derniere_saturation=patient.oxygene_saturation,
            dernier_poids=patient.poids_kg,
            derniere_taille=patient.taille_cm,
            dernier_imc=patient.imc,
            resume_clinique=resume_clinique,
            examens_realises=examens_realises,
            traitements_en_cours=traitements_en_cours,
            statut='ENVOYE',
            created_by=current_user.id
        )
        
        db.session.add(reference)
        db.session.commit()
        
        flash('✅ Référence créée avec succès', 'success')
        return redirect(url_for('imprimer_reference', id=reference.id))
    
    return render_template('references/nouvelle.html',
                         patients=patients,
                         search=search)

@app.route('/api/patient/<int:patient_id>/constantes')
@login_required
def api_patient_constantes(patient_id):
    """Récupérer les constantes d'un patient (pour AJAX)"""
    from models import Patient
    
    patient = Patient.query.get_or_404(patient_id)
    
    return jsonify({
        'tension': patient.tension_arterielle,
        'temperature': patient.temperature_c,
        'pouls': patient.pulse_bpm,
        'saturation': patient.oxygene_saturation,
        'poids': patient.poids_kg,
        'taille': patient.taille_cm,
        'imc': patient.imc
    })
@app.route('/api/patients/search')
@login_required
def api_patients_search():
    """Recherche de patients pour autocomplétion"""
    q = request.args.get('q', '').strip()
    if len(q) < 2:
        return jsonify([])

    patients = Patient.query.filter(
        Patient.id_structure == current_user.id_structure,
        Patient.archived == False,
        _patient_recherche_conditions(q)
    ).limit(20).all()
    
    result = []
    for p in patients:
        result.append({
            'id': p.id,
            'nom': p.nom,
            'prenom': p.prenom,
            'telephone': p.telephone,
            'id_padded': f"{p.id:05d}"
        })
    
    return jsonify(result)
@app.context_processor
def utility_processor():
    from models import PermissionTemp
    from datetime import datetime
    
    def has_temp_permission(permission):
        if not current_user.is_authenticated:
            return False
        
        # Vérifier les permissions temporaires
        temp_perm = PermissionTemp.query.filter(
            PermissionTemp.user_id == current_user.id,
            PermissionTemp.permission == permission,
            PermissionTemp.actif == True,
            PermissionTemp.date_fin > datetime.utcnow()
        ).first()
        
        return temp_perm is not None
    
    return dict(has_temp_permission=has_temp_permission)
# ==================== GESTION DES SALLES ====================

# ==================== GESTION DES SALLES ====================

@app.route('/salles')
@login_required
def liste_salles():
    """Liste des salles par service"""
    from models import Service, Salle, Lit  # ⭐ AJOUTER CET IMPORT
    
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    services = Service.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    return render_template('salles/liste.html', services=services)


@app.route('/salles/service/ajouter', methods=['GET', 'POST'])
@login_required
def ajouter_service():
    """Ajouter un service"""
    from models import Service  # ⭐ AJOUTER CET IMPORT
    
    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    if request.method == 'POST':
        nom = request.form.get('nom')
        description = request.form.get('description')
        
        if not nom:
            flash('Le nom du service est obligatoire', 'danger')
            return redirect(url_for('ajouter_service'))
        
        service = Service(
            structure_id=current_user.id_structure,
            nom=nom,
            description=description
        )
        db.session.add(service)
        db.session.commit()
        
        flash(f'Service "{nom}" créé avec succès', 'success')
        return redirect(url_for('liste_salles'))
    
    return render_template('salles/ajouter_service.html')


@app.route('/salles/salle/ajouter', methods=['GET', 'POST'])
@login_required
def ajouter_salle():
    """Ajouter une salle"""
    from models import Service, Salle, Lit  # ⭐ AJOUTER CES IMPORTS
    
    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    services = Service.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    if request.method == 'POST':
        service_id = request.form.get('service_id')
        nom = request.form.get('nom')
        type_salle = request.form.get('type_salle')
        nombre_lits = request.form.get('nombre_lits', type=int)
        prix_journalier = request.form.get('prix_journalier', type=float)
        description = request.form.get('description')
        # ⭐ Renseignés automatiquement par le sélecteur "Tarif GHP" côté JS
        # (un type de salle du catalogue GHP -> les 3 paliers d'un coup) —
        # jamais de saisie libre pour ces 3 champs, voir ajouter_salle.html.
        acte_ghp_semaine1 = (request.form.get('acte_ghp_semaine1') or '').strip() or None
        acte_ghp_semaine2 = (request.form.get('acte_ghp_semaine2') or '').strip() or None
        acte_ghp_semaine3 = (request.form.get('acte_ghp_semaine3') or '').strip() or None

        if not service_id or not nom or not type_salle or not nombre_lits:
            flash('Tous les champs obligatoires doivent être remplis', 'danger')
            return redirect(url_for('ajouter_salle'))

        salle = Salle(
            service_id=int(service_id),
            nom=nom,
            type_salle=type_salle,
            nombre_lits=nombre_lits,
            prix_journalier=prix_journalier,
            description=description,
            acte_ghp_semaine1=acte_ghp_semaine1,
            acte_ghp_semaine2=acte_ghp_semaine2,
            acte_ghp_semaine3=acte_ghp_semaine3
        )
        db.session.add(salle)
        db.session.flush()
        
        # Créer les lits
        for i in range(nombre_lits):
            lit = Lit(
                salle_id=salle.id,
                numero=chr(65 + i)  # A, B, C, D, ...
            )
            db.session.add(lit)
        
        db.session.commit()

        flash(f'Salle "{nom}" créée avec {nombre_lits} lits', 'success')
        if not acte_ghp_semaine1:
            flash(
                f'Aucun tarif GHP choisi pour "{nom}" — la facturation automatique à la '
                f'clôture d\'une hospitalisation ne fonctionnera pas tant que ce n\'est pas '
                f'configuré (modifiez la salle pour le faire).',
                'warning'
            )
        return redirect(url_for('liste_salles'))
    
    return render_template('salles/ajouter_salle.html', services=services)


@app.route('/salles/salle/<int:id>')
@login_required
def detail_salle(id):
    """Détail d'une salle avec ses lits"""
    from models import Salle, Lit, Hospitalisation, Patient  # ⭐ AJOUTER CES IMPORTS
    
    salle = Salle.query.get_or_404(id)
    
    if salle.service.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_salles'))
    
    lits = Lit.query.filter_by(salle_id=id).all()
    
    # Récupérer les infos des patients pour les lits occupés
    for lit in lits:
        if lit.statut == 'occupe' and lit.hospitalisation_id:
            hospitalisation = Hospitalisation.query.get(lit.hospitalisation_id)
            if hospitalisation:
                lit.patient = hospitalisation.patient
    
    return render_template('salles/detail_salle.html', salle=salle, lits=lits)


@app.route('/salles/salle/<int:id>/modifier', methods=['GET', 'POST'])
@login_required
def modifier_salle(id):
    """Modifier une salle — notamment son tarif GHP (paliers d'hospitalisation),
    seul moyen désormais de le configurer/corriger après création (l'ancien
    onglet Paramétrage AMU ne fait plus que les seuils de jours, communs à
    la structure)."""
    from models import Service, Salle

    salle = Salle.query.get_or_404(id)
    if salle.service.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_salles'))

    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_salle', id=id))

    services = Service.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()

    if request.method == 'POST':
        service_id = request.form.get('service_id')
        nom = request.form.get('nom')
        type_salle = request.form.get('type_salle')
        prix_journalier = request.form.get('prix_journalier', type=float)
        description = request.form.get('description')
        acte_ghp_semaine1 = (request.form.get('acte_ghp_semaine1') or '').strip() or None
        acte_ghp_semaine2 = (request.form.get('acte_ghp_semaine2') or '').strip() or None
        acte_ghp_semaine3 = (request.form.get('acte_ghp_semaine3') or '').strip() or None

        if not service_id or not nom or not type_salle:
            flash('Tous les champs obligatoires doivent être remplis', 'danger')
            return redirect(url_for('modifier_salle', id=id))

        salle.service_id = int(service_id)
        salle.nom = nom
        salle.type_salle = type_salle
        salle.prix_journalier = prix_journalier
        salle.description = description
        salle.acte_ghp_semaine1 = acte_ghp_semaine1
        salle.acte_ghp_semaine2 = acte_ghp_semaine2
        salle.acte_ghp_semaine3 = acte_ghp_semaine3
        db.session.commit()

        flash(f'Salle "{nom}" mise à jour', 'success')
        if not acte_ghp_semaine1:
            flash(
                f'Aucun tarif GHP choisi pour "{nom}" — la facturation automatique à la '
                f'clôture d\'une hospitalisation ne fonctionnera pas tant que ce n\'est pas configuré.',
                'warning'
            )
        return redirect(url_for('detail_salle', id=id))

    return render_template('salles/modifier_salle.html', salle=salle, services=services)


@app.route('/hospitalisation/<int:id>/assigner-lit', methods=['POST'])
@login_required
def assigner_lit(id):
    """Assigner un lit à une hospitalisation"""
    from models import Hospitalisation, Lit
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    lit_id = request.form.get('lit_id', type=int)
    
    if not lit_id:
        flash('Veuillez sélectionner un lit', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    lit = Lit.query.get(lit_id)
    
    if not lit or lit.statut != 'disponible':
        flash('Ce lit n\'est pas disponible', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    # Occuper le lit
    lit.occuper(id)
    hospitalisation.lit_id = lit.id
    hospitalisation.chambre = lit.salle.nom
    hospitalisation.lit = lit.numero
    
    db.session.commit()
    
    flash(f'Lit {lit.salle.nom} - {lit.numero} attribué avec succès', 'success')
    return redirect(url_for('detail_hospitalisation', id=id))

@app.route('/api/salles/disponibles')
@login_required
def api_salles_disponibles():
    from models import Salle, Service, Lit  # ⭐ TOUS LES IMPORTS
    
    service_nom = request.args.get('service', '')
    
    if not service_nom:
        return jsonify([])
    
    # Récupérer les salles du service
    salles = Salle.query.join(Service).filter(
        Service.nom == service_nom,
        Service.structure_id == current_user.id_structure,
        Salle.actif == True
    ).all()
    
    result = []
    for salle in salles:
        # Compter les lits disponibles
        lits_disponibles = Lit.query.filter_by(
            salle_id=salle.id,
            statut='disponible'
        ).count()
        
        result.append({
            'id': salle.id,
            'nom': salle.nom,
            'type': salle.type_salle,
            'lits_disponibles': lits_disponibles,
            'prix': salle.prix_journalier
        })
    
    return jsonify(result)

@app.route('/api/lits/disponibles')
@login_required
def api_lits_disponibles():
    """Récupérer les lits disponibles d'une salle"""
    from models import Lit
    
    salle_id = request.args.get('salle_id', type=int)
    
    if not salle_id:
        return jsonify([])
    
    lits = Lit.query.filter_by(
        salle_id=salle_id,
        statut='disponible'
    ).all()
    
    result = []
    for lit in lits:
        result.append({
            'id': lit.id,
            'numero': lit.numero
        })
    
    return jsonify(result)
# ==================== ANTÉCÉDENTS PATIENT ====================

@app.route('/patient/<int:patient_id>/antecedents')
@login_required
def patient_antecedents(patient_id):
    from models import Patient, AntecedentPatient
    
    patient = Patient.query.get_or_404(patient_id)
    antecedents = AntecedentPatient.query.filter_by(patient_id=patient_id).order_by(AntecedentPatient.date_recueil.desc()).all()
    
    # ⭐ Récupérer les paramètres de retour
    return_to = request.args.get('return_to')
    consultation_id = request.args.get('consultation_id')
    
    return render_template('patients/antecedents.html',
                         patient=patient,
                         antecedents=antecedents,
                         return_to=return_to,
                         consultation_id=consultation_id,
                         current_user=current_user)

@app.route('/patient/<int:patient_id>/antecedent/ajouter', methods=['POST'])
@login_required
def ajouter_antecedent(patient_id):
    from models import AntecedentPatient
    from datetime import datetime, timezone
    
    patient = Patient.query.get_or_404(patient_id)
    
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    type_antecedent = request.form.get('type_antecedent')
    type_precision = request.form.get('type_precision')
    description = request.form.get('description')
    severite = request.form.get('severite')
    date_debut = request.form.get('date_debut')
    date_fin = request.form.get('date_fin')
    traitement = request.form.get('traitement')
    notes = request.form.get('notes')
    actif = request.form.get('actif') == 'on'
    
    antecedent = AntecedentPatient(
        patient_id=patient_id,
        type_antecedent=type_antecedent,
        type_precision=type_precision if type_antecedent == 'AUTRE' else None,
        description=description,
        severite=severite if severite else None,
        date_debut=datetime.strptime(date_debut, '%Y-%m-%d').date() if date_debut else None,
        date_fin=datetime.strptime(date_fin, '%Y-%m-%d').date() if date_fin else None,
        traitement=traitement if traitement else None,
        notes=notes if notes else None,
        actif=actif,
        recueilli_par=current_user.id,
        date_recueil=datetime.now(timezone.utc)
    )
    
    db.session.add(antecedent)
    db.session.commit()
    
    # ⭐ AJOUTER UN PARAMÈTRE POUR DIFFÉRENCIER AJAX
    format = request.args.get('format')
    if format == 'json':
        return jsonify({
            'success': True,
            'message': 'Antécédent ajouté avec succès',
            'antecedent_id': antecedent.id
        })
    
    flash('✅ Antécédent ajouté avec succès', 'success')
    
    return_to = request.form.get('return_to')
    if return_to == 'pre_consultation' or current_user.role == 'infirmier':
        return redirect(url_for('infirmier_pre_consultation', patient_id=patient_id))
    else:
        return redirect(url_for('patient_antecedents', patient_id=patient_id))


@app.route('/antecedent/<int:id>/modifier', methods=['POST'])
@login_required
def modifier_antecedent(id):
    """Modifier un antécédent (médecin ou infirmier)"""
    from models import AntecedentPatient
    
    antecedent = AntecedentPatient.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Mise à jour des champs
    antecedent.type_antecedent = request.form.get('type_antecedent')
    antecedent.description = request.form.get('description')
    date_debut = request.form.get('date_debut')
    date_fin = request.form.get('date_fin')
    antecedent.actif = request.form.get('actif') == 'on'
    antecedent.severite = request.form.get('severite')
    antecedent.traitement = request.form.get('traitement')
    antecedent.notes = request.form.get('notes')
    antecedent.modified_by = current_user.id
    antecedent.modified_at = datetime.utcnow()
    
    if date_debut:
        antecedent.date_debut = datetime.strptime(date_debut, '%Y-%m-%d')
    if date_fin:
        antecedent.date_fin = datetime.strptime(date_fin, '%Y-%m-%d')
    
    db.session.commit()
    
    flash('✅ Antécédent modifié avec succès', 'success')
    return redirect(url_for('patient_antecedents', patient_id=antecedent.patient_id))


@app.route('/antecedent/<int:id>/supprimer', methods=['POST'])
@login_required
def supprimer_antecedent(id):
    from models import AntecedentPatient
    
    antecedent = AntecedentPatient.query.get_or_404(id)
    patient_id = antecedent.patient_id
    
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    db.session.delete(antecedent)
    db.session.commit()
    
    flash('✅ Antécédent supprimé avec succès', 'success')
    
    # ⭐ REDIRECTION : Si c'est un infirmier, retour à la pré-consultation
    if current_user.role == 'infirmier':
        return redirect(url_for('infirmier_pre_consultation', patient_id=patient_id))
    else:
        return redirect(url_for('patient_antecedents', patient_id=patient_id))


@app.route('/patient/<int:patient_id>/habitudes_vie', methods=['POST'])
@login_required
def modifier_habitudes_vie(patient_id):
    from flask import jsonify
    
    patient = Patient.query.get_or_404(patient_id)
    
    patient.tabac = request.form.get('tabac')
    patient.alcool = request.form.get('alcool')
    patient.groupe_sanguin = request.form.get('groupe_sanguin')
    patient.medecin_traitant = request.form.get('medecin_traitant')
    patient.mutuelle = request.form.get('mutuelle')
    patient.allaitement = request.form.get('allaitement') == 'on'
    patient.grossesse = request.form.get('grossesse') == 'on'
    
    db.session.commit()
    
    # ⭐ SI C'EST UNE REQUÊTE AJAX, RETOURNER JSON
    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return jsonify({'success': True, 'message': 'Habitudes de vie mises à jour'})
    
    # SINON, REDIRIGER NORMALEMENT
    flash('Habitudes de vie mises à jour avec succès', 'success')
    
    return_to = request.form.get('return_to') or request.args.get('return_to')
    consultation_id = request.form.get('consultation_id') or request.args.get('consultation_id')
    
    if return_to == 'consultation' and consultation_id:
        return redirect(url_for('consultation_detail', id=consultation_id))
    else:
        return redirect(url_for('patient_antecedents', patient_id=patient_id))

# Dans app.py
@app.route('/api/patient/<int:patient_id>/habitudes_vie')
@login_required
def api_patient_habitudes_vie(patient_id):
    """Récupère les habitudes de vie d'un patient"""
    from models import Patient
    
    patient = Patient.query.get_or_404(patient_id)
    
    return jsonify({
        'tabac': patient.tabac,
        'alcool': patient.alcool,
        'allaitement': patient.allaitement,
        'grossesse': patient.grossesse,
        'groupe_sanguin': patient.groupe_sanguin,
        'mutuelle': patient.mutuelle,
        'medecin_traitant': patient.medecin_traitant
    })

@app.route('/api/patient/<int:patient_id>/antecedents')
@login_required
def api_patient_antecedents(patient_id):
    """API pour récupérer les antécédents (pour le formulaire consultation)"""
    from models import AntecedentPatient
    
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        return jsonify([])
    
    antecedents = AntecedentPatient.query.filter_by(
        patient_id=patient_id,
        actif=True
    ).order_by(AntecedentPatient.date_recueil.desc()).all()
    
    result = []
    for a in antecedents:
        result.append({
            'id': a.id,
            'type': a.type_antecedent,
            'description': a.description,
            'date_debut': a.date_debut.strftime('%d/%m/%Y') if a.date_debut else None,
            'severite': a.severite,
            'traitement': a.traitement,
            'recueilli_par': f"{a.recueillant.prenom} {a.recueillant.nom}" if a.recueillant else 'Inconnu'
        })
    
    return jsonify(result)

@app.route('/patient/<int:id>/constante/ajouter', methods=['POST'])
@login_required
def ajouter_constante_patient(id):
    """Ajouter une nouvelle constante pour un patient (infirmier)"""
    from models import Patient
    from datetime import datetime
    
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    patient = Patient.query.get_or_404(id)
    
    temperature = request.form.get('temperature')
    tension = request.form.get('tension')
    pouls = request.form.get('pouls')
    saturation = request.form.get('saturation')
    poids = request.form.get('poids')
    taille = request.form.get('taille')
    imc = request.form.get('imc')
    
    # ✅ Mettre à jour les constantes (écraser les anciennes)
    if temperature:
        patient.temperature_c = float(temperature)
    if tension:
        patient.tension_arterielle = tension
    if pouls:
        patient.pulse_bpm = int(pouls)
    if saturation:
        patient.oxygene_saturation = int(saturation)
    if poids:
        patient.poids_kg = float(poids)
    if taille:
        patient.taille_cm = float(taille)
    if imc:
        patient.imc = float(imc)
    
    patient.updated_at = datetime.utcnow()
    db.session.commit()
    
    flash('✅ Nouvelles constantes enregistrées avec succès', 'success')
    return redirect(url_for('patient_detail', id=patient.id))

# ==================== ANALYSES DE RÉFÉRENCE ====================

import csv
import os
from flask import jsonify  # ⭐ IMPORTANT

_analyses_cache = None
_analyses_last_update = None

def charger_analyses_reference():
    """Charge les analyses depuis le fichier CSV (id, nom)"""
    global _analyses_cache, _analyses_last_update
    
    if _analyses_cache and _analyses_last_update:
        from datetime import datetime
        if (datetime.now() - _analyses_last_update).seconds < 3600:
            return _analyses_cache
    
    try:
        analyses_file = os.path.join(os.path.dirname(__file__), 'analyses_reference.csv')
        
        if not os.path.exists(analyses_file):
            print(f"Fichier {analyses_file} non trouve")
            return ['NFS', 'Glycemie', 'CRP', 'Radiographie']
        
        analyses_list = []
        with open(analyses_file, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            for row in reader:
                if row and len(row) >= 2:
                    nom = row[1].strip()
                    if nom:
                        analyses_list.append(nom)
                elif row and row[0].strip():
                    analyses_list.append(row[0].strip())
        
        _analyses_cache = analyses_list
        from datetime import datetime
        _analyses_last_update = datetime.now()
        
        print(f"{len(analyses_list)} analyses chargees depuis le fichier")
        return analyses_list
    except Exception as e:
        print(f"Erreur chargement analyses: {e}")
        return []

def search_analyses(search_term, limit=20):
    """Recherche des analyses par nom"""
    if not search_term or len(search_term) < 2:
        return []
    
    all_analyses = charger_analyses_reference()
    search_term = search_term.lower().strip()
    
    results = []
    for analyse in all_analyses:
        if search_term in analyse.lower():
            results.append({'nom': analyse})
            if len(results) >= limit:
                break
    
    return results

@app.route('/api/analyses/search')
@login_required
def api_analyses_search():
    """API de recherche d'analyses"""
    term = request.args.get('term', '')
    if len(term) < 2:
        return jsonify([])
    
    results = search_analyses(term)
    return jsonify(results)

@app.route('/api/analyses/ajouter', methods=['POST'])
@login_required
def api_analyses_ajouter():
    """Ajouter une nouvelle analyse au fichier CSV"""
    if current_user.role not in ['admin_structure', 'medecin']:
        return jsonify({'success': False, 'error': 'Non autorise'}), 403
    
    nom = request.json.get('nom', '').strip()
    if not nom:
        return jsonify({'success': False, 'error': 'Nom requis'}), 400
    
    analyses = charger_analyses_reference()
    if nom in analyses:
        return jsonify({'success': False, 'error': 'Deja existante'}), 400
    
    try:
        analyses_file = os.path.join(os.path.dirname(__file__), 'analyses_reference.csv')
        with open(analyses_file, 'a', encoding='utf-8') as f:
            f.write(f'\n{nom}')
        
        global _analyses_cache, _analyses_last_update
        _analyses_cache = None
        _analyses_last_update = None
        
        return jsonify({'success': True, 'message': f'Analyse "{nom}" ajoutee'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


import pdfkit
import tempfile
import os

@app.route('/patient/<int:id>/pdf')
@login_required
def patient_pdf(id):
    """Générer le dossier patient en PDF"""
    from models import Patient, Consultation, Prescription, Hospitalisation, AnalyseDemande, Reference, AntecedentPatient
    from datetime import datetime
    
    patient = Patient.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Récupérer les données
    consultations = Consultation.query.filter_by(id_patient=patient.id).order_by(
        Consultation.date_consultation.desc()
    ).all()
    
    prescriptions = Prescription.query.filter_by(id_patient=patient.id).order_by(
        Prescription.date_prescription.desc()
    ).all()
    
    hospitalisations = Hospitalisation.query.filter_by(patient_id=patient.id).order_by(
        Hospitalisation.date_debut.desc()
    ).all()
    
    analyses = AnalyseDemande.query.filter_by(patient_id=patient.id).order_by(
        AnalyseDemande.date_demande.desc()
    ).all()
    
    references = Reference.query.filter_by(patient_id=patient.id).order_by(
        Reference.date_reference.desc()
    ).all()
    
    antecedents = AntecedentPatient.query.filter_by(
        patient_id=patient.id,
        actif=True
    ).all()
    
    age = None
    if patient.date_naissance:
        today = datetime.utcnow().date()
        age = today.year - patient.date_naissance.year - ((today.month, today.day) < (patient.date_naissance.month, patient.date_naissance.day))
    
    # Rendre le template HTML
    html_content = render_template('patients/pdf.html',
                                 patient=patient,
                                 age=age,
                                 consultations=consultations,
                                 prescriptions=prescriptions,
                                 hospitalisations=hospitalisations,
                                 analyses=analyses,
                                 references=references,
                                 antecedents=antecedents,
                                 now=datetime.utcnow())
    
    # Chemin vers wkhtmltopdf
    path_wkhtmltopdf = r'C:\Program Files\wkhtmltopdf\bin\wkhtmltopdf.exe'
    
    # Options de configuration
    options = {
        'page-size': 'A4',
        'margin-top': '1.5cm',
        'margin-bottom': '1.5cm',
        'margin-left': '1.5cm',
        'margin-right': '1.5cm',
        'encoding': 'UTF-8',
        'enable-local-file-access': None,
        'no-stop-slow-scripts': None
    }
    
    config = pdfkit.configuration(wkhtmltopdf=path_wkhtmltopdf)
    
    # Générer le PDF
    pdf_file = pdfkit.from_string(html_content, False, options=options, configuration=config)
    
    # Créer un fichier temporaire
    with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as f:
        f.write(pdf_file)
        temp_path = f.name
    
    return send_file(
        temp_path,
        as_attachment=True,
        download_name=f'Dossier_Patient_{patient.prenom}_{patient.nom}_{datetime.utcnow().strftime("%Y%m%d")}.pdf',
        mimetype='application/pdf'
    )
@app.route('/patient/<int:id>/pdf-impression')
@login_required
def patient_pdf_impression(id):
    """Version imprimable du dossier patient (pour PDF via navigateur)"""
    from models import Patient, Consultation, Prescription, Hospitalisation, AnalyseDemande, Reference, AntecedentPatient
    from datetime import datetime
    
    patient = Patient.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    consultations = Consultation.query.filter_by(id_patient=patient.id).order_by(
        Consultation.date_consultation.desc()
    ).all()
    
    prescriptions = Prescription.query.filter_by(id_patient=patient.id).order_by(
        Prescription.date_prescription.desc()
    ).all()
    
    hospitalisations = Hospitalisation.query.filter_by(patient_id=patient.id).order_by(
        Hospitalisation.date_debut.desc()
    ).all()
    
    analyses = AnalyseDemande.query.filter_by(patient_id=patient.id).order_by(
        AnalyseDemande.date_demande.desc()
    ).all()
    
    references = Reference.query.filter_by(patient_id=patient.id).order_by(
        Reference.date_reference.desc()
    ).all()
    
    antecedents = AntecedentPatient.query.filter_by(
        patient_id=patient.id,
        actif=True
    ).all()
    
    age = None
    if patient.date_naissance:
        today = datetime.utcnow().date()
        age = today.year - patient.date_naissance.year - ((today.month, today.day) < (patient.date_naissance.month, patient.date_naissance.day))
    
    return render_template('patients/pdf_impression.html',
                         patient=patient,
                         age=age,
                         consultations=consultations,
                         prescriptions=prescriptions,
                         hospitalisations=hospitalisations,
                         analyses=analyses,
                         references=references,
                         antecedents=antecedents,
                         now=datetime.utcnow())
# ==================== EXAMEN PHYSIQUE ====================

_SECTIONS_EXAMEN = None

def get_sections_examen():
    """Récupère les sections de l'examen physique"""
    global _SECTIONS_EXAMEN
    
    if _SECTIONS_EXAMEN is not None:
        return _SECTIONS_EXAMEN
    
    _SECTIONS_EXAMEN = [
        {
            'nom': 'Général',
            'icone': 'fa-user',
            'fr': 'Patient conscient, orienté, collaborant.\nBon état général.',
            'en': 'Patient conscious, oriented, cooperative.\nGood general condition.'
        },
        {
            'nom': 'Neurologique',
            'icone': 'fa-brain',
            'fr': 'Motricité et sensibilité conservées.\nRéflexes ostéotendineux présents et symétriques.\nPas de déficit neurologique.\nPas de trouble de la marche ou de l\'équilibre.',
            'en': 'Motor and sensory functions preserved.\nOsteotendinous reflexes present and symmetrical.\nNo neurological deficit.\nNo gait or balance disorders.'
        },
        {
            'nom': 'Cardiovasculaire',
            'icone': 'fa-heart',
            'fr': 'Bruits du cœur réguliers, rythme sinusal régulier.\nPas de souffle cardiaque.\nPulsations périphériques présentes et symétriques.\nPas d\'œdème des membres inférieurs.',
            'en': 'Regular heart sounds, regular sinus rhythm.\nNo heart murmur.\nPeripheral pulses present and symmetrical.\nNo lower limb edema.'
        },
        {
            'nom': 'Respiratoire',
            'icone': 'fa-lungs',
            'fr': 'Auscultation pulmonaire normale, murmure vésiculaire bien perçu.\nPas de bruits anormaux (crépitants, sibilants).\nPas de douleur thoracique à la respiration.',
            'en': 'Normal lung auscultation, vesicular breath sounds well heard.\nNo abnormal sounds (crackles, wheezes).\nNo chest pain on respiration.'
        },
        {
            'nom': 'Digestif',
            'icone': 'fa-stomach',
            'fr': 'Abdomen souple, non douloureux à la palpation.\nBruits hydro-aériques présents.\nPas de masse, pas de défense.\nPas de douleur à la décompression.',
            'en': 'Soft abdomen, non-tender on palpation.\nBowel sounds present.\nNo mass, no guarding.\nNo pain on decompression.'
        },
        {
            'nom': 'Splénoganglionnaire',
            'icone': 'fa-blood',
            'fr': 'Pas de splénomégalie palpable.\nPas de polyadénopathie périphérique palpable.\nAires ganglionnaires libres.',
            'en': 'No palpable splenomegaly.\nNo palpable peripheral lymphadenopathy.\nLymph node areas clear.'
        },
        {
            'nom': 'Urogénital',
            'icone': 'fa-kidney',
            'fr': 'Examen urogénital normal.\nPas de douleur à la palpation des fosses lombaires.\nPas de globe vésical.\nOrganes génitaux externes normaux.',
            'en': 'Normal urogenital examination.\nNo pain on palpation of the lumbar fossae.\nNo urinary retention.\nNormal external genitalia.'
        },
        {
            'nom': 'Odonto-stomatologique',
            'icone': 'fa-tooth',
            'fr': 'Cavité buccale normale, muqueuse buccale saine.\nPas de lésion, pas d\'infection.\nDents en bon état.\nPas de mobilité dentaire anormale.',
            'en': 'Normal oral cavity, healthy oral mucosa.\nNo lesions, no infection.\nTeeth in good condition.\nNo abnormal tooth mobility.'
        },
        {
            'nom': 'Dermatologique',
            'icone': 'fa-hand',
            'fr': 'Peau normale, pas de lésion, pas d\'éruption.\nMuqueuses sèches et normales.\nPas de prurit.\nOngles normaux.',
            'en': 'Normal skin, no lesions, no rash.\nMucous membranes dry and normal.\nNo pruritus.\nNormal nails.'
        },
        {
            'nom': 'Locomoteur (Ostéo-articulaire)',
            'icone': 'fa-bone',
            'fr': 'Amplitudes articulaires complètes.\nPas de déformation, pas de douleur à la mobilisation.\nPas de limitation de mouvement.\nPas de raideur.',
            'en': 'Complete joint ranges of motion.\nNo deformity, no pain on mobilization.\nNo limitation of movement.\nNo stiffness.'
        },
        {
            'nom': 'Oto-rhino-laryngologique',
            'icone': 'fa-ear-deaf',
            'fr': 'Conduits auditifs externes libres, tympans normaux.\nFosses nasales libres, muqueuse normale.\nPharynx normal.\nPas de douleur à la mastication.',
            'en': 'External auditory canals clear, normal tympanic membranes.\nNasal passages clear, normal mucosa.\nNormal pharynx.\nNo pain on mastication.'
        },
        {
            'nom': 'Endocrinien',
            'icone': 'fa-flask',
            'fr': 'Pas de goitre palpable à la palpation cervicale.\nPas de signe d\'hypo ou hyperthyroïdie.\nPas de trouble de la croissance ou du développement.',
            'en': 'No palpable goiter on cervical palpation.\nNo signs of hypo or hyperthyroidism.\nNo growth or developmental disorders.'
        },
        {
            'nom': 'Psychiatrique',
            'icone': 'fa-brain',
            'fr': 'Humeur stable, contact facile et approprié.\nPas de trouble du comportement, pas d\'idées délirantes.\nPas de trouble de l\'humeur.\nPas d\'anxiété ou de dépression.',
            'en': 'Stable mood, easy and appropriate contact.\nNo behavioral disorders, no delusional ideas.\nNo mood disorders.\nNo anxiety or depression.'
        },
        {
            'nom': 'Autre à préciser',
            'icone': 'fa-plus-circle',
            'fr': 'Section personnalisée à ajouter selon les besoins de l\'examen.',
            'en': 'Custom section to add according to the needs of the examination.'
        }
    ]
    
    return _SECTIONS_EXAMEN

@app.route('/api/examen-physique/sections')
@login_required
def api_sections_examen():
    """Récupère les sections de l'examen physique"""
    lang = request.args.get('lang', 'fr')
    sections = get_sections_examen()
    
    result = []
    for s in sections:
        result.append({
            'nom': s['nom'],
            'icone': s['icone'],
            'texte': s['fr'] if lang == 'fr' else s['en']
        })
    
    return jsonify(result)


def _rendre_sections_examen_physique(examen, lang='fr'):
    """Reconstruit, côté serveur, la liste des sections de l'examen
    physique avec leur texte effectif et un indicateur 'modifie' — pour
    affichage en LECTURE SEULE (détail consultation/hospitalisation), avec
    surlignage gris des sections modifiées (patron : "les parties
    modifiées dans l'examen physique apparaissent sous une couleur
    grise"). Référence de comparaison : sections_origine (instantané figé
    à la copie, hospitalisation uniquement) si présent, sinon le
    catalogue par défaut (cas normal d'une consultation) — même logique
    que sectionsOriginales côté JS (examen_physique.html)."""
    import json

    if not examen:
        return []

    catalogue = get_sections_examen()

    modifiees = {}
    if examen.sections_modifiees:
        try:
            modifiees = json.loads(examen.sections_modifiees)
        except (ValueError, TypeError):
            modifiees = {}

    origine = {}
    if getattr(examen, 'sections_origine', None):
        try:
            origine = json.loads(examen.sections_origine)
        except (ValueError, TypeError):
            origine = {}

    sections = []
    for i, s in enumerate(catalogue):
        idx = str(i)
        defaut = s['fr'] if lang == 'fr' else s['en']
        reference = origine.get(idx, defaut)
        texte = modifiees.get(idx, reference if origine else defaut)
        modifie = idx in modifiees and modifiees[idx] != reference
        sections.append({'nom': s['nom'], 'icone': s['icone'], 'texte': texte, 'modifie': modifie})

    return sections


@app.route('/consultation/<int:id>/examen-physique')
@login_required
def examen_physique(id):
    """Page de l'examen physique"""
    from models import Consultation, Patient, ExamenPhysique
    
    consultation = Consultation.query.get_or_404(id)
    patient = Patient.query.get(consultation.id_patient)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Récupérer l'examen existant
    examen = ExamenPhysique.query.filter_by(consultation_id=id).first()
    
    return render_template('consultations/examen_physique.html',
                         consultation=consultation,
                         patient=patient,
                         examen=examen)
@app.route('/patient/<int:patient_id>/examen-physique/ajouter')
@login_required
def examen_physique_ajouter(patient_id):
    from models import Patient, Consultation
    from datetime import datetime
    
    patient = Patient.query.get_or_404(patient_id)
    
    if patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # ⭐ VÉRIFIER SI UNE CONSULTATION TEMPORAIRE EXISTE DÉJÀ
    consultation_temp = Consultation.query.filter_by(
        id_patient=patient.id,
        is_temporary=True
    ).first()
    
    if consultation_temp:
        consultation = consultation_temp
        print(f"📋 Consultation temporaire existante: #{consultation.id}")
    else:
        consultation = Consultation(
            id_patient=patient.id,
            id_medecin=current_user.id,
            date_consultation=datetime.utcnow(),
            is_temporary=True,
            statut='en_cours'
        )
        db.session.add(consultation)
        db.session.commit()
        print(f"✅ Nouvelle consultation temporaire créée avec ID: {consultation.id}")
    
    # ⭐ REDIRIGER VERS LA PAGE D'EXAMEN AVEC L'ID DANS L'URL
    return redirect(url_for('examen_physique_page', consultation_id=consultation.id))


@app.route('/examen-physique/<int:consultation_id>')
@login_required
def examen_physique_page(consultation_id):
    """Affiche la page d'examen physique avec l'ID de consultation"""
    from models import Consultation, Patient, ExamenPhysique
    
    consultation = Consultation.query.get_or_404(consultation_id)
    patient = Patient.query.get(consultation.id_patient)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Récupérer l'examen existant
    examen = ExamenPhysique.query.filter_by(consultation_id=consultation_id).first()
    
    return render_template('consultations/examen_physique.html',
                         patient=patient,
                         consultation=consultation,
                         examen=examen)

@app.route('/patient/<int:patient_id>/examen-physique/creer-temporaire', methods=['POST'])
@login_required
def creer_consultation_temporaire(patient_id):
    """Crée une consultation temporaire et retourne l'ID"""
    from models import Patient, Consultation
    from flask import jsonify
    from datetime import datetime
    
    patient = Patient.query.get_or_404(patient_id)
    
    if patient.id_structure != current_user.id_structure:
        return jsonify({'success': False, 'message': 'Accès non autorisé'}), 403
    
    # Vérifier si une consultation temporaire existe
    consultation_temp = Consultation.query.filter_by(
        id_patient=patient.id,
        is_temporary=True
    ).first()
    
    if consultation_temp:
        consultation = consultation_temp
    else:
        consultation = Consultation(
            id_patient=patient.id,
            id_medecin=current_user.id,
            date_consultation=datetime.utcnow(),
            is_temporary=True,
            statut='en_cours'
        )
        db.session.add(consultation)
        db.session.commit()
    
    return jsonify({
        'success': True,
        'consultation_id': consultation.id,
        'patient_id': patient.id
    })

@app.route('/consultation/<int:id>/examen-physique/enregistrer', methods=['POST'])
@login_required
def enregistrer_examen_physique(id):
    from models import Consultation, ExamenPhysique
    from datetime import datetime
    from flask import jsonify
    import json
    import re
    
    print(f"🟢 Enregistrement examen physique pour consultation ID: {id}")
    
    consultation = Consultation.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'message': 'Accès non autorisé'}), 403
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    try:
        examen_complet = request.form.get('examen_complet', '')
        sections_modifiees = request.form.get('sections_modifiees', '{}')
        
        print(f"📝 Examen complet reçu: {len(examen_complet)} caractères")
        
        examen_complet = nettoyer_examen_complet(examen_complet)
        
        examen = ExamenPhysique.query.filter_by(consultation_id=id).first()
        
        if examen:
            examen.examen_complet = examen_complet
            examen.sections_modifiees = sections_modifiees
            examen.modified_at = datetime.utcnow()
            print(f"🔄 Examen physique #{examen.id} mis à jour")
        else:
            examen = ExamenPhysique(
                consultation_id=id,
                examen_complet=examen_complet,
                sections_modifiees=sections_modifiees,
                created_by=current_user.id
            )
            db.session.add(examen)
            print(f"✅ Nouvel examen physique créé pour consultation #{id}")
        
        # Mettre à jour les notes cliniques
        if examen_complet and examen_complet.strip():
            if consultation.notes_cliniques:
                if "--- EXAMEN PHYSIQUE ---" not in consultation.notes_cliniques:
                    consultation.notes_cliniques = consultation.notes_cliniques + f"\n\n--- EXAMEN PHYSIQUE ---\n{examen_complet}"
                else:
                    pattern = r'--- EXAMEN PHYSIQUE ---\n.*?(?=\n---|$)'
                    consultation.notes_cliniques = re.sub(pattern, f"--- EXAMEN PHYSIQUE ---\n{examen_complet}", consultation.notes_cliniques, flags=re.DOTALL)
            else:
                consultation.notes_cliniques = f"--- EXAMEN PHYSIQUE ---\n{examen_complet}"
        
        db.session.commit()
        
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({
                'success': True,
                'message': 'Examen enregistré avec succès',
                'examen_id': examen.id,
                'consultation_id': consultation.id,
                'examen_complet': examen_complet  # ⭐ AJOUTER CECI
            })
        
        flash('✅ Examen physique enregistré avec succès', 'success')
        return redirect(url_for('consultation_detail', id=id))
        
    except Exception as e:
        db.session.rollback()
        print(f"❌ Erreur: {e}")
        import traceback
        traceback.print_exc()
        
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'message': str(e)}), 500
        
        flash(f'❌ Erreur: {str(e)}', 'danger')
        return redirect(url_for('consultation_detail', id=id))


@app.route('/hospitalisation/<int:id>/examen-physique')
@login_required
def hospitalisation_examen_physique(id):
    """Examen physique d'une hospitalisation — même éditeur que pour une
    consultation (patron : "il faut prévoir l'examen physique préremplie
    à modifier exactement comme dans consultations"). Pré-rempli à la
    première visite avec l'examen physique de la consultation d'origine
    (hospitalisation.consultation_id), sinon celui de la dernière
    consultation du patient ; ensuite librement modifiable, avec
    surlignage gris des sections changées depuis l'admission (voir
    _rendre_sections_examen_physique)."""
    from models import Hospitalisation, ExamenPhysique, Consultation
    import json

    hospitalisation = Hospitalisation.query.get_or_404(id)
    patient = hospitalisation.patient

    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    examen = ExamenPhysique.query.filter_by(hospitalisation_id=id).first()

    if not examen:
        # ⭐ Chercher l'examen physique source : celui de la consultation
        # d'origine en priorité, sinon la dernière consultation du patient
        # qui en a un.
        examen_source = None
        if hospitalisation.consultation_id:
            examen_source = ExamenPhysique.query.filter_by(consultation_id=hospitalisation.consultation_id).first()
        if not examen_source:
            derniere_consultation = (
                Consultation.query
                .filter_by(id_patient=hospitalisation.patient_id)
                .order_by(Consultation.date_consultation.desc())
                .all()
            )
            for c in derniere_consultation:
                examen_source = ExamenPhysique.query.filter_by(consultation_id=c.id).first()
                if examen_source:
                    break

        if examen_source:
            examen = ExamenPhysique(
                hospitalisation_id=id,
                sections_modifiees=examen_source.sections_modifiees,
                sections_origine=examen_source.sections_modifiees or '{}',
                examen_complet=examen_source.examen_complet,
                created_by=current_user.id
            )
        else:
            examen = ExamenPhysique(
                hospitalisation_id=id,
                sections_modifiees='{}',
                sections_origine='{}',
                created_by=current_user.id
            )
        db.session.add(examen)
        db.session.commit()

    return render_template('consultations/examen_physique.html',
                         patient=patient,
                         consultation=None,
                         hospitalisation=hospitalisation,
                         examen=examen)


@app.route('/hospitalisation/<int:id>/examen-physique/enregistrer', methods=['POST'])
@login_required
def enregistrer_examen_physique_hospitalisation(id):
    """Équivalent de enregistrer_examen_physique() pour une hospitalisation
    — ne touche jamais sections_origine (instantané figé à la copie)."""
    from models import Hospitalisation, ExamenPhysique
    from datetime import datetime

    hospitalisation = Hospitalisation.query.get_or_404(id)

    if current_user.role not in ['admin_structure', 'medecin']:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'message': 'Accès non autorisé'}), 403
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))

    try:
        examen_complet = request.form.get('examen_complet', '')
        sections_modifiees = request.form.get('sections_modifiees', '{}')

        examen_complet = nettoyer_examen_complet(examen_complet)

        examen = ExamenPhysique.query.filter_by(hospitalisation_id=id).first()

        if examen:
            examen.examen_complet = examen_complet
            examen.sections_modifiees = sections_modifiees
            examen.modified_at = datetime.utcnow()
        else:
            examen = ExamenPhysique(
                hospitalisation_id=id,
                examen_complet=examen_complet,
                sections_modifiees=sections_modifiees,
                sections_origine='{}',
                created_by=current_user.id
            )
            db.session.add(examen)

        db.session.commit()

        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({
                'success': True,
                'message': 'Examen enregistré avec succès',
                'examen_id': examen.id,
                'hospitalisation_id': hospitalisation.id,
                'examen_complet': examen_complet
            })

        flash('✅ Examen physique enregistré avec succès', 'success')
        return redirect(url_for('detail_hospitalisation', id=id))

    except Exception as e:
        db.session.rollback()
        print(f"❌ Erreur: {e}")
        import traceback
        traceback.print_exc()

        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'message': str(e)}), 500

        flash(f'❌ Erreur: {str(e)}', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))


def nettoyer_examen_complet(examen_complet):
    """Supprime les sections vides ou non modifiées"""
    if not examen_complet:
        return ''
    
    lines = examen_complet.split('\n')
    result = []
    skip_next = False
    
    for line in lines:
        if '═══════════════════════════════════════════════════' in line:
            result.append(line)
            continue
        
        if 'section(s) active(s)' in line or 'modification(s)' in line:
            result.append(line)
            continue
        
        if '--- AUTRE À PRÉCISER ---' in line:
            skip_next = True
            continue
        
        if skip_next:
            if not line.strip() or line.strip() == '':
                skip_next = False
            continue
        
        if '---' in line and 'AUTRE À PRÉCISER' not in line:
            result.append(line)
            continue
        
        if line.strip() and 'Section personnalisée' not in line:
            result.append(line)
    
    return '\n'.join(result)


@app.route('/cleanup-temp-consultations')
@login_required
def cleanup_temp_consultations():
    """Supprime les consultations temporaires sans examen"""
    from models import Consultation, ExamenPhysique
    
    if current_user.role not in ['admin_structure']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Trouver les consultations temporaires sans examen
    temp_consultations = Consultation.query.filter_by(statut='en_cours').all()
    deleted = 0
    
    for c in temp_consultations:
        examen = ExamenPhysique.query.filter_by(consultation_id=c.id).first()
        if not examen:
            db.session.delete(c)
            deleted += 1
    
    db.session.commit()
    flash(f'{deleted} consultation(s) temporaire(s) nettoyée(s)', 'success')
    return redirect(url_for('dashboard'))


# ==================== SYNCHRONISATION GHP ====================

import requests
import uuid
from datetime import datetime

@app.route('/sync/ghp')
@login_required
def sync_ghp_config():
    """Page de configuration de la synchronisation GHP"""
    from models import Structure, StructureMapping
    
    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    structures = Structure.query.filter_by(
        id_structure=current_user.id_structure,
        actif=True
    ).all() if current_user.role == 'super_admin' else Structure.query.filter_by(id=current_user.id_structure).all()
    
    mappings = StructureMapping.query.filter_by(local_structure_id=current_user.id_structure).all()
    
    return render_template('sync/mapping.html', 
                         structures=structures,
                         mappings=mappings)


@app.route('/sync/ghp', methods=['POST'])
@login_required
def sync_ghp_save():
    """Enregistrer une configuration de mapping"""
    from models import StructureMapping
    
    if current_user.role != 'admin_structure':
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    local_structure_id = request.form.get('local_structure_id', type=int)
    source_structure_id = request.form.get('source_structure_id', type=int)
    api_url = request.form.get('api_url')
    api_key = request.form.get('api_key')
    
    if not local_structure_id or not source_structure_id:
        flash('Tous les champs sont obligatoires', 'danger')
        return redirect(url_for('sync_ghp_config'))
    
    # Vérifier si le mapping existe déjà
    mapping = StructureMapping.query.filter_by(
        local_structure_id=local_structure_id,
        source_structure_id=source_structure_id
    ).first()
    
    if mapping:
        mapping.api_url = api_url
        mapping.api_key = api_key
        mapping.actif = True
    else:
        mapping = StructureMapping(
            local_structure_id=local_structure_id,
            source_structure_id=source_structure_id,
            api_url=api_url,
            api_key=api_key,
            source_name='ghp'
        )
        db.session.add(mapping)
    
    db.session.commit()
    
    flash('✅ Configuration enregistrée avec succès', 'success')
    return redirect(url_for('sync_ghp_config'))

# ============================================================
# FONCTIONS DE NORMALISATION
# ============================================================

def normalize_assurance_type(value):
    """
    Normalise le type d'assurance pour standardiser les valeurs
    """
    if not value:
        return 'NON_ASSURÉ'
    
    # Convertir en string et mettre en minuscules
    value = str(value).lower().strip()
    
    # Mapping des valeurs
    mapping = {
        # AMU-CNSS
        'amu_cnss': 'AMU-CNSS',
        'amu-cnss': 'AMU-CNSS',
        'amucnss': 'AMU-CNSS',
        'amu cnss': 'AMU-CNSS',
        'cnss': 'AMU-CNSS',
        
        # AMU-INAM
        'amu_inam': 'AMU-INAM',
        'amu-inam': 'AMU-INAM',
        'amuinam': 'AMU-INAM',
        'amu inam': 'AMU-INAM',
        'inam': 'AMU-INAM',
        
        # Autre assurance
        'autre_assurance': 'AUTRE_ASSURANCE',
        'autre-assurance': 'AUTRE_ASSURANCE',
        'autre assurance': 'AUTRE_ASSURANCE',
        'autre': 'AUTRE_ASSURANCE',
        'other': 'AUTRE_ASSURANCE',
        
        # Non assuré
        'non_assure': 'NON_ASSURÉ',
        'non-assure': 'NON_ASSURÉ',
        'non assure': 'NON_ASSURÉ',
        'nonassure': 'NON_ASSURÉ',
        'non': 'NON_ASSURÉ',
        'aucune': 'NON_ASSURÉ',
        '': 'NON_ASSURÉ',
    }
    
    # Vérifier si la valeur existe dans le mapping
    if value in mapping:
        return mapping[value]
    
    # Si la valeur contient 'amu' ou 'assurance', essayer de deviner
    if 'amu' in value or 'assurance' in value:
        if 'cnss' in value or 'inam' in value:
            # Essayer de trouver le type
            if 'cnss' in value:
                return 'AMU-CNSS'
            elif 'inam' in value:
                return 'AMU-INAM'
    
    # Si rien ne correspond, retourner la valeur en majuscules
    return value.upper().replace('_', '-')

def sync_patients_from_ghp(structure_mapping):
    """
    Synchronise UNIQUEMENT les informations patient depuis GHP
    Version avec normalisation des données
    """
    from models import Patient, Utilisateur
    from datetime import datetime
    import uuid
    import requests
    
    try:
        token = structure_mapping.api_key
        
        if not token:
            print("❌ Token manquant dans le mapping")
            return {'cree': 0, 'mis_a_jour': 0, 'erreur': 0, 'message': 'Token manquant'}
        
        url = f"{structure_mapping.api_url}/api/sync/patients"
        params = {'token': token}
        
        print(f"🔄 Synchronisation depuis: {url}")
        
        response = requests.get(url, params=params, timeout=60)
        
        if response.status_code != 200:
            print(f"❌ Erreur API GHP: {response.status_code}")
            return {'cree': 0, 'mis_a_jour': 0, 'erreur': 1, 'message': f'Erreur API: {response.status_code}'}
        
        data = response.json()
        patients = data.get('patients', [])
        
        print(f"✅ {len(patients)} patients récupérés depuis GHP")
        
        compteur = {'cree': 0, 'mis_a_jour': 0, 'erreur': 0}
        
        # ⭐⭐⭐ AUCUN MÉDECIN RÉFÉRENT N'EST ASSIGNÉ ICI ⭐⭐⭐
        # Le médecin référent sera assigné lors de la première consultation
        print("ℹ️ Aucun médecin référent assigné automatiquement - assignation lors de la première consultation")
        
        for p_data in patients:
            try:
                source_id = p_data.get('ID')
                if not source_id:
                    compteur['erreur'] += 1
                    print(f"❌ Patient sans ID ignoré")
                    continue
                
                # NORMALISATION DES DONNÉES
                raw_type = p_data.get('type_assurance') or p_data.get('TypeAssurance') or 'non_assure'
                type_assurance = normalize_assurance_type(raw_type)
                
                taux_prise_charge = p_data.get('taux_assurance') or p_data.get('taux_prise_charge') or 0
                numero_assure = p_data.get('numero_assure') or p_data.get('NumeroAssure') or p_data.get('num_assure') or ''
                
                assurance2_nom = p_data.get('assurance2_nom') or p_data.get('Assurance2Nom') or ''
                if assurance2_nom:
                    assurance2_nom = assurance2_nom.upper().strip()
                
                taux_assurance2 = p_data.get('taux_assurance2') or p_data.get('TauxAssurance2') or 0
                numero_assure2 = p_data.get('numero_assure2') or p_data.get('NumeroAssure2') or ''
                
                personne_a_prevenir_nom = p_data.get('personne_a_prevenir_nom') or p_data.get('PersonneAPrevenirNom') or ''
                personne_a_prevenir_telephone = p_data.get('personne_a_prevenir_telephone') or p_data.get('PersonneAPrevenirTelephone') or ''
                personne_a_prevenir_relation = p_data.get('personne_a_prevenir_relation') or p_data.get('PersonneAPrevenirRelation') or ''
                
                # Date de naissance
                date_naissance = p_data.get('date_naissance')
                if date_naissance and isinstance(date_naissance, str):
                    try:
                        date_naissance = datetime.strptime(date_naissance, '%Y-%m-%d').date()
                    except:
                        date_naissance = None
                elif isinstance(date_naissance, datetime):
                    date_naissance = date_naissance.date()
                
                # Chercher si le patient existe déjà
                patient = Patient.query.filter_by(
                    patient_source_id=str(source_id),
                    source_structure_id=structure_mapping.source_structure_id,
                    id_structure=structure_mapping.local_structure_id
                ).first()
                
                if patient:
                    # 📝 MISE À JOUR DU PATIENT EXISTANT
                    print(f"📝 Mise à jour: {p_data.get('nom')} {p_data.get('prenom')} (ID GHP: {source_id})")
                    
                    patient.nom = p_data.get('nom') or ''
                    patient.prenom = p_data.get('prenom') or ''
                    patient.telephone = str(p_data.get('telephone') or '')
                    patient.adresse = p_data.get('adresse') or ''
                    patient.date_naissance = date_naissance
                    
                    # ASSURANCE PRINCIPALE
                    patient.type_assurance = type_assurance
                    patient.taux_prise_charge = str(taux_prise_charge) if taux_prise_charge else None
                    patient.numero_assure = str(numero_assure) if numero_assure else ''
                    
                    # ASSURANCE 2
                    patient.assurance2_nom = assurance2_nom if assurance2_nom else None
                    patient.taux_assurance2 = float(taux_assurance2) if taux_assurance2 else None
                    patient.numero_assure2 = str(numero_assure2) if numero_assure2 else ''
                    
                    # PERSONNE À PRÉVENIR
                    patient.personne_a_prevenir_nom = personne_a_prevenir_nom
                    patient.personne_a_prevenir_telephone = personne_a_prevenir_telephone
                    patient.personne_a_prevenir_relation = personne_a_prevenir_relation
                    
                    # ⭐⭐⭐ ON NE TOUCHE PAS AU MÉDECIN RÉFÉRENT ⭐⭐⭐
                    # Le médecin référent reste celui qui a été assigné lors de la première consultation
                    print(f"   👨‍⚕️ Médecin référent actuel: ID {patient.id_medecin_referent if patient.id_medecin_referent else 'Aucun'}")
                    
                    patient.synced_at = datetime.utcnow()
                    patient.synced_from = 'ghp'
                    compteur['mis_a_jour'] += 1
                    
                else:
                    # ➕ CRÉATION D'UN NOUVEAU PATIENT
                    print(f"➕ Création: {p_data.get('nom')} {p_data.get('prenom')} (ID GHP: {source_id})")
                    
                    patient = Patient(
                        id_structure=structure_mapping.local_structure_id,
                        uuid=str(uuid.uuid4()),
                        patient_source_id=str(source_id),
                        source_structure_id=structure_mapping.source_structure_id,
                        source_name='ghp',
                        
                        # Identité
                        nom=p_data.get('nom') or '',
                        prenom=p_data.get('prenom') or '',
                        telephone=str(p_data.get('telephone') or ''),
                        adresse=p_data.get('adresse') or '',
                        date_naissance=date_naissance,
                        
                        # ASSURANCE PRINCIPALE
                        type_assurance=type_assurance,
                        taux_prise_charge=str(taux_prise_charge) if taux_prise_charge else None,
                        numero_assure=str(numero_assure) if numero_assure else '',
                        
                        # ASSURANCE 2
                        assurance2_nom=assurance2_nom if assurance2_nom else None,
                        taux_assurance2=float(taux_assurance2) if taux_assurance2 else None,
                        numero_assure2=str(numero_assure2) if numero_assure2 else '',
                        
                        # PERSONNE À PRÉVENIR
                        personne_a_prevenir_nom=personne_a_prevenir_nom,
                        personne_a_prevenir_telephone=personne_a_prevenir_telephone,
                        personne_a_prevenir_relation=personne_a_prevenir_relation,
                        
                        # Autres champs
                        lieu_naissance=p_data.get('lieu_naissance') or '',
                        sexe=p_data.get('sexe') or '',
                        email=p_data.get('email') or '',
                        profession=p_data.get('profession') or '',
                        
                        # ⭐⭐⭐ PAS DE MÉDECIN RÉFÉRENT À LA CRÉATION ⭐⭐⭐
                        id_medecin_referent=None,  # Sera assigné lors de la première consultation
                        
                        # Statut
                        statut_medical='PREMIERE_VISITE',
                        archived=False,
                        
                        # Métadonnées
                        synced_at=datetime.utcnow(),
                        synced_from='ghp'
                    )
                    
                    db.session.add(patient)
                    compteur['cree'] += 1
                    
                    print(f"   ℹ️ Patient créé sans médecin référent - sera assigné lors de la première consultation")
                
                db.session.flush()
                
            except Exception as e:
                compteur['erreur'] += 1
                print(f"❌ Erreur patient {p_data.get('ID')}: {e}")
                import traceback
                traceback.print_exc()
                db.session.rollback()
                continue
        
        # Mettre à jour la date de dernière synchronisation
        structure_mapping.last_sync = datetime.utcnow()
        db.session.commit()
        
        message = f"✅ Sync terminée: {compteur['cree']} créés, {compteur['mis_a_jour']} mis à jour, {compteur['erreur']} erreurs"
        print(message)
        
        return {
            'cree': compteur['cree'],
            'mis_a_jour': compteur['mis_a_jour'],
            'erreur': compteur['erreur'],
            'message': message,
            'total': len(patients)
        }
        
    except requests.exceptions.RequestException as e:
        print(f"❌ Erreur de connexion à GHP: {e}")
        return {
            'cree': 0,
            'mis_a_jour': 0,
            'erreur': 1,
            'message': f'Erreur de connexion: {str(e)}'
        }
    except Exception as e:
        print(f"❌ Erreur générale: {e}")
        import traceback
        traceback.print_exc()
        db.session.rollback()
        return {
            'cree': 0,
            'mis_a_jour': 0,
            'erreur': 1,
            'message': f'Erreur: {str(e)}'
        }


@app.route('/api/sync/patients/<int:mapping_id>', methods=['POST'])
@login_required
def api_sync_patients(mapping_id):
    """Déclencher la synchronisation des patients depuis GHP"""
    from models import StructureMapping
    
    if current_user.role not in ['admin_structure', 'super_admin']:
        return jsonify({'success': False, 'message': 'Non autorisé'}), 403
    
    mapping = StructureMapping.query.get_or_404(mapping_id)
    
    if mapping.local_structure_id != current_user.id_structure:
        return jsonify({'success': False, 'message': 'Accès non autorisé'}), 403
    
    if not mapping.api_key:
        return jsonify({
            'success': False,
            'message': 'Token manquant dans la configuration'
        }), 400
    
    try:
        # ⭐ APPELER LA FONCTION DE SYNCHRONISATION
        resultat = sync_patients_from_ghp(mapping)
        
        if resultat.get('erreur', 0) > 0 and resultat.get('cree', 0) == 0 and resultat.get('mis_a_jour', 0) == 0:
            return jsonify({
                'success': False,
                'message': resultat.get('message', 'Erreur lors de la synchronisation'),
                'details': resultat
            }), 500
        
        return jsonify({
            'success': True,
            'message': resultat.get('message', 'Synchronisation terminée'),
            'details': {
                'cree': resultat.get('cree', 0),
                'mis_a_jour': resultat.get('mis_a_jour', 0),
                'erreur': resultat.get('erreur', 0),
                'total': resultat.get('total', 0)
            }
        })
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'message': f'Erreur: {str(e)}'
        }), 500

@app.route('/api/webhook/patient-created', methods=['POST'])
def webhook_patient_created():
    from models import StructureMapping
    from threading import Thread
    import os
    
    try:
        webhook_secret = os.environ.get('WEBHOOK_SECRET', 'mon_secret_webhook_123456')
        token = request.headers.get('X-Webhook-Token')
        
        if token != webhook_secret:
            return jsonify({'success': False, 'message': 'Token invalide'}), 401
        
        data = request.json
        if not data:
            return jsonify({'success': False, 'message': 'Données JSON manquantes'}), 400
        
        patient_id = data.get('patient_id')
        structure_id = data.get('structure_id')
        
        if not patient_id or not structure_id:
            return jsonify({'success': False, 'message': 'patient_id et structure_id sont obligatoires'}), 400
        
        mapping = StructureMapping.query.filter_by(
            local_structure_id=structure_id,
            actif=True
        ).first()
        
        if not mapping:
            return jsonify({
                'success': False,
                'message': f'Configuration GHP non trouvée pour la structure {structure_id}'
            }), 404
        
        def sync_in_background():
            with app.app_context():
                try:
                    print(f"⚡ Webhook: Sync immédiate patient {patient_id}")
                    resultat = sync_patients_from_ghp(mapping)
                    
                    if resultat.get('cree', 0) > 0:
                        print(f"✅ Patient {patient_id} synchronisé immédiatement")
                    else:
                        print(f"⚠️ Patient {patient_id} déjà existant ou non trouvé")
                except Exception as e:
                    print(f"❌ Erreur webhook patient {patient_id}: {e}")
                    import traceback
                    traceback.print_exc()
        
        Thread(target=sync_in_background).start()
        
        print(f"📡 Webhook: Sync déclenchée pour patient {patient_id}")
        
        return jsonify({
            'success': True,
            'message': f'Synchronisation déclenchée en arrière-plan pour patient {patient_id}'
        })
        
    except Exception as e:
        print(f"❌ Erreur webhook: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'message': f'Erreur interne: {str(e)}'
        }), 500

# ═══════════════════════════════════════════
# ROUTE DE TEST DU WEBHOOK
# ═══════════════════════════════════════════

# app.py - Ajouter cette route

@app.route('/api/webhook/test', methods=['GET', 'POST'])
def webhook_test():
    """
    Route de test pour vérifier que le webhook est accessible
    """
    if request.method == 'POST':
        data = request.json or {}
        return jsonify({
            'success': True,
            'message': 'Webhook accessible',
            'received_data': data,
            'headers': dict(request.headers)
        })
    else:
        return jsonify({
            'success': True,
            'message': '✅ Webhook accessible',
            'instructions': 'Envoyer une requête POST avec patient_id et structure_id',
            'example': {
                'url': '/api/webhook/patient-created',
                'method': 'POST',
                'headers': {
                    'X-Webhook-Token': 'mon_secret_webhook_123456',
                    'Content-Type': 'application/json'
                },
                'body': {
                    'patient_id': 123,
                    'structure_id': 1
                }
            }
        })

@app.route('/api/sync/mapping/<int:id>', methods=['DELETE'])
@login_required
def api_delete_mapping(id):
    """Supprimer un mapping"""
    from models import StructureMapping
    
    mapping = StructureMapping.query.get_or_404(id)
    
    if mapping.local_structure_id != current_user.id_structure:
        return jsonify({'success': False, 'message': 'Non autorisé'}), 403
    
    db.session.delete(mapping)
    db.session.commit()
    
    return jsonify({'success': True})
@app.route('/patient/<int:patient_id>/update-habitudes', methods=['POST'])
@login_required
def patient_update_habitudes(patient_id):
    """Mettre à jour les habitudes de vie du patient"""
    from models import Patient
    from flask import jsonify
    
    patient = Patient.query.get_or_404(patient_id)
    
    # Vérifier les permissions
    if current_user.role not in ['admin_structure', 'medecin', 'infirmier']:
        return jsonify({'success': False, 'message': 'Accès non autorisé'}), 403
    
    try:
        # Récupérer les données
        tabac = request.form.get('tabac')
        alcool = request.form.get('alcool')
        allaitement = request.form.get('allaitement')
        grossesse = request.form.get('grossesse')
        groupe_sanguin = request.form.get('groupe_sanguin')
        mutuelle = request.form.get('mutuelle')
        medecin_traitant = request.form.get('medecin_traitant')
        
        # Mettre à jour
        patient.tabac = tabac if tabac else None
        patient.alcool = alcool if alcool else None
        patient.allaitement = allaitement == 'Oui'
        patient.grossesse = grossesse == 'Oui'
        patient.groupe_sanguin = groupe_sanguin if groupe_sanguin else None
        patient.mutuelle = mutuelle if mutuelle else None
        patient.medecin_traitant = medecin_traitant if medecin_traitant else None
        
        db.session.commit()
        
        # ⭐ RETOURNER JSON (pour l'appel AJAX)
        return jsonify({'success': True, 'message': 'Habitudes de vie mises à jour avec succès'})
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)}), 500


@app.route('/api/medicaments/disponibles')
@login_required
def api_medicaments_disponibles():
    """
    Récupère les médicaments depuis Google Sheets (via GHP)
    """
    from models import StructureMapping
    import requests
    
    # Récupérer le mapping GHP
    mapping = StructureMapping.query.filter_by(
        local_structure_id=current_user.id_structure,
        actif=True
    ).first()
    
    if not mapping:
        return jsonify([])
    
    try:
        # Appeler l'API de GHP pour récupérer les médicaments
        url = f"{mapping.api_url}/api/medicaments"
        params = {'token': mapping.api_key}
        
        response = requests.get(url, params=params, timeout=10)
        
        if response.status_code != 200:
            return jsonify([])
        
        data = response.json()
        medicaments = data.get('medicaments', [])
        
        # Filtrer et formater pour le médecin
        result = []
        for m in medicaments:
            # ⭐ Le médecin voit seulement le nom et le stock
            result.append({
                'id': m.get('ID'),
                'nom': m.get('nom', ''),
                'quantite_stock': m.get('quantite_stock', 0)
            })
        
        # Trier par nom
        result.sort(key=lambda x: x['nom'])
        
        return jsonify(result)
        
    except Exception as e:
        print(f"❌ Erreur récupération médicaments: {e}")
        return jsonify([])

@app.route('/api/actes-types/disponibles')
@login_required
def api_actes_types_disponibles():
    """
    Récupère le catalogue d'actes depuis GHP (même principe que
    /api/medicamentos/disponibles) — pour la recherche d'actes posés.
    """
    from models import StructureMapping
    import requests

    mapping = StructureMapping.query.filter_by(
        local_structure_id=current_user.id_structure,
        actif=True
    ).first()

    if not mapping:
        print(f"⚠️ Aucun mapping GHP trouvé pour la structure {current_user.id_structure}")
        return jsonify([])

    try:
        url = f"{mapping.api_url}/api/actes/disponibles"
        params = {'token': mapping.api_key}

        response = requests.get(url, params=params, timeout=15)
        if response.status_code != 200:
            print(f"❌ Erreur GHP (actes): {response.status_code} - {response.text[:100]}")
            return jsonify([])

        data = response.json()
        actes = data.get('actes', [])
        # ⭐ prix/pbr transmis en plus du nom (si présents côté GHP) — la
        # recherche d'actes posés les ignore, l'aperçu de facturation
        # d'hospitalisation les utilise pour l'estimation avant envoi.
        result = [
            {'nom': a.get('nom', ''), 'prix': a.get('prix', 0), 'pbr': a.get('pbr', 0)}
            for a in actes if a.get('nom')
        ]
        result.sort(key=lambda x: x['nom'])

        print(f"✅ {len(result)} actes disponibles chargés")
        return jsonify(result)

    except requests.exceptions.Timeout:
        print("❌ Timeout lors de la récupération des actes")
        return jsonify([])
    except requests.exceptions.ConnectionError:
        print("❌ Erreur de connexion à GHP")
        return jsonify([])
    except Exception as e:
        print(f"❌ Erreur récupération actes: {e}")
        import traceback
        traceback.print_exc()
        return jsonify([])


@app.route('/api/medicamentos/disponibles')
@login_required
def api_medicamentos_disponibles():
    """
    Récupère les médicaments disponibles depuis GHP
    Le médecin voit seulement : id, nom, quantite_stock
    """
    from models import StructureMapping
    import requests
    
    # Récupérer le mapping GHP
    mapping = StructureMapping.query.filter_by(
        local_structure_id=current_user.id_structure,
        actif=True
    ).first()
    
    if not mapping:
        print(f"⚠️ Aucun mapping GHP trouvé pour la structure {current_user.id_structure}")
        return jsonify([])
    
    try:
        # Appeler l'API de GHP pour récupérer les médicaments
        url = f"{mapping.api_url}/api/medicamentos"
        params = {'token': mapping.api_key}
        
        print(f"📡 Récupération des médicaments depuis: {url}")
        print(f"   Structure source: {mapping.source_structure_id}")
        
        response = requests.get(url, params=params, timeout=15)
        
        if response.status_code != 200:
            print(f"❌ Erreur GHP: {response.status_code} - {response.text[:100]}")
            return jsonify([])
        
        data = response.json()
        medicamentos = data.get('medicamentos', [])
        
        # Formater pour le médecin (seulement nom + stock)
        result = []
        for m in medicamentos:
            if m.get('nom'):
                result.append({
                    'id': m.get('ID'),
                    'nom': m.get('nom', ''),
                    'quantite_stock': m.get('quantite_stock', 0)
                })
        
        result.sort(key=lambda x: x['nom'])
        
        print(f"✅ {len(result)} médicaments disponibles chargés")
        return jsonify(result)
        
    except requests.exceptions.Timeout:
        print("❌ Timeout lors de la récupération des médicaments")
        return jsonify([])
    except requests.exceptions.ConnectionError:
        print("❌ Erreur de connexion à GHP")
        return jsonify([])
    except Exception as e:
        print(f"❌ Erreur récupération médicaments: {e}")
        import traceback
        traceback.print_exc()
        return jsonify([])

@app.route('/api/sync/prescriptions', methods=['POST'])
@login_required
def api_sync_prescriptions_to_ghp():
    """
    Envoie les prescriptions vers GHP
    """
    from models import Prescription, Patient, StructureMapping
    from datetime import datetime
    import requests

    try:
        # ⭐ Récupérer le mapping GHP
        mapping = StructureMapping.query.filter_by(
            local_structure_id=current_user.id_structure,
            actif=True
        ).first()

        if not mapping:
            return jsonify({'success': False, 'message': 'Configuration GHP non trouvée'}), 400

        # ⭐ Récupérer les prescriptions non synchronisées — UNIQUEMENT
        # celles dont le patient appartient à la structure de cet
        # utilisateur (sinon, avec plusieurs structures actives sur ce
        # déploiement, on enverrait aussi les prescriptions des AUTRES
        # structures avec le mapping/token de celle-ci).
        prescriptions = (
            Prescription.query
            .join(Patient, Prescription.id_patient == Patient.id)
            .filter(
                Prescription.synced_at.is_(None),
                Prescription.statut == 'active',
                Patient.id_structure == current_user.id_structure,
            )
            .all()
        )
        
        if not prescriptions:
            return jsonify({'success': True, 'message': 'Aucune prescription à synchroniser'})
        
        # ⭐ Formater les données
        data = []
        for p in prescriptions:
            # 🔥 Déterminer le type automatiquement
            type_presc = 'medicament'  # Par défaut
            
            # Si c'est un acte (ex: contient des mots-clés)
            mots_actes = ['examen', 'radio', 'scan', 'echo', 'analyse', 'test', 'biopsie', 'radiographie', 'irm']
            if p.medicament and any(mot in p.medicament.lower() for mot in mots_actes):
                type_presc = 'acte'
            
            # Ou si c'est un médicament (par défaut)
            presc_data = {
                'id': p.id,
                'patient_id': p.id_patient,
                'patient_nom': p.patient.nom if p.patient else '',
                'patient_prenom': p.patient.prenom if p.patient else '',
                'medicament': p.medicament or '',
                'type_prescription': type_presc,  # 🔥 AJOUT DU TYPE
                'dosage': p.dosage or '',
                'forme': p.forme or '',
                'quantite': p.quantite or '1',
                'duree_jours': p.duree_jours or 0,
                'frequence': p.frequence or '',
                'instructions': p.instructions or '',
                'date_prescription': p.date_prescription.isoformat() if p.date_prescription else datetime.now().isoformat(),
                'prescripteur': p.prescripteur or ''
            }
            data.append(presc_data)
        
        # ⭐ Envoyer vers GHP
        url = f"{mapping.api_url}/api/prescriptions"
        params = {'token': mapping.api_key}
        
        print(f"📡 Envoi de {len(data)} prescriptions vers GHP")
        
        response = requests.post(
            url,
            json={'prescriptions': data},
            params=params,
            timeout=30
        )
        
        if response.status_code == 200:
            # ⭐ Marquer comme synchronisées
            for p in prescriptions:
                p.synced_at = datetime.utcnow()
            db.session.commit()
            
            return jsonify({
                'success': True,
                'message': f'✅ {len(data)} prescriptions envoyées'
            })
        else:
            return jsonify({
                'success': False,
                'message': f'Erreur GHP: {response.status_code}',
                'response': response.text[:500]
            }), 500
            
    except Exception as e:
        print(f"❌ Erreur sync prescriptions: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'message': str(e)}), 500

# ================================================================
# ════════════════════════════════════════════════════════════════
# SECTION : GESTION DES TEMPLATES (PROTOCOLES, ORDONNANCES, EXAMENS)
# ════════════════════════════════════════════════════════════════
# ================================================================

# ================================================================
# 1. DASHBOARD DES TEMPLATES
# ================================================================

@app.route('/templates')
@login_required
def templates_dashboard():
    """Dashboard des templates médicaux"""
    from models import ProtocoleSoins, OrdonnanceType, ExamenType
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    structure_id = current_user.id_structure
    
    # Statistiques
    nb_protocoles = ProtocoleSoins.query.filter_by(structure_id=structure_id, actif=True).count()
    nb_ordonnances = OrdonnanceType.query.filter_by(structure_id=structure_id, actif=True).count()
    nb_examens = ExamenType.query.filter_by(structure_id=structure_id, actif=True).count()
    nb_inactifs = ProtocoleSoins.query.filter_by(structure_id=structure_id, actif=False).count() + \
                  OrdonnanceType.query.filter_by(structure_id=structure_id, actif=False).count() + \
                  ExamenType.query.filter_by(structure_id=structure_id, actif=False).count()
    
    # Derniers templates créés
    derniers_protocoles = ProtocoleSoins.query.filter_by(structure_id=structure_id).order_by(
        ProtocoleSoins.created_at.desc()
    ).limit(5).all()
    
    return render_template('templates/dashboard.html',
                         nb_protocoles=nb_protocoles,
                         nb_ordonnances=nb_ordonnances,
                         nb_examens=nb_examens,
                         nb_inactifs=nb_inactifs,
                         derniers_protocoles=derniers_protocoles)


# ================================================================
# 2. PROTOCOLES DE SOINS - CRUD
# ================================================================

@app.route('/templates/protocoles')
@login_required
def liste_protocoles():
    """Liste des protocoles de soins"""
    from models import ProtocoleSoins
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    protocoles = ProtocoleSoins.query.filter_by(
        structure_id=current_user.id_structure
    ).order_by(ProtocoleSoins.nom).all()
    
    return render_template('templates/protocoles/liste.html', protocoles=protocoles)


def _lire_fichier_source(champ='fichier_source'):
    """Lit le fichier optionnel joint au formulaire (docx/pdf importé) et
    renvoie (nom, mime, data) ou (None, None, None) — patron : "le fichier
    importé ne s'affiche pas dans le dossier du patient" : l'import ne
    gardait jusqu'ici que le texte extrait (voir api_import_fichier_texte),
    jamais le fichier lui-même. Utilisé par les 6 routes de création/
    modification de protocole/ordonnance-type/examen-type."""
    fichier = request.files.get(champ)
    if not fichier or not fichier.filename:
        return None, None, None
    return fichier.filename, fichier.mimetype, fichier.read()


def _fichier_source_info(source_type, source_id, categorie):
    """Résout la provenance (source_type/source_id, voir Ordonnance et
    ExamenPrescrit) d'une prescription vers le fichier importé du modèle
    dont elle vient — directement (source_type='template') ou via le
    ProtocoleSoins associé (source_type='protocole') — pour que ce fichier
    réapparaisse enfin dans le dossier du patient. categorie : 'ordonnance'
    ou 'examen'. Renvoie {'nom':..., 'url':...} ou None.

    ⭐ Pour un protocole : le fichier importé sur le PROTOCOLE lui-même
    (document général couvrant ordonnance+examens) est prioritaire s'il
    existe ; à défaut, on retombe sur celui de l'ordonnance-type/examen-type
    associé — sinon un fichier importé directement sur le protocole
    n'apparaissait nulle part (patron : "avec ceci si on importe un fichier
    de protocole dans le dossier il sera bien visible ?" — non, avant ce
    correctif)."""
    if not source_id:
        return None
    from models import OrdonnanceType, ExamenType, ProtocoleSoins

    modele = None
    if source_type == 'template':
        modele = (OrdonnanceType if categorie == 'ordonnance' else ExamenType).query.get(source_id)
    elif source_type == 'protocole':
        protocole = ProtocoleSoins.query.get(source_id)
        if protocole:
            if protocole.fichier_nom:
                modele = protocole
            else:
                modele = protocole.ordonnance_type if categorie == 'ordonnance' else protocole.examen_type

    if not modele or not modele.fichier_nom:
        return None

    if source_type == 'protocole' and modele.__class__.__name__ == 'ProtocoleSoins':
        endpoint = 'api_telecharger_fichier_protocole'
    else:
        endpoint = 'api_telecharger_fichier_ordonnance_type' if categorie == 'ordonnance' else 'api_telecharger_fichier_examen_type'
    return {'nom': modele.fichier_nom, 'url': url_for(endpoint, id=modele.id)}


def _telecharger_fichier_source(modele, nom_defaut):
    """Sert fichier_data/fichier_nom/fichier_mime d'un ProtocoleSoins/
    OrdonnanceType/ExamenType — même patron que api_telecharger_modele_resultat."""
    from flask import Response
    if not modele.fichier_data:
        return "Aucun fichier importé pour cet élément.", 404
    return Response(
        modele.fichier_data, mimetype=modele.fichier_mime or 'application/octet-stream',
        headers={'Content-Disposition': f'attachment; filename="{modele.fichier_nom or nom_defaut}"'}
    )


@app.route('/templates/protocole/<int:id>/fichier')
@login_required
def api_telecharger_fichier_protocole(id):
    from models import ProtocoleSoins
    protocole = ProtocoleSoins.query.filter_by(id=id, structure_id=current_user.id_structure).first()
    if not protocole:
        return "Protocole introuvable", 404
    return _telecharger_fichier_source(protocole, 'protocole')


@app.route('/templates/ordonnance/<int:id>/fichier')
@login_required
def api_telecharger_fichier_ordonnance_type(id):
    from models import OrdonnanceType
    ordonnance = OrdonnanceType.query.filter_by(id=id, structure_id=current_user.id_structure).first()
    if not ordonnance:
        return "Ordonnance type introuvable", 404
    return _telecharger_fichier_source(ordonnance, 'ordonnance')


@app.route('/templates/examen/<int:id>/fichier')
@login_required
def api_telecharger_fichier_examen_type(id):
    from models import ExamenType
    examen = ExamenType.query.filter_by(id=id, structure_id=current_user.id_structure).first()
    if not examen:
        return "Examen type introuvable", 404
    return _telecharger_fichier_source(examen, 'examen')


@app.route('/templates/protocole/ajouter', methods=['GET', 'POST'])
@login_required
def ajouter_protocole():
    """Ajouter un nouveau protocole de soins"""
    from models import ProtocoleSoins, OrdonnanceType, ExamenType
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    # Récupérer les ordonnances et examens types pour les associations
    ordonnances = OrdonnanceType.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    examens = ExamenType.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    if request.method == 'POST':
        nom = request.form.get('nom', '').strip()
        description = request.form.get('description', '').strip()
        ordonnance_type_id = request.form.get('ordonnance_type_id')
        examen_type_id = request.form.get('examen_type_id')
        
        if not nom or not description:
            flash('Le nom et la description sont obligatoires', 'danger')
            return redirect(url_for('ajouter_protocole'))
        
        fichier_nom, fichier_mime, fichier_data = _lire_fichier_source()

        protocole = ProtocoleSoins(
            structure_id=current_user.id_structure,
            nom=nom,
            description=description,
            ordonnance_type_id=int(ordonnance_type_id) if ordonnance_type_id else None,
            examen_type_id=int(examen_type_id) if examen_type_id else None,
            fichier_nom=fichier_nom,
            fichier_mime=fichier_mime,
            fichier_data=fichier_data,
            created_by=current_user.id,
            actif=True
        )

        db.session.add(protocole)
        db.session.commit()

        _pousser_protocole_ghp(
            'protocole_soins', 'ProtocoleSoins', protocole.id, protocole.structure_id,
            titre=nom, description=description,
            contenu=_generer_contenu_protocole_soins(nom, description), actif=True,
        )

        flash(f'Protocole "{nom}" créé avec succès', 'success')
        return redirect(url_for('liste_protocoles'))
    
    return render_template('templates/protocoles/ajouter.html',
                         ordonnances=ordonnances,
                         examens=examens)


@app.route('/templates/protocole/<int:id>/modifier', methods=['GET', 'POST'])
@login_required
def modifier_protocole(id):
    """Modifier un protocole de soins"""
    from models import ProtocoleSoins, OrdonnanceType, ExamenType
    
    protocole = ProtocoleSoins.query.get_or_404(id)
    
    if protocole.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_protocoles'))
    
    ordonnances = OrdonnanceType.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    examens = ExamenType.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    if request.method == 'POST':
        protocole.nom = request.form.get('nom', '').strip()
        protocole.description = request.form.get('description', '').strip()
        protocole.ordonnance_type_id = request.form.get('ordonnance_type_id', type=int) or None
        protocole.examen_type_id = request.form.get('examen_type_id', type=int) or None
        protocole.actif = request.form.get('actif') == 'on'
        protocole.updated_at = datetime.utcnow()

        fichier_nom, fichier_mime, fichier_data = _lire_fichier_source()
        if fichier_nom:
            protocole.fichier_nom = fichier_nom
            protocole.fichier_mime = fichier_mime
            protocole.fichier_data = fichier_data

        db.session.commit()

        _pousser_protocole_ghp(
            'protocole_soins', 'ProtocoleSoins', protocole.id, protocole.structure_id,
            titre=protocole.nom, description=protocole.description,
            contenu=_generer_contenu_protocole_soins(protocole.nom, protocole.description),
            actif=protocole.actif,
        )

        flash(f'Protocole "{protocole.nom}" modifié avec succès', 'success')
        return redirect(url_for('liste_protocoles'))
    
    return render_template('templates/protocoles/modifier.html',
                         protocole=protocole,
                         ordonnances=ordonnances,
                         examens=examens)


@app.route('/templates/protocole/<int:id>/supprimer', methods=['POST'])
@login_required
def supprimer_protocole(id):
    """Supprimer un protocole de soins"""
    from models import ProtocoleSoins
    
    protocole = ProtocoleSoins.query.get_or_404(id)
    
    if protocole.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_protocoles'))
    
    nom = protocole.nom
    _pousser_protocole_ghp(
        'protocole_soins', 'ProtocoleSoins', protocole.id, protocole.structure_id,
        titre=nom, action='archive',
    )
    db.session.delete(protocole)
    db.session.commit()

    flash(f'Protocole "{nom}" supprimé avec succès', 'success')
    return redirect(url_for('liste_protocoles'))


# ================================================================
# 3. ORDONNANCES TYPES - CRUD
# ================================================================

@app.route('/templates/ordonnances')
@login_required
def liste_ordonnances():
    """Liste des ordonnances types"""
    from models import OrdonnanceType
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    ordonnances = OrdonnanceType.query.filter_by(
        structure_id=current_user.id_structure
    ).order_by(OrdonnanceType.nom).all()
    
    return render_template('templates/ordonnances/liste.html', ordonnances=ordonnances)


@app.route('/templates/ordonnance/ajouter', methods=['GET', 'POST'])
@login_required
def ajouter_ordonnance():
    """Ajouter une nouvelle ordonnance type"""
    from models import OrdonnanceType
    import json
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    if request.method == 'POST':
        nom = request.form.get('nom', '').strip()
        description = request.form.get('description', '').strip()
        medicaments_json = request.form.get('medicaments_json', '[]')
        
        if not nom:
            flash('Le nom est obligatoire', 'danger')
            return redirect(url_for('ajouter_ordonnance'))
        
        # Vérifier que le JSON est valide
        try:
            medicaments = json.loads(medicaments_json)
        except:
            flash('Format des médicaments invalide', 'danger')
            return redirect(url_for('ajouter_ordonnance'))
        
        fichier_nom, fichier_mime, fichier_data = _lire_fichier_source()

        ordonnance = OrdonnanceType(
            structure_id=current_user.id_structure,
            nom=nom,
            description=description,
            medicaments=medicaments_json,
            fichier_nom=fichier_nom,
            fichier_mime=fichier_mime,
            fichier_data=fichier_data,
            created_by=current_user.id,
            actif=True
        )

        db.session.add(ordonnance)
        db.session.commit()

        _pousser_protocole_ghp(
            'ordonnance_type', 'OrdonnanceType', ordonnance.id, ordonnance.structure_id,
            titre=nom, description=description,
            contenu=_generer_contenu_ordonnance(nom, medicaments),
            medicaments=medicaments, actif=True,
        )

        flash(f'Ordonnance "{nom}" créée avec succès', 'success')
        return redirect(url_for('liste_ordonnances'))
    
    return render_template('templates/ordonnances/ajouter.html')


@app.route('/templates/ordonnance/<int:id>/modifier', methods=['GET', 'POST'])
@login_required
def modifier_ordonnance(id):
    """Modifier une ordonnance type"""
    from models import OrdonnanceType
    import json
    
    ordonnance = OrdonnanceType.query.get_or_404(id)
    
    if ordonnance.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_ordonnances'))
    
    if request.method == 'POST':
        ordonnance.nom = request.form.get('nom', '').strip()
        ordonnance.description = request.form.get('description', '').strip()
        ordonnance.medicaments = request.form.get('medicaments_json', '[]')
        ordonnance.actif = request.form.get('actif') == 'on'
        ordonnance.updated_at = datetime.utcnow()

        fichier_nom, fichier_mime, fichier_data = _lire_fichier_source()
        if fichier_nom:
            ordonnance.fichier_nom = fichier_nom
            ordonnance.fichier_mime = fichier_mime
            ordonnance.fichier_data = fichier_data

        db.session.commit()

        try:
            medicaments = json.loads(ordonnance.medicaments or '[]')
        except Exception:
            medicaments = []
        _pousser_protocole_ghp(
            'ordonnance_type', 'OrdonnanceType', ordonnance.id, ordonnance.structure_id,
            titre=ordonnance.nom, description=ordonnance.description,
            contenu=_generer_contenu_ordonnance(ordonnance.nom, medicaments),
            medicaments=medicaments, actif=ordonnance.actif,
        )

        flash(f'Ordonnance "{ordonnance.nom}" modifiée avec succès', 'success')
        return redirect(url_for('liste_ordonnances'))
    
    return render_template('templates/ordonnances/modifier.html', ordonnance=ordonnance)


@app.route('/templates/ordonnance/<int:id>/supprimer', methods=['POST'])
@login_required
def supprimer_ordonnance(id):
    """Supprimer une ordonnance type"""
    from models import OrdonnanceType
    
    ordonnance = OrdonnanceType.query.get_or_404(id)
    
    if ordonnance.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_ordonnances'))
    
    nom = ordonnance.nom
    _pousser_protocole_ghp(
        'ordonnance_type', 'OrdonnanceType', ordonnance.id, ordonnance.structure_id,
        titre=nom, action='archive',
    )
    db.session.delete(ordonnance)
    db.session.commit()

    flash(f'Ordonnance "{nom}" supprimée avec succès', 'success')
    return redirect(url_for('liste_ordonnances'))


# ================================================================
# 4. EXAMENS TYPES - CRUD
# ================================================================

@app.route('/templates/examens')
@login_required
def liste_examens():
    """Liste des examens types"""
    from models import ExamenType
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    examens = ExamenType.query.filter_by(
        structure_id=current_user.id_structure
    ).order_by(ExamenType.nom).all()
    
    return render_template('templates/examens/liste.html', examens=examens)


@app.route('/templates/examen/ajouter', methods=['GET', 'POST'])
@login_required
def ajouter_examen():
    """Ajouter un nouvel examen type"""
    from models import ExamenType
    import json
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('dashboard'))
    
    if request.method == 'POST':
        nom = request.form.get('nom', '').strip()
        nature = request.form.get('nature', '').strip()
        motif = request.form.get('motif', '').strip()
        description = request.form.get('description', '').strip()
        examens_json = request.form.get('examens_json', '[]')
        
        if not nom or not nature:
            flash('Le nom et la nature sont obligatoires', 'danger')
            return redirect(url_for('ajouter_examen'))
        
        try:
            examens = json.loads(examens_json)
        except:
            flash('Format des examens invalide', 'danger')
            return redirect(url_for('ajouter_examen'))
        
        fichier_nom, fichier_mime, fichier_data = _lire_fichier_source()

        examen = ExamenType(
            structure_id=current_user.id_structure,
            nom=nom,
            nature=nature,
            motif=motif,
            description=description,
            examens=examens_json,
            fichier_nom=fichier_nom,
            fichier_mime=fichier_mime,
            fichier_data=fichier_data,
            created_by=current_user.id,
            actif=True
        )

        db.session.add(examen)
        db.session.commit()

        _pousser_protocole_ghp(
            'bulletin_examen', 'ExamenType', examen.id, examen.structure_id,
            titre=nom, description=description,
            contenu=_generer_contenu_bulletin(nom, motif, examens),
            examens=examens, actif=True,
        )

        flash(f'Examen type "{nom}" créé avec succès', 'success')
        return redirect(url_for('liste_examens'))
    
    return render_template('templates/examens/ajouter.html')


@app.route('/templates/examen/<int:id>/modifier', methods=['GET', 'POST'])
@login_required
def modifier_examen(id):
    """Modifier un examen type"""
    from models import ExamenType
    import json
    
    examen = ExamenType.query.get_or_404(id)
    
    if examen.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_examens'))
    
    if request.method == 'POST':
        examen.nom = request.form.get('nom', '').strip()
        examen.nature = request.form.get('nature', '').strip()
        examen.motif = request.form.get('motif', '').strip()
        examen.description = request.form.get('description', '').strip()
        examen.examens = request.form.get('examens_json', '[]')
        examen.actif = request.form.get('actif') == 'on'
        examen.updated_at = datetime.utcnow()

        fichier_nom, fichier_mime, fichier_data = _lire_fichier_source()
        if fichier_nom:
            examen.fichier_nom = fichier_nom
            examen.fichier_mime = fichier_mime
            examen.fichier_data = fichier_data

        db.session.commit()

        try:
            examens_list = json.loads(examen.examens or '[]')
        except Exception:
            examens_list = []
        _pousser_protocole_ghp(
            'bulletin_examen', 'ExamenType', examen.id, examen.structure_id,
            titre=examen.nom, description=examen.description,
            contenu=_generer_contenu_bulletin(examen.nom, examen.motif, examens_list),
            examens=examens_list, actif=examen.actif,
        )

        flash(f'Examen type "{examen.nom}" modifié avec succès', 'success')
        return redirect(url_for('liste_examens'))
    
    return render_template('templates/examens/modifier.html', examen=examen)


@app.route('/templates/examen/<int:id>/supprimer', methods=['POST'])
@login_required
def supprimer_examen(id):
    """Supprimer un examen type"""
    from models import ExamenType
    
    examen = ExamenType.query.get_or_404(id)
    
    if examen.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_examens'))
    
    nom = examen.nom
    _pousser_protocole_ghp(
        'bulletin_examen', 'ExamenType', examen.id, examen.structure_id,
        titre=nom, action='archive',
    )
    db.session.delete(examen)
    db.session.commit()

    flash(f'Examen type "{nom}" supprimé avec succès', 'success')
    return redirect(url_for('liste_examens'))

# ================================================================
# HOSPITALISATION - GESTION DU PROTOCOLE, ORDONNANCE ET EXAMENS
# ================================================================

@app.route('/hospitalisation/<int:id>/appliquer-protocole', methods=['POST'])
@login_required
def appliquer_protocole(id):
    """Appliquer un protocole de soins à une hospitalisation"""
    from models import Hospitalisation, ProtocoleSoins, OrdonnanceType, ExamenType, ExamenPrescrit
    import json
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    protocole_id = request.form.get('protocole_id', type=int)
    
    if not protocole_id:
        flash('Veuillez sélectionner un protocole', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    protocole = ProtocoleSoins.query.get_or_404(protocole_id)
    
    if protocole.structure_id != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    try:
        # Appliquer le protocole
        hospitalisation.protocole_id = protocole.id
        hospitalisation.updated_at = datetime.utcnow()
        
        # Si le protocole a une ordonnance associée, la copier
        medicaments_pour_ghp = []
        if protocole.ordonnance_type_id:
            ordonnance_type = OrdonnanceType.query.get(protocole.ordonnance_type_id)
            if ordonnance_type:
                anciens_medicaments = []
                if hospitalisation.ordonnance_prescite:
                    try:
                        anciens_medicaments = json.loads(hospitalisation.ordonnance_prescite)
                    except Exception:
                        anciens_medicaments = []
                try:
                    nouveaux_medicaments = json.loads(ordonnance_type.medicaments) if ordonnance_type.medicaments else []
                except Exception:
                    nouveaux_medicaments = []
                hospitalisation.ordonnance_prescite = ordonnance_type.medicaments
                # ⭐ Miroir + synchro GHP — même trou que modifier/creer_ordonnance_hospitalisation :
                # une ordonnance copiée depuis un protocole ne partait jamais vers GHP.
                medicaments_pour_ghp = _items_nouveaux(nouveaux_medicaments, anciens_medicaments)
                if medicaments_pour_ghp:
                    prescriptions_creees = _creer_prescriptions_miroir(
                        patient_id=hospitalisation.patient_id,
                        prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
                        items=medicaments_pour_ghp,
                        type_prescription='medicament'
                    )
                    # ⭐ Traçabilité côté infirmier : "cette dose vient du
                    # protocole X" (patron : "à un protocole on peut
                    # assigner une ordonnance et examens à faire").
                    for p in prescriptions_creees:
                        p.protocole_id = protocole.id

        # Si le protocole a des examens associés, les créer
        examens_pour_ghp = []
        if protocole.examen_type_id:
            examen_type = ExamenType.query.get(protocole.examen_type_id)
            if examen_type:
                # Créer l'examen prescrit
                examen_prescrit = ExamenPrescrit(
                    hospitalisation_id=hospitalisation.id,
                    patient_id=hospitalisation.patient_id,
                    medecin_id=current_user.id,
                    examen_type_id=examen_type.id,
                    nature=examen_type.nature,
                    motif=examen_type.motif,
                    description=examen_type.description,
                    examens=examen_type.examens,
                    statut='EN_ATTENTE',
                    date_prescription=datetime.utcnow(),
                    # ⭐ Traçabilité (champs déjà prévus pour ça, jusqu'ici
                    # jamais renseignés par cette route) : "cet examen vient
                    # du protocole X".
                    source_type='protocole',
                    source_id=protocole.id,
                    source_nom=protocole.nom
                )
                db.session.add(examen_prescrit)
                db.session.flush()
                try:
                    examens_pour_ghp = json.loads(examen_type.examens) if examen_type.examens else []
                except Exception:
                    examens_pour_ghp = []
                # ⭐ Pont vers /analyses (labo/radio) + synchro GHP — jusqu'ici
                # un protocole avec examen ne faisait NI l'un NI l'autre,
                # contrairement aux autres façons de prescrire un examen.
                _creer_analyses_demandees_depuis_examen_prescrit(
                    examen_prescrit, examens_pour_ghp, current_user.id_structure
                )
                examens_prescriptions_creees = _creer_prescriptions_miroir(
                    patient_id=hospitalisation.patient_id,
                    prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
                    items=examens_pour_ghp,
                    type_prescription='acte'
                )
                for p in examens_prescriptions_creees:
                    p.protocole_id = protocole.id

        db.session.commit()
        if examens_pour_ghp or medicaments_pour_ghp:
            _envoyer_prescriptions_ghp_immediat()

        flash(f'Protocole "{protocole.nom}" appliqué avec succès', 'success')
        
    except Exception as e:
        db.session.rollback()
        flash(f'Erreur : {str(e)}', 'danger')
    
    return redirect(url_for('detail_hospitalisation', id=id))


@app.route('/hospitalisation/<int:id>/ordonnance/modifier', methods=['POST'])
@login_required
def modifier_ordonnance_hospitalisation(id):
    """Modifier une ordonnance existante (crée une nouvelle version)"""
    from models import Hospitalisation
    import json
    from datetime import datetime
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    if hospitalisation.statut != 'actif':
        flash('Impossible de modifier une hospitalisation clôturée', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    medicaments_json = request.form.get('medicaments_json', '[]')
    motif_modification = request.form.get('motif_modification', 'Modification')
    
    try:
        medicaments = json.loads(medicaments_json)
    except:
        flash('Format des médicaments invalide', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    if not medicaments:
        flash('Veuillez ajouter au moins un médicament', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    try:
        # Récupérer l'historique existant
        historique = []
        if hospitalisation.ordonnance_historique:
            try:
                historique = json.loads(hospitalisation.ordonnance_historique)
            except:
                historique = []
        
        # Sauvegarder l'ancienne ordonnance dans l'historique
        anciens_medicaments = []
        if hospitalisation.ordonnance_prescite:
            ancienne_version = json.loads(hospitalisation.ordonnance_prescite)
            anciens_medicaments = ancienne_version
            historique.append({
                'version': hospitalisation.ordonnance_version or 1,
                'date': datetime.utcnow().isoformat(),
                'medicaments': ancienne_version,
                'prescrit_par': current_user.id,
                'prescrit_par_nom': f"{current_user.prenom} {current_user.nom}",
                'motif': motif_modification
            })

        # Sauvegarder l'historique
        hospitalisation.ordonnance_historique = json.dumps(historique, ensure_ascii=False)

        # Incrémenter la version
        hospitalisation.ordonnance_version = (hospitalisation.ordonnance_version or 0) + 1

        # Mettre à jour la nouvelle ordonnance
        hospitalisation.ordonnance_prescite = medicaments_json
        hospitalisation.updated_at = datetime.utcnow()
        hospitalisation.created_by = current_user.id

        # ⭐ Miroir Prescription pour la synchronisation GHP — mêmes principes
        # que creer_ordonnance_hospitalisation : seuls les médicaments
        # réellement nouveaux par rapport à la version précédente sont
        # renvoyés (évite les doublons à chaque modification/réimpression).
        # Avant ce correctif, une ordonnance modifiée en hospitalisation
        # n'apparaissait jamais dans "Prescriptions reçues" côté GHP.
        nouveaux = _items_nouveaux(medicaments, anciens_medicaments)
        if nouveaux:
            _creer_prescriptions_miroir(
                patient_id=hospitalisation.patient_id,
                prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
                items=nouveaux,
                type_prescription='medicament'
            )

        db.session.commit()
        if nouveaux:
            _envoyer_prescriptions_ghp_immediat()

        flash(f'✅ Ordonnance modifiée - Version {hospitalisation.ordonnance_version} créée', 'success')
        
    except Exception as e:
        db.session.rollback()
        flash(f'❌ Erreur : {str(e)}', 'danger')

    return redirect(url_for('detail_hospitalisation', id=id))


@app.route('/hospitalisation/<int:id>/ordonnance-sortie', methods=['POST'])
@login_required
def definir_ordonnance_sortie(id):
    """Ordonnance de sortie — médicaments à poursuivre par le patient
    APRÈS l'hospitalisation, distincte de l'ordonnance de séjour
    (hospitalisation.ordonnance_prescite). Un seul document final (pas de
    versionnement séparé) : ré-enregistrer ce champ le corrige simplement.
    Rédigeable dès que l'hospitalisation n'est plus active — écrire
    l'ordonnance de sortie pendant le séjour n'a pas de sens."""
    from models import Hospitalisation
    import json
    from datetime import datetime

    hospitalisation = Hospitalisation.query.get_or_404(id)

    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))

    if hospitalisation.patient.id_structure != current_user.id_structure:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('liste_hospitalisations'))

    if hospitalisation.statut == 'actif':
        flash('L\'ordonnance de sortie se rédige une fois l\'hospitalisation clôturée', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))

    medicaments_json = request.form.get('medicaments_json', '[]')
    try:
        medicaments = json.loads(medicaments_json)
    except Exception:
        flash('Format des médicaments invalide', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))

    if not medicaments:
        flash('Veuillez ajouter au moins un médicament', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))

    try:
        anciens_medicaments = []
        if hospitalisation.ordonnance_sortie:
            try:
                anciens_medicaments = json.loads(hospitalisation.ordonnance_sortie)
            except Exception:
                anciens_medicaments = []

        hospitalisation.ordonnance_sortie = medicaments_json
        hospitalisation.ordonnance_sortie_date = datetime.utcnow()
        hospitalisation.ordonnance_sortie_par = current_user.id

        # ⭐ Miroir Prescription pour la synchronisation GHP — même principe
        # que l'ordonnance de séjour (modifier_ordonnance_hospitalisation) :
        # uniquement les médicaments réellement nouveaux par rapport à une
        # correction précédente, pour éviter les doublons.
        nouveaux = _items_nouveaux(medicaments, anciens_medicaments)

        db.session.commit()

        if nouveaux:
            _creer_prescriptions_miroir(
                patient_id=hospitalisation.patient_id,
                prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
                items=nouveaux,
                type_prescription='medicament'
            )
            db.session.commit()
            _envoyer_prescriptions_ghp_immediat()

        flash('✅ Ordonnance de sortie enregistrée', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'❌ Erreur : {str(e)}', 'danger')

    return redirect(url_for('detail_hospitalisation', id=id))


@app.route('/hospitalisation/<int:id>/ordonnance-sortie/imprimer')
@login_required
def imprimer_ordonnance_sortie(id):
    """Version imprimable de l'ordonnance de sortie — remise au patient."""
    from models import Hospitalisation, Structure
    import json

    hospitalisation = Hospitalisation.query.get_or_404(id)
    structure = Structure.query.get(current_user.id_structure)
    patient = hospitalisation.patient

    if not hospitalisation.ordonnance_sortie:
        flash('Aucune ordonnance de sortie pour cette hospitalisation', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))

    try:
        medicaments = json.loads(hospitalisation.ordonnance_sortie)
    except Exception:
        medicaments = []

    prescripteur = (
        f"{hospitalisation.prescripteur_sortie.prenom} {hospitalisation.prescripteur_sortie.nom}"
        if hospitalisation.prescripteur_sortie
        else f"{current_user.prenom} {current_user.nom}"
    )

    return render_template('impressions/ordonnance_sortie.html',
                         hospitalisation=hospitalisation,
                         structure=structure,
                         patient=patient,
                         medicaments=medicaments,
                         prescripteur=prescripteur,
                         date_ordonnance=hospitalisation.ordonnance_sortie_date)


@app.route('/hospitalisation/<int:id>/ordonnance/historique')
@login_required
def historique_ordonnances_hospitalisation(id):
    from models import Hospitalisation
    import json
    from datetime import datetime
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    
    # Récupérer l'historique
    historique = []
    if hospitalisation.ordonnance_historique:
        try:
            historique = json.loads(hospitalisation.ordonnance_historique)
            # Formater les dates
            for v in historique:
                if 'date' in v and v['date']:
                    try:
                        dt = datetime.fromisoformat(v['date'].replace('Z', '+00:00'))
                        v['date_formatted'] = dt.strftime('%d/%m/%Y %H:%M')
                    except:
                        v['date_formatted'] = v['date']
                else:
                    v['date_formatted'] = '-'
        except:
            historique = []
    
    # Récupérer l'ordonnance actuelle
    ordonnance_actuelle_medicaments = []
    if hospitalisation.ordonnance_prescite:
        try:
            ordonnance_actuelle_medicaments = json.loads(hospitalisation.ordonnance_prescite)
        except:
            ordonnance_actuelle_medicaments = []
    
    return render_template('hospitalisations/historique_ordonnances.html',
                         hospitalisation=hospitalisation,
                         historique=historique,
                         ordonnance_prescite=hospitalisation.ordonnance_prescite,
                         ordonnance_actuelle_medicaments=ordonnance_actuelle_medicaments,
                         version_actuelle=hospitalisation.ordonnance_version or 1)

@app.route('/hospitalisation/<int:id>/ordonnance/version/<int:version>/imprimer')
@login_required
def imprimer_ordonnance_version(id, version):
    """Imprimer une version spécifique de l'ordonnance"""
    from models import Hospitalisation, Structure
    import json
    from datetime import datetime
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    structure = Structure.query.get(current_user.id_structure)
    
    # ⭐ RÉCUPÉRER LE PATIENT
    patient = hospitalisation.patient
    
    # ⭐ SI C'EST LA VERSION ACTUELLE
    if version == hospitalisation.ordonnance_version:
        medicaments = []
        if hospitalisation.ordonnance_prescite:
            try:
                medicaments = json.loads(hospitalisation.ordonnance_prescite)
            except:
                medicaments = []
        
        prescripteur = f"{hospitalisation.createur.prenom} {hospitalisation.createur.nom}" if hospitalisation.createur else current_user.prenom + " " + current_user.nom
        date_version = hospitalisation.updated_at.strftime('%d/%m/%Y %H:%M') if hospitalisation.updated_at else datetime.now().strftime('%d/%m/%Y %H:%M')
        
        return render_template('impressions/ordonnance_version.html',
                             hospitalisation=hospitalisation,
                             structure=structure,
                             patient=patient,
                             medicaments=medicaments,
                             version=version,
                             prescripteur=prescripteur,
                             date_version=date_version,
                             now=datetime.utcnow())
    
    # ⭐ SI C'EST UNE VERSION ANCIENNE
    historique = []
    if hospitalisation.ordonnance_historique:
        try:
            historique = json.loads(hospitalisation.ordonnance_historique)
        except:
            historique = []
    
    # Chercher la version
    version_data = None
    for v in historique:
        if v.get('version') == version:
            version_data = v
            break
    
    if not version_data:
        flash('Version non trouvée', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    medicaments = version_data.get('medicaments', [])
    prescripteur = version_data.get('prescrit_par_nom', 'Inconnu')
    date_version = version_data.get('date_formatted', version_data.get('date', datetime.now().strftime('%d/%m/%Y %H:%M')))
    
    return render_template('impressions/ordonnance_version.html',
                         hospitalisation=hospitalisation,
                         structure=structure,
                         patient=patient,
                         medicaments=medicaments,
                         version=version,
                         prescripteur=prescripteur,
                         date_version=date_version,
                         now=datetime.utcnow())

@app.route('/hospitalisation/<int:id>/ordonnance/creer', methods=['POST'])
@login_required
def creer_ordonnance_hospitalisation(id):
    """Créer une nouvelle ordonnance pour une hospitalisation"""
    from models import Hospitalisation
    import json
    from datetime import datetime
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    if hospitalisation.statut != 'actif':
        flash('Impossible de créer une ordonnance sur une hospitalisation clôturée', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    medicaments_json = request.form.get('medicaments_json', '[]')
    
    try:
        medicaments = json.loads(medicaments_json)
    except:
        flash('Format des médicaments invalide', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    if not medicaments:
        flash('Veuillez ajouter au moins un médicament', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    try:
        # ⭐ SAUVEGARDER L'ANCIENNE ORDONNANCE DANS L'HISTORIQUE (SI ELLE EXISTE)
        historique = []
        if hospitalisation.ordonnance_historique:
            try:
                historique = json.loads(hospitalisation.ordonnance_historique)
            except:
                historique = []
        
        # Si une ancienne ordonnance existe, la sauvegarder dans l'historique
        anciens_medicaments = []
        if hospitalisation.ordonnance_prescite:
            ancienne_version = json.loads(hospitalisation.ordonnance_prescite)
            anciens_medicaments = ancienne_version
            historique.append({
                'version': hospitalisation.ordonnance_version or 1,
                'date': datetime.utcnow().isoformat(),
                'medicaments': ancienne_version,
                'prescrit_par': current_user.id,
                'prescrit_par_nom': f"{current_user.prenom} {current_user.nom}",
                'motif': 'Nouvelle prescription'
            })

        # Sauvegarder l'historique
        hospitalisation.ordonnance_historique = json.dumps(historique, ensure_ascii=False)

        # Incrémenter la version
        hospitalisation.ordonnance_version = (hospitalisation.ordonnance_version or 0) + 1

        # Mettre à jour la nouvelle ordonnance
        hospitalisation.ordonnance_prescite = medicaments_json
        hospitalisation.updated_at = datetime.utcnow()
        hospitalisation.created_by = current_user.id

        # ⭐ Miroir Prescription pour la synchronisation GHP — uniquement les
        # médicaments nouveaux par rapport à la version précédente (le champ
        # ordonnance_prescite est réécrit en entier à chaque version, pas un
        # historique ligne par ligne comme Prescription).
        nouveaux = _items_nouveaux(medicaments, anciens_medicaments)
        if nouveaux:
            _creer_prescriptions_miroir(
                patient_id=hospitalisation.patient_id,
                prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
                items=nouveaux,
                type_prescription='medicament'
            )

        db.session.commit()
        if nouveaux:
            _envoyer_prescriptions_ghp_immediat()

        flash(f'✅ Nouvelle prescription enregistrée (Version {hospitalisation.ordonnance_version})', 'success')
        
    except Exception as e:
        db.session.rollback()
        flash(f'❌ Erreur : {str(e)}', 'danger')
    
    return redirect(url_for('detail_hospitalisation', id=id))

@app.route('/patient/<int:patient_id>/dossier/imprimer')
@login_required
def imprimer_dossier_patient(patient_id):
    from models import Patient, Consultation, Prescription, AntecedentPatient, Hospitalisation, AnalyseDemande, ExamenPhysique, Utilisateur, ExamenPrescrit, AvisExterne, EvolutionPatient, ConstanteVitale, HospitalisationMedecin, HospitalisationInfirmier, NoteAdmission, Engagement
    from datetime import datetime
    import json
    
    patient = Patient.query.get_or_404(patient_id)
    
    if current_user.role not in ['admin_structure', 'medecin', 'super_admin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('patient_detail', id=patient_id))
    
    # ============================================================ #
    # 1. CONSULTATIONS
    # ============================================================ #
    consultations = Consultation.query.filter_by(
        id_patient=patient.id
    ).order_by(Consultation.date_consultation.desc()).all()
    
    consultations_data = []
    for consultation in consultations:
        medecin = None
        if consultation.id_medecin:
            medecin = Utilisateur.query.get(consultation.id_medecin)
        
        prescriptions = Prescription.query.filter_by(
            id_patient=patient.id,
            id_consultation=consultation.id,
            type_prescription='medicament'
        ).all()
        
        examens_prescrits = ExamenPrescrit.query.filter_by(
            consultation_id=consultation.id,
            est_active=True
        ).all()
        
        analyses = AnalyseDemande.query.filter_by(
            consultation_id=consultation.id
        ).all()
        
        examen_physique = ExamenPhysique.query.filter_by(
            consultation_id=consultation.id
        ).first()
        
        consultations_data.append({
            'consultation': consultation,
            'medecin': medecin,
            'prescriptions': prescriptions,
            'examens_prescrits': examens_prescrits,
            'analyses': analyses,
            'examen_physique': examen_physique
        })
    
    # ============================================================ #
    # 2. HOSPITALISATIONS
    # ============================================================ #
    hospitalisations = Hospitalisation.query.filter_by(
        patient_id=patient.id
    ).order_by(Hospitalisation.date_debut.desc()).all()
    
    hospitalisations_data = []
    for hosp in hospitalisations:
        medecins = HospitalisationMedecin.query.filter_by(
            hospitalisation_id=hosp.id,
            actif=True
        ).all()
        
        infirmiers = HospitalisationInfirmier.query.filter_by(
            hospitalisation_id=hosp.id,
            actif=True
        ).all()
        
        constantes = ConstanteVitale.query.filter_by(
            hospitalisation_id=hosp.id
        ).order_by(ConstanteVitale.date_prise.desc()).limit(20).all()
        
        evolutions = EvolutionPatient.query.filter_by(
            hospitalisation_id=hosp.id
        ).order_by(EvolutionPatient.date_evolution.desc()).all()
        
        avis_externes = AvisExterne.query.filter_by(
            hospitalisation_id=hosp.id
        ).order_by(AvisExterne.date_demande.desc()).all()
        
        note_active = None
        if hosp.note_admission_active_id:
            note_active = NoteAdmission.query.get(hosp.note_admission_active_id)
        
        ordonnance_medicaments = []
        if hosp.ordonnance_prescite:
            try:
                ordonnance_medicaments = json.loads(hosp.ordonnance_prescite)
            except:
                ordonnance_medicaments = []
        
        examens_hosp = ExamenPrescrit.query.filter_by(
            hospitalisation_id=hosp.id,
            est_active=True
        ).all()
        
        hospitalisations_data.append({
            'hospitalisation': hosp,
            'medecins': medecins,
            'infirmiers': infirmiers,
            'constantes': constantes,
            'evolutions': evolutions,
            'avis_externes': avis_externes,
            'note_active': note_active,
            'ordonnance_medicaments': ordonnance_medicaments,
            'examens_prescrits': examens_hosp,
            'protocole': hosp.protocole
        })
    
    # ============================================================ #
    # 3. ENGAGEMENTS (DNR, REFUS_TRAITEMENT, SORTIE_AVIS, AUTRE)
    # ============================================================ #
    engagements = Engagement.query.filter_by(
        patient_id=patient.id
    ).order_by(Engagement.date_creation.desc()).all()
    
    # ============================================================ #
    # 4. ANTÉCÉDENTS
    # ============================================================ #
    antecedents = AntecedentPatient.query.filter_by(
        patient_id=patient.id,
        actif=True
    ).order_by(AntecedentPatient.date_recueil.desc()).all()
    
    age = None
    if patient.date_naissance:
        today = datetime.utcnow().date()
        age = today.year - patient.date_naissance.year - ((today.month, today.day) < (patient.date_naissance.month, patient.date_naissance.day))

    # ============================================================ #
    # 5. SOINS ADMINISTRÉS (journal de soins — actes réellement effectués)
    # ============================================================ #
    from models import ActePose
    soins_poses = ActePose.query.filter_by(
        patient_id=patient.id, statut='actif'
    ).order_by(ActePose.date_pose.desc()).all()

    return render_template('impressions/dossier_patient.html',
                         patient=patient,
                         age=age,
                         consultations_data=consultations_data,
                         hospitalisations_data=hospitalisations_data,
                         engagements=engagements,
                         antecedents=antecedents,
                         soins_poses=soins_poses,
                         now=datetime.utcnow())

@app.route('/hospitalisation/<int:id>/examen/ajouter', methods=['POST'])
@login_required
def ajouter_examen_prescrit(id):
    """Ajouter un examen prescrit à une hospitalisation"""
    from models import Hospitalisation, ExamenPrescrit, ExamenType, AnalyseDemande
    import json
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    examen_type_id = request.form.get('examen_type_id', type=int)
    nature = request.form.get('nature', '').strip()
    motif = request.form.get('motif', '').strip()
    description = request.form.get('description', '').strip()
    examens_json = request.form.get('examens_json', '[]')
    
    try:
        # Récupérer la liste des examens
        examens = json.loads(examens_json) if examens_json else []
        
        if not examens and not examen_type_id:
            flash('Veuillez ajouter au moins un examen', 'danger')
            return redirect(url_for('detail_hospitalisation', id=id))
        
        if examen_type_id:
            examen_type = ExamenType.query.get(examen_type_id)
            if examen_type:
                # Utiliser les données du template
                nature = examen_type.nature
                motif = examen_type.motif or motif
                description = examen_type.description or description
                examens = json.loads(examen_type.examens) if examen_type.examens else []
                examens_json = examen_type.examens
                
                examen_prescrit = ExamenPrescrit(
                    hospitalisation_id=hospitalisation.id,
                    patient_id=hospitalisation.patient_id,
                    medecin_id=current_user.id,
                    examen_type_id=examen_type.id,
                    version=1,
                    est_active=True,
                    nature=nature,
                    motif=motif,
                    description=description,
                    examens=examens_json,
                    statut='EN_ATTENTE',
                    date_prescription=datetime.utcnow()
                )
                db.session.add(examen_prescrit)
                db.session.flush()
        else:
            # Création manuelle
            examen_prescrit = ExamenPrescrit(
                hospitalisation_id=hospitalisation.id,
                patient_id=hospitalisation.patient_id,
                medecin_id=current_user.id,
                version=1,
                est_active=True,
                nature=nature,
                motif=motif,
                description=description,
                examens=examens_json,
                statut='EN_ATTENTE',
                date_prescription=datetime.utcnow()
            )
            db.session.add(examen_prescrit)
            db.session.flush()
        
        # ⭐⭐⭐ CRÉER LES ANALYSES DEMANDÉES (pour le laborantin) ⭐⭐⭐
        analyses_creees = 0
        try:
            for nom_analyse in examens:
                if nom_analyse and nom_analyse.strip():
                    analyse = AnalyseDemande(
                        hospitalisation_id=hospitalisation.id,
                        consultation_id=None,
                        patient_id=hospitalisation.patient_id,
                        structure_id=current_user.id_structure,
                        type_analyse=nature,  # BIOLOGIE, IMAGERIE, AUTRE
                        nom_analyse=nom_analyse.strip(),
                        description=description,
                        prescrit_par=current_user.id,
                        # ⭐ Lien fiable vers l'ExamenPrescrit d'origine — voir
                        # saisir_resultats_examen, qui matchait jusqu'ici sur
                        # nom_analyse==nature (ne matchait presque jamais).
                        examen_prescrit_id=examen_prescrit.id,
                        statut='EN_ATTENTE',
                        date_demande=datetime.utcnow()
                    )
                    db.session.add(analyse)
                    analyses_creees += 1
            print(f"✅ {analyses_creees} analyse(s) créée(s) pour le laborantin (hospitalisation)")
        except Exception as e:
            print(f"⚠️ Erreur création analyses: {e}")

        # ⭐ Miroir Prescription pour la synchronisation GHP
        _creer_prescriptions_miroir(
            patient_id=hospitalisation.patient_id,
            prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
            items=examens,
            type_prescription='acte'
        )

        db.session.commit()
        _envoyer_prescriptions_ghp_immediat()
        flash(f'✅ Examen prescrit et {analyses_creees} analyse(s) envoyée(s) au laborantin', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'❌ Erreur : {str(e)}', 'danger')
    
    return redirect(url_for('detail_hospitalisation', id=id))


@app.route('/hospitalisation/<int:id>/examen/<int:examen_id>/resultats', methods=['POST'])
@login_required
def saisir_resultats_examen(id, examen_id):
    """Saisir les résultats d'un examen prescrit"""
    from models import Hospitalisation, ExamenPrescrit, AnalyseDemande
    from datetime import datetime
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    examen = ExamenPrescrit.query.get_or_404(examen_id)
    
    # ⭐ PERMISSIONS : Laborantin, Radiologue, Medecin, Admin
    if current_user.role not in ['admin_structure', 'medecin', 'laborantin', 'radiologue']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    resultats = request.form.get('resultats', '').strip()
    statut = request.form.get('statut', 'TERMINE')
    
    if not resultats:
        flash('Veuillez saisir les résultats', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    # ⭐ METTRE À JOUR L'EXAMEN PRESCRIT
    examen.resultats = resultats
    examen.statut = statut
    examen.laborantin_id = current_user.id
    examen.date_resultats = datetime.utcnow()
    examen.updated_at = datetime.utcnow()
    
    # ⭐ METTRE À JOUR LES ANALYSES DEMANDÉES CORRESPONDANTES — matching par
    # examen_prescrit_id (fiable) plutôt que l'ancien nom_analyse==nature
    # (qui comparait un nom d'examen à "BIOLOGIE"/"IMAGERIE", donc ne
    # matchait quasiment jamais). Un ExamenPrescrit peut porter plusieurs
    # examens (un par AnalyseDemande liée) — même résultat reporté partout.
    analyses = AnalyseDemande.query.filter_by(examen_prescrit_id=examen.id).all()
    for analyse in analyses:
        analyse.resultats = resultats
        analyse.statut = statut
        analyse.date_resultats = datetime.utcnow()
        analyse.resultats_par = current_user.id
    
    # ⭐ METTRE À JOUR LES RÉSULTATS DANS L'HOSPITALISATION
    if statut == 'TERMINE':
        if examen.nature == 'BIOLOGIE':
            if hospitalisation.notes_admission:
                hospitalisation.notes_admission += f"\n\n--- RÉSULTATS BIOLOGIE ---\n{resultats}"
            else:
                hospitalisation.notes_admission = f"--- RÉSULTATS BIOLOGIE ---\n{resultats}"
        elif examen.nature == 'IMAGERIE':
            if hospitalisation.notes_admission:
                hospitalisation.notes_admission += f"\n\n--- RÉSULTATS IMAGERIE ---\n{resultats}"
            else:
                hospitalisation.notes_admission = f"--- RÉSULTATS IMAGERIE ---\n{resultats}"
    
    db.session.commit()
    
    flash('✅ Résultats enregistrés avec succès', 'success')
    return redirect(url_for('detail_hospitalisation', id=id))


@app.route('/hospitalisation/<int:id>/examen/<int:examen_id>/supprimer', methods=['POST'])
@login_required
def supprimer_examen_prescrit(id, examen_id):
    """Supprimer un examen prescrit"""
    from models import ExamenPrescrit
    
    examen = ExamenPrescrit.query.get_or_404(examen_id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    db.session.delete(examen)
    db.session.commit()
    
    flash('Examen supprimé avec succès', 'success')
    return redirect(url_for('detail_hospitalisation', id=id))


@app.route('/hospitalisation/<int:id>/ordonnance/imprimer')
@login_required
def imprimer_ordonnance(id):
    """Imprimer l'ordonnance médicale"""
    from models import Hospitalisation, Structure
    import json
    from datetime import datetime
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    structure = Structure.query.get(current_user.id_structure)
    
    if not hospitalisation.ordonnance_prescite:
        flash('Aucune ordonnance à imprimer', 'warning')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    medicaments = json.loads(hospitalisation.ordonnance_prescite) if hospitalisation.ordonnance_prescite else []
    
    return render_template('impressions/ordonnance.html',
                         hospitalisation=hospitalisation,
                         structure=structure,
                         patient=hospitalisation.patient,  # ⭐ PASSER LE PATIENT
                         medicaments=medicaments,
                         now=datetime.utcnow())


@app.route('/hospitalisation/<int:id>/examen/<int:examen_id>/imprimer')
@login_required
def imprimer_examen(id, examen_id):
    """Imprimer la demande d'examens"""
    from models import Hospitalisation, ExamenPrescrit, Structure
    import json
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    examen = ExamenPrescrit.query.get_or_404(examen_id)
    structure = Structure.query.get(current_user.id_structure)
    
    examens_list = json.loads(examen.examens) if examen.examens else []
    
    return render_template('impressions/examen.html',
                         hospitalisation=hospitalisation,
                         structure=structure,
                         examen=examen,
                         examens_list=examens_list,
                         now=datetime.utcnow())

# ================================================================
# API - TEMPLATES (pour sélection dynamique)
# ================================================================

@app.route('/api/protocole/<int:id>')
@login_required
def api_get_protocole(id):
    """Récupère les détails d'un protocole pour l'aperçu dynamique"""
    from models import ProtocoleSoins
    
    protocole = ProtocoleSoins.query.get_or_404(id)
    
    if protocole.structure_id != current_user.id_structure:
        return jsonify({'error': 'Accès non autorisé'}), 403
    
    return jsonify({
        'id': protocole.id,
        'nom': protocole.nom,
        'description': protocole.description,
        'ordonnance_type': protocole.ordonnance_type.nom if protocole.ordonnance_type else None,
        'examen_type': protocole.examen_type.nom if protocole.examen_type else None
    })


@app.route('/api/examen-type/<int:id>')
@login_required
def api_get_examen_type(id):
    """Récupère les détails d'un examen type pour le pré-remplissage"""
    from models import ExamenType
    import json
    
    examen_type = ExamenType.query.get_or_404(id)
    
    if examen_type.structure_id != current_user.id_structure:
        return jsonify({'error': 'Accès non autorisé'}), 403
    
    examens_list = json.loads(examen_type.examens) if examen_type.examens else []
    
    return jsonify({
        'id': examen_type.id,
        'nom': examen_type.nom,
        'nature': examen_type.nature,
        'motif': examen_type.motif,
        'description': examen_type.description,
        'examens': examens_list
    })


@app.route('/api/ordonnance-type/<int:id>')
@login_required
def api_get_ordonnance_type(id):
    """Récupère les détails d'une ordonnance type pour le pré-remplissage"""
    from models import OrdonnanceType
    import json
    
    ordonnance = OrdonnanceType.query.get_or_404(id)
    
    if ordonnance.structure_id != current_user.id_structure:
        return jsonify({'error': 'Accès non autorisé'}), 403
    
    medicaments = json.loads(ordonnance.medicaments) if ordonnance.medicaments else []
    
    return jsonify({
        'id': ordonnance.id,
        'nom': ordonnance.nom,
        'description': ordonnance.description,
        'medicaments': medicaments
    })


@app.route('/api/hospitalisation/<int:id>/examen/<int:examen_id>')
@login_required
def api_get_examen_prescrit(id, examen_id):
    """Récupère les détails d'un examen prescrit"""
    from models import Hospitalisation, ExamenPrescrit
    import json
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    examen = ExamenPrescrit.query.get_or_404(examen_id)
    
    if hospitalisation.patient.id_structure != current_user.id_structure:
        return jsonify({'error': 'Accès non autorisé'}), 403
    
    examens_list = json.loads(examen.examens) if examen.examens else []
    
    return jsonify({
        'id': examen.id,
        'nature': examen.nature,
        'motif': examen.motif,
        'description': examen.description,
        'examens': examens_list,
        'statut': examen.statut,
        'resultats': examen.resultats,
        'date_prescription': examen.date_prescription.strftime('%d/%m/%Y %H:%M') if examen.date_prescription else None
    })

# ================================================================
# CONSULTATION - ORDONNANCES (AVEC HISTORIQUE)
# ================================================================

@app.route('/consultation/<int:id>/ordonnance/creer', methods=['GET', 'POST'])
@login_required
def creer_ordonnance_consultation(id):
    """Créer une nouvelle ordonnance pour une consultation"""
    from models import Consultation, Ordonnance, OrdonnanceType, ProtocoleSoins
    import json
    
    consultation = Consultation.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('consultation_detail', id=id))
    
    # Récupérer les templates et protocoles disponibles
    ordonnances_types = OrdonnanceType.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    protocoles = ProtocoleSoins.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    if request.method == 'POST':
        source_type = request.form.get('source_type', 'manuel')
        source_id = request.form.get('source_id', type=int)
        medicaments_json = request.form.get('medicaments_json', '[]')
        motif_modification = request.form.get('motif_modification', 'Création')
        
        # Valider le JSON
        try:
            medicaments = json.loads(medicaments_json)
        except:
            flash('Format des médicaments invalide', 'danger')
            return redirect(url_for('creer_ordonnance_consultation', id=id))
        
        if not medicaments:
            flash('Veuillez ajouter au moins un médicament', 'danger')
            return redirect(url_for('creer_ordonnance_consultation', id=id))
        
        # Compter les ordonnances existantes
        nb_ordonnances = consultation.ordonnances.count()
        nouvelle_version = nb_ordonnances + 1
        
        # Récupérer le nom de la source
        source_nom = None
        if source_type == 'template' and source_id:
            template = OrdonnanceType.query.get(source_id)
            if template:
                source_nom = template.nom
        elif source_type == 'protocole' and source_id:
            protocole = ProtocoleSoins.query.get(source_id)
            if protocole:
                source_nom = protocole.nom
                
                # Si c'est un protocole, appliquer aussi les examens
                if protocole.examen_type_id:
                    from models import ExamenPrescrit, ExamenType
                    examen_type = ExamenType.query.get(protocole.examen_type_id)
                    if examen_type:
                        examen_prescrit = ExamenPrescrit(
                            consultation_id=consultation.id,
                            patient_id=consultation.id_patient,
                            medecin_id=current_user.id,
                            examen_type_id=examen_type.id,
                            version=1,
                            est_active=True,
                            nature=examen_type.nature,
                            motif=examen_type.motif,
                            description=examen_type.description,
                            examens=examen_type.examens,
                            source_type='protocole',
                            source_id=protocole.id,
                            source_nom=protocole.nom,
                            statut='EN_ATTENTE',
                            date_prescription=datetime.utcnow()
                        )
                        db.session.add(examen_prescrit)
                        db.session.flush()
                        try:
                            examens_protocole_pour_ghp = json.loads(examen_type.examens) if examen_type.examens else []
                        except Exception:
                            examens_protocole_pour_ghp = []
                        # ⭐ Pont vers /analyses (labo/radio) + synchro GHP pour
                        # cet examen — jusqu'ici, un protocole appliqué via
                        # "créer ordonnance" ne faisait NI l'un NI l'autre pour
                        # l'examen (seuls les médicaments de l'ordonnance,
                        # plus bas, étaient synchronisés).
                        _creer_analyses_demandees_depuis_examen_prescrit(
                            examen_prescrit, examens_protocole_pour_ghp, current_user.id_structure
                        )
                        _creer_prescriptions_miroir(
                            patient_id=consultation.id_patient,
                            prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
                            items=examens_protocole_pour_ghp,
                            type_prescription='acte',
                            id_consultation=consultation.id
                        )
        
        # Désactiver les anciennes ordonnances
        consultation.ordonnances.update({'est_active': False})
        
        # Créer la nouvelle ordonnance
        ordonnance = Ordonnance(
            consultation_id=consultation.id,
            version=nouvelle_version,
            est_active=True,
            medicaments=medicaments_json,
            source_type=source_type,
            source_id=source_id,
            source_nom=source_nom,
            redige_par=current_user.id,
            date_redaction=datetime.utcnow(),
            motif_modification=motif_modification
        )
        
        db.session.add(ordonnance)
        db.session.flush()

        # Mettre à jour la consultation
        consultation.ordonnance_active_id = ordonnance.id

        # Si source_type est 'protocole', enregistrer le protocole
        if source_type == 'protocole' and source_id:
            consultation.protocole_applique_id = source_id

        # ⭐ Miroir Prescription pour la synchronisation GHP (voir commentaire
        # au-dessus de _creer_prescriptions_miroir) — sans ça, cette
        # ordonnance ne partait jamais vers "Prescriptions reçues" côté GHP.
        _creer_prescriptions_miroir(
            patient_id=consultation.id_patient,
            prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
            items=medicaments,
            type_prescription='medicament',
            id_consultation=consultation.id
        )

        db.session.commit()
        _envoyer_prescriptions_ghp_immediat()

        flash(f'✅ Ordonnance (version {nouvelle_version}) créée avec succès', 'success')
        return redirect(url_for('consultation_detail', id=id))
    
    return render_template('consultations/modals/ordonnance_creer.html',
                         consultation=consultation,
                         ordonnances_types=ordonnances_types,
                         protocoles=protocoles,
                         now=datetime.utcnow())


@app.route('/consultation/<int:id>/ordonnance/<int:ordonnance_id>/modifier', methods=['GET', 'POST'])
@login_required
def modifier_ordonnance_consultation(id, ordonnance_id):
    """Modifier une ordonnance existante (crée une nouvelle version)"""
    from models import Consultation, Ordonnance
    import json
    
    consultation = Consultation.query.get_or_404(id)
    ordonnance_old = Ordonnance.query.get_or_404(ordonnance_id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('consultation_detail', id=id))
    
    if request.method == 'POST':
        medicaments_json = request.form.get('medicaments_json', '[]')
        motif_modification = request.form.get('motif_modification', 'Modification')
        
        try:
            medicaments = json.loads(medicaments_json)
        except:
            flash('Format des médicaments invalide', 'danger')
            return redirect(url_for('consultation_detail', id=id))
        
        if not medicaments:
            flash('Veuillez ajouter au moins un médicament', 'danger')
            return redirect(url_for('consultation_detail', id=id))
        
        # Compter les ordonnances existantes
        nb_ordonnances = consultation.ordonnances.count()
        nouvelle_version = nb_ordonnances + 1
        
        # Désactiver les anciennes ordonnances
        consultation.ordonnances.update({'est_active': False})
        
        # Créer la nouvelle ordonnance
        ordonnance = Ordonnance(
            consultation_id=consultation.id,
            version=nouvelle_version,
            est_active=True,
            medicaments=medicaments_json,
            source_type='modification',
            source_id=ordonnance_old.id,
            source_nom=f"Version {ordonnance_old.version}",
            redige_par=current_user.id,
            date_redaction=datetime.utcnow(),
            motif_modification=motif_modification
        )
        
        db.session.add(ordonnance)
        db.session.flush()

        # Mettre à jour la consultation
        consultation.ordonnance_active_id = ordonnance.id

        # ⭐ Miroir Prescription pour GHP — uniquement les médicaments
        # réellement nouveaux par rapport à l'ancienne version, pour ne pas
        # renvoyer en double ce qui a déjà été prescrit/synchronisé.
        anciens_medicaments = ordonnance_old.get_medicaments_list()
        nouveaux = _items_nouveaux(medicaments, anciens_medicaments)
        if nouveaux:
            _creer_prescriptions_miroir(
                patient_id=consultation.id_patient,
                prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
                items=nouveaux,
                type_prescription='medicament',
                id_consultation=consultation.id
            )

        db.session.commit()
        if nouveaux:
            _envoyer_prescriptions_ghp_immediat()

        flash(f'✅ Ordonnance modifiée (version {nouvelle_version})', 'success')
        return redirect(url_for('consultation_detail', id=id))
    
    return render_template('consultations/modals/ordonnance_modifier.html',
                         consultation=consultation,
                         ordonnance=ordonnance_old,
                         now=datetime.utcnow())


@app.route('/consultation/<int:id>/ordonnance/imprimer')
@login_required
def imprimer_ordonnance_consultation(id):
    """Imprimer l'ordonnance active"""
    from models import Consultation, Structure
    import json
    
    consultation = Consultation.query.get_or_404(id)
    structure = Structure.query.get(current_user.id_structure)
    patient = consultation.patient
    
    ordonnance_active = consultation.ordonnance_active
    
    if not ordonnance_active:
        flash('Aucune ordonnance active à imprimer', 'warning')
        return redirect(url_for('consultation_detail', id=id))
    
    medicaments = ordonnance_active.get_medicaments_list()
    
    return render_template('impressions/ordonnance_consultation.html',
                         consultation=consultation,
                         patient=patient,
                         structure=structure,
                         ordonnance=ordonnance_active,
                         medicaments=medicaments,
                         now=datetime.utcnow())


@app.route('/consultation/<int:id>/ordonnance/historique')
@login_required
def historique_ordonnances_consultation(id):
    from models import Consultation, Ordonnance  # ⭐ AJOUTE CET IMPORT
    
    consultation = Consultation.query.get_or_404(id)
    
    ordonnances = consultation.ordonnances.order_by(
        Ordonnance.version.desc()
    ).all()
    
    return render_template('consultations/historique_ordonnances.html',
                         consultation=consultation,
                         ordonnances=ordonnances)

# ================================================================
# CONSULTATION - EXAMENS PRESCRITS
# ================================================================

@app.route('/consultation/<int:id>/examen/ajouter', methods=['POST'])
@login_required
def ajouter_examen_prescrit_consultation(id):
    from models import Consultation, ExamenPrescrit, ExamenType
    import json
    
    consultation = Consultation.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('consultation_detail', id=id))
    
    examen_type_id = request.form.get('examen_type_id', type=int)
    nature = request.form.get('nature', '').strip()
    motif = request.form.get('motif', '').strip()
    description = request.form.get('description', '').strip()
    examens_json = request.form.get('examens_json', '[]')
    
    try:
        examens_pour_ghp = []
        if examen_type_id:
            examen_type = ExamenType.query.get(examen_type_id)
            if examen_type:
                examen_prescrit = ExamenPrescrit(
                    # ⭐ hospitalisation_id = None pour consultation
                    hospitalisation_id=None,
                    consultation_id=consultation.id,
                    patient_id=consultation.id_patient,
                    medecin_id=current_user.id,
                    examen_type_id=examen_type.id,
                    nature=examen_type.nature,
                    motif=examen_type.motif,
                    description=examen_type.description,
                    examens=examen_type.examens,
                    statut='EN_ATTENTE',
                    date_prescription=datetime.utcnow()
                )
                db.session.add(examen_prescrit)
                try:
                    examens_pour_ghp = json.loads(examen_type.examens) if examen_type.examens else []
                except Exception:
                    examens_pour_ghp = []
        else:
            # Création manuelle
            examens = json.loads(examens_json) if examens_json else []
            examen_prescrit = ExamenPrescrit(
                # ⭐ hospitalisation_id = None pour consultation
                hospitalisation_id=None,
                consultation_id=consultation.id,
                patient_id=consultation.id_patient,
                medecin_id=current_user.id,
                nature=nature,
                motif=motif,
                description=description,
                examens=examens_json,
                statut='EN_ATTENTE',
                date_prescription=datetime.utcnow()
            )
            db.session.add(examen_prescrit)
            examens_pour_ghp = examens

        db.session.flush()
        # ⭐ Pont vers /analyses (labo/radio) — jusqu'ici ce chemin
        # (consultation, prescription directe) mirait bien vers GHP mais
        # ne créait aucune AnalyseDemande, donc restait invisible pour
        # laborantin/radiologue.
        _creer_analyses_demandees_depuis_examen_prescrit(
            examen_prescrit, examens_pour_ghp, current_user.id_structure
        )

        # ⭐ Miroir Prescription pour la synchronisation GHP
        _creer_prescriptions_miroir(
            patient_id=consultation.id_patient,
            prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
            items=examens_pour_ghp,
            type_prescription='acte',
            id_consultation=consultation.id
        )

        db.session.commit()
        _envoyer_prescriptions_ghp_immediat()
        flash('Examen prescrit ajoute avec succès', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erreur : {str(e)}', 'danger')

    return redirect(url_for('consultation_detail', id=id))


@app.route('/consultation/<int:id>/examen/<int:examen_id>/supprimer', methods=['POST'])
@login_required
def supprimer_examen_prescrit_consultation(id, examen_id):
    """Supprimer un examen prescrit d'une consultation"""
    from models import ExamenPrescrit
    
    examen = ExamenPrescrit.query.get_or_404(examen_id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('consultation_detail', id=id))
    
    db.session.delete(examen)
    db.session.commit()
    
    flash('✅ Examen supprimé avec succès', 'success')
    return redirect(url_for('consultation_detail', id=id))


# ================================================================
# CONSULTATION - EXAMENS PRESCRITS (AVEC HISTORIQUE)
# ================================================================

@app.route('/consultation/<int:id>/examen/creer', methods=['GET', 'POST'])
@login_required
def creer_examen_consultation(id):
    """Créer une nouvelle demande d'examens pour une consultation"""
    from models import Consultation, ExamenPrescrit, ExamenType, ProtocoleSoins, AnalyseDemande
    import json
    
    consultation = Consultation.query.get_or_404(id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('consultation_detail', id=id))
    
    examens_types = ExamenType.query.filter_by(
        structure_id=current_user.id_structure,
        actif=True
    ).all()
    
    if request.method == 'POST':
        source_type = request.form.get('source_type', 'manuel')
        source_id = request.form.get('source_id', type=int)
        
        nature = request.form.get('nature', '').strip()
        motif = request.form.get('motif', '').strip()
        description = request.form.get('description', '').strip()
        examens_json = request.form.get('examens_json', '[]')
        
        try:
            examens = json.loads(examens_json)
        except:
            flash('Format des examens invalide', 'danger')
            return redirect(url_for('creer_examen_consultation', id=id))
        
        if not examens:
            flash('Veuillez ajouter au moins un examen', 'danger')
            return redirect(url_for('creer_examen_consultation', id=id))
        
        # Récupérer le nom de la source
        source_nom = None
        if source_type == 'template' and source_id:
            template = ExamenType.query.get(source_id)
            if template:
                source_nom = template.nom
                nature = template.nature
                motif = template.motif
                description = template.description
        elif source_type == 'protocole' and source_id:
            protocole = ProtocoleSoins.query.get(source_id)
            if protocole and protocole.examen_type_id:
                template = ExamenType.query.get(protocole.examen_type_id)
                if template:
                    source_nom = protocole.nom
                    nature = template.nature
                    motif = template.motif
                    description = template.description
        
        # Compter les examens existants
        nb_examens = consultation.examens_prescrits.count()
        nouvelle_version = nb_examens + 1
        
        # Créer l'examen prescrit
        examen_prescrit = ExamenPrescrit(
            consultation_id=consultation.id,
            patient_id=consultation.id_patient,
            medecin_id=current_user.id,
            version=nouvelle_version,
            est_active=True,
            nature=nature,
            motif=motif,
            description=description,
            examens=examens_json,
            source_type=source_type,
            source_id=source_id,
            source_nom=source_nom,
            statut='EN_ATTENTE',
            date_prescription=datetime.utcnow()
        )
        db.session.add(examen_prescrit)
        db.session.flush()  # Pour obtenir l'ID si besoin
        
        # ⭐⭐⭐ CRÉER LES ANALYSES DEMANDÉES (pour le laborantin) ⭐⭐⭐
        analyses_creees = 0
        try:
            for nom_analyse in examens:
                if nom_analyse.strip():
                    analyse = AnalyseDemande(
                        consultation_id=consultation.id,
                        patient_id=consultation.id_patient,
                        structure_id=current_user.id_structure,
                        type_analyse=nature,  # BIOLOGIE, IMAGERIE, AUTRE
                        nom_analyse=nom_analyse.strip(),
                        description=description,
                        prescrit_par=current_user.id,
                        # ⭐ Lien fiable vers l'ExamenPrescrit d'origine — voir
                        # saisir_resultats_examen.
                        examen_prescrit_id=examen_prescrit.id,
                        statut='EN_ATTENTE',
                        date_demande=datetime.utcnow()
                    )
                    db.session.add(analyse)
                    analyses_creees += 1
            print(f"✅ {analyses_creees} analyse(s) créée(s) pour le laborantin")
        except Exception as e:
            print(f"⚠️ Erreur création analyses: {e}")

        # ⭐ Miroir Prescription pour la synchronisation GHP
        _creer_prescriptions_miroir(
            patient_id=consultation.id_patient,
            prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
            items=examens,
            type_prescription='acte',
            id_consultation=consultation.id
        )

        db.session.commit()
        _envoyer_prescriptions_ghp_immediat()

        flash(f'✅ Demande d\'examens créée avec succès ({analyses_creees} analyse(s) envoyée(s) au laborantin)', 'success')
        return redirect(url_for('consultation_detail', id=id))
    
    return render_template('consultations/modals/examen_creer.html',
                         consultation=consultation,
                         examens_types=examens_types,
                         now=datetime.utcnow())

# ================================================================
# CONSULTATION - EXAMENS (AVEC HISTORIQUE)
# ================================================================

@app.route('/consultation/<int:id>/examen/<int:examen_id>/modifier', methods=['GET', 'POST'])
@login_required
def modifier_examen_consultation(id, examen_id):
    """Modifier un examen prescrit (crée une nouvelle version)"""
    from models import Consultation, ExamenPrescrit
    import json
    
    consultation = Consultation.query.get_or_404(id)
    examen_old = ExamenPrescrit.query.get_or_404(examen_id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('consultation_detail', id=id))
    
    if request.method == 'POST':
        nature = request.form.get('nature', '').strip()
        motif = request.form.get('motif', '').strip()
        description = request.form.get('description', '').strip()
        examens_json = request.form.get('examens_json', '[]')
        
        try:
            examens = json.loads(examens_json)
        except:
            flash('Format des examens invalide', 'danger')
            return redirect(url_for('consultation_detail', id=id))
        
        if not examens:
            flash('Veuillez ajouter au moins un examen', 'danger')
            return redirect(url_for('consultation_detail', id=id))
        
        # Compter les examens existants
        nb_examens = consultation.examens_prescrits.count()
        nouvelle_version = nb_examens + 1
        
        # Désactiver les anciens examens
        consultation.examens_prescrits.update({'est_active': False})
        
        # Creer le nouvel examen
        examen = ExamenPrescrit(
            consultation_id=consultation.id,
            patient_id=consultation.id_patient,
            medecin_id=current_user.id,
            version=nouvelle_version,
            est_active=True,
            nature=nature,
            motif=motif,
            description=description,
            examens=examens_json,
            source_type='modification',
            source_id=examen_old.id,
            source_nom=f"Version {examen_old.version}",
            statut='EN_ATTENTE',
            date_prescription=datetime.utcnow()
        )
        
        db.session.add(examen)
        db.session.flush()

        # ⭐ Miroir Prescription pour GHP — uniquement les examens réellement
        # nouveaux par rapport à l'ancienne version.
        anciens_examens = examen_old.get_examens_list()
        nouveaux = _items_nouveaux(examens, anciens_examens)
        if nouveaux:
            # ⭐ Pont vers /analyses (labo/radio) — même principe que le
            # miroir GHP juste en dessous : uniquement les items réellement
            # nouveaux, pour ne pas créer un doublon d'AnalyseDemande pour
            # un examen déjà présent dans la version précédente.
            _creer_analyses_demandees_depuis_examen_prescrit(
                examen, nouveaux, current_user.id_structure
            )
            _creer_prescriptions_miroir(
                patient_id=consultation.id_patient,
                prescripteur_nom=f"{current_user.prenom} {current_user.nom}",
                items=nouveaux,
                type_prescription='acte',
                id_consultation=consultation.id
            )

        db.session.commit()
        if nouveaux:
            _envoyer_prescriptions_ghp_immediat()

        flash(f'Demande d\'examens modifiee (version {nouvelle_version})', 'success')
        return redirect(url_for('consultation_detail', id=id))
    
    return render_template('consultations/modals/examen_modifier.html',
                         consultation=consultation,
                         examen=examen_old)

@app.route('/consultation/<int:id>/examen/<int:examen_id>/imprimer')
@login_required
def imprimer_examen_consultation(id, examen_id):
    """Imprimer une demande d'examens"""
    from models import Consultation, ExamenPrescrit, Structure
    import json
    
    consultation = Consultation.query.get_or_404(id)
    examen = ExamenPrescrit.query.get_or_404(examen_id)
    structure = Structure.query.get(current_user.id_structure)
    patient = consultation.patient
    
    examens_list = examen.get_examens_list()
    
    return render_template('impressions/examen_consultation.html',
                         consultation=consultation,
                         patient=patient,
                         structure=structure,
                         examen=examen,
                         examens_list=examens_list,
                         now=datetime.utcnow())
@app.route('/hospitalisation/<int:id>/reference/ajouter', methods=['GET', 'POST'])
@login_required
def ajouter_reference_hospitalisation(id):
    """Ajouter une référence depuis une hospitalisation"""
    from models import Hospitalisation, Reference, Patient
    
    hospitalisation = Hospitalisation.query.get_or_404(id)
    patient = Patient.query.get(hospitalisation.patient_id)
    
    if current_user.role not in ['admin_structure', 'medecin']:
        flash('Accès non autorisé', 'danger')
        return redirect(url_for('detail_hospitalisation', id=id))
    
    if request.method == 'POST':
        motif = request.form.get('motif')
        diagnostic = request.form.get('diagnostic')
        centre_reference = request.form.get('centre_reference')
        service_reference = request.form.get('service_reference')
        medecin_referent = request.form.get('medecin_referent')
        resume_clinique = request.form.get('resume_clinique')
        examens_realises = request.form.get('examens_realises')
        traitements_en_cours = request.form.get('traitements_en_cours')
        
        if not motif or not centre_reference:
            flash('Le motif et le centre de référence sont obligatoires', 'danger')
            return redirect(url_for('ajouter_reference_hospitalisation', id=id))
        
        reference = Reference(
            patient_id=patient.id,
            hospitalisation_id=hospitalisation.id,
            structure_id=current_user.id_structure,
            motif=motif,
            diagnostic=diagnostic or hospitalisation.motif,
            centre_reference=centre_reference,
            service_reference=service_reference,
            medecin_referent=medecin_referent,
            derniere_tension=patient.tension_arterielle,
            derniere_temperature=patient.temperature_c,
            derniere_pulse=patient.pulse_bpm,
            derniere_saturation=patient.oxygene_saturation,
            dernier_poids=patient.poids_kg,
            derniere_taille=patient.taille_cm,
            dernier_imc=patient.imc,
            resume_clinique=resume_clinique,
            examens_realises=examens_realises,
            traitements_en_cours=traitements_en_cours,
            statut='ENVOYE',
            created_by=current_user.id
        )
        
        db.session.add(reference)
        db.session.commit()
        
        flash('✅ Référence créée avec succès', 'success')
        return redirect(url_for('imprimer_reference', id=reference.id))
    
    return render_template('hospitalisations/ajouter_reference.html',
                         hospitalisation=hospitalisation,
                         patient=patient)






if __name__ == '__main__':
    app.run(debug=True)